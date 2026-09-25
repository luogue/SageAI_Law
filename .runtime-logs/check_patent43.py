import json
from pathlib import Path
from docx import Document
from review_v2.engine import review_contract_v2
from review_v2.user_profile import UserRiskProfile

doc = Document(next(Path('F:/ai/Sage/test_data').glob('43*.docx')))
text = '\n'.join(p.text for p in doc.paragraphs) + '\n' + '\n'.join('\t'.join(c.text for c in r.cells) for t in doc.tables for r in t.rows)
result = review_contract_v2(text, 'technology', use_llm=False, include_low_risk=True, profile=UserRiskProfile(user_identity='受让方'))
print(json.dumps({'release': result.get('release_quality_gate'), 'role': result.get('perspective_resolution'), 'items': [{k: i.get(k) for k in ('item_id','coverage_status','risk_level','risk_group','reasoning_summary','coverage_elements','action_priority')} for i in result['items'] if i['item_id'] in ('TC-06-01','TC-06-02','TC-06-03','TC-06-04','TC-12-01')], 'findings': [{k: f.get(k) for k in ('risk_title','review_ids','severity','final_action')} for f in result['final_dispositions']]}, ensure_ascii=False, indent=2))
