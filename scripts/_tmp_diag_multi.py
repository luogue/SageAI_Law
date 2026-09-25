"""多候选诊断：对 93/33/22/16 四份技术咨询合同跑引擎，找出复现"假 MISSING"的样本。"""
import io
import hashlib
import sys
from pathlib import Path

ROOT = Path(r"F:/ai/Sage")
sys.path.insert(0, str(ROOT / "backend"))
sys.path.insert(0, str(ROOT))

from docx import Document
from review_v2.engine import review_contract_v2
from review_v2.user_profile import UserRiskProfile

OUT = ROOT / "scripts" / "_tmp_diag_multi.out.txt"

CANDIDATES = [
    "93 技术咨询合同.docx",
    "33 技术咨询服务合同.docx",
    "22 技术咨询.docx",
    "16 技术咨询合同.docx",
]

TARGETS = {
    "TC-01-03", "TC-08-01", "TC-08-02", "TC-08-03", "TC-08-04",
    "TC-11-01", "TC-11-02", "TC-11-03", "TC-11-04", "TC-12-02",
}

buf = io.StringIO()


def w(s=""):
    buf.write(s + "\n")


def docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for tb in doc.tables:
        for row in tb.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts)


for name in CANDIDATES:
    p = ROOT / "test_data" / name
    if not p.exists():
        w(f"\n### {name} 不存在")
        continue
    TEXT = docx_text(p)
    sha = hashlib.sha256(TEXT.encode("utf-8")).hexdigest()
    w("\n" + "#" * 70)
    w(f"### {name}   len={len(TEXT)}  sha256={sha[:16]}")
    w("#" * 70)
    result = review_contract_v2(
        TEXT, "technology", use_llm=False, include_low_risk=True,
        profile=UserRiskProfile(user_identity="甲方"),
    )
    gate = result.get("release_quality_gate", {})
    w(f"gate={gate.get('status')} failures={len(gate.get('failures') or [])}")
    for f in (gate.get("failures") or []):
        w(f"   FAIL: {f}")
    total_conflict = 0
    for item in result.get("items", []):
        iid = item.get("item_id")
        if iid not in TARGETS:
            continue
        covered = {r["element_id"] for r in item.get("coverage_elements", []) if r.get("status") == "COVERED"}
        trace = item.get("decision_trace", {}) or {}
        missing = {r["element_id"] for r in trace.get("element_absence_searches", []) if r.get("result") == "NOT_FOUND"}
        conflict = covered & missing
        status = item.get("status")
        w(f"\n [{iid}] {status}")
        w(f"    COVERED: {sorted(covered)}")
        w(f"    NOT_FOUND: {sorted(missing)}")
        if conflict:
            total_conflict += len(conflict)
            w(f"    *** 冲突(假MISSING): {sorted(conflict)}")
    w(f"\n >>> {name} 假MISSING冲突总数: {total_conflict}")

OUT.write_text(buf.getvalue(), encoding="utf-8")
print("written", OUT)
