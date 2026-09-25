"""临时排查2：列出技术咨询类样本修改时间+完整文本，输出到 utf-8 文件。"""
import io
from pathlib import Path

ROOT = Path(r"F:/ai/Sage")
OUT = ROOT / "scripts" / "_tmp_inspect2.out.txt"

def docx_text(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    # 表格
    for tb in doc.tables:
        for row in tb.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts)

buf = io.StringIO()

def w(s=""):
    buf.write(s + "\n")

w("=" * 70)
w("A. test_data 技术咨询/服务类样本（按修改时间）")
w("=" * 70)
rows = []
for p in sorted((ROOT / "test_data").glob("*.docx")):
    n = p.name
    if any(k in n for k in ["技术咨询", "技术服务"]):
        rows.append((p.stat().st_mtime, n, p.stat().st_size))
for mtime, n, size in sorted(rows, key=lambda r: -r[0]):
    import datetime
    ts = datetime.datetime.fromtimestamp(mtime).strftime("%Y-%m-%d %H:%M")
    w(f"{ts}  {size:>8}  {n}")

for name in ["33 技术咨询服务合同.docx", "22 技术咨询.docx", "16 技术咨询合同.docx"]:
    p = ROOT / "test_data" / name
    w("\n" + "=" * 70)
    w(f"B. 全文: {name}")
    w("=" * 70)
    if p.exists():
        try:
            w(docx_text(p))
        except Exception as e:
            w(f"(失败 {e})")

OUT.write_text(buf.getvalue(), encoding="utf-8")
print("written", OUT)
