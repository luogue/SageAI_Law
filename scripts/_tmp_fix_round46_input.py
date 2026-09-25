"""固定 Round46 输入：33 技术咨询合同 -> artifact + 测试 fixture。"""
import hashlib
import shutil
from pathlib import Path

ROOT = Path(r"F:/ai/Sage")
SRC = ROOT / "test_data" / "33 技术咨询服务合同.docx"
ART = ROOT / "artifacts" / "technology_pass_semantics" / "round46_consulting_generalization"
FIX = ROOT / "backend" / "tests" / "fixtures" / "round46_consulting_generalization.txt"


def docx_text(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for tb in doc.tables:
        for row in tb.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts)


ART.mkdir(parents=True, exist_ok=True)
dst = ART / "input.docx"
if not dst.exists():
    shutil.copy2(SRC, dst)

TEXT = docx_text(SRC)
(ART / "input.txt").write_text(TEXT, encoding="utf-8")
FIX.write_text(TEXT, encoding="utf-8")

sha_docx = hashlib.sha256(SRC.read_bytes()).hexdigest()
sha_txt = hashlib.sha256(TEXT.encode("utf-8")).hexdigest()

meta = {
    "round": "Round46",
    "input_source": "test_data/33 技术咨询服务合同.docx",
    "input_docx_sha256": sha_docx,
    "input_text_sha256": sha_txt,
    "artifact_input": str(dst.relative_to(ROOT)).replace("\\", "/"),
    "fixture": "backend/tests/fixtures/round46_consulting_generalization.txt",
    "text_length": len(TEXT),
}
import json
(ART / "input-manifest.json").write_text(
    json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")

print("docx sha:", sha_docx)
print("text sha:", sha_txt)
print("fixture written:", FIX)
print("manifest:", ART / "input-manifest.json")
