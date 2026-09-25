"""Round46 复现诊断：固定输入 93 技术咨询合同，输出各 TC 的 element/coverage/SEARCH_AUDIT 状态。

只做诊断，不修改任何业务代码。
"""
import io
import hashlib
import shutil
import sys
from pathlib import Path

ROOT = Path(r"F:/ai/Sage")
sys.path.insert(0, str(ROOT / "backend"))
sys.path.insert(0, str(ROOT))

from docx import Document
from review_v2.clause_parser import parse_contract_clauses
from review_v2.consulting_element_facts import elements as consulting_elements
from review_v2.technology_service_elements import elements as technology_elements
from review_v2.confidentiality_remedies import canonical_confidentiality_remedies
from review_v2.engine import review_contract_v2
from review_v2.user_profile import UserRiskProfile

OUT = ROOT / "scripts" / "_tmp_diag_round46.out.txt"
ROUND_DIR = ROOT / "artifacts" / "technology_pass_semantics" / "round46_consulting_generalization"
SRC = ROOT / "test_data" / "93 技术咨询合同.docx"

buf = io.StringIO()


def w(s=""):
    buf.write(s + "\n")


# ---------- 1. 固定输入 ----------
ROUND_DIR.mkdir(parents=True, exist_ok=True)
dst_docx = ROUND_DIR / "input.docx"
if not dst_docx.exists():
    shutil.copy2(SRC, dst_docx)
sha = hashlib.sha256(SRC.read_bytes()).hexdigest()

doc = Document(str(SRC))
parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
for tb in doc.tables:
    for row in tb.rows:
        cells = [c.text.strip() for c in row.cells]
        line = " | ".join(c for c in cells if c)
        if line:
            parts.append(line)
TEXT = "\n".join(parts)
(ROUND_DIR / "input.txt").write_text(TEXT, encoding="utf-8")
txt_sha = hashlib.sha256(TEXT.encode("utf-8")).hexdigest()

w("=" * 70)
w("Round46 输入固定")
w("=" * 70)
w(f"源样本: {SRC.name}")
w(f"docx sha256: {sha}")
w(f"text sha256: {txt_sha}")
w(f"文本长度: {len(TEXT)}")
w(f"复制至: {dst_docx}")

# ---------- 2. Element 层诊断 ----------
clauses = parse_contract_clauses(TEXT)


def indexed(rows):
    return {row["element_id"]: row for row in rows}


FAMILY_MAP = {
    "TC-01-03": ("consulting", "consulting_subject", {"consulting_project", "consulting_scope"}),
    "TC-08-01": ("consulting", "consulting_quality", {"consulting_scope", "report_contents", "acceptance_standard"}),
    "TC-08-02": ("consulting", "consulting_decision", {"qualified_consulting_decision_risk", "consultant_fault_exception"}),
    "TC-08-03": ("consulting", "consulting_inputs", {"truthfulness_completeness"}),
    "TC-08-04": ("consulting", "consulting_remedy", {"defect_trigger", "correction_period", "correction_remedy"}),
    "TC-11-01": ("technology", "confidentiality_controls", {"third_party_disclosure_restriction", "security_measures"}),
    "TC-11-02": ("technology", "confidentiality_disclosure", {"independent_development_exception"}),
    "TC-11-03": ("technology", "material_exit", {
        "processing_deadline", "retained_backup_use_restriction", "confidential_copy_deletion",
        "confidential_backup_deletion", "confidential_disposal_confirmation", "confidential_material_destruction",
    }),
    "TC-11-04": ("remedies", None, {"breach_damage_liability", "stop_disclosure", "stop_use", "mitigation_obligation", "incident_notification"}),
    "TC-12-02": ("technology", "service_claims", {
        "third_party_claim_trigger", "claim_handling_responsibility", "provider_infringement_liability",
        "direct_damage", "defense_cost_allocation", "notification_requirement", "notification_deadline",
        "necessary_assistance", "termination_remedy", "refund_remedy", "settlement_authority",
        "replacement_remedy", "license_cure", "defense_control", "lawyer_selection",
    }),
}

w("\n" + "=" * 70)
w("Element 层状态（93 合同）")
w("=" * 70)
remedy_rows = indexed(canonical_confidentiality_remedies(clauses))
for tc in sorted(FAMILY_MAP):
    kind, family, keys = FAMILY_MAP[tc]
    if kind == "consulting":
        rows = indexed(consulting_elements(family, clauses))
    elif kind == "technology":
        rows = indexed(technology_elements(family, clauses))
    else:
        rows = remedy_rows
    w(f"\n[{tc}] (source={kind}/{family or 'remedies'})")
    for key in sorted(keys):
        r = rows.get(key)
        if r is None:
            w(f"   {key}: <无此 element>")
        else:
            st = r.get("status")
            ev = r.get("supporting_evidence") or []
            w(f"   {key}: {st}  evidence={len(ev)}")

# ---------- 3. 引擎层：coverage vs SEARCH_AUDIT 假 MISSING ----------
TARGETS = set(FAMILY_MAP.keys())
w("\n" + "=" * 70)
w("引擎层：coverage COVERED ∩ SEARCH_AUDIT NOT_FOUND（假 MISSING 检测）")
w("=" * 70)

for party in ["甲方", "乙方"]:
    w(f"\n--- 立场 {party} ---")
    try:
        result = review_contract_v2(
            TEXT, "technology", use_llm=False, include_low_risk=True,
            profile=UserRiskProfile(user_identity=party),
        )
    except Exception as e:
        w(f"  引擎异常: {type(e).__name__}: {e}")
        continue
    gate = result.get("release_quality_gate", {})
    w(f"  gate={gate.get('status')}  failures={len(gate.get('failures') or [])}")
    for item in result.get("items", []):
        iid = item.get("item_id")
        if iid not in TARGETS:
            continue
        covered = {r["element_id"] for r in item.get("coverage_elements", []) if r.get("status") == "COVERED"}
        trace = item.get("decision_trace", {}) or {}
        missing = {r["element_id"] for r in trace.get("element_absence_searches", []) if r.get("result") == "NOT_FOUND"}
        conflict = covered & missing
        w(f"\n  [{iid}] status={item.get('status')}")
        w(f"     COVERED({len(covered)}): {sorted(covered)}")
        w(f"     NOT_FOUND({len(missing)}): {sorted(missing)}")
        if conflict:
            w(f"     *** 假MISSING冲突: {sorted(conflict)}")
        # 所有 coverage element 明细
        det = ", ".join(f"{r['element_id']}={r.get('status')}" for r in item.get("coverage_elements", []))
        w(f"     明细: {det}")

OUT.write_text(buf.getvalue(), encoding="utf-8")
print("written", OUT)
