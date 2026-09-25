"""Run the actual utility fixture through the report pipeline and save evidence."""
import json
from pathlib import Path
from collections import Counter
from docx import Document
from review_v2.engine import review_contract_v2

def main():
    document = Document("test_data/62 供用水电气热力合同.docx")
    text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
    text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    result = review_contract_v2(text, "utility_supply", use_llm=False)
    path = Path("artifacts/utility_semantic_boundary_acceptance.json")
    path.parent.mkdir(exist_ok=True)
    path.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    path.with_suffix(".md").write_text(result["result"], encoding="utf-8")
    print("gate", json.dumps({k: result["release_quality_gate"].get(k) for k in ("status", "failures")}, ensure_ascii=False))
    print("counts", Counter(i["coverage_status"] for i in result["items"]))
    for item in result["items"]:
        if item["item_id"] in {"US-03-04", "US-06-04", "US-11-03"}:
            print(json.dumps({k:item.get(k) for k in ["item_id", "coverage_status", "reasoning_summary", "evidence_location", "evidence_clauses", "applicability"]}, ensure_ascii=False))

if __name__ == "__main__":
    main()
