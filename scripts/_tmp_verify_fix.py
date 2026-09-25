"""验证修复：33 合同目标 element 是否转为 COVERED，且 Round41-45 fixture 不退化。"""
import io
import sys
from pathlib import Path

ROOT = Path(r"F:/ai/Sage")
sys.path.insert(0, str(ROOT / "backend"))
sys.path.insert(0, str(ROOT))

from docx import Document
from review_v2.clause_parser import parse_contract_clauses
from review_v2.consulting_element_facts import elements as consulting_elements

OUT = ROOT / "scripts" / "_tmp_verify_fix.out.txt"
buf = io.StringIO()


def w(s=""):
    buf.write(s + "\n")


def docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for tb in doc.tables:
        for row in tb.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts)


CHECK = {
    "consulting_decision": ["qualified_consulting_decision_risk", "consultant_fault_exception"],
    "consulting_remedy": ["defect_trigger", "correction_period", "correction_remedy"],
    "consulting_quality": ["consulting_scope", "report_contents", "acceptance_standard", "delivery_form"],
    "consulting_inputs": ["truthfulness_completeness"],
    "consulting_subject": ["consulting_project", "consulting_scope"],
}


def report(name, TEXT):
    w("\n" + "=" * 66)
    w(f"{name}")
    w("=" * 66)
    clauses = parse_contract_clauses(TEXT)
    for family, keys in CHECK.items():
        try:
            rows = {r["element_id"]: r for r in consulting_elements(family, clauses)}
        except Exception as e:
            w(f"  [{family}] 异常 {e}")
            continue
        for key in keys:
            r = rows.get(key)
            if r is None:
                w(f"  {family}.{key}: <无>")
            else:
                ev = r.get("supporting_evidence") or []
                q = ev[0]["quote"][:70].replace("\n", " ") if ev else ""
                w(f"  {family}.{key}: {r['status']}  | {q}")


# 33 新技术咨询合同（Round46 目标）
report("33 技术咨询服务合同.docx (Round46 目标)",
       docx_text(ROOT / "test_data" / "33 技术咨询服务合同.docx"))

# Round41-45 fixtures 回归（不得退化）
fx = ROOT / "backend" / "tests" / "fixtures"
for name in ["round41_consulting.txt", "round42_consulting.txt", "round43_consulting.txt",
             "round43_consulting_630k.txt", "round44_consulting_projection.txt",
             "round45_consulting_projection.txt", "round36_consulting.txt"]:
    p = fx / name
    if p.exists():
        report(name, p.read_text(encoding="utf-8"))

OUT.write_text(buf.getvalue(), encoding="utf-8")
print("written", OUT)
