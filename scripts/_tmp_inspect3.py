"""临时排查3：提取 93 技术咨询合同.docx / 94 技术服务合同.docx 全文。"""
import io
from pathlib import Path

ROOT = Path(r"F:/ai/Sage")
OUT = ROOT / "scripts" / "_tmp_inspect3.out.txt"


def docx_text(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for tb in doc.tables:
        for row in tb.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts)


buf = io.StringIO()
for name in ["93 技术咨询合同.docx", "94 技术服务合同.docx"]:
    p = ROOT / "test_data" / name
    buf.write("=" * 70 + "\n")
    buf.write(f"{name}\n")
    buf.write("=" * 70 + "\n")
    if p.exists():
        try:
            buf.write(docx_text(p) + "\n")
        except Exception as e:
            buf.write(f"(失败 {e})\n")
    else:
        buf.write("(不存在)\n")

OUT.write_text(buf.getvalue(), encoding="utf-8")
print("written", OUT)
