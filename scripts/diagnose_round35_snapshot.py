import sys
from pathlib import Path
sys.path[:0]=['backend','.']
from docx import Document
from review_v2.clause_parser import parse_contract_clauses
from review_v2.engine import _register_atomic_family_facts
from review_v2.contract_facts import extract_contract_fact_graph,canonical_fact_key
from review_v2.application_transfer_elements import FAMILIES
d=Document(next(Path('test_data').glob('42*.docx')))
t='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in r.cells) for tab in d.tables for r in tab.rows)
cs=parse_contract_clauses(t);g=extract_contract_fact_graph(cs)
for f in FAMILIES:_register_atomic_family_facts(g,f,cs)
for family,es in g.atomic_family_snapshots.items():
 for e in es:
  facts=g.facts.get(canonical_fact_key(e.get('source_fact_key','')),[])
  for s in e.get('supporting_evidence',[]):
   if not any(s['clause_id'] in f.source_clause_ids and s['quote'] in f.source_spans for f in facts):
    print(family,e['element_id'],e['status'],repr(s['quote']),[(f.value,f.source_spans) for f in facts])
