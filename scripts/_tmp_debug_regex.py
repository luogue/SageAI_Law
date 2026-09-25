"""调试：33 合同 defect_trigger / correction_period 正则为何未命中。"""
import io
import re
import sys
from pathlib import Path

ROOT = Path(r"F:/ai/Sage")
sys.path.insert(0, str(ROOT / "backend"))
sys.path.insert(0, str(ROOT))

from docx import Document
from review_v2.clause_parser import parse_contract_clauses
from review_v2.obligation_detail_elements import _segments, _DURATION
from review_v2.proposition_scope import arrangement_matches

OUT = ROOT / "scripts" / "_tmp_debug_regex.out.txt"
buf = io.StringIO()


def w(s=""):
    buf.write(s + "\n")


def docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for tb in doc.tables:
        for row in tb.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts)


TEXT = docx_text(ROOT / "test_data" / "33 技术咨询服务合同.docx")
clauses = parse_contract_clauses(TEXT)

# 导入修改后的实际 predicate
from review_v2.consulting_element_facts import FAMILIES

defect_pat = dict((r[0], r[2]) for r in FAMILIES["consulting_remedy"])["defect_trigger"]
corr_pat = dict((r[0], r[2]) for r in FAMILIES["consulting_remedy"])["correction_period"]

w("defect_trigger 正则(截断):")
w("  " + defect_pat[:400])
w("\ncorrection_period 正则(截断):")
w("  " + corr_pat[:400])

w("\n" + "=" * 66)
w("包含关键词的 segments 及匹配情况")
w("=" * 66)
for clause, heading, quote in _segments(clauses):
    text = re.sub(r"\s+", "", quote)
    if not re.search(r"满意|改正|弥补|期限", text):
        continue
    w(f"\nclause={clause.clause_id} block={getattr(clause,'block_type','')} title={getattr(clause,'title','')}")
    w(f"  text: {text[:160]}")
    d_match = list(arrangement_matches(defect_pat, text))
    c_match = list(arrangement_matches(corr_pat, text))
    w(f"  defect_trigger matched={bool(d_match)}")
    w(f"  correction_period matched={bool(c_match)}")
    # 单独测 30天 的期限
    m = re.search(_DURATION + r"(?:内|的期限|期限)[^。；]{0,20}(?:补充|补正|修正|修订|改正|修改|弥补)", text)
    w(f"  直接正则(期限分支)={bool(m)} -> {m[0] if m else ''}")
    m2 = re.search(r"(?:咨询成果|咨询报告|报告|成果|服务).{0,25}(?:不符合|未达到|不满意|不能令[^。；]{0,8}满意)", text)
    w(f"  直接正则(服务+不满意)={bool(m2)} -> {m2[0] if m2 else ''}")

OUT.write_text(buf.getvalue(), encoding="utf-8")
print("written", OUT)
