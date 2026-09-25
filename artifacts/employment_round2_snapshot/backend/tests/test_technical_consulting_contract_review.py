import io
from pathlib import Path

import pytest
import main as legacy_main
from docx import Document
from fastapi.testclient import TestClient

from auth.routes import get_current_user
from main import app
from review_v2.engine import TECHNICAL_CONSULTING_EVIDENCE_IMPLEMENTATION
from review_v2.engine import review_contract_v2
from review_v2.clause_parser import parse_contract_clauses
from review_v2.user_profile import UserRiskProfile
from review_v2.technology_service_analysis import (
    analyze_payment_schedule,
    analyze_timeline,
    analyze_technology_relationship_risks,
)


CONSULTING_CONTRACT = """
技术咨询合同
项目名称：智能仓储管理系统技术可行性及架构优化咨询项目
委托方（甲方）：杭州启明科技有限公司
受托方（乙方）：北京智策技术咨询有限公司
签署日期：2020年7月15日

第一条 咨询内容、咨询要求和咨询方式
1、咨询内容：乙方对甲方现有智能仓储管理系统进行技术可行性分析和系统架构评估，分析系统架构、性能瓶颈、接口方案及技术风险。
2、咨询要求：乙方形成技术咨询报告，内容包括技术改造建议、推荐技术架构、实施路线、风险、成本及周期估算。
3、乙方不承担软件开发、技术实施或成果制造义务。

第二条 项目进度
乙方应按双方确认的咨询工作计划完成调研并提交咨询成果。

第三条 甲方协作事项
1、甲方提供系统架构图、数据库结构、ERP/WMS接口文档和系统性能数据。
2、甲方安排技术人员访谈，提供测试环境、系统访问权限并协助说明业务流程。
3、甲方保证资料真实、完整，并通过加密电子文件及现场交接方式提供。

第四条 技术咨询报酬和支付方式
1、技术咨询报酬总额为180,000元（含税），分三期支付：合同生效后5个工作日、提交咨询报告初稿后5个工作日、最终验收通过后10个工作日。
2、收款账户、开户银行和帐号以本条所列信息为准。

第五条 保密
1、保密内容包括甲方的系统架构、源代码、接口资料、业务数据、仓储运营数据及其他未公开技术和经营信息，以及乙方的非公开咨询方法、分析模型和技术方案。
2、涉密人员范围为参与本项目的双方员工及外部顾问。
3、保密期限为合同有效期内及合同终止后3年。
4、违反保密义务造成对方损失的，应赔偿实际损失。

第六条 联系与协调
双方通过指定项目联系人协调咨询工作。

第七条 咨询成果和验收
1、乙方提交电子版及纸质版技术咨询报告。
2、咨询报告应完整覆盖第一条约定内容，并就架构、性能瓶颈、改造路线和主要风险提出明确意见。
3、甲方组织技术人员书面验收。

第八条 咨询成果瑕疵救济
因乙方原因导致咨询成果明显不符合合同约定内容的，乙方应在甲方要求的合理期限内免费修改。

第九条 咨询决策风险
甲方按照乙方符合本合同约定标准和方式完成的技术咨询工作成果作出决策并予以实施所造成的损失，按以下第1种方式处理：
1、乙方不承担责任。
2、乙方承担部分责任。具体承担方式为：无。

第十条 新技术成果
甲方利用咨询成果完成的新的技术成果归甲方所有；乙方利用甲方资料和工作条件完成的新的技术成果归乙方所有。

第十五条 技术文件
经双方书面盖章或确认的电子邮件确认后，下列文件为合同组成部分：《项目技术需求说明书》《现有系统技术资料清单》《项目咨询工作计划》《咨询成果验收确认单》。
"""


@pytest.fixture(scope="module")
def consulting_review() -> dict:
    return review_contract_v2(
        CONSULTING_CONTRACT,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/委托方"),
    )


def _items(output: dict) -> dict:
    return {item["item_id"]: item for item in output["items"]}


def _consulting_data_review(data_clause: str = "") -> dict:
    text = f"""
技术咨询合同
甲方为委托方，乙方为咨询方。
第一条 乙方负责开展生产线技术评估并向甲方提交技术咨询报告。
第二条 技术咨询费为10000元。
{data_clause}
"""
    return review_contract_v2(
        text,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/委托方"),
    )


def test_consulting_context_and_role_are_single_source(consulting_review: dict):
    classification = consulting_review["deterministic_analysis"]["technology_classification"]
    context = consulting_review["report_context"]

    assert classification["primary_subtype"] == "TECHNICAL_CONSULTING"
    assert context["transaction_structure"] == "B2B_TECHNICAL_CONSULTING"
    assert context["usage_scenario_label"] == "技术咨询"
    assert context["contract_type_label"] == "技术合同—技术咨询合同"
    assert consulting_review["perspective_resolution"]["display_role"] == "甲方 / 委托方"
    assert consulting_review["result"].startswith("# 技术咨询合同审查报告")
    for stale_role in ("专利权受让方", "让与方", "被许可方", "许可方", "委托开发方", "合作开发方"):
        assert stale_role not in consulting_review["result"]


def test_consulting_tc08_rules_and_direct_evidence(consulting_review: dict):
    by_id = _items(consulting_review)
    assert [by_id[f"TC-08-0{index}"]["coverage_status"] for index in range(1, 5)] == ["通过"] * 4
    assert by_id["TC-08-05"]["coverage_status"] == "通过"
    assert by_id["TC-08-05"]["risk_level"] == "低"
    assert by_id["TC-08-05"]["action_priority"] == "NONE"

    locations = {
        item_id: " ".join(str(row.get("location") or "") for row in by_id[item_id]["evidence_clauses"])
        for item_id in ("TC-08-01", "TC-08-02", "TC-08-03", "TC-08-04")
    }
    quality_text = " ".join(
        str(row.get("quote") or row.get("text") or "") for row in by_id["TC-08-01"]["evidence_clauses"]
    )
    assert "咨询内容" in quality_text and "完整覆盖" in quality_text
    assert "第九条" in locations["TC-08-02"] and "第十一条" not in locations["TC-08-02"]
    assert "第三条" in locations["TC-08-03"]
    assert "第八条" in locations["TC-08-04"]


def test_consulting_subtype_applicability_is_isolated(consulting_review: dict):
    by_id = _items(consulting_review)
    assert by_id["TC-01-04"]["coverage_status"] == "不适用"
    for item_id in ("TC-10-01", "TC-10-02", "TC-10-03", "TC-10-04"):
        assert by_id[item_id]["coverage_status"] == "不适用"
        assert "当前技术咨询交易" in by_id[item_id]["applicability"]["reason"]
        assert "当前技术服务交易" not in by_id[item_id]["applicability"]["reason"]
    for item_id in ("TC-03-04", "TC-12-02", "TC-12-04", "TC-13-02", "TC-13-03"):
        assert by_id[item_id]["coverage_status"] == "不适用"


def test_consulting_confidentiality_uses_shared_atomic_evidence(consulting_review: dict):
    by_id = _items(consulting_review)
    expected = {
        "TC-11-01": ("部分覆盖", "中"),
        "TC-11-02": ("部分覆盖", "低"),
        "TC-11-03": ("缺失", "低"),
        "TC-11-04": ("部分覆盖", "中"),
    }
    for item_id, (coverage, level) in expected.items():
        assert by_id[item_id]["coverage_status"] == coverage
        assert by_id[item_id]["risk_level"] == level

    tc1104 = {row["element_id"]: row for row in by_id["TC-11-04"]["coverage_elements"]}
    assert tc1104["breach_damage_liability"]["status"] == "COVERED"
    assert tc1104["stop_disclosure"]["status"] == "MISSING"
    assert tc1104["return_delete_destroy"]["status"] == "MISSING"
    assert tc1104["specific_confidentiality_penalty"]["status"] == "NOT_APPLICABLE"

    evidence = by_id["TC-11-04"]["evidence_clauses"]
    assert "第五条第4项" in evidence[0]["location"]
    assert "违反保密义务" in evidence[0]["quote"]
    assert "实际损失" in evidence[0]["quote"]
    assert evidence[0]["semantic_evidence_type"] == "EXISTING_CONFIDENTIALITY_DAMAGE_LIABILITY"
    assert evidence[0]["semantic_support_level"] == "DIRECT"
    assert any(
        row.get("semantic_support_level") == "SUPPORTING" and "第五条第1项" in row["location"]
        for row in evidence[1:]
    )


def test_consulting_attachments_and_release_gate(consulting_review: dict):
    by_id = _items(consulting_review)
    attachment = by_id["TC-13-01"]
    assert attachment["coverage_status"] == "通过"
    assert all("第十五条" in str(row.get("location") or "") for row in attachment["evidence_clauses"])
    assert consulting_review["review_evidence_element_gate"]["mismatch_count"] == 0
    assert consulting_review["release_quality_gate"]["status"] == "REPORT_READY"
    assert consulting_review["system_metadata"]["technology_technical_consulting"] == "FROZEN_REGRESSION"
    assert consulting_review["result"]


def test_consulting_signing_law_context_is_not_false_positive(consulting_review: dict):
    by_id = _items(consulting_review)
    expected = {
        "TC-08-02": ("第359条", "第881条"),
        "TC-11-04": ("第113条", "第584条"),
    }
    for item_id, (historical_article, current_article) in expected.items():
        context = by_id[item_id]["decision_trace"]["legal_time_context"]
        assert context["signed_at"] == "2020-07-15"
        assert context["law_at_signing"] == "中华人民共和国合同法"
        assert historical_article in context["historical_reference"]
        assert current_article in context["current_reference"]
        assert context["status"] == "HISTORICAL_AT_SIGNING"
        assert context["display_text"] in consulting_review["result"]
    assert not any(finding.get("risk_type") == "legal_reference" for finding in consulting_review["risk_findings"])


def test_modern_consulting_contract_does_not_display_repealed_contract_law():
    modern_contract = CONSULTING_CONTRACT.replace("2020年7月15日", "2024年7月15日")
    output = review_contract_v2(
        modern_contract,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/委托方"),
    )
    by_id = _items(output)
    for item_id, article in (("TC-08-02", "第881条"), ("TC-11-04", "第584条")):
        context = by_id[item_id]["decision_trace"]["legal_time_context"]
        assert context["status"] == "CURRENT_LAW_AT_SIGNING"
        assert context["display_text"] == f"《中华人民共和国民法典》{article}"
    assert "签署时法律依据：《中华人民共和国合同法》" not in output["result"]


def test_consulting_party_b_role_is_not_development_role():
    output = review_contract_v2(
        CONSULTING_CONTRACT,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="乙方/受托方"),
    )
    assert output["perspective_resolution"]["display_role"] == "乙方 / 受托方"
    assert output["report_context"]["perspective"] == "technical_consultant"


def test_consulting_subtype_is_inferred_from_transaction_substance_without_explicit_title():
    title_neutral_contract = CONSULTING_CONTRACT.replace("技术咨询合同", "智能仓储项目协议", 1)
    output = review_contract_v2(
        title_neutral_contract,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/委托方"),
    )
    classification = output["deterministic_analysis"]["technology_classification"]
    by_id = _items(output)

    assert classification["primary_subtype"] == "TECHNICAL_CONSULTING"
    assert classification["primary_contract_type"] == "technology_consulting"
    assert "title:technical_consulting" not in classification["classification_evidence"]
    assert {
        "core_obligation:analysis_evaluation_advice",
        "deliverable:consulting_report_or_opinion",
        "consideration:consulting_fee",
    }.issubset(set(classification["classification_evidence"]))
    assert output["contract_context"]["contract_subtype"] == "TECHNICAL_CONSULTING"
    assert output["report_context"]["contract_type_label"] == "技术合同—技术咨询合同"
    assert all(by_id[f"TC-08-0{index}"]["coverage_status"] != "不适用" for index in range(1, 6))
    assert by_id["TC-01-03"]["coverage_status"] == "通过"
    assert by_id["TC-01-04"]["coverage_status"] == "不适用"
    assert all(by_id[f"TC-10-0{index}"]["coverage_status"] == "不适用" for index in range(1, 5))
    assert by_id["TC-11-01"]["coverage_status"] != "缺失"
    assert by_id["TC-11-02"]["coverage_status"] != "缺失"
    for item_id in ("TC-03-04", "TC-12-02", "TC-12-04", "TC-13-02"):
        assert by_id[item_id]["coverage_status"] == "不适用"


CONSULTING_ELEMENT_EVIDENCE_CONTRACT = """
项目名称：新能源汽车动力电池生产线节能优化技术咨询项目
甲方（委托方）：测试甲方
乙方（咨询方）：测试乙方
第一条 咨询事项
乙方针对动力电池生产线能耗数据、主要耗能设备和工艺流程进行分析评估，设计节能改造方案并提交最终技术咨询报告。
第二条 咨询报酬
甲方支付咨询报酬人民币20万元。
第三条 咨询资料
甲方应向乙方提供能耗数据、生产线图纸、设计资料、适用规范和其他咨询所需信息，并保证资料真实、完整。
第四条 保密
保密资料包括双方在项目中接触的商业秘密、技术秘密、图纸、设计资料、规范和未公开信息。
双方承担保密义务，未经对方书面同意不得向第三方披露，但已公开信息和依法必须披露的信息除外。
在任何情形下，本条规定的保密义务应持续有效。
第五条 咨询报告瑕疵救济
甲方对报告不满意的，可以通知乙方整改；乙方应在收到通知后10日内完成修改并重新提交。
整改后仍不能达到验收标准的，甲方可以解除合同，乙方应退还已付款并按年利率4%支付利息。
"""


def test_consulting_elements_and_evidence_are_status_separated_and_contract_grounded():
    output = review_contract_v2(
        CONSULTING_ELEMENT_EVIDENCE_CONTRACT,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/委托方"),
    )
    by_id = _items(output)

    assert by_id["TC-01-03"]["coverage_status"] == "通过"
    assert "专利" not in by_id["TC-01-03"]["narrative_state"]["judgment"]
    assert by_id["TC-13-01"]["coverage_status"] == "不适用"

    materials = by_id["TC-08-03"]
    assert materials["coverage_status"] in {"通过", "部分覆盖"}
    assert "第三条" in str(materials["evidence_clauses"][0]["location"])
    assert "甲方应向乙方提供" in materials["evidence_clauses"][0]["raw_text"]

    remedy = by_id["TC-08-04"]
    remedy_elements = {row["element_id"]: row["status"] for row in remedy["coverage_elements"]}
    assert remedy_elements == {
        "defect_trigger": "COVERED",
        "correction_period": "COVERED",
        "correction_remedy": "COVERED",
        "termination_remedy": "COVERED",
        "refund_remedy": "COVERED",
    }

    confidentiality = by_id["TC-11-02"]
    state = confidentiality["narrative_state"]
    assert confidentiality["coverage_status"] == "部分覆盖"
    assert state["covered_elements"]
    assert state["missing_elements"]
    assert not set(state["covered_elements"]) & set(state["missing_elements"])
    assert "持续有效" in " ".join(row.get("raw_text", "") for row in confidentiality["evidence_clauses"])
    assert "终止后三年" not in output["result"]
    assert "尚未覆盖：双方保密" not in output["result"]


CONSULTING_CROSS_SECTION_EVIDENCE_CONTRACT = """
技术咨询合同
甲方（委托方）：测试甲方
乙方（咨询方）：测试乙方
第一条 咨询事项
乙方就甲方生产系统进行技术分析、评估并提交技术咨询报告。
第三条 技术咨询报酬和费用
3.1 合同总价为500,000美元，其中分项一合同价120,000美元、分项二合同价130,000美元、分项三合同价110,000美元、分项四合同价140,000美元。
3.2 甲方按30%、40%和30%的付款比例分期支付，各期付款期限均为收到相应咨询报告后15个工作日内。
3.3 甲方付款银行及乙方收款银行、收款账户以本条所列信息为准；双方各自承担本方银行费用。
3.4 超出本合同咨询范围的额外服务，由双方另行协商费用并书面确认调整后的服务价格。
第五条 保密
双方对履行本合同过程中知悉的技术资料、经营信息及其他保密信息均负有保密义务；未经对方书面同意，不得向任何第三方泄露，也不得用于本项目之外。
5.1 合同终止后，双方继续受本条保密义务约束。
第七条 咨询成果异议和改正
7.2 甲方对咨询服务不满意的，应通知乙方；乙方应在15日内改正或弥补，在改正或弥补完成前甲方有权暂停付款。
第十条 验收不合格救济
10.3 咨询报告未达到最低验收标准的，甲方有权解除合同并要求乙方退还全部已付款项。
"""


CONSULTING_THEME_AGGREGATION_CONTRACT = """
技术咨询合同
合同号：TC-TEST-2026-0830
签订日期：2026年8月30日
委托方（甲方）：测试委托方
咨询方（乙方）：测试咨询方
咨询方国家：德国
技术咨询项目：智能制造生产线升级及数字化改造项目
第一条 合同内容
1.1 咨询内容及咨询范围：乙方对智能制造生产线自动化、能耗优化及数字化控制系统进行技术可行性分析和系统架构评估，并提交最终技术咨询报告。
1.5 技术咨询服务期限：12个月。
1.6 最终技术咨询报告提交期限：11个月。
第二条 咨询资料和协作
甲方应在合同生效后5个工作日内以加密电子文件向乙方提供系统架构图、运行数据和技术资料，安排技术人员访谈并提供必要访问权限；甲方保证所提供资料真实、完整、准确。
第三条 价格与支付
3.1 合同总价币种：美元（USD），合同总价为500,000美元。
分项一合同价：120,000美元；分项二合同价：130,000美元；分项三合同价：110,000美元；分项四合同价：140,000美元。
3.3 委托方付款银行、咨询方收款银行和咨询方账户以本条所列信息为准。
3.3.1 合同总价比例20%，金额100,000美元，审核后15天付款。
3.3.2 分项一合同价付款比例80%，金额96,000美元，审核后15天付款。
3.3.3 分项二合同价付款比例80%，金额104,000美元，审核后15天付款。
3.3.4 分项三合同价付款比例80%，金额88,000美元，审核后15天付款。
3.3.5 分项四合同价付款比例60%，金额84,000美元，审核后15天付款。
3.3.6 分项四合同价付款比例20%，金额28,000美元，审核后15天付款。
第四条 交付
分项一技术咨询报告于合同生效后3个月内提交；分项二于6个月内提交；分项三于9个月内提交；分项四及最终技术咨询报告于11个月内提交。
第七条 保证
7.2 整改或弥补期限：15天。
7.3 保证期：最终验收后12个月。
第十条 违约和合同解除
10.1 迟延交付违约金分阶段计算，迟延交付违约金总额上限为合同总价的10%。
10.3 超过规定交付期限60天仍不能交付的，退还已付款并按年利率6%支付利息。
10.4 一般违约整改期限为30天。
第十五条 合同有效期
合同有效期为3年。
附件一：技术咨询服务范围
现有智能制造生产线技术评估、自动化设备及控制系统升级方案、PLC及工业控制系统优化方案、生产线数字化架构设计、能耗监测及节能优化方案、最终总体技术改造方案和实施阶段技术咨询支持。
附件二：进度安排
第1个月现场调研和资料收集；第2至3个月提交分项一报告；第4至6个月提交分项二报告；第7至9个月提交分项三报告；第10至11个月提交分项四及最终报告；第12个月提供技术答疑、验收及整改支持。
附件三：人员安排
项目总负责人、自动化控制专家、机械系统专家、数字化系统专家、能源管理专家和项目协调人员共6人。
"""


COMPLETE_CONSULTING_CONFIDENTIALITY = """
第十六条 完整保密机制
保密信息包括双方技术秘密、商业秘密和未公开技术资料，涉密人员范围为参与本项目的员工和顾问。
双方承担严格保密义务，保密信息仅用于本项目，未经披露方书面同意不得复制，也不得向任何第三方披露。
双方应采取访问控制、加密和留痕等安全措施，本条在合同终止后继续有效五年。
已公开信息、接收方从有权第三方合法取得的信息、接收方独立开发的信息不属于保密信息；法律法规要求披露、司法程序要求披露或监管要求披露时，接收方应在披露前书面通知披露方。
项目结束或合同终止后，接收方应返还、删除或销毁全部保密资料及其副本、备份文件，并书面确认处置结果。
任何一方违反保密义务的，应赔偿对方实际损失，并立即停止披露、立即停止使用，采取补救措施防止继续扩散。
"""


def test_consulting_cross_section_evidence_recall_and_binding():
    output = review_contract_v2(
        CONSULTING_CROSS_SECTION_EVIDENCE_CONTRACT,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/委托方"),
    )
    by_id = _items(output)

    remedy = by_id["TC-08-04"]
    remedy_text = " ".join(
        str(row.get("quote") or row.get("text") or "") for row in remedy["evidence_clauses"]
    )
    remedy_locations = " ".join(str(row.get("location") or "") for row in remedy["evidence_clauses"])
    assert remedy["coverage_status"] in {"通过", "部分覆盖"}
    assert "第七条" in remedy_locations and "第十条" in remedy_locations
    assert all(term in remedy_text for term in ("不满意", "改正", "暂停付款", "未达到最低验收标准", "退还全部已付款项"))

    payment = by_id["TC-08-05"]
    payment_elements = {row["element_id"]: row for row in payment["coverage_elements"]}
    payment_locations = " ".join(str(row.get("location") or "") for row in payment["evidence_clauses"])
    assert payment["coverage_status"] == "通过"
    assert "第三条" in payment_locations
    for element_id in (
        "consulting_fee_total", "itemized_prices", "payment_milestones",
        "collection_account", "bank_expense_allocation", "out_of_scope_repricing",
    ):
        assert payment_elements[element_id]["status"] == "COVERED", element_id
        assert payment_elements[element_id]["evidence_ids"], element_id

    tc1101 = by_id["TC-11-01"]
    tc1101_elements = {row["element_id"]: row for row in tc1101["coverage_elements"]}
    assert tc1101["coverage_status"] != "缺失"
    for element_id in (
        "confidential_information_scope", "confidentiality_obligation",
        "purpose_limitation", "third_party_disclosure_restriction", "confidentiality_duration",
    ):
        assert tc1101_elements[element_id]["status"] == "COVERED", element_id
        assert tc1101_elements[element_id]["evidence_ids"], element_id

    tc1102 = by_id["TC-11-02"]
    tc1102_elements = {row["element_id"]: row for row in tc1102["coverage_elements"]}
    assert tc1102["coverage_status"] != "缺失"
    assert tc1102_elements["confidentiality_period"]["status"] == "COVERED"
    assert tc1102_elements["confidentiality_period"]["evidence_ids"]

    tc1104 = by_id["TC-11-04"]
    tc1104_elements = {row["element_id"]: row for row in tc1104["coverage_elements"]}
    assert tc1104["coverage_status"] == "部分覆盖"
    assert tc1104_elements["confidentiality_obligation"]["status"] == "COVERED"
    assert tc1104_elements["stop_disclosure"]["status"] == "MISSING"
    assert tc1104_elements["stop_use"]["status"] == "MISSING"
    assert tc1104_elements["mitigation_obligation"]["status"] == "MISSING"
    assert "未明确双方基本保密义务" not in tc1104["risk_point_description"]
    assert "返还、删除或销毁" not in tc1104["suggested_clause_example"]


def test_consulting_report_projects_seven_atomic_findings_into_three_themes():
    output = review_contract_v2(
        CONSULTING_THEME_AGGREGATION_CONTRACT,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/委托方"),
    )
    by_id = _items(output)
    themes = {theme["title"]: theme for theme in output["risk_themes"]}

    assert output["report_context"]["contract_type_label"] == "技术合同—技术咨询合同"
    assert len(output["risk_findings"]) == 7
    assert output["finding_count"] == 7
    assert output["risk_theme_count"] == output["aggregated_risk_count"] == 3
    assert set(themes) == {
        "咨询成果质量与验收机制不完整",
        "咨询决策风险承担",
        "保密机制不完整",
    }
    assert set(themes["咨询成果质量与验收机制不完整"]["source_review_ids"]) == {"TC-08-01", "TC-08-04"}
    assert set(themes["保密机制不完整"]["source_review_ids"]) == {
        "TC-11-01", "TC-11-02", "TC-11-03", "TC-11-04",
    }
    assert themes["咨询成果质量与验收机制不完整"]["final_action"] == "SHOULD_CHANGE"
    assert themes["咨询决策风险承担"]["final_action"] == "OPTIONAL_IMPROVEMENT"
    assert themes["保密机制不完整"]["final_action"] == "OPTIONAL_IMPROVEMENT"
    assert output["risk_theme_summary"] == {
        "total": 3,
        "must_change": 0,
        "should_change": 1,
        "optional_improvement": 2,
        "manual_review": 0,
    }
    assert output["risk_finding_summary"]["低风险"] == 7
    assert by_id["TC-08-01"]["risk_severity"] == "LOW"
    assert by_id["TC-08-01"]["action_priority"] == "B_PRACTICAL"
    assert by_id["TC-08-04"]["risk_severity"] == "LOW"
    assert by_id["TC-08-04"]["action_priority"] == "C_OPTIMIZATION"
    assert by_id["TC-08-02"]["action_priority"] == "C_OPTIMIZATION"
    for item_id in ("TC-11-01", "TC-11-02", "TC-11-03", "TC-11-04"):
        assert by_id[item_id]["risk_severity"] == "LOW"
        assert by_id[item_id]["action_priority"] == "C_OPTIMIZATION"
    assert "风险主题数（聚合后）：3" in output["result"]
    assert "底层审查发现数：7" in output["result"]
    assert "建议处理风险：1" in output["result"]
    assert "可选优化主题数（聚合后）：2" in output["result"]
    assert output["overall_assessment"]["signability"] == "可以签署，建议进一步完善"
    assert "总体可签署性：可以签署，建议进一步完善" in output["result"]
    tc0804_finding = next(
        finding for finding in output["risk_findings"]
        if "TC-08-04" in finding["source_review_ids"]
    )
    assert tc0804_finding["recommendation"] == "建议明确咨询成果不符合约定或验收标准时的异议触发条件。"
    assert "明确明确" not in tc0804_finding["recommendation"]
    assert "建议补充明确咨询成果不符合约定或验收标准的异议触发条件明确" not in output["result"]
    assert "明确明确" not in output["result"]
    assert output["report_disposition_projection_gate"]["status"] == "REPORT_READY"
    assert output["report_integrity_gate"]["status"] == "REPORT_READY"


def test_missing_consulting_data_provision_is_low_risk_but_should_change():
    output = _consulting_data_review()
    by_id = _items(output)
    item = by_id["TC-08-03"]
    finding = next(
        row for row in output["risk_findings"]
        if "TC-08-03" in row["source_review_ids"]
    )

    assert item["coverage_status"] == "缺失"
    assert item["risk_severity"] == "LOW"
    assert item["action_priority"] == "B_PRACTICAL"
    assert item["risk_point_description"] == (
        "当前未完整明确委托方应提供的合理资料和协作范围，以及咨询方发现资料缺陷后的书面通知、"
        "影响说明、因果关系证明和顺延、费用调整或免责仅限实际受影响范围的边界。"
    )
    assert "真实性、完整性责任" not in finding["issue"]
    assert finding["recommendation"] == (
        "建议将委托方的资料提供和协作义务限定于开展咨询工作所合理必要的范围；"
        "咨询方发现资料缺失、矛盾或不适用时，应及时书面通知并说明具体影响。"
        "咨询方主张工期顺延、费用调整或免责的，应证明资料问题与受影响工作之间的因果关系，"
        "且顺延、调整或免责仅限于实际受影响范围。"
    )
    assert "建议明确委托方真实性、完整性及迟延责任" not in finding["recommendation"]
    disposition = next(
        row for row in output["final_dispositions"]
        if row["risk_id"] == finding["finding_id"]
    )
    assert disposition["final_action"] == "SHOULD_CHANGE"
    assert finding["issue"] in output["result"]
    assert finding["recommendation"] in output["result"]


def test_simple_consulting_data_provision_is_partial_not_missing():
    output = _consulting_data_review(
        "第三条 甲方应向乙方提供开展咨询所需的技术资料和相关数据。"
    )

    item = _items(output)["TC-08-03"]
    finding = next(
        row for row in output["risk_findings"]
        if "TC-08-03" in row["source_review_ids"]
    )
    truthfulness = next(
        row for row in item["coverage_elements"]
        if row["element_id"] == "truthfulness_completeness"
    )

    assert item["coverage_status"] == "部分覆盖"
    assert truthfulness["status"] == "MISSING"
    assert truthfulness["risk_relevant"] is False
    assert truthfulness["user_disposition"] == "KEEP_CURRENT"
    assert all(
        row.get("element_id") != "truthfulness_completeness"
        for row in finding["final_missing_elements"]
    )
    assert "资料真实、完整或准确义务" not in finding["issue"]
    assert "及时书面通知" in finding["recommendation"]
    assert "因果关系" in finding["recommendation"]
    assert "实际受影响范围" in finding["recommendation"]


def test_complete_consulting_data_provision_is_pass():
    output = _consulting_data_review(
        "第三条 甲方应在合同生效后3日内以电子文件方式向乙方提供设备图纸、技术资料和业务数据，"
        "并提供现场调研、测试环境和访问权限，保证资料真实、完整；乙方发现缺漏应通知甲方补充，"
        "甲方迟延提供或资料缺陷影响咨询工作的，相关延期和损失由甲方承担。"
    )

    assert _items(output)["TC-08-03"]["coverage_status"] == "通过"


def test_tc0803_priority_override_does_not_upgrade_other_low_missing_rules():
    output = _consulting_data_review()
    by_id = _items(output)

    assert by_id["TC-08-02"]["risk_severity"] == "LOW"
    assert by_id["TC-08-02"]["action_priority"] == "C_OPTIMIZATION"
    for item_id in ("TC-11-01", "TC-11-02", "TC-11-03", "TC-11-04"):
        if by_id[item_id]["coverage_status"] in {"缺失", "部分覆盖", "风险条款"}:
            assert by_id[item_id]["risk_severity"] == "LOW"
            assert by_id[item_id]["action_priority"] == "C_OPTIMIZATION"


def test_consulting_without_confidentiality_has_four_atomic_findings_but_one_theme():
    output = review_contract_v2(
        CONSULTING_THEME_AGGREGATION_CONTRACT,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/委托方"),
    )
    confidentiality_findings = [
        finding for finding in output["risk_findings"]
        if set(finding["source_review_ids"]) & {"TC-11-01", "TC-11-02", "TC-11-03", "TC-11-04"}
    ]
    confidentiality_themes = [
        theme for theme in output["risk_themes"] if theme["title"] == "保密机制不完整"
    ]

    assert len(confidentiality_findings) == 4
    assert len(confidentiality_themes) == 1
    assert confidentiality_themes[0]["underlying_finding_count"] == 4


def test_complete_consulting_confidentiality_produces_no_confidentiality_theme():
    output = review_contract_v2(
        CONSULTING_CONTRACT + COMPLETE_CONSULTING_CONFIDENTIALITY,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/委托方"),
    )
    by_id = _items(output)

    assert all(by_id[item_id]["coverage_status"] == "通过" for item_id in (
        "TC-11-01", "TC-11-02", "TC-11-03", "TC-11-04",
    ))
    assert not any(theme["title"] == "保密机制不完整" for theme in output["risk_themes"])
    assert not any(
        set(finding["source_review_ids"]) & {"TC-11-01", "TC-11-02", "TC-11-03", "TC-11-04"}
        for finding in output["risk_findings"]
    )
    assert output["release_quality_gate"]["status"] == "REPORT_READY"


@pytest.mark.parametrize("text", [
    """技术服务合同\n甲方委托乙方提供设备维护、故障排查和技术支持，甲方按月支付技术服务费。""",
    """软件委托开发合同\n甲方委托乙方研发软件并交付源代码和开发成果，甲方支付开发费用。""",
    """合作开发合同\n甲乙双方共同投入人员和经费，共同研究开发新技术并共享研发成果。""",
    """专利权转让合同\n乙方将专利权完整转让给甲方，双方办理专利权转让登记，甲方支付转让价款。""",
    """专利实施许可合同\n乙方许可甲方在中国境内实施专利，许可期限五年，甲方支付许可使用费。""",
])
def test_consulting_theme_projection_does_not_attach_to_other_technology_subtypes(text: str):
    output = review_contract_v2(text, "technology", use_llm=False)
    by_id = _items(output)

    assert output["deterministic_analysis"]["technology_classification"]["primary_subtype"] != "TECHNICAL_CONSULTING"
    assert all(by_id[item_id]["coverage_status"] == "不适用" for item_id in (
        "TC-08-01", "TC-08-02", "TC-08-03", "TC-08-04", "TC-08-05",
    ))
    assert not any(
        ((theme.get("aggregation_signals") or {}).get("user_visible_risk_cluster") or {}).get("scope") == "technical_consulting"
        for theme in output["risk_themes"]
    )


def test_vague_consulting_material_obligation_has_specific_attachment_gaps():
    output = review_contract_v2(
        CONSULTING_ELEMENT_EVIDENCE_CONTRACT + "\n第六条 乙方应提供相关技术资料。",
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/委托方"),
    )
    item = _items(output)["TC-13-01"]
    missing_ids = {
        row["element_id"] for row in item["coverage_elements"] if row["status"] == "MISSING"
    }

    assert item["coverage_status"] == "部分覆盖"
    assert missing_ids == {
        "material_names_and_types",
        "material_version_and_quantity",
        "material_delivery_form",
        "material_receipt_confirmation",
    }
    assert "技术资料附件”事项" not in item["narrative_state"]["judgment"]


CONSULTING_CONFIDENTIALITY_BINDING_CONTRACT = CONSULTING_ELEMENT_EVIDENCE_CONTRACT + """
第六条 保密范围与义务
保密资料包括与本项目有关的经营情况、财务情况、商业秘密、技术秘密、图纸和其他未公开信息。
双方应采取适当措施严格保密，未经信息拥有方书面同意，不得使用或向第三方泄露从对方获得的保密信息。
公众已知信息以及法律法规或司法程序要求披露的信息不属于违反保密义务，但依法披露前应尽合理努力通知信息拥有方。
在任何情形下，本条规定的保密义务应持续有效。
第七条 保密违约救济
任何一方违反保密义务的，守约方可以解除合同，并可采取法律允许的其他补救措施。
"""


def test_consulting_confidentiality_elements_bind_to_actual_clause_evidence():
    output = review_contract_v2(
        CONSULTING_CONFIDENTIALITY_BINDING_CONTRACT,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/委托方"),
    )
    by_id = _items(output)

    tc1101 = by_id["TC-11-01"]
    tc1101_elements = {row["element_id"]: row for row in tc1101["coverage_elements"]}
    assert tc1101["coverage_status"] == "部分覆盖"
    for element_id in (
        "confidential_information_scope",
        "confidentiality_obligation",
        "third_party_disclosure_restriction",
        "confidentiality_duration",
    ):
        assert tc1101_elements[element_id]["status"] == "COVERED"
        assert tc1101_elements[element_id]["evidence_ids"]
    assert tc1101_elements["purpose_limitation"]["status"] == "PARTIALLY_COVERED"
    assert tc1101_elements["purpose_limitation"]["evidence_ids"]
    tc1101_finding = next(
        finding for finding in output["risk_findings"]
        if "TC-11-01" in finding["source_review_ids"]
    )
    assert "保密信息仅可用于本合同约定的项目目的" in tc1101_finding["issue"]
    assert any(
        row["element_id"] == "purpose_limitation"
        for row in tc1101_finding["final_missing_elements"]
    )
    confidentiality_theme = next(
        theme for theme in output["risk_themes"]
        if theme["title"] == "保密机制不完整"
    )
    assert "项目目的使用" in confidentiality_theme["issue"]

    tc1102 = by_id["TC-11-02"]
    tc1102_elements = {row["element_id"]: row for row in tc1102["coverage_elements"]}
    assert tc1102["coverage_status"] == "部分覆盖"
    assert tc1102_elements["confidential_information_scope"]["status"] == "COVERED"
    assert tc1102_elements["confidentiality_period"]["status"] == "COVERED"
    assert tc1102_elements["public_information_exception"]["status"] == "COVERED"
    assert tc1102_elements["required_disclosure_exception"]["status"] == "PARTIALLY_COVERED"
    assert tc1102_elements["lawful_third_party_source_exception"]["status"] == "MISSING"
    assert tc1102_elements["independent_development_exception"]["status"] == "MISSING"
    assert "终止后三年" not in output["result"]

    tc1104 = by_id["TC-11-04"]
    tc1104_elements = {row["element_id"]: row for row in tc1104["coverage_elements"]}
    assert tc1104["coverage_status"] == "部分覆盖"
    assert tc1104_elements["confidentiality_obligation"]["status"] == "COVERED"
    assert tc1104_elements["confidentiality_obligation"]["evidence_ids"]
    assert tc1104_elements["breach_damage_liability"]["status"] == "PARTIALLY_COVERED"
    assert tc1104_elements["breach_damage_liability"]["evidence_ids"]
    assert tc1104_elements["stop_use"]["status"] == "MISSING"
    assert tc1104_elements["return_delete_destroy"]["status"] == "MISSING"

    for item_id in ("TC-11-01", "TC-11-02", "TC-11-04"):
        item = by_id[item_id]
        for element in item["coverage_elements"]:
            if element["status"] in {"COVERED", "PARTIALLY_COVERED"}:
                assert element["evidence_ids"], (item_id, element["element_id"])
        assert not (
            item["coverage_status"] in {"通过", "部分覆盖"}
            and item["evidence_location"] == "合同正文及附件中未找到对应约定"
        )


def test_consulting_material_direct_obligation_ranks_before_collection_progress():
    output = review_contract_v2(
        CONSULTING_ELEMENT_EVIDENCE_CONTRACT,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/委托方"),
    )
    item = _items(output)["TC-08-03"]
    first = item["evidence_clauses"][0]["raw_text"]
    assert "甲方应向乙方提供" in first
    assert "资料收集" not in first


def test_contract_section_parser_keeps_confidentiality_section_visible_with_offsets():
    text = """
第一条 服务
乙方提供技术咨询服务。
第二条 付款
甲方应支付咨询报酬。
第三条 责任
违约方应承担违约责任。
第四条 保密
双方应对商业秘密和技术秘密严格保密，未经书面同意不得向第三方披露。本条保密义务持续有效。
第五条 保证
乙方保证咨询报告符合约定。
"""
    clauses = parse_contract_clauses(text)
    article_clauses = {
        clause.location.split("第")[1].split("条")[0]: clause
        for clause in clauses
        if clause.location.startswith("第") and "条" in clause.location
    }

    for article in ("一", "二", "三", "四", "五"):
        assert article in article_clauses
        clause = article_clauses[article]
        assert clause.start_offset >= 0
        assert clause.end_offset > clause.start_offset
    confidentiality = article_clauses["四"]
    assert "商业秘密" in confidentiality.text
    assert "不得向第三方披露" in confidentiality.text
    assert "持续有效" in confidentiality.text


def test_production_file_endpoint_uses_active_consulting_evidence_implementation():
    document = Document()
    for line in CONSULTING_CONFIDENTIALITY_BINDING_CONTRACT.strip().splitlines():
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)

    app.dependency_overrides[get_current_user] = lambda: {"id": 1, "phone": "test"}
    try:
        with TestClient(app) as client:
            response = client.post(
                "/review-v2/file",
                files={
                    "file": (
                        "consulting-production-path.docx",
                        buffer.getvalue(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                data={
                    "contract_type": "technology",
                    "include_low_risk": "true",
                    "use_llm": "false",
                    "profile_json": '{"user_identity":"甲方/委托方"}',
                },
            )
    finally:
        app.dependency_overrides.pop(get_current_user, None)

    assert response.status_code == 200
    payload = response.json()
    trace = payload["production_path_trace"]
    assert trace["entrypoint"] == "routers.review_v2.review_v2_file"
    assert trace["engine_function"] == "review_v2.engine.review_contract_v2"
    assert trace["evidence_implementation"] == TECHNICAL_CONSULTING_EVIDENCE_IMPLEMENTATION
    assert trace["cache"]["status"] == "MISS"
    assert trace["input_visibility"]["confidentiality_terms_present"] is True
    assert trace["pipeline_id"] == "canonical-review-v2-production-6.4"
    assert trace["classifier_version"] == "technical-transaction-classifier-v3-single-source"
    assert trace["classification_runtime"]["actual_execution_count"] == 1
    assert trace["classification_runtime"]["invocation_count"] >= 1
    assert trace["classification_runtime"]["cache_hit_count"] == trace["classification_runtime"]["invocation_count"] - 1
    lifecycle = trace["classification_lifecycle"]
    assert lifecycle["raw_classification"]["primary_subtype"] == "TECHNICAL_CONSULTING"
    assert lifecycle["normalized_classification"]["contract_subtype"] == "TECHNICAL_CONSULTING"
    assert lifecycle["review_context"]["contract_subtype"] == "TECHNICAL_CONSULTING"
    assert lifecycle["rule_router"]["contract_subtype"] == "TECHNICAL_CONSULTING"
    assert lifecycle["renderer"]["contract_subtype"] == "TECHNICAL_CONSULTING"
    assert lifecycle["renderer"]["scene"] == "技术咨询"
    assert lifecycle["renderer"]["contract_type_label"] == "技术合同—技术咨询合同"

    by_id = {item["item_id"]: item for item in payload["items"]}
    for item_id in ("TC-11-01", "TC-11-02", "TC-11-03", "TC-11-04", "TC-08-03"):
        assert by_id[item_id]["decision_trace"]["implementation_identity"]["version"] == TECHNICAL_CONSULTING_EVIDENCE_IMPLEMENTATION
    assert by_id["TC-11-01"]["coverage_status"] == "部分覆盖"
    assert by_id["TC-11-02"]["coverage_status"] == "部分覆盖"
    assert by_id["TC-11-03"]["coverage_status"] == "缺失"
    assert by_id["TC-11-04"]["coverage_status"] == "部分覆盖"
    assert "甲方应向乙方提供" in by_id["TC-08-03"]["evidence_clauses"][0]["raw_text"]
    assert by_id["TC-01-03"]["coverage_status"] == "通过"
    assert all(by_id[f"TC-08-0{index}"]["coverage_status"] != "不适用" for index in range(1, 6))
    assert all(by_id[f"TC-10-0{index}"]["coverage_status"] == "不适用" for index in range(1, 5))
    assert by_id["TC-13-01"]["coverage_status"] != "缺失"


def test_legacy_batch_entry_rejects_confident_technical_contract(monkeypatch):
    import security.middleware as security_middleware
    import utils.security as security_utils

    document = Document()
    for line in CONSULTING_CONTRACT.strip().splitlines():
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)

    class AllowRateLimiter:
        def is_allowed(self, *args, **kwargs):
            return True, ""

    monkeypatch.setattr(legacy_main, "get_user_bonus_credits", lambda user_id: 3)
    monkeypatch.setattr(legacy_main.blacklist_manager, "check", lambda *args, **kwargs: False)
    monkeypatch.setattr(legacy_main, "is_auto_banned", lambda **kwargs: (False, ""))
    monkeypatch.setattr(legacy_main, "get_rate_limiter", lambda: AllowRateLimiter())
    monkeypatch.setattr(security_middleware, "check_blacklist", lambda *args, **kwargs: None)
    monkeypatch.setattr(security_middleware, "get_rate_limiter", lambda: AllowRateLimiter())
    monkeypatch.setattr(security_utils, "check_auto_ban", lambda *args, **kwargs: False)
    monkeypatch.setenv("ENABLE_DUPLICATE_CHECK", "false")
    app.dependency_overrides[get_current_user] = lambda: 1
    try:
        with TestClient(app) as client:
            response = client.post(
                "/api/batch/review",
                files={
                    "files": (
                        "consulting-must-use-canonical.docx",
                        buffer.getvalue(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                data={"additional_instructions": ""},
            )
    finally:
        app.dependency_overrides.pop(get_current_user, None)

    assert response.status_code == 400
    detail = response.json()["detail"]
    rejected = detail["rejected_files"][0]
    assert rejected["error_code"] == "CANONICAL_TECHNICAL_PIPELINE_REQUIRED"
    assert rejected["canonical_endpoint"] == "/review-v2/file"
    assert rejected["detected_subtype"] == "TECHNICAL_CONSULTING"


def test_production_file_endpoint_is_deterministic_across_fresh_no_cache_runs():
    document = Document()
    for line in CONSULTING_CONTRACT.strip().splitlines():
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    file_bytes = buffer.getvalue()

    app.dependency_overrides[get_current_user] = lambda: {"id": 1, "phone": "test"}
    try:
        with TestClient(app) as client:
            payloads = []
            for _ in range(2):
                response = client.post(
                    "/review-v2/file",
                    files={
                        "file": (
                            "consulting-no-cache.docx",
                            file_bytes,
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        )
                    },
                    data={
                        "contract_type": "technology",
                        "include_low_risk": "true",
                        "use_llm": "false",
                        "profile_json": '{"user_identity":"甲方/委托方"}',
                    },
                )
                assert response.status_code == 200
                payloads.append(response.json())
    finally:
        app.dependency_overrides.pop(get_current_user, None)

    first, second = payloads
    assert first["production_path_trace"]["cache"]["status"] == "MISS"
    assert second["production_path_trace"]["cache"]["status"] == "MISS"
    assert first["production_path_trace"]["document_hash"] == second["production_path_trace"]["document_hash"]
    assert first["report_context"] == second["report_context"]
    assert first["result"] == second["result"]
    first_items = {item["item_id"]: item for item in first["items"]}
    second_items = {item["item_id"]: item for item in second["items"]}
    for item_id in ("TC-01-03", "TC-08-01", "TC-08-02", "TC-08-03", "TC-08-04", "TC-08-05", "TC-10-01", "TC-11-01", "TC-11-02", "TC-11-04", "TC-13-01"):
        assert first_items[item_id]["coverage_status"] == second_items[item_id]["coverage_status"]


def test_consulting_eur_mixed_payment_bases_and_direction_are_deterministic():
    text = """
技术咨询合同
合同总价为EUR 1,200,000。
分项一的合同价为EUR 300,000；分项二的合同价为EUR 300,000；分项三的合同价为EUR 300,000；分项四的合同价为EUR 300,000。
合同总价的20%，即EUR 240,000，支付给咨询方。
分项一合同价的80%，即EUR 240,000，支付给咨询方。
分项二合同价的80%，即EUR 240,000，支付给委托方。
分项三合同价的80%，即EUR 240,000，支付给咨询方。
分项四合同价的80%，即EUR 240,000，支付给咨询方。
分项四合同价的20%，即EUR 60,000，支付给咨询方。
"""
    payment = analyze_payment_schedule(text)

    assert payment["status"] == "INCONSISTENT"
    assert payment["contract_total"] == 1_200_000
    assert payment["amount_total"] == 1_260_000
    assert payment["percentage_total"] == 105
    assert payment["amount_delta"] == 60_000
    assert payment["currency"] == "EUR"
    assert payment["payment_fact_count"] == 6
    assert payment["dominant_payee"] == "咨询方"
    assert [row["payee"] for row in payment["payment_direction_outliers"]] == ["委托方"]
    whole_contract_node = payment["nodes"][0]
    assert whole_contract_node["payment_group_label"] in {"合同总价", "合同总价款"}


def test_consulting_relative_duration_cap_and_weekly_penalty_are_shared_facts():
    text = """
技术咨询合同
技术咨询服务自合同生效之日起8个月内完成。
合同生效后9个月内提交分项四报告，10个月内提交最终报告，11个月内完成最终验收。
咨询方因服务疏忽造成人身损害和财产损失的，应承担赔偿责任。
咨询方对本合同的任何和所有责任都限定在合同总价之内。
迟延第一至第四周每周支付合同总价2%，第五至第八周每周3%，第九周起每周5%。
迟延交付违约金总额不得超过合同总价的30%。
"""
    timeline = analyze_timeline(text)
    risks = analyze_technology_relationship_risks(text)

    conflicts = [
        row for row in timeline["conflicts"]
        if row["type"] == "MILESTONE_AFTER_PERFORMANCE_DURATION"
    ]
    assert timeline["performance_duration_months"] == 8
    assert {row["milestone_offset_months"] for row in conflicts} == {9, 10, 11}
    assert risks["liability_cap_validity"]["present"] is True
    assert risks["liability_cap_validity"]["analysis"]["personal_injury_within_broad_cap"] is True
    penalty = risks["potential_excessive_liquidated_damages"]
    assert penalty["present"] is True
    weekly = penalty["analysis"]["weekly_penalties"]
    assert sorted({rate for row in weekly for rate in row["rates_percent_per_week"]}) == [2, 3, 5]
    assert max(row["max_rate_percent_per_week"] for row in weekly) == 5
    assert penalty["analysis"]["cumulative_penalty_caps"][0]["cap_percent"] == 30


def test_latest_consulting_docx_full_regression_has_no_subtype_or_evidence_regression():
    path = Path(__file__).resolve().parents[2] / "test_data" / "33 技术咨询服务合同.docx"
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    output = review_contract_v2(
        text,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/委托方"),
    )
    by_id = _items(output)
    titles = {row["title"] for row in output["risk_findings"]}
    report = output["result"]

    assert output["deterministic_analysis"]["technology_classification"]["primary_subtype"] == "TECHNICAL_CONSULTING"
    assert all(by_id[item_id]["coverage_status"] == "不适用" for item_id in (
        "TC-09-01", "TC-09-02", "TC-09-03", "TC-09-04",
    ))
    assert by_id["TC-01-03"]["coverage_status"] == "通过"
    assert by_id["TC-08-01"]["coverage_status"] == "通过"
    assert "交付形式明确" in by_id["TC-08-01"]["narrative_state"]["judgment"]
    assert "资料提供时间和方式明确" in by_id["TC-08-03"]["narrative_state"]["judgment"]
    tc1104 = {row["element_id"]: row["status"] for row in by_id["TC-11-04"]["coverage_elements"]}
    assert tc1104["termination_remedy"] == "COVERED"
    assert "合同解除救济" in by_id["TC-11-04"]["risk_point_description"]
    assert by_id["TC-13-04"]["coverage_status"] != "不适用"

    assert {
        "付款比例、节点金额与合同总价不一致",
        "服务期限与交付验收节点冲突",
        "依法不得限责情形仍适用责任上限的效力风险",
        "违约金标准偏高及司法调减风险",
        "技术进出口管制附件",
    } <= titles
    for marker in ("1,260,000", "105%", "60,000", "支付给委托方", "第9、10、11个月"):
        assert marker in report
    assert "延至第11个月" in report
    assert "延至第10个月" not in report
    assert "可选优化主题数（聚合后）：" in report
    assert "- 可选优化数：" not in report

    for item_id in ("GEN-LIAB-02", "GEN-LIAB-03"):
        assert by_id[item_id]["economic_impact"] == "medium"
        assert by_id[item_id]["economic_impact_detail"]["level"] == "medium"
    findings_by_group = {row["risk_group"]: row for row in output["risk_findings"]}
    assert findings_by_group["generic_liability_cap_validity"]["economic_impact"]["level"] == "medium"
    liability_finding = findings_by_group["generic_liability_cap_validity"]
    assert "造成对方人身损害" in liability_finding["consequence"]
    assert "因故意或者重大过失造成对方财产损失" in liability_finding["consequence"]
    assert "因过错造成的财产损失" not in report
    assert "保留当前审查方依法追偿的权利" in liability_finding["recommendation"]
    penalty_impact = findings_by_group["generic_potential_excessive_liquidated_damages"]["economic_impact"]
    assert penalty_impact["level"] == "medium"
    assert penalty_impact["amount_estimate"] == 360_000
    assert "潜在可主张金额为EUR 360,000，占合同总价30%" in report
    assert "该金额由相对方向当前审查方支付，不属于当前审查方的资金暴露" in report
    assert "直接资金暴露为EUR 360,000" not in report
    assert "保留当前审查方可主张的违约救济" in findings_by_group[
        "generic_potential_excessive_liquidated_damages"
    ]["recommendation"]

    confidentiality = findings_by_group["technical_secret_confidentiality"]
    assert "明确保密信息仅可用于本合同约定的项目目的" in confidentiality["recommendation"]
    assert "补充明确未" not in confidentiality["recommendation"]
    assert "第874条" not in report
    assert "第884条" not in report
    assert output["release_quality_gate"]["status"] == "REPORT_READY"
