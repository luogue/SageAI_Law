import io
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from auth.routes import get_current_user
from main import app


ROOT = Path(__file__).resolve().parents[2]
SOURCE_CONTRACT = ROOT / "test_data" / "8、技术服务.docx"
SOURCE_CONSULTING_CONTRACT = ROOT / "test_data" / "16 技术咨询合同.docx"


def _source_text() -> str:
    document = Document(str(SOURCE_CONTRACT))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = [
        "\t".join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
    ]
    return "\n".join([*paragraphs, *tables])


def _consulting_variant(text: str) -> str:
    return (
        text.replace("技术服务合同", "技术咨询合同")
        .replace("技术服务", "技术咨询")
        .replace("服务内容", "咨询内容")
        .replace("服务方式", "咨询方式")
    )


def _docx_bytes(text: str) -> bytes:
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _review_file(file_bytes: bytes, filename: str) -> dict:
    app.dependency_overrides[get_current_user] = lambda: {"id": 1, "phone": "test"}
    try:
        with TestClient(app) as client:
            response = client.post(
                "/review-v2/file",
                files={
                    "file": (
                        filename,
                        file_bytes,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                data={
                    "contract_type": "technology",
                    "include_low_risk": "true",
                    "use_llm": "false",
                    "profile_json": '{"user_identity":"甲方/委托方"}',
                },
            )
    finally:
        app.dependency_overrides.pop(get_current_user, None)
    assert response.status_code == 200
    return response.json()


@pytest.fixture(scope="module")
def consulting_relationship_review() -> dict:
    return _review_file(
        _docx_bytes(_consulting_variant(_source_text())),
        "relationship-risk-consulting.docx",
    )


def _groups(payload: dict) -> set[str]:
    return {finding["risk_group"] for finding in payload["risk_findings"]}


def test_real_loader_parser_pipeline_keeps_general_relationship_risks_for_consulting(
    consulting_relationship_review: dict,
):
    payload = consulting_relationship_review
    assert payload["contract_context"]["contract_subtype"] == "TECHNICAL_CONSULTING"
    assert payload["report_context"]["scenario_label"] == "技术咨询"
    assert {
        "payment_formula_inconsistency",
        "technical_timeline_conflict",
        "data_retention_warranty_conflict",
        "personal_information_processing_gap",
        "technical_unilateral_termination_change",
        "technical_uncapped_cumulative_liability",
        "technical_background_ip_overtransfer",
    } <= _groups(payload)

    trace = payload["observability"]["risk_recall_lifecycle"]
    assert trace["subtype_independent"] is True
    assert trace["candidate_count"] >= 7
    assert trace["dropped_groups"] == []
    assert trace["candidate_count"] == trace["final_count"]

    by_id = {item["item_id"]: item for item in payload["items"]}
    for item_id in ("TC-04-01", "TC-05-01", "TC-06-01", "TC-07-01"):
        assert by_id[item_id]["coverage_status"] == "不适用"
    assert by_id["TC-03-04"]["coverage_status"] != "不适用"
    assert by_id["TC-13-02"]["coverage_status"] != "不适用"

    confidentiality_groups = {
        "technical_secret_confidentiality",
        "confidentiality_breach_liability",
        "confidentiality_period_and_scope",
        "confidential_material_return_destroy",
    }
    assert confidentiality_groups <= _groups(payload)

    payment_trace = payload["observability"]["payment_risk_lifecycle"]
    events = {event["event"]: event for event in payment_trace["events"]}
    assert payment_trace["status"] == "FINAL_RISK"
    assert events["[RULE_EXECUTED]"]["invocation_count"] >= 1
    assert events["[CANDIDATE_CREATED]"]["created"] is True
    assert events["[CANDIDATE_FILTERED]"]["status"] == "NOT_FILTERED"
    assert events["[CANDIDATE_MERGED]"]["merged"] is True
    assert events["[FINAL_RISK]"]["present"] is True
    assert events["[FINAL_RISK]"]["rendered"] is True


def test_current_consulting_fixture_keeps_amount_only_node_without_false_formula_risk():
    payload = _review_file(SOURCE_CONSULTING_CONTRACT.read_bytes(), SOURCE_CONSULTING_CONTRACT.name)
    payment = payload["deterministic_analysis"]["payment_formula"]
    trace = payload["observability"]["payment_risk_lifecycle"]
    events = {event["event"]: event for event in trace["events"]}

    assert payload["contract_context"]["contract_subtype"] == "TECHNICAL_CONSULTING"
    assert payment["status"] == "INSUFFICIENT_DATA"
    assert len(payment["nodes"]) == 1
    assert payment["nodes"][0]["amount"] == 54000
    assert payment["nodes"][0]["percentage"] is None
    assert trace["status"] == "NO_CANDIDATE"
    assert events["[CANDIDATE_CREATED]"]["created"] is False
    assert "arithmetic_status=INSUFFICIENT_DATA" in events["[CANDIDATE_CREATED]"]["reason"]
    assert events["[FINAL_RISK]"]["present"] is False


def test_payment_percentage_mutation_removes_only_formula_conflict():
    text = _consulting_variant(_source_text())
    text = text.replace(
        "合同总额的25%，即人民币120,000元",
        "合同总额的20%，即人民币96,000元",
    )
    payload = _review_file(_docx_bytes(text), "payment-consistent-consulting.docx")
    formula = payload["deterministic_analysis"]["payment_formula"]
    assert formula["percentage_total"] == 100
    assert formula["amount_total"] == formula["contract_total"] == 480000
    assert formula["status"] == "CONSISTENT"
    assert "payment_formula_inconsistency" not in _groups(payload)


def test_changed_total_with_same_105_percent_still_triggers_formula_conflict():
    text = _consulting_variant(_source_text())
    replacements = {
        "480,000": "960,000",
        "144,000": "288,000",
        "192,000": "384,000",
        "120,000": "240,000",
        "48,000": "96,000",
    }
    for original, replacement in replacements.items():
        text = text.replace(original, replacement)
    payload = _review_file(_docx_bytes(text), "payment-scaled-consulting.docx")
    formula = payload["deterministic_analysis"]["payment_formula"]
    assert formula["contract_total"] == 960000
    assert formula["percentage_total"] == 105
    assert formula["amount_delta"] == 48000
    assert "payment_formula_inconsistency" in _groups(payload)


def test_post_signature_timeline_mutation_removes_timeline_conflict():
    text = _consulting_variant(_source_text()).replace("2026年9月15日", "2026年8月15日")
    payload = _review_file(_docx_bytes(text), "timeline-consistent-consulting.docx")
    assert payload["deterministic_analysis"]["timeline"]["conflicts"] == []
    assert "technical_timeline_conflict" not in _groups(payload)


def test_removing_personal_data_facts_removes_processing_risk():
    text = _consulting_variant(_source_text())
    text = text.replace(
        "甲方向乙方提供的数据可能包括师生姓名、工号、电话号码、电子邮箱、科研项目信息、项目经费信息及系统操作日志。",
        "甲方向乙方仅提供不含可识别自然人信息的汇总统计资料。",
    )
    text = (
        text.replace("电子信箱", "项目联系渠道")
        .replace("电子邮箱：", "项目联系渠道：")
        .replace("系统操作日志", "系统统计指标")
    )
    payload = _review_file(_docx_bytes(text), "no-personal-data-consulting.docx")
    assert "personal_information_processing_gap" not in _groups(payload)
