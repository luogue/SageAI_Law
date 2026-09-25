from pathlib import Path

from docx import Document

from review_v2.domain_resolution import resolve_contract_domain
from review_v2.engine import review_contract_v2
from review_v2.clause_parser import parse_contract_clauses


ROOT = Path(__file__).resolve().parents[2]


def _docx_text(path: Path) -> str:
    document = Document(path)
    rows = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            value = " | ".join(cell.text.strip() for cell in row.cells)
            if value.strip(" |"):
                rows.append(value)
    return "\n".join(rows)


def _hb003_output():
    path = next((ROOT / "test_data").glob("6*.docx"))
    return review_contract_v2(_docx_text(path), "sale", use_llm=False)


def test_sale_hb003_routes_to_residential_real_estate_and_isolates_goods_reviews():
    output = _hb003_output()
    context = output["contract_context"]
    assert context["contract_family"] == "SALE"
    assert context["subject_matter"] == "REAL_ESTATE"
    assert context["transaction_nature"] == "OWNERSHIP_TRANSFER"
    assert context["asset_usage"] == "RESIDENTIAL"
    assert context["contract_subtype"] == "RESIDENTIAL_SECONDARY_REAL_ESTATE_SALE"
    assert context["user_role"] == "buyer"
    assert output["contract_type_name"] == "住宅二手房地产买卖合同"
    owners = output["observability"]["review_registry_debug"]["domain_owners"]
    assert owners == ["sale/common", "sale/real_estate"]
    assert output["observability"]["review_registry_debug"]["prohibited_reviews"] == []


def test_sale_hb003_extracts_material_real_estate_facts_and_passes_domain_gates():
    output = _hb003_output()
    facts = output["contract_facts"]["real_estate"]
    assert facts["seller_name"] == "张伟民"
    assert facts["buyer_name"] == "李若晴"
    assert facts["subject"]["certificate_number"] == "沪（2020）浦字不动产权第012345号"
    assert facts["subject"]["registered_owner"] == "张伟民"
    assert facts["subject"]["address"] == "上海市浦东新区张杨路1888弄8号1202室"
    assert facts["subject"]["building_area_sqm"] == 96.42
    assert facts["subject"]["land_term_start"] == "2002-06-01"
    assert facts["subject"]["land_term_end"] == "2072-05-31"
    assert facts["financing"] == {
        "total_price": 6880000.0, "deposit": 200000.0, "first_payment": 1800000.0,
        "cumulative_payment": 2000000.0, "loan_amount": 3500000.0,
        "balance": 1380000.0, "household_holdback": 100000.0,
    }
    assert facts["mortgage"]["release_deadline"] == "2026-09-05"
    assert facts["mortgage"]["payment_suspension_trigger"] == "2026-09-15"
    assert facts["mortgage"]["termination_grace_days"] == 10
    assert facts["delivery"]["delivery_date"] == "2026-09-30"
    assert facts["delivery"]["inspection_deadline"] == "3日内"
    assert facts["residential"]["household_count"] == 2
    assert facts["residential"]["spouse_name"] == "王静怡"
    gates = output["release_quality_gate"]["gate_results"]
    assert gates["DOMAIN_SUBTYPE_GATE"] == "PASS"
    assert gates["DOMAIN_REVIEW_ISOLATION_GATE"] == "PASS"
    assert gates["DOMAIN_MATERIAL_COVERAGE_GATE"] == "PASS"
    assert gates["NARRATIVE_TRUTH_GATE"] == "PASS"
    assert output["release_quality_gate"]["status"] == "REPORT_READY"


def test_real_estate_domain_requires_multiple_independent_signals():
    weak = "房屋内使用的机器设备买卖合同，设备名称、规格型号、数量及交货地点详见附件。"
    weak_resolution = resolve_contract_domain("sale", weak, parse_contract_clauses(weak))
    assert weak_resolution.subject_matter != "REAL_ESTATE"

    no_title_words = "住宅买卖合同。双方约定转移登记并到房地产交易中心办理过户，房屋连同土地使用权一并转让。"
    resolution = resolve_contract_domain("sale", no_title_words, parse_contract_clauses(no_title_words))
    assert resolution.subject_matter == "REAL_ESTATE"
    assert len({row.signal_type for row in resolution.evidence}) >= 2


def test_real_estate_pack_uses_domain_metadata_not_fixture_values():
    text = """住宅房地产买卖合同
    出卖人（甲方）：甲某；买受人（乙方）：乙某。
    不动产权证号为苏（2025）宁字不动产权第999号，登记权利人为甲某。
    房地产坐落南京市某区某路1号，建筑面积88平方米，土地用途为住宅用地。
    双方于登记机构办理转移登记和预告登记，房屋交付时办理钥匙交接。
    """
    output = review_contract_v2(text, "sale", use_llm=False)
    assert output["contract_context"]["subject_matter"] == "REAL_ESTATE"
    assert "sale/goods" not in output["observability"]["review_registry_debug"]["domain_owners"]
