from pathlib import Path
import re

import pytest
from docx import Document

from review_v2.engine import review_contract_v2
from review_v2.user_profile import UserRiskProfile


ROOT = Path(__file__).resolve().parents[2]
FIXTURE = ROOT / "test_data" / "36 技术开发（合作）合同.docx"


def _document_text(path: Path) -> str:
    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = [
        "\t".join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
    ]
    return "\n".join(paragraphs + tables)


@pytest.fixture(scope="module")
def cooperative_review() -> dict:
    return review_contract_v2(
        _document_text(FIXTURE),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/合作开发方"),
    )


def _items(review: dict) -> dict:
    return {row["item_id"]: row for row in review["items"]}


def _elements(item: dict) -> dict:
    return {row["element_id"]: row for row in item["coverage_elements"]}


def test_cooperative_investment_and_work_facts_are_atomically_resolved(cooperative_review: dict):
    item = _items(cooperative_review)["TC-05-01"]
    elements = _elements(item)

    for element_id in (
        "research_task_allocation",
        "development_schedule_and_milestones",
        "funding_investment",
        "technical_and_personnel_investment",
        "data_equipment_environment_investment",
    ):
        assert elements[element_id]["status"] == "COVERED"

    assert elements["party_investment_total_consistency"]["status"] == "COVERED"
    assert elements["nominal_investment_payment_consistency"]["status"] == "CONFLICTING"
    assert elements["equity_investment_ratio_basis"]["status"] == "PARTIALLY_COVERED"
    analysis = item["decision_trace"]["cooperative_investment_consistency"]
    assert analysis["stated_project_total"] == 2_200_000
    assert analysis["declared_investment_sum"] == 2_200_000
    assert analysis["payment_conflicts"] == [{
        "party": "甲方",
        "nominal_investment": 1_200_000,
        "payment_total": 1_250_000,
        "difference": 50_000,
        "nodes": [400_000, 400_000, 450_000],
    }]
    assert analysis["equity_mismatch_is_illegality"] is False

    absence = item["decision_trace"]["specialized_search_audit"]["missing_elements"]
    assert not any(
        covered in missing
        for covered in ("研发任务", "研发期限", "资金投入", "数据、设备")
        for missing in absence
    )


def test_final_result_patent_license_and_revenue_rights_are_compared(cooperative_review: dict):
    item = _items(cooperative_review)["TC-05-02"]
    elements = _elements(item)

    assert item["coverage_status"] == "风险条款"
    assert elements["final_result_ownership"]["status"] == "COVERED"
    assert elements["patent_application_right"]["status"] == "CONFLICTING"
    assert elements["patent_right_after_grant"]["status"] == "CONFLICTING"
    assert elements["external_license_right"]["status"] == "ADVERSE"
    assert elements["external_license_revenue"]["status"] == "ADVERSE"
    assert "共同所有" in item["risk_point_description"]
    assert "收益归该许可方单独所有" in item["risk_point_description"]


def test_existing_failure_and_exit_terms_are_adverse_not_missing(cooperative_review: dict):
    item = _items(cooperative_review)["TC-05-04"]
    elements = _elements(item)

    assert item["coverage_status"] == "风险条款"
    for element_id in (
        "failure_payment_allocation",
        "stage_result_handover",
        "technical_risk_determination",
        "reviewer_exit_cost_and_ip",
        "counterparty_exit_liability",
    ):
        assert elements[element_id]["status"] == "ADVERSE"
    assert "已付款不退" in item["risk_point_description"]
    assert "未付费用仍需继续支付" in item["risk_point_description"]


def test_governance_deadlock_silence_and_contact_authority_are_one_control_risk(cooperative_review: dict):
    item = _items(cooperative_review)["TECH-SHARED-GOV-01"]
    elements = _elements(item)

    assert item["issue_type"] == "CONFLICTING"
    assert elements["major_matter_voting"]["status"] == "COVERED"
    assert elements["deadlock_resolution"]["status"] == "CONFLICTING"
    assert elements["silence_as_major_change_consent"]["status"] == "ADVERSE"
    assert elements["project_contact_authority"]["status"] == "ADVERSE"
    trace = item["decision_trace"]["cooperative_governance_analysis"]
    assert all(trace.values())


def test_component_control_and_loss_direction_are_separate_from_claim_procedure(cooperative_review: dict):
    item = _items(cooperative_review)["TC-12-02"]
    elements = _elements(item)

    assert elements["component_selection_control"]["status"] == "COVERED"
    assert elements["component_license_cost_allocation"]["status"] == "ADVERSE"
    assert elements["component_infringement_loss_allocation"]["status"] == "ADVERSE"
    assert elements["claim_notice"]["status"] == "MISSING"
    assert elements["defense_and_settlement_control"]["status"] == "MISSING"
    evidence_text = "\n".join(row.get("text") or "" for row in item["evidence_clauses"])
    assert all(term in evidence_text for term in ("第三方软件", "许可费用", "甲方承担50%"))


def test_result_reuse_data_reuse_and_follow_on_improvement_remain_distinct(cooperative_review: dict):
    by_id = _items(cooperative_review)
    result_rights = by_id["TC-05-03"]
    data_use = by_id["TECH-SHARED-DATA-01"]
    improvement = by_id["TC-10-02"]

    assert result_rights["risk_group"] == "cooperative_result_rights_and_commercialization"
    assert data_use["risk_group"] == "technology_data_secondary_use"
    assert improvement["risk_group"] == "improvement_sharing_license"
    assert _elements(data_use)["project_purpose_limitation"]["status"] == "ADVERSE"
    assert _elements(data_use)["post_termination_retention_and_deletion"]["status"] == "ADVERSE"
    assert _elements(improvement)["improvement_ownership_boundary"]["status"] == "COVERED"
    assert _elements(improvement)["improvement_disclosure"]["status"] == "ADVERSE"
    improvement_evidence = "\n".join(row.get("text") or "" for row in improvement["evidence_clauses"])
    assert "归实际完成改进的一方独立所有" in improvement_evidence
    assert "共同改进成果归合作各方共同" not in cooperative_review["result"]


def test_reviewer_funded_equipment_transfer_is_an_independent_adverse_risk(cooperative_review: dict):
    item = _items(cooperative_review)["TECH-SHARED-ASSET-01"]
    elements = _elements(item)

    assert elements["reviewer_funded_asset_ownership"]["status"] == "ADVERSE"
    assert elements["jointly_funded_asset_ownership"]["status"] == "ADVERSE"
    assert elements["return_or_value_compensation"]["status"] == "ADVERSE"
    assert "甲方资金购买" in item["risk_point_description"]


def test_complex_five_day_deemed_acceptance_preserves_latent_defect_rights(cooperative_review: dict):
    item = _items(cooperative_review)["TECH-SHARED-12"]
    elements = _elements(item)

    assert item["issue_type"] == "ADVERSE"
    assert elements["short_deemed_acceptance"]["status"] == "ADVERSE"
    assert elements["continuous_performance_observation"]["status"] == "COVERED"
    assert elements["latent_defect_and_warranty_carveout"]["status"] == "MISSING"
    assert item["decision_trace"]["deemed_acceptance_analysis"] == {
        "deadline_working_days": 5,
        "complex_project": True,
        "sustained_observation_required": True,
        "latent_defect_carveout": False,
    }


def test_domestic_cooperation_does_not_trigger_cross_border_or_personal_result_rules(cooperative_review: dict):
    by_id = _items(cooperative_review)
    assert by_id["TC-13-04"]["coverage_status"] == "不适用"
    assert by_id["TC-03-01"]["coverage_status"] == "不适用"
    assert by_id["TC-03-02"]["coverage_status"] == "不适用"
    assert "签署日期：2026年9月2日\n本合同" not in "\n".join(
        row.get("text") or "" for row in by_id["TC-13-04"].get("evidence_clauses") or []
    )


def test_shared_governance_and_creator_facts_survive_final_projection(cooperative_review: dict):
    by_id = _items(cooperative_review)
    investment = by_id["TC-05-01"]
    creator = by_id["TC-03-03"]

    major = _elements(investment)["major_decision_process"]
    assert major["status"] == "COVERED"
    assert major["required_fact_key"] == "technology.governance.major_decision"
    assert major["source_clause_ids"] == ["C0007"]
    assert not any(
        "重大" in value and "决策" in value
        for value in investment["decision_trace"]["specialized_search_audit"]["missing_elements"]
    )

    assert creator["coverage_status"] == "部分覆盖"
    assert creator["applicability"]["applicable"] is True
    assert _elements(creator)["creator_attribution"]["status"] == "COVERED"
    assert _elements(creator)["creator_list_determination"]["status"] == "COVERED"
    assert _elements(creator)["creator_reward_arrangement"]["status"] == "MISSING"


def test_attachment_reference_and_attachment_content_are_distinct_facts(cooperative_review: dict):
    by_id = _items(cooperative_review)
    for review_id in ("TC-03-04", "TC-13-02"):
        item = by_id[review_id]
        elements = _elements(item)
        assert item["coverage_status"] == "部分覆盖"
        assert elements["ip_proof_attachment_arrangement"]["status"] == "COVERED"
        assert elements["actual_ip_proof_attachment"]["status"] == "REFERENCED_BUT_NOT_PROVIDED"
        assert "C0029" in elements["actual_ip_proof_attachment"]["source_clause_ids"]
        assert "附件安排并非缺失" in item["risk_point_description"]

    materials = _elements(by_id["TC-13-01"])
    assert materials["body_development_scope"]["status"] == "COVERED"
    assert materials["body_research_work_allocation"]["status"] == "COVERED"
    assert materials["requirements_specification"]["status"] == "REFERENCED_BUT_NOT_PROVIDED"
    assert materials["detailed_module_schedule"]["status"] == "REFERENCED_BUT_NOT_PROVIDED"

    acceptance = _elements(by_id["TC-13-03"])
    assert acceptance["body_acceptance_metrics"]["status"] == "COVERED"
    assert acceptance["test_report"]["status"] == "COVERED"
    assert acceptance["test_attachment_arrangement"]["status"] == "COVERED"
    assert acceptance["actual_test_attachment_content"]["status"] == "REFERENCED_BUT_NOT_PROVIDED"


def test_confidentiality_evidence_is_not_expanded_beyond_the_quote(cooperative_review: dict):
    item = _items(cooperative_review)["TC-11-04"]
    evidence = next(row for row in item["evidence_clauses"] if row["clause_id"] == "C0015")

    assert "保密义务" in evidence["support_summary"]
    assert "第三方披露" in evidence["support_summary"]
    assert "复制" not in evidence["support_summary"]
    assert "超范围使用" not in evidence["support_summary"]


def test_unrelated_breach_remedies_do_not_activate_infringement_coordination(cooperative_review: dict):
    item = _items(cooperative_review)["TC-12-04"]

    assert item["coverage_status"] == "不适用"
    assert item["coverage_presence"] == "NOT_APPLICABLE"
    assert item["coverage_elements"] == []
    assert item["evidence_clauses"] == []
    assert item["applicability"]["machine_reason"] == "SAME_LOSS_GENERAL_REMEDY_NOT_FOUND"
    assert "迟延付款、退出" in item["applicability"]["reason"]


def test_registry_bound_risk_evidence_reuses_final_review_sources(cooperative_review: dict):
    provenance = next(
        finding for finding in cooperative_review["risk_findings"]
        if finding["risk_group"] == "technology_ip_rights_provenance"
    )
    body_rows = [
        row for row in provenance["evidence_bundle"]["items"]
        if row.get("source_type") != "SEARCH_AUDIT"
    ]

    assert {row.get("clause_id") for row in body_rows} == {"C0019", "C0020", "C0029"}
    assert all(row.get("supporting_element_ids") for row in body_rows)
    assert all(row.get("canonical_fact_keys") for row in body_rows)
    assert cooperative_review["review_evidence_element_gate"]["status"] == "REPORT_READY"
    assert cooperative_review["risk_evidence_ownership_gate"]["status"] == "REPORT_READY"
    assert cooperative_review["release_quality_gate"]["status"] == "REPORT_READY"


def test_creator_reward_is_optional_internal_governance_with_correct_legal_basis(cooperative_review: dict):
    item = _items(cooperative_review)["TC-03-03"]

    assert item["legal_basis_index"] == "《民法典》第849条"
    assert item["risk_level"] == "低"
    assert item["recommended_action"] == "OPTIONAL"
    assert _elements(item)["creator_reward_arrangement"]["status"] == "MISSING"


def test_referenced_attachments_are_material_pending_not_contract_defects(cooperative_review: dict):
    by_id = _items(cooperative_review)
    for review_id in ("TC-03-04", "TC-13-01", "TC-13-02", "TC-13-03"):
        item = by_id[review_id]
        state = item["decision_trace"]["attachment_material_state"]
        assert item["recommended_action"] == "MANUAL_REVIEW"
        assert state["status"] == "REFERENCED_BUT_NOT_PROVIDED"
        assert state["requires_contract_edit"] is False
        assert not any(
            row.get("status") == "MISSING" and row.get("risk_relevant") is not False
            for row in item["coverage_elements"]
        )


def test_retention_clause_is_adverse_not_missing_and_search_audit_stays_factual(cooperative_review: dict):
    by_id = _items(cooperative_review)
    retention = by_id["TC-11-03"]
    assert retention["risk_status"] == "ADVERSE"
    assert _elements(retention)["post_termination_material_retention"]["status"] == "ADVERSE"
    assert retention["decision_trace"]["specialized_search_audit"]["required_positive_facts"] == []

    acceptance = by_id["TC-13-03"]
    audit = acceptance["decision_trace"]["specialized_search_audit"]
    assert not any("尚未" in value or "无法核验" in value for value in audit["required_positive_facts"])


def test_article_870_is_kept_for_a_license_but_not_for_cooperative_development(cooperative_review: dict):
    cooperative = _items(cooperative_review)["TC-13-02"]
    assert "第870条" not in cooperative["legal_basis_index"]
    assert "第870条" in cooperative["decision_trace"]["legal_basis_applicability"]["removed_scene_inapplicable_articles"]

    license_path = ROOT / "test_data" / "13、技术许可.docx"
    license_review = review_contract_v2(_document_text(license_path), "technology", use_llm=False)
    license_item = _items(license_review)["TC-13-02"]
    assert "第870条" in license_item["legal_basis_index"]
    assert license_item["decision_trace"]["legal_basis_applicability"]["removed_scene_inapplicable_articles"] == []


def test_search_audit_queries_are_positive_objective_facts(cooperative_review: dict):
    audit_rows = [
        row
        for finding in cooperative_review["risk_findings"]
        for row in finding.get("evidence_bundle", {}).get("items", [])
        if row.get("source_type") == "SEARCH_AUDIT"
    ]
    targets = [str(value) for row in audit_rows for value in row.get("search_targets") or []]

    assert "非履约使用须取得数据提供方事先书面同意" in targets
    assert "未将数据提供方事先书面同意作为非履约使用前提" not in targets
    assert not any(re.match(r"^(?:未|尚未|无法|缺失|不足)", value) for value in targets)


def test_manual_review_records_are_executable_and_have_no_empty_formal_advice(cooperative_review: dict):
    records = {row["item_id"]: row for row in cooperative_review["manual_review_items"]}
    by_id = _items(cooperative_review)

    for review_id in ("TC-03-04", "TC-13-01", "TC-13-02", "TC-13-03"):
        instruction = records[review_id]["review_instruction"]
        assert "请提供" in instruction
        assert "核验" in instruction
        assert "本事项适用" not in instruction
        display = by_id[review_id]["decision_trace"]["recommendation_display_state"]
        assert display["formal_recommendations"] == []
        assert display["review_instructions"] == [instruction]

    assert "背景知识产权清单" in records["TC-03-04"]["review_instruction"]
    assert "测试设备、环境和统计口径" in records["TC-13-03"]["review_instruction"]
    assert not re.search(r"(?:【)?正式建议(?:】)?：\s*(?:\n|$)", cooperative_review["result"])


def test_unprovided_attachment_narrative_does_not_assert_content_is_missing(cooperative_review: dict):
    by_id = _items(cooperative_review)
    for review_id in ("TC-03-04", "TC-13-01", "TC-13-02", "TC-13-03"):
        judgment = by_id[review_id]["narrative_state"]["judgment"]
        assert judgment == "正文存在部分相关约定，但附件未提供，当前无法核验完整性。"
        assert "尚无法核验已有部分约定" not in judgment

    acceptance = by_id["TC-13-03"]
    assert "合同尚未约定验收争议复测或第三方检测机制" not in acceptance["risk_point_description"]
