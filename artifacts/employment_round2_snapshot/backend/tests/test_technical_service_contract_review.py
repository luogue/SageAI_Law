import json
from pathlib import Path

from docx import Document

from review_v2.engine import COVERAGE_NOT_APPLICABLE, review_contract_v2
from review_v2.technology_service_analysis import analyze_payment_schedule, analyze_timeline
from review_v2.user_profile import UserRiskProfile


ROOT = Path(__file__).resolve().parents[2]
FIXTURE = ROOT / "test_data" / "8、技术服务.docx"
GOLD = Path(__file__).parent / "fixtures" / "technical_contracts" / "technical_service_data_governance.gold.json"


def _fixture_text() -> str:
    document = Document(str(FIXTURE))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = ["\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    return "\n".join(paragraphs + tables)


def _review() -> dict:
    return review_contract_v2(
        _fixture_text(),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(
            user_identity="乙方/技术服务提供方",
            contract_scenario="高校科研数据治理技术服务",
        ),
    )


def test_payment_formula_analyzer_handles_generic_inconsistent_schedules():
    excessive = analyze_payment_schedule(
        "合同总价为200000元。首期支付30%，即60000元；第二期支付60%，即120000元；质保金15%，即30000元。"
    )
    assert excessive["percentage_total"] == 105
    assert excessive["amount_total"] == 210000
    assert excessive["amount_delta"] == 10000
    assert excessive["status"] == "INCONSISTENT"

    insufficient = analyze_payment_schedule(
        "合同总价为100000元。首期支付20%，即20000元；验收后支付70%，即70000元。"
    )
    assert insufficient["percentage_total"] == 90
    assert insufficient["amount_delta"] == -10000

    wrong_amount = analyze_payment_schedule(
        "合同总价为100000元。首期支付50%，即50000元；尾款支付50%，即45000元。"
    )
    assert wrong_amount["percentage_total"] == 100
    assert wrong_amount["amount_total"] == 95000
    assert wrong_amount["node_mismatches"]

    included_retention = analyze_payment_schedule(
        "合同总价为100000元。支付90%，即90000元；尾款包含10%质保金，即10000元。"
    )
    assert included_retention["percentage_total"] == 100
    assert included_retention["status"] == "CONSISTENT"


def test_timeline_analyzer_detects_pre_effective_performance_generically():
    timeline = analyze_timeline(
        "服务期限为2027年3月1日至2027年6月30日。委托方应在2027年3月5日前提供资料。合同于2027年3月15日签署并生效。"
    )
    assert timeline["effective_date"] == "2027-03-15"
    assert timeline["service_start_date"] == "2027-03-01"
    assert timeline["data_delivery_deadline"] == "2027-03-05"
    assert {row["type"] for row in timeline["conflicts"]} == {
        "SERVICE_BEFORE_EFFECTIVE",
        "DATA_DUE_BEFORE_EFFECTIVE",
    }


def test_fixture_classification_context_and_applicability():
    output = _review()
    context = output["contract_context"]
    assert context["contract_family"] == "TECHNOLOGY"
    assert context["contract_subtype"] == "TECHNICAL_SERVICE"
    assert context["transaction_structure"] == "B2B_TECHNICAL_SERVICE"
    assert context["user_role"] == "service_provider"
    assert output["report_context"]["contract_type_label"] == "技术服务合同"
    assert output["report_context"]["scenario_label"] == "高校科研数据治理技术服务"

    by_id = {item["item_id"]: item for item in output["items"]}
    for item_id in ("TC-03-01", "TC-03-02", "TC-03-03", "TC-05-01", "TC-05-02", "TC-05-03", "TC-05-04", "TC-12-01"):
        assert by_id[item_id]["coverage_status"] == COVERAGE_NOT_APPLICABLE
    assert by_id["TC-13-03"]["coverage_status"] == "通过"


def test_fixture_separates_coverage_from_adverse_risk_and_finds_gold_risks():
    output = _review()
    by_id = {item["item_id"]: item for item in output["items"]}
    assert by_id["TC-09-02"]["coverage_presence"] == "PRESENT"
    assert by_id["TC-09-02"]["risk_status"] == "ADVERSE"
    assert by_id["TC-09-02"]["recommended_action"] == "MUST_FIX"
    assert by_id["TC-12-04"]["coverage_presence"] == "PRESENT"
    assert by_id["TC-12-04"]["risk_status"] == "ADVERSE"

    groups = {finding["risk_group"] for finding in output["risk_findings"]}
    assert {
        "technical_acceptance_payment_control",
        "technical_background_ip_overtransfer",
        "technical_uncapped_cumulative_liability",
        "technical_unilateral_termination_change",
        "personal_information_processing_gap",
        "payment_formula_inconsistency",
    } <= groups
    assert "technical_timeline_conflict" in groups


def test_fixture_formula_timeline_evidence_and_report_quality():
    output = _review()
    formula = output["deterministic_analysis"]["payment_formula"]
    assert formula["percentage_total"] == 105
    assert formula["amount_total"] == 504000
    assert formula["contract_total"] == 480000
    assert formula["amount_delta"] == 24000

    conflict_types = {row["type"] for row in output["deterministic_analysis"]["timeline"]["conflicts"]}
    assert {"SERVICE_BEFORE_EFFECTIVE", "DATA_DUE_BEFORE_EFFECTIVE"} <= conflict_types

    report = output["result"]
    assert "附件二" in report
    assert "出卖人" not in report
    assert "买受人" not in report
    assert "可追溯到具体缺失要素" not in report
    assert "未定位到单一明确条款" not in report
    assert output["release_quality_gate"]["status"] == "REPORT_READY"


def test_fixture_matches_frozen_gold_semantics():
    gold = json.loads(GOLD.read_text(encoding="utf-8"))
    output = _review()
    context = output["contract_context"]
    assert context["contract_family"] == gold["classification"]["contract_family"]
    assert context["contract_subtype"] == gold["classification"]["contract_subtype"]
    assert context["transaction_structure"] == gold["classification"]["transaction_structure"]
    assert context["user_role"] == gold["classification"]["review_role"]

    groups = {finding["risk_group"] for finding in output["risk_findings"]}
    assert set(gold["required_risk_groups"]) <= groups
    assert set(gold["required_secondary_risk_groups"]) <= groups
    by_id = {item["item_id"]: item for item in output["items"]}
    assert all(by_id[item_id]["coverage_status"] == COVERAGE_NOT_APPLICABLE for item_id in gold["not_applicable_review_ids"])
    assert all(by_id[item_id]["coverage_status"] == "通过" for item_id in gold["required_attachment_review_ids"])
    assert all(token not in output["result"] for token in gold["forbidden_report_tokens"])

    finding_summary = output["risk_finding_summary"]
    assert len(output["risk_findings"]) == sum(finding_summary[level] for level in ("重大风险", "中等风险", "低风险"))
    disposition_summary = output["final_disposition_summary"]
    assert disposition_summary["total"] == len(output["final_dispositions"]) == len(output["risk_findings"])
    assert output["overall_assessment"]["core_risk_count"] == finding_summary["重大风险"]
    assert output["overall_assessment"]["priority_modify_count"] == disposition_summary["actionable"]
    assert output["overall_assessment"]["optional_count"] == disposition_summary["optional_improvement"]
