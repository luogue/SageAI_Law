from pathlib import Path
from types import SimpleNamespace

from docx import Document

import review_v2.engine as engine
from review_v2.canonical_claims import CanonicalClaim
from review_v2.recommendation_impact import evaluate_recommendations


ROOT = Path(__file__).resolve().parents[2]


def _shared_claim(hierarchy):
    return CanonicalClaim(
        claim_id="document.precedence", subject="contract_documents", predicate="precedence",
        object={"hierarchy": hierarchy}, qualifier="contract_clause", polarity="POSITIVE",
        status="ASSERTED", fact_ids=["commercial.document_precedence"], evidence_ids=["CCE-C1"],
        source_clause_ids=["C1"], source_spans=["文件效力顺序"], source_locations=["第十三条"],
        extraction_method="shared_claim_rule", normalization="semantic_elements", confidence="high",
        consumer_review_ids=["A", "B"],
    )


def _review(review_id, semantic_key, status):
    return SimpleNamespace(
        item_id=review_id,
        coverage_elements=[{"claim_id": "document.precedence", "element_id": semantic_key, "status": status}],
        decision_trace={},
    )


def _hb4_text():
    document = Document(next((ROOT / "test_data").glob("7*.docx")))
    rows = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        rows.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
    return "\n".join(row for row in rows if row.strip(" |"))


def test_same_claim_same_semantic_element_contradiction_fails_closed():
    result = engine._cross_review_shared_claim_state(
        [_review("A", "attachment_effect", "COVERED"), _review("B", "attachment_effect", "MISSING")],
        {"document.precedence": _shared_claim({"attachment_effect": True})},
    )
    assert result["gate_results"]["CROSS_REVIEW_SHARED_CLAIM_GATE"] == "FAIL"
    assert result["contradiction_count"] == 1


def test_same_claim_different_semantic_elements_may_have_different_coverage():
    result = engine._cross_review_shared_claim_state(
        [_review("A", "attachment_effect", "COVERED"), _review("B", "order_priority", "MISSING")],
        {"document.precedence": _shared_claim({"attachment_effect": True, "order_priority": False})},
    )
    assert result["gate_results"]["CROSS_REVIEW_SHARED_CLAIM_GATE"] == "PASS"
    assert result["contradiction_count"] == 0


def test_hb4_reuses_document_precedence_and_propagates_default_law():
    result = engine.review_contract_v2(_hb4_text(), "house_lease", use_llm=False)
    items = {item["item_id"]: item for item in result["items"]}
    shared = next(row for row in result["observability"]["shared_claim_consumer_debug"] if row["claim_id"] == "document.precedence")
    assert result["release_quality_gate"]["status"] == "REPORT_READY"
    assert shared["consumer_reviews"] == ["S-12-04", "FS-14"]
    assert shared["contradictions"] == []
    assert all(element["claim_id"] == "document.precedence" for review_id in ("S-12-04", "FS-14") for element in items[review_id]["coverage_elements"])
    assert result["default_law_provenance_gate"]["gate_results"]["DEFAULT_LAW_PROVENANCE_GATE"] == "PASS"
    assert result["default_law_provenance_gate"]["break_count"] == 0
    assert "第605条" in result["result"]


def test_uncertain_legal_freshness_cannot_produce_clear_default(monkeypatch):
    monkeypatch.setattr(engine, "validate_legal_citation", lambda *args, **kwargs: {"status": "UNCERTAIN", "support_type": "UNCERTAIN"})
    item = SimpleNamespace(
        coverage_status="部分覆盖", risk_group="sale_risk_buyer_delay_acceptance",
        node_id="sale_risk_buyer_delay_acceptance",
    )
    assessment = engine._default_law_assessment(item, "sale", "买受人无正当理由拒收并迟延受领。")
    assert assessment["status"] == "UNCERTAIN_DEFAULT"
    assert assessment["applicability_confidence"] == "low"


def test_recommendation_impact_blocks_remedy_narrowing_but_allows_exact_statutory_mirror():
    item = SimpleNamespace(
        item_id="R1", node_id="sale_expectation_loss",
        coverage_elements=[{"status": "COVERED", "proposition": "买方可以依法主张可得利益损失。"}],
        decision_trace={"effective_gap": {"statutory_protection_scope": "法律保护依法可主张的可得利益。"}},
    )
    finding = SimpleNamespace(finding_id="F1", source_review_ids=["R1"], risk_group="sale_expectation_loss")
    narrowing = SimpleNamespace(
        recommendation_id="REC-1", target_finding_ids=["F1"],
        proposed_change="将可得利益赔偿封顶为合同价款的10%。", suggested_text="",
    )
    mirror = SimpleNamespace(
        recommendation_id="REC-2", target_finding_ids=["F1"],
        proposed_change="按照民法典第584条现行法律规则明确可预见性标准，不缩窄买方依法享有的赔偿范围。", suggested_text="",
    )
    accepted, impacts, gate = evaluate_recommendations([narrowing, mirror], [finding], [item], "BUYER")
    assert [row.recommendation_id for row in accepted] == ["REC-2"]
    assert impacts[0].narrows_remedy is True and impacts[0].gate_result == "BLOCKED"
    assert impacts[1].weakens_protection is False and impacts[1].gate_result == "PASS"
    assert gate["gate_results"]["RECOMMENDATION_NON_DEGRADATION_GATE"] == "PASS"
