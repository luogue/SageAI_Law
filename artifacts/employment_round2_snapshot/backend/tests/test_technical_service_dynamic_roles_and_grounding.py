from pathlib import Path

from docx import Document

from review_v2.engine import (
    COVERAGE_NOT_APPLICABLE,
    ReviewItemResult,
    _absence_proposition_evidence,
    _apply_factual_grounding_gate,
    review_contract_v2,
)
from review_v2.technology_service_analysis import (
    analyze_payment_schedule,
    analyze_personal_information_processing,
    analyze_technical_service_core_facts,
    analyze_technology_relationship_risks,
    analyze_timeline,
)
from review_v2.user_profile import UserRiskProfile


ROOT = Path(__file__).resolve().parents[2]
FIXTURE = ROOT / "test_data" / "23 技术服务.docx"
SECOND_FIXTURE = ROOT / "test_data" / "24 技术服务.docx"


def _fixture_text() -> str:
    document = Document(str(FIXTURE))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = ["\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    return "\n".join(paragraphs + tables)


def _review() -> dict:
    return review_contract_v2(
        _fixture_text(),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方", contract_scenario="软件技术服务"),
    )


def _second_fixture_text() -> str:
    document = Document(str(SECOND_FIXTURE))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = ["\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    return "\n".join(paragraphs + tables)


def _second_review() -> dict:
    return review_contract_v2(
        _second_fixture_text(),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方", contract_scenario="ERP性能诊断技术服务"),
    )


def test_contract_defined_service_roles_override_party_letter_defaults():
    output = _review()
    context = output["contract_context"]
    perspective = context["perspective_resolution"]

    assert context["contract_subtype"] == "TECHNICAL_SERVICE"
    assert context["transaction_role"] == "SERVICE_PROVIDER"
    assert context["user_role"] == "service_provider"
    assert perspective["party_side"] == "PARTY_A"
    assert perspective["legal_role"] == "SERVICE_PROVIDER"
    assert perspective["contract_display_role"] == "服务提供方"
    assert perspective["display_role"] == "甲方 / 服务提供方"
    assert "审查立场：甲方 / 服务提供方" in output["result"]
    assert "甲方 / 委托方" not in output["result"]


def test_service_acceptance_and_new_result_core_facts_are_not_keyword_passes():
    text = _fixture_text()
    facts = analyze_technical_service_core_facts(text)
    assert facts["acceptance"]["mechanism_present"] is True
    assert facts["acceptance"]["standard_deferred"] is True
    assert facts["acceptance"]["deadline_defined"] is True
    assert facts["acceptance"]["written_objection"] is True
    assert facts["acceptance"]["deemed_acceptance"] is True
    assert facts["acceptance"]["remedy_present"] is True
    assert facts["acceptance"]["remedy_deadline_specific"] is False
    assert facts["service_generated_result"]["scope_defined"] is True
    assert facts["service_generated_result"]["ownership_deferred"] is True
    assert facts["service_generated_result"]["ownership_finalized"] is False

    output = _review()
    by_id = {item["item_id"]: item for item in output["items"]}
    acceptance = by_id["TC-09-02"]
    ownership = by_id["TC-09-03"]

    assert acceptance["coverage_status"] == "部分覆盖"
    acceptance_elements = {row["element_id"]: row["status"] for row in acceptance["coverage_elements"]}
    assert acceptance_elements["acceptance_process"] == "COVERED"
    assert acceptance_elements["acceptance_deadline"] == "COVERED"
    assert acceptance_elements["written_objection_mechanism"] == "COVERED"
    assert acceptance_elements["deemed_acceptance_arrangement"] == "COVERED"
    assert acceptance_elements["acceptance_standard_finality"] == "PARTIALLY_COVERED"
    assert acceptance_elements["rectification_deadline"] == "PARTIALLY_COVERED"

    assert ownership["coverage_status"] == "部分覆盖"
    ownership_elements = {row["element_id"]: row["status"] for row in ownership["coverage_elements"]}
    assert ownership_elements["service_result_scope"] == "COVERED"
    assert ownership_elements["service_result_ownership_finality"] == "PARTIALLY_COVERED"
    assert ownership_elements["interim_result_use"] == "COVERED"
    assert "另行协商" in ownership["risk_point_description"]
    assert "已覆盖，无需修改" not in ownership["action_suggestion"]


def test_service_ip_proof_and_personal_information_rules_are_fact_triggered():
    output = _review()
    by_id = {item["item_id"]: item for item in output["items"]}

    assert by_id["TC-03-04"]["coverage_status"] == COVERAGE_NOT_APPLICABLE
    assert by_id["TC-13-02"]["coverage_status"] == COVERAGE_NOT_APPLICABLE
    assert not any(row["risk_group"] == "technology_ip_rights_provenance" for row in output["risk_findings"])
    assert not any(row["risk_group"] == "personal_information_processing_gap" for row in output["risk_findings"])

    generic = analyze_technology_relationship_risks(
        "服务方可访问业务数据、系统数据、接口数据及系统账号；项目结束后删除或返还。"
    )
    assert generic["personal_information_processing_gap"]["present"] is False

    explicit = analyze_technology_relationship_risks(
        "服务方受托处理员工姓名、手机号和身份证号，并在项目结束后删除个人信息。"
    )
    assert explicit["personal_information_processing_gap"]["present"] is True
    assert {"姓名", "手机号", "身份证"} <= set(
        explicit["personal_information_processing_gap"]["analysis"]["matched_categories"]
    )


def test_final_report_contains_no_ungrounded_service_facts():
    output = _review()
    report = output["result"]
    forbidden = (
        "当前技术服务仅以SLA",
        "月度服务记录",
        "最终书面验收确认履行质量",
        "姓名、工号、电话、邮箱及系统日志",
        "专利证书、专利申请文件",
    )
    assert all(token not in report for token in forbidden)

    test_item = next(item for item in output["items"] if item["item_id"] == "TC-13-03")
    assert test_item["coverage_status"] == COVERAGE_NOT_APPLICABLE
    reason = str((test_item.get("applicability") or {}).get("reason") or "")
    assert "SLA" not in reason
    assert "响应时间" not in reason
    assert "月度服务记录" not in reason


def test_factual_grounding_gate_blocks_an_unbound_concrete_contract_fact():
    item = ReviewItemResult(
        item_id="TC-TEST",
        node_id="test",
        title="测试审查项",
        risk_level="中",
        status="待修改",
        risk_point_description="合同已经约定SLA和月度服务记录。",
        suggested_clause_example="建议沿用现有SLA。",
        legal_basis_index="",
        legal_consequence_if_missing="现有SLA可能引发争议。",
        applicable_if="",
        matched_terms=[],
        coverage_status="部分覆盖",
        issue_type="PARTIAL",
        evidence_clauses=[],
        evidence_location="",
        reasoning_summary="合同存在SLA。",
    )
    item.decision_trace["factual_grounding_claims"] = [{
        "claim_id": "UNGROUNDED-SLA",
        "required_term_groups": [["SLA"], ["月度服务记录"]],
        "evidence_clause_ids": [],
        "fields": [
            "risk_point_description", "reasoning_summary", "legal_consequence_if_missing",
            "suggested_clause_example",
        ],
    }]

    checked = _apply_factual_grounding_gate([item], "合同仅约定一般技术支持。")
    assert checked[0].status == "UNKNOWN"
    assert checked[0].issue_type == "MANUAL_REVIEW"
    assert checked[0].needs_manual_review is True
    assert checked[0].decision_trace["factual_grounding_gate"]["status"] == "NEEDS_REVIEW"
    assert "SLA" not in checked[0].risk_point_description


def test_provider_side_secondary_facts_are_recorded_without_forcing_risks():
    facts = analyze_technical_service_core_facts(_fixture_text())
    assert facts["liability"]["third_party_platform_exception"] is True
    assert facts["liability"]["indirect_loss_exclusion"] is True
    assert facts["liability"]["provider_liability_cap"] is True
    assert facts["change_control"]["principle_only"] is True
    assert facts["termination_settlement"]["actual_workload_basis"] is True
    assert facts["termination_settlement"]["workload_determination_unclear"] is True

    output = _review()
    assert output["deterministic_analysis"]["technical_service_core_facts"] == facts


PERSPECTIVE_CASE_TEXT = """
软件技术服务合同
甲方（服务提供方）为乙方（服务接受方）提供软件维护、故障排查和技术支持服务。
服务期限一年，服务费人民币十万元。服务成果按双方确认的标准书面验收；不合格的，甲方负责整改。
甲方保证其提供的程序代码不侵犯第三方知识产权。因甲方提供的程序代码侵犯第三方知识产权并导致乙方遭受损失的，由甲方承担相应责任。
因乙方提供的软件、数据、图片、文档或商标导致知识产权纠纷的，由乙方自行承担责任。
双方对履约中知悉的商业秘密承担保密义务。
"""


def _perspective_review(party: str) -> dict:
    return review_contract_v2(
        PERSPECTIVE_CASE_TEXT,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity=party, contract_scenario="软件技术服务"),
    )


def test_third_party_claim_direction_changes_between_provider_and_recipient():
    provider = _perspective_review("甲方")
    recipient = _perspective_review("乙方")
    provider_item = next(item for item in provider["items"] if item["item_id"] == "TC-12-02")
    recipient_item = next(item for item in recipient["items"] if item["item_id"] == "TC-12-02")
    provider_finding = next(
        row for row in provider["risk_findings"] if "TC-12-02" in row["source_review_ids"]
    )
    recipient_finding = next(
        row for row in recipient["risk_findings"] if "TC-12-02" in row["source_review_ids"]
    )
    provider_disposition = next(
        row for row in provider["final_dispositions"] if "TC-12-02" in row["review_ids"]
    )
    recipient_disposition = next(
        row for row in recipient["final_dispositions"] if "TC-12-02" in row["review_ids"]
    )

    # Coverage is a contract fact and remains the same; risk and drafting
    # direction change with the reviewed party.
    assert provider_item["coverage_status"] == recipient_item["coverage_status"] == "部分覆盖"
    assert provider_item["risk_level"] == "低"
    assert provider_finding["severity"] == "LOW"
    assert provider_finding["priority"] == "C_OPTIMIZATION"
    assert provider_item["decision_trace"]["risk_to_current_party"]["risk_scope"] == "PROCEDURAL_CONTROL_ONLY"
    assert provider_item["decision_trace"]["recommendation_direction_check"]["expands_current_party_liability"] is False
    assert "未经服务提供方书面同意" in provider_finding["recommendation"]
    assert "退款或赔偿救济" not in provider_finding["recommendation"]
    assert provider_disposition["client_effect"] == "FAVORABLE"
    assert provider_disposition["final_action"] == "OPTIONAL_IMPROVEMENT"
    assert provider_disposition["blocked_recommendation_ids"] == []
    assert "未经服务提供方书面同意" in provider_disposition["recommendation"]

    assert recipient_item["risk_level"] == "中"
    assert recipient_finding["severity"] == "MEDIUM"
    assert recipient_finding["priority"] == "B_PRACTICAL"
    assert recipient_item["decision_trace"]["risk_to_current_party"]["risk_scope"] == "BASE_PROTECTION_AND_REMEDY"
    assert "解除权及退款或价款返还" in recipient_finding["recommendation"]
    assert "替代救济" not in recipient_finding["recommendation"]
    assert "责任承担" not in recipient_finding["recommendation"]
    assert recipient_disposition["client_effect"] == "FAVORABLE"
    assert recipient_disposition["final_action"] == "SHOULD_CHANGE"


def test_service_result_consequence_survives_finding_aggregation():
    output = _review()
    finding = next(
        row for row in output["risk_findings"] if "TC-09-03" in row["source_review_ids"]
    )
    assert "成果权属" in finding["consequence"]
    assert "交付使用范围" in finding["consequence"]
    assert "停用、替换成本" not in finding["consequence"]
    assert "项目定制成果的知识产权归属" in finding["recommendation"]


def test_second_service_payment_keeps_amount_only_tail_node_in_arithmetic():
    payment = analyze_payment_schedule(_second_fixture_text())

    assert payment["status"] == "INCONSISTENT"
    assert payment["contract_total"] == 500_000
    assert [node["amount"] for node in payment["nodes"]] == [200_000, 150_000, 125_000, 75_000]
    assert [node["percentage"] for node in payment["nodes"]] == [40, 30, 25, None]
    assert [node["effective_percentage"] for node in payment["nodes"]] == [40, 30, 25, 15]
    assert payment["amount_total"] == 550_000
    assert payment["percentage_total"] == 110
    assert payment["amount_delta"] == 50_000
    assert payment["nodes"][-1]["percentage_inferred"] is True
    assert payment["nodes"][-1]["payment_type"] == "final"


def test_payment_nodes_allow_amount_percentage_and_dynamic_remainder_independently():
    amount_only = analyze_payment_schedule("合同总价为100000元。项目验收合格后支付尾款10000元。")
    percentage_only = analyze_payment_schedule("合同总价为100000元。第一期验收后支付30%。")
    remainder = analyze_payment_schedule(
        "合同总价为100000元。第一期验收后支付30%，即30000元；最终验收后支付剩余款项。"
    )

    assert amount_only["nodes"][0]["amount"] == 10_000
    assert amount_only["nodes"][0]["percentage"] is None
    assert amount_only["nodes"][0]["effective_percentage"] == 10
    assert percentage_only["nodes"][0]["percentage"] == 30
    assert percentage_only["nodes"][0]["amount"] is None
    assert percentage_only["nodes"][0]["effective_amount"] == 30_000
    assert remainder["nodes"][-1]["dynamic_expression"] == "REMAINDER"
    assert remainder["nodes"][-1]["effective_amount"] == 70_000
    assert remainder["percentage_total"] == 100
    assert remainder["status"] == "CONSISTENT"


def test_second_service_existing_result_ownership_is_not_reported_missing():
    facts = analyze_technical_service_core_facts(_second_fixture_text())["service_generated_result"]
    assert facts["scope_defined"] is True
    assert facts["ownership_finalized"] is True
    assert facts["ownership_deferred"] is False
    assert facts["ownership_party"] == "乙方"
    assert facts["recipient_internal_use_only"] is True
    assert facts["provider_reuse_allowed"] is True

    output = _second_review()
    item = next(row for row in output["items"] if row["item_id"] == "TC-09-03")
    assert item["coverage_status"] != "缺失"
    assert item["coverage_presence"] == "PRESENT"
    assert "缺失" not in item["risk_point_description"]
    assert any(row["risk_group"] == "service_new_result_ownership" for row in output["risk_findings"])


def test_second_service_timeline_is_deterministic_and_binds_only_direct_dates():
    timeline = analyze_timeline(_second_fixture_text())
    assert timeline["service_start_date"] == "2026-10-10"
    assert timeline["contract_sign_date"] == "2026-10-20"
    assert timeline["effective_date"] is None
    assert timeline["comparison_basis"] == "SIGNING_DATE"
    assert timeline["conflict_types"] == ["SERVICE_BEFORE_SIGNING"]

    output = _second_review()
    item = next(row for row in output["items"] if row["item_id"] == "TECH-SHARED-03")
    finding = next(row for row in output["risk_findings"] if row["risk_group"] == "technical_timeline_conflict")
    quotes = " ".join(str(row.get("quoted_text") or "") for row in finding["evidence_bundle"]["items"])
    assert item["needs_manual_review"] is False
    assert finding["evidence_bundle"]["entailment_status"] == "PASSED"
    assert item["title"] == "履行开始日期早于合同签订日期"
    assert "合同生效" not in item["risk_point_description"]
    assert any(
        "contract_signing_date" in (row.get("fact_types") or [row.get("fact_type")])
        for row in finding["evidence_bundle"]["items"]
    )
    assert "2026年10月10日" in quotes
    assert "2026年10月20日" in quotes
    assert "甲方资料提供期限" not in item["risk_point_description"]


def test_second_service_personal_information_uses_strong_processing_facts_and_direct_evidence():
    analysis = analyze_technology_relationship_risks(_second_fixture_text())["personal_information_processing_gap"]["analysis"]
    assert analysis["trigger_strength"] == "STRONG"
    assert all(analysis["strong_facts"][key] for key in (
        "production_database_access",
        "personal_data_access",
        "data_transfer_to_provider",
        "data_copy_to_provider_server",
        "data_storage_by_provider",
        "data_analysis_by_provider",
        "real_data_processing",
    ))

    output = _second_review()
    item = next(row for row in output["items"] if row["item_id"] == "TECH-SHARED-02")
    finding = next(row for row in output["risk_findings"] if row["risk_group"] == "personal_information_processing_gap")
    fact_types = {
        fact_type
        for row in finding["evidence_bundle"]["items"]
        for fact_type in row.get("fact_types") or [row.get("fact_type")]
    }
    assert item["needs_manual_review"] is False
    assert finding["evidence_bundle"]["entailment_status"] == "PASSED"
    assert {
        "production_database_access", "data_transfer_to_provider",
        "data_copy_to_provider_server", "data_analysis_by_provider",
    } <= fact_types
    assert "信息种类" not in finding["evidence_bundle"]["missing_element_types"]
    absence_text = " ".join(
        str(row.get("quoted_text") or "")
        for row in finding["evidence_bundle"]["items"]
        if row.get("evidence_role") == "ABSENCE_EVIDENCE"
    )
    assert "未发现完整约定：信息种类" not in absence_text


def test_second_service_scene_liability_limit_and_acceptance_disconnect_are_grounded():
    output = _second_review()
    assert output["report_context"]["scenario_label"] == "ERP系统性能诊断与数据库优化技术服务"
    assert output["contract_context"]["usage_scenario"] == "ENTERPRISE_INFORMATION_SYSTEM_PERFORMANCE_OPTIMIZATION_SERVICE"

    by_group = {row["risk_group"]: row for row in output["risk_findings"]}
    liability = by_group["service_provider_liability_cap_exclusion"]
    acceptance = by_group["service_performance_acceptance_disconnect"]
    assert liability["severity"] == "MEDIUM"
    assert liability["title"] == "服务方赔偿责任上限及免责范围过宽"
    assert all(term in liability["issue"] for term in ("已支付技术服务费", "业务中断损失", "第三方索赔"))
    assert liability["economic_impact"]["level"] == "medium"
    assert acceptance["severity"] == "MEDIUM"
    assert acceptance["title"] == "量化服务目标与验收标准脱节"
    assert set(acceptance["source_review_ids"]) == {"TC-09-02", "TECH-SHARED-09"}
    # The generic liquidated-damages validity screen now correctly adds the
    # contract's daily 0.5% clause as a separate, evidence-backed risk.
    assert len(output["risk_findings"]) == 10
    assert output["release_quality_gate"]["status"] == "REPORT_READY"

    quality_acceptance = next(row for row in output["items"] if row["item_id"] == "TC-09-02")
    quality_elements = {row["element_id"]: row["status"] for row in quality_acceptance["coverage_elements"]}
    assert quality_acceptance["coverage_status"] == "部分覆盖"
    assert quality_acceptance["action_suggestion"] != "已覆盖，无需修改。"
    assert quality_elements == {
        "basic_acceptance_framework": "COVERED",
        "performance_target_acceptance_link": "PARTIALLY_COVERED",
    }
    assert quality_acceptance["decision_trace"]["specialized_risk_consistency"]["test_method_details_excluded"] is True

    personal = by_group["personal_information_processing_gap"]
    assert personal["title"] == "个人信息委托处理条款不完整"
    assert "个人信息委托处理条款缺失" not in output["result"]
    assert "个人信息委托处理机制不足" not in output["result"]

    attachment = next(row for row in output["items"] if row["item_id"] == "TC-13-03")
    assert attachment["risk_level"] == "低"
    assert "测试方法、次数和统计口径" in attachment["risk_point_description"]
    assert "争议复测机制" in attachment["risk_point_description"]
    assert "目标与验收条件" not in attachment["risk_point_description"]


def test_service_provider_does_not_receive_counterparty_liability_cap_risk():
    contract = """
    技术服务合同
    甲方（服务提供方）：某技术公司
    乙方（服务接受方）：某客户公司
    甲方可访问乙方生产数据库并进行数据库参数调整和SQL优化。
    甲方承担的全部赔偿责任最高不超过乙方已经支付的技术服务费用。
    对于乙方利润损失、业务中断损失、第三方索赔及其他间接损失，甲方不承担责任。
    """
    output = review_contract_v2(
        contract,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方", contract_scenario="技术服务"),
    )
    item = next(row for row in output["items"] if row["item_id"] == "TECH-SHARED-08")
    assert output["perspective_resolution"]["canonical_role"] == "SERVICE_PROVIDER"
    assert item["coverage_status"] == "通过"
    assert item["decision_trace"]["risk_to_current_party"]["result"] == "FAVORABLE_SUPPRESSED"
    assert not any(
        row["risk_group"] == "service_provider_liability_cap_exclusion"
        for row in output["risk_findings"]
    )


def test_signing_date_only_becomes_effective_date_with_an_explicit_condition():
    timeline = analyze_timeline(
        "技术服务合同。本合同履行期限自2026年10月10日起。"
        "本合同自双方签字盖章之日起生效。签订日期：2026年10月20日。"
    )
    assert timeline["contract_sign_date"] == "2026-10-20"
    assert timeline["effective_date"] == "2026-10-20"
    assert timeline["effective_date_basis"] == "SIGNATURE_DATE_PLUS_EFFECTIVE_CONDITION"
    assert timeline["comparison_basis"] == "EFFECTIVE_DATE"


def test_timeline_labels_do_not_capture_neighboring_period_end_dates():
    timeline = analyze_timeline(
        "技术服务合同。签订日期：2026年8月18日。"
        "生效期限：2026年9月1日至2031年8月31日。"
        "服务期限：2026年9月1日至2027年6月30日。"
    )
    assert timeline["contract_sign_date"] == "2026-08-18"
    assert timeline["effective_date"] == "2026-09-01"
    assert timeline["service_start_date"] == "2026-09-01"
    assert timeline["conflicts"] == []


def test_structured_personal_categories_are_partial_scope_not_absent_scope():
    analysis = analyze_personal_information_processing(_second_fixture_text())
    assert {"员工账号信息", "订单数据"} <= set(analysis["personal_data_categories"])
    assert analysis["control_coverage"]["scope"] == "PARTIALLY_COVERED"
    assert "scope" not in analysis["missing_controls"]


def test_second_service_attachment_acceptance_and_confidentiality_are_scenario_specific():
    output = _second_review()
    by_id = {row["item_id"]: row for row in output["items"]}

    assert output["release_quality_gate"]["status"] == "REPORT_READY"

    material = by_id["TC-13-01"]
    assert material["coverage_status"] == "部分覆盖"
    assert "诊断报告" in material["suggested_clause_example"]
    assert "优化建议书" in material["suggested_clause_example"]
    assert "测试结果记录" in material["suggested_clause_example"]
    assert all(token not in material["suggested_clause_example"] for token in ("源代码", "图纸", "样品", "培训材料"))

    acceptance = by_id["TC-13-03"]
    acceptance_elements = {row["element_id"]: row["status"] for row in acceptance["coverage_elements"]}
    assert acceptance["coverage_status"] == "部分覆盖"
    assert acceptance_elements["quantitative_performance_targets"] == "COVERED"
    assert acceptance_elements["test_method"] == "MISSING"
    assert acceptance_elements["acceptance_dispute_retest"] == "MISSING"
    assert "量化指标完全缺失" not in acceptance["risk_point_description"]

    confidentiality = by_id["TC-11-02"]
    assert confidentiality["coverage_status"] == "通过"
    assert all(row["status"] == "COVERED" for row in confidentiality["coverage_elements"])
    assert confidentiality["action_suggestion"] == "已覆盖，无需修改。"

    assert by_id["TECH-SHARED-02"]["needs_manual_review"] is False
    assert by_id["TECH-SHARED-03"]["needs_manual_review"] is False


def test_service_search_audit_and_material_attachment_follow_current_subtype_facts():
    output = _review()
    items = {item["item_id"]: item for item in output["items"]}
    liability_projection = str({
        "item": items["TC-12-04"],
        "findings": [
            row for row in output["risk_findings"] if "TC-12-04" in row["source_review_ids"]
        ],
    })
    assert "被许可方" not in liability_projection
    assert "超范围实施" not in liability_projection

    attachment = items["TC-13-01"]
    assert attachment["coverage_status"] == "部分覆盖"
    attachment_elements = {row["element_id"]: row["status"] for row in attachment["coverage_elements"]}
    assert attachment_elements["attachment_formation_mechanism"] == "COVERED"
    assert attachment_elements["actual_attachment_schedule"] == "MISSING"
    assert "未显示附件级技术资料清单已经实际形成" in attachment["risk_point_description"]


def test_search_audit_preserves_scalar_scope_and_searches_positive_facts():
    evidence, audit = _absence_proposition_evidence(
        risk_group="test_positive_search_audit",
        searched_scope="合同正文及全部附件",
        required_elements=[
            "量化服务目标与最终验收合格条件之间的衔接不完整",
            "未达到量化指标时的整改义务",
        ],
    )

    assert audit["searched_scope"] == ["合同正文及全部附件"]
    assert audit["query_semantics"] == "POSITIVE_REQUIRED_FACT"
    assert audit["required_elements"] == [
        "量化服务目标与最终验收合格条件之间的衔接",
        "未达到量化指标时的整改义务",
    ]
    assert "已检索合同正文及全部附件" in evidence["quoted_text"]
    assert "合、同、正、文" not in evidence["quoted_text"]
    assert "不完整" not in " ".join(evidence["search_targets"])


def test_second_service_search_audits_classification_and_screening_are_semantically_consistent():
    output = _second_review()
    by_id = {row["item_id"]: row for row in output["items"]}
    by_group = {row["risk_group"]: row for row in output["risk_findings"]}

    assert len(output["risk_findings"]) == 10
    assert output["release_quality_gate"]["status"] == "REPORT_READY"
    assert {
        row["risk_group"] for row in output["risk_findings"]
    } == {
        "payment_formula_inconsistency",
        "personal_information_processing_gap",
        "service_provider_liability_cap_exclusion",
        "service_new_result_ownership",
        "technical_timeline_conflict",
        "service_performance_acceptance_disconnect",
        "confidential_material_return_destroy",
        "technical_material_attachment",
        "acceptance_test_attachment",
        "generic_potential_excessive_liquidated_damages",
    }

    audit_texts = []
    audit_targets = {}
    for finding in output["risk_findings"]:
        for evidence in finding["evidence_bundle"]["items"]:
            if evidence.get("source_type") != "SEARCH_AUDIT":
                continue
            audit_texts.append(str(evidence.get("quoted_text") or ""))
            audit_targets[finding["risk_group"]] = evidence.get("search_targets") or []
    rendered_audits = "\n".join(audit_texts)
    assert "合、同、正、文" not in rendered_audits
    assert "量化服务目标与最终验收合格条件之间的衔接不完整" not in rendered_audits
    assert "实际交付成果尚可通过清单进一步固定版本和验收基线" not in rendered_audits
    assert "量化性能指标被明确列为最终验收合格的必要条件" in audit_targets[
        "service_performance_acceptance_disconnect"
    ]
    material_targets = audit_targets["technical_material_attachment"]
    assert any("诊断报告" in target and "优化建议书" in target for target in material_targets)
    assert any("测试结果记录" in target and "记录格式" in target for target in material_targets)
    assert all("已有部分约定" not in target for target in material_targets)

    for review_id in ("TC-01-01", "TC-01-02", "TC-01-03"):
        evidence = by_id[review_id]["evidence_clauses"]
        assert evidence[0]["location"] == "一、服务内容、方式和要求"
        assert all(term in evidence[0]["text"] for term in ("性能诊断", "数据库分析", "系统优化"))

    report = output["result"]
    assert "审查判断：未发现“非法垄断技术”对应的不利条款或安排。" in report
    assert "审查判断：未发现“侵害他人技术成果”对应的不利条款或安排。" in report
    assert "合同已覆盖非法垄断技术要求" not in report
    assert "合同已覆盖侵害他人技术成果要求" not in report

    liability = by_id["TECH-SHARED-08"]
    assessment = liability["decision_trace"]["liability_limit_legal_assessment"]
    assert liability["legal_basis_index"] == "《中华人民共和国民法典》第506条、第584条"
    assert assessment["validity_conclusion"] == "NOT_DETERMINED"
    assert "不足以将全部责任上限直接认定为无效" in assessment["reason"]
    liability_detail = report[report.index("服务方赔偿责任上限及免责范围过宽（TECH-SHARED-08）"):]
    assert "《中华人民共和国民法典》第506条" in liability_detail
    assert "《中华人民共和国民法典》第584条" in liability_detail
    assert "该责任上限无效" not in by_group["service_provider_liability_cap_exclusion"]["issue"]

    ownership_detail = report[report.index("服务中新成果及使用权安排偏向服务方（TC-09-03）"):]
    assert "【实际法律暴露】：" in ownership_detail
    assert "尚未形成经过双方主体比较的方向性结论" not in ownership_detail
    assert "该条款对当前审查方形成不利或限制性安排" not in ownership_detail
    assert "尚未形成经过双方主体比较的方向性结论" not in report

    personal_detail = report[report.index("个人信息委托处理条款不完整（TECH-SHARED-02）"):]
    assert (
        "【缺失的法律后果】：数据处理边界和安全责任不清，可能产生个人信息合规、数据泄露和项目停摆风险。"
        in personal_detail
    )
    assert all(stutter not in report for stutter in (
        "不完整不清", "缺失不清", "未明确不清", "不足不清", "未约定不清",
    ))


def test_service_confidentiality_period_and_scope_bind_direct_confidentiality_clauses():
    output = _review()
    item = next(row for row in output["items"] if row["item_id"] == "TC-11-02")
    evidence_text = "\n".join(
        str(row.get("text") or row.get("quote") or "") for row in item["evidence_clauses"]
    )
    elements = {row["element_id"]: row["status"] for row in item["coverage_elements"]}

    assert item["coverage_status"] == "通过"
    assert item["action_priority"] == "NONE"
    assert all(elements[element_id] == "COVERED" for element_id in (
        "confidential_information_scope",
        "mutual_confidentiality_responsibility",
        "third_party_disclosure_restriction",
        "statutory_disclosure_exception",
        "post_termination_survival",
        "confidentiality_period",
    ))
    assert "商业秘密、技术资料、客户信息、经营信息及其他未公开信息" in evidence_text
    assert "任何一方不得向第三方披露" in evidence_text
    assert "司法机关或行政机关依法要求披露的除外" in evidence_text
    assert "合同终止后继续有效" in evidence_text
    assert "双方保密义务长期有效" in evidence_text
    assert "仅可在履行本合同所必需的范围内使用乙方提供的数据" not in evidence_text
    assert not any("TC-11-02" in row["source_review_ids"] for row in output["risk_findings"])


def test_service_provider_confidentiality_breach_narrative_excludes_liability_expansion():
    output = _review()
    item = next(row for row in output["items"] if row["item_id"] == "TC-11-04")
    finding = next(row for row in output["risk_findings"] if "TC-11-04" in row["source_review_ids"])
    disposition = next(row for row in output["final_dispositions"] if "TC-11-04" in row["review_ids"])
    elements = {row["element_id"]: row for row in item["coverage_elements"]}
    report = output["result"]
    section_start = report.index("TC-11-04）")
    section_end = report.index("TC-12-01）", section_start)
    section = report[section_start:section_end]

    # Coverage extraction retains the absent compensation fields for internal
    # audit, while party-direction filtering keeps them out of the provider's
    # user-facing risk and recommendation.
    for element_id in ("direct_loss", "lost_profit", "attorney_fee", "litigation_cost"):
        assert elements[element_id]["status"] == "MISSING"
        assert elements[element_id]["risk_relevant"] is False
        assert elements[element_id]["direction"] == "NEUTRAL"
    for element_id in ("stop_disclosure", "stop_use", "mitigation_obligation"):
        assert elements[element_id]["status"] == "MISSING"
        assert elements[element_id]["risk_relevant"] is True

    assert item["decision_trace"]["risk_to_current_party"]["risk_scope"] == "INCIDENT_CONTAINMENT_ONLY"
    assert item["decision_trace"]["recommendation_direction_check"]["expands_current_party_liability"] is False
    assert "停止披露" in finding["issue"]
    assert "停止使用" in finding["issue"]
    assert "防止损失或披露继续扩大" in finding["issue"]
    assert "可得利益损失" not in finding["issue"]
    assert "合理律师费" not in finding["issue"]
    assert "诉讼费用" not in finding["issue"]
    assert "无需重复创设同一赔偿义务" in disposition["recommendation"]
    assert "赔偿范围包括可得利益损失" not in section
    assert "赔偿范围包括合理律师费" not in section
    assert "赔偿范围包括诉讼费用" not in section
