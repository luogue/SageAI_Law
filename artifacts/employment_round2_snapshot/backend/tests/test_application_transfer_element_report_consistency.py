"""Final atomic state is authoritative for absence searches and report prose."""
from copy import deepcopy
from pathlib import Path
import asyncio
import re

import pytest
from docx import Document

from review_v2.engine import (
    ReviewItemResult, _build_narrative_state, _evidence_constrained_recommendation,
    _finding_recommendation, _normalize_technology_transaction_terminology,
    _objective_fact_search_targets,
    _operative_technology_clauses, _reconcile_item_element_precedence,
    _reconcile_item_search_audit_with_positive_evidence, _set_multi_element_state,
    _uncovered_element_label, review_contract_v2,
    _apply_application_right_transfer_facts, _recommendation_action_for_element,
)
from review_v2.technology_patent_transfer_analysis import extract_transfer_registration_facts, extract_transfer_authority_facts
from review_v2.technology_contract_checklist import get_technology_checklist
from review_v2.user_profile import UserRiskProfile


def review(elements=(), status="部分覆盖"):
    return ReviewItemResult(
        item_id="SYNTHETIC", node_id="test", title="办理事项", risk_level="低",
        status=status, coverage_status=status, issue_type="PARTIAL",
        risk_point_description="旧风险文本", suggested_clause_example="保留现有办理安排。",
        legal_basis_index="", legal_consequence_if_missing="", applicable_if="本项适用",
        matched_terms=[], evidence_clauses=[{"clause_id": "E1", "text": "甲方负责办理登记。", "source_type": "SOURCE_CLAUSE"}],
        evidence_location="第一条", reasoning_summary="", applicability={"status": "APPLICABLE"},
        coverage_elements=deepcopy(list(elements)),
    )


@pytest.mark.parametrize("party", ["甲方", "乙方", "受让方", "让与方", "买方", "卖方"])
@pytest.mark.parametrize("days", ["60日", "十五个工作日", "3个月"])
def test_registration_actor_and_timing_are_independent_of_contract_roles(party, days):
    clauses = _operative_technology_clauses(
        f"第四条 转让登记\n由{party}负责在签署后{days}内办理专利申请权转让登记。"
        "乙方应提供必要文件并协助办理。"
    )
    facts = extract_transfer_registration_facts(clauses)
    assert set(facts) == {"registration_responsible_party", "registration_deadline", "registration_cooperation"}
    assert all(row["clause_id"] == clauses[0].clause_id for rows in facts.values() for row in rows)


def test_registration_payment_timing_and_unrelated_assistance_do_not_fill_gaps():
    text = (
        "第一条 转让登记\n甲方负责办理专利申请权转让登记。\n"
        "第二条 付款\n转让登记完成后10日内支付价款。\n"
        "第三条 资料交付\n乙方应提供技术文件并协助验收。"
    )
    assert set(extract_transfer_registration_facts(_operative_technology_clauses(text))) == {"registration_responsible_party"}


@pytest.mark.parametrize("text", [
    "甲方不负责办理专利申请权转让登记。",
    "专利申请权转让登记的主体及60日办理期限未约定。",
    "专利申请权转让登记由双方另行协商办理。",
])
def test_negative_or_unallocated_registration_is_not_positive(text):
    assert not extract_transfer_registration_facts(_operative_technology_clauses("第一条 登记\n" + text))


def test_source_bound_positive_beats_later_absence_and_recomputes_pass():
    row = review([
        {"element_id": "filing_owner", "status": "COVERED", "proposition": "办理主体", "evidence_ids": ["E1"]},
        {"element_id": "filing_owner", "status": "MISSING", "proposition": "办理主体"},
    ], "缺失")
    row.decision_trace["specialized_search_audit"] = {"missing_elements": ["办理主体"]}
    missing = _reconcile_item_element_precedence(row)
    _reconcile_item_search_audit_with_positive_evidence(row, missing)
    assert not missing
    assert row.coverage_status == "通过"
    assert row.decision_trace["specialized_search_audit"]["missing_elements"] == []
    assert row.decision_trace["specialized_search_audit"]["search_status"] == "FOUND"
    assert not _build_narrative_state(row)["missing_elements"]


@pytest.mark.parametrize("evidence_ids", [[], ["SEARCH-1"], ["not-a-source"]])
def test_unbound_or_search_audit_is_not_positive_source(evidence_ids):
    row = review([
        {"element_id": "filing_owner", "status": "COVERED", "proposition": "办理主体", "evidence_ids": evidence_ids},
        {"element_id": "filing_owner", "status": "MISSING", "proposition": "办理主体"},
    ], "缺失")
    row.evidence_clauses.append({"clause_id": "SEARCH-1", "source_type": "SEARCH_AUDIT", "text": "检索已完成"})
    assert _reconcile_item_element_precedence(row)[0]["status"] == "MISSING"
    assert row.coverage_status == "缺失"


def test_similar_words_do_not_erase_a_distinct_atomic_deadline_gap():
    row = review([
        {"element_id": "filing_owner", "status": "COVERED", "proposition": "专利申请权转让登记办理主体", "evidence_ids": ["E1"]},
        {"element_id": "filing_deadline", "status": "MISSING", "proposition": "专利申请权转让登记办理期限"},
    ])
    row.decision_trace.update({"explicit_missing_elements": ["合同已明确覆盖办理主体"],
                               "specialized_search_audit": {"missing_elements": ["专利申请权转让登记办理主体"]}})
    missing = _reconcile_item_element_precedence(row)
    _reconcile_item_search_audit_with_positive_evidence(row, missing)
    assert row.decision_trace["specialized_search_audit"]["missing_elements"] == ["专利申请权转让登记办理期限"]
    narrative = _build_narrative_state(row)
    assert narrative["missing_elements"] == ["专利申请权转让登记办理期限"]
    assert "办理主体" not in narrative["issue"] + narrative["recommendation"]


@pytest.mark.parametrize("fact", ["费用分担方式", "发生技术困难时的协作义务", "投入不足后的补足机制"])
def test_unkeyed_final_gap_keeps_its_objective_search_target(fact):
    assert _objective_fact_search_targets({"element_id": "unkeyed", "proposition": fact + "尚未明确"}) == [fact]


def test_positive_projection_does_not_change_not_applicable_routing():
    row = review([{"element_id": "core", "status": "COVERED", "proposition": "核心要求", "evidence_ids": ["E1"]}], "不适用")
    row.applicability = {"status": "NOT_APPLICABLE", "applicable": False}
    _reconcile_item_element_precedence(row)
    assert row.coverage_status == "不适用"


def test_uncertain_evidence_does_not_get_a_free_pass():
    row = review([
        {"element_id": "core", "status": "COVERED", "proposition": "核心要求", "evidence_ids": ["E1"]},
        {"element_id": "proof", "status": "UNCERTAIN", "proposition": "独立证明文件"},
    ])
    assert _reconcile_item_element_precedence(row)
    assert row.coverage_status == "部分覆盖"


@pytest.mark.parametrize("proposition", [
    "费用凭证已列入交付材料。", "权利负担已有披露。", "文件版本已有安排。",
    "合同已明确覆盖通知期限。", "已有程序文件已列入附件清单。",
])
def test_missing_labels_never_reuse_completed_coverage_claims(proposition):
    row = review()
    _set_multi_element_state(row, [{"element_id": "proof", "status": "MISSING", "proposition": proposition,
                                    "bound_evidence_ids": []}], [])
    _reconcile_item_element_precedence(row)
    label = _uncovered_element_label(row.coverage_elements[0])
    narrative = _build_narrative_state(row)
    assert not re.search(r"已列入|已有披露|已有安排|已明确覆盖", label + narrative["issue"] + narrative["recommendation"])
    assert row.decision_trace["final_missing_element_ids"] == ["proof"]


@pytest.mark.parametrize("advice", [
    "保留现有登记办理安排；明确期限。", "建议在现有安排基础上，补充通知。",
    "建议在已有责任分配基础之上补充费用承担。",
])
def test_missing_advice_cannot_invent_existing_arrangement(advice):
    row = review([{"element_id": "notice", "status": "MISSING", "proposition": "通知期限"}], "缺失")
    row.evidence_clauses = []
    assert _evidence_constrained_recommendation(row, advice) == "建议补充明确通知期限。"
    assert "保留现有" not in _finding_recommendation("third_party_ip_claim_and_remedy", row)


@pytest.mark.parametrize("primary,expected", [
    ("TECHNOLOGY_TRANSFER", "受让方通知让与方"),
    ("TECHNOLOGY_LICENSE", "被许可方通知许可方"),
    ("COOPERATIVE_DEVELOPMENT", "接受相关技术的合作方通知技术提供方"),
])
def test_transaction_role_copy_preserves_source_quotes(primary, expected):
    row = review()
    row.suggested_clause_example = "被许可方通知许可方"
    row.evidence_clauses[0]["quote"] = "既有被许可方与许可方另行协商。"
    source = deepcopy(row.evidence_clauses)
    _normalize_technology_transaction_terminology([row], [], {"primary_subtype": primary})
    assert row.suggested_clause_example == expected
    assert row.evidence_clauses == source


@pytest.mark.parametrize("primary", ["TECHNOLOGY_TRANSFER", "TECHNOLOGY_LICENSE", "COOPERATIVE_DEVELOPMENT"])
def test_generic_applicability_is_not_subject_to_role_localization(primary):
    row = review()
    original = next(cp for cp in get_technology_checklist() if cp.item_id == "TC-12-01").applicable_if
    row.applicable_if = original
    for _ in range(2):
        _normalize_technology_transaction_terminology([row], [], {"primary_subtype": primary})
    assert row.applicable_if == original
    assert "受让人或被许可人" in original
    assert "受让人、受让人" not in row.applicable_if


@pytest.mark.parametrize("statement", [
    "保证其有权转让本项专利申请权", "确认其有权处分本项专利申请权",
    "保证其合法拥有本项专利申请权并有权转让", "保证对本项专利申请权具有完整处分权",
])
@pytest.mark.parametrize("header,provider", [
    ("受让方（甲方）：收购公司\n让与方（乙方）：研发公司", "乙方"),
    ("甲方（让与方）：研发公司\n乙方（受让方）：收购公司", "甲方"),
    ("受让方（甲方）：收购公司\n转让方（丙方）：研发公司", "丙方"),
    ("", "让与方"),
])
def test_authority_synonyms_share_an_element_without_assuming_party_b(header, provider, statement):
    text = header + "\n第六条 权利保证\n" + provider + statement + "。"
    clauses = _operative_technology_clauses(text)
    row = review(status="缺失")
    row.item_id = "TC-06-03"
    row.decision_trace["specialized_search_audit"] = {"missing_elements": ["让与方具有专利申请权处分或转让权的保证"]}
    _apply_application_right_transfer_facts([row], clauses, {"tertiary_subtype": "PATENT_APPLICATION_RIGHT_TRANSFER"})
    missing = _reconcile_item_element_precedence(row)
    _reconcile_item_search_audit_with_positive_evidence(row, missing)
    elements = {e["element_id"]: e for e in row.coverage_elements}
    authority = elements["transferor_title_and_authority"]
    assert authority["status"] == "COVERED"
    assert authority["evidence_ids"]
    assert authority["evidence_level"] == "CONTRACTUAL_REPRESENTATION"
    assert row.coverage_status == "部分覆盖"
    assert {e["element_id"] for e in missing} == {"application_coownership_disclosure", "application_pledge_disclosure", "application_seizure_disclosure", "application_other_restrictions_disclosure", "third_party_claim_handling"}
    assert "处分或转让权的保证" not in "；".join(row.decision_trace["specialized_search_audit"]["missing_elements"])
    recommendation = _finding_recommendation(row.risk_group, row)
    assert "处分或转让权的保证" not in recommendation
    assert "质押" in recommendation and "抗辩" in recommendation


def test_exact_self_representation_is_positive_even_in_an_isolated_clause():
    facts = extract_transfer_authority_facts(_operative_technology_clauses("第六条\n乙方保证其有权转让本项专利申请权。"))
    assert facts and facts[0]["actor"] == "乙方"
    assert "乙方保证其有权转让本项专利申请权" in facts[0]["quote"]


def test_application_authority_cannot_borrow_a_granted_patent_representation():
    clauses = _operative_technology_clauses("让与方（乙方）：公司B\n第六条 权利保证\n乙方保证有权转让另一项已授权专利权。")
    row = review(status="缺失")
    row.item_id = "TC-06-03"
    _apply_application_right_transfer_facts([row], clauses, {"tertiary_subtype": "PATENT_APPLICATION_RIGHT_TRANSFER"})
    assert next(e for e in row.coverage_elements if e["element_id"] == "transferor_title_and_authority")["status"] == "MISSING"


@pytest.mark.parametrize("statement", [
    "乙方不保证其有权转让本项专利申请权。",
    "乙方无法保证其具有本项专利申请权的完整处分权。",
    "乙方对本项专利申请权不具有完整处分权。",
    "若乙方有权转让本项专利申请权，双方另行签约。",
    "如果乙方保证其有权转让本项专利申请权，甲方愿意受让。",
    "乙方是否有权处分本项专利申请权有待核实。",
    "乙方有权许可本项专利申请权涉及的技术。",
    "乙方保证其有权转让设备，专利申请权由甲方保证有权处分。",
    "甲方保证其有权转让本项专利申请权。",
])
def test_authority_requires_affirmative_matching_actor_and_asset(statement):
    clauses = _operative_technology_clauses("受让方（甲方）：公司A\n让与方（乙方）：公司B\n第六条 权利保证\n" + statement)
    assert not extract_transfer_authority_facts(clauses)


def test_final_element_policy_ignores_stale_recommendation_and_covered_input():
    row = review([
        {"element_id": "authority", "status": "COVERED", "proposition": "处分权保证", "evidence_ids": ["E1"]},
        {"element_id": "notice", "status": "MISSING", "proposition": "通知期限"},
        {"element_id": "delivery", "status": "PARTIALLY_COVERED", "proposition": "已有交付义务，但交付范围尚未明确", "residual_gap": "交付资料范围"},
    ])
    row.decision_trace["recommendation_element_projection"] = {"policy": "FINAL_UNRESOLVED_ELEMENTS"}
    _reconcile_item_element_precedence(row)
    # Even a stale final-id input cannot upgrade a covered row to a drafting gap.
    row.decision_trace["final_missing_element_ids"].append("authority")
    advice = _evidence_constrained_recommendation(row, "建议补充处分权保证、通知期限、交付义务和支付义务。")
    assert advice == "建议补充明确通知期限；补充明确交付资料范围。"
    assert _recommendation_action_for_element(row.coverage_elements[0])[0] == ""


def test_final_element_policy_does_not_fabricate_advice_for_an_empty_delta():
    row = review([{"element_id": "core", "status": "COVERED", "proposition": "核心义务", "evidence_ids": ["E1"]}])
    row.decision_trace.update({"recommendation_element_projection": {"policy": "FINAL_UNRESOLVED_ELEMENTS"}, "final_missing_element_ids": []})
    assert _evidence_constrained_recommendation(row, "建议补充核心义务。") == ""


@pytest.fixture(scope="module")
def result():
    path = next((Path(__file__).resolve().parents[2] / "test_data").glob("44*.docx"))
    document = Document(path)
    text = "\n".join(p.text for p in document.paragraphs)
    text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for table in document.tables for row in table.rows)
    return review_contract_v2(text, "technology", use_llm=False, include_low_risk=True,
                              profile=UserRiskProfile(user_identity="甲方/受让方"))


def test_current_contract_registration_and_confidentiality_final_state(result):
    by_id = {row["item_id"]: row for row in result["items"]}
    for key in ("TC-06-04", "TC-11-01"):
        row = by_id[key]
        assert row["coverage_status"] == "通过"
        assert row["decision_trace"]["final_missing_element_ids"] == []
        assert all(element["status"] == "COVERED" and element["evidence_ids"] for element in row["coverage_elements"])
        assert not row["narrative_state"]["missing_elements"]
        assert not row["narrative_state"]["issue"]
    assert "第十一条" in by_id["TC-06-04"]["evidence_location"]
    assert not result["release_quality_gate"]["failures"]
    assert result["release_quality_gate"]["status"] == "REPORT_READY"
    assert result["deterministic_analysis"]["technology_classification"]["tertiary_subtype"] == "PATENT_APPLICATION_RIGHT_TRANSFER"


def test_current_contract_missing_lists_and_active_role_text_are_consistent(result):
    for row in result["items"]:
        if row["coverage_status"] == "不适用":
            continue
        narrative = row["narrative_state"]
        assert not re.search(r"已列入|已有披露|已明确覆盖", "；".join(narrative.get("missing_elements", [])))
        text = " ".join(str(row.get(key) or "") for key in (
            "risk_point_description", "legal_consequence_if_missing", "suggested_clause_example", "action_suggestion"))
        assert not re.search(r"被许可[方人]|许可[方人]", text), (row["item_id"], text)
        covered = {element["element_id"] for element in row["coverage_elements"] if element["status"] == "COVERED"}
        assert not covered.intersection(row["decision_trace"].get("final_missing_element_ids", []))


def test_current_contract_claim_context_does_not_invent_protective_duties(result):
    response = next(row for row in result["items"] if row["item_id"] == "TC-12-02")
    elements = {row["element_id"]: row for row in response["coverage_elements"]}
    assert response["coverage_status"] == "部分覆盖"
    assert elements["claim_trigger"]["status"] == "COVERED"
    assert elements["defense_cost_allocation"]["status"] == "MISSING"
    for key in ("replacement_or_workaround", "license_cure", "refund_or_termination_remedy"):
        assert elements[key]["status"] == "MISSING"


def test_current_authority_gap_and_recommendations_are_consistent_in_all_sections(result):
    row = next(row for row in result["items"] if row["item_id"] == "TC-06-03")
    assert row["coverage_status"] == "部分覆盖"
    assert row["decision_trace"]["final_missing_element_ids"] == ["application_coownership_disclosure", "application_pledge_disclosure", "application_seizure_disclosure", "application_other_restrictions_disclosure", "third_party_claim_handling"]
    assert "第六条" in row["evidence_location"]
    assert row["coverage_elements"][0]["status"] == "COVERED"
    for text in (row["risk_point_description"], row["suggested_clause_example"], row["narrative_state"]["issue"]):
        assert "处分或转让权的保证" not in text
        assert "质押" in text and "抗辩" in text
    finding = next(f for f in result["risk_findings"] if "TC-06-03" in f["source_review_ids"])
    assert "处分或转让权的保证" not in finding["recommendation"]
    assert "处分或转让权的保证" not in finding["issue"]
    assert {e["element_id"] for e in finding["final_missing_elements"]} == {"application_coownership_disclosure", "application_pledge_disclosure", "application_seizure_disclosure", "application_other_restrictions_disclosure", "third_party_claim_handling"}
    for recommendation in result["recommendations"]:
        if finding["finding_id"] not in recommendation["source_finding_ids"]:
            continue
        assert "处分或转让权的保证" not in recommendation["suggested_text"]
        assert set(recommendation["recommendation_input_elements"]) == {"application_coownership_disclosure", "application_pledge_disclosure", "application_seizure_disclosure", "application_other_restrictions_disclosure", "third_party_claim_handling"}
    for line in result["result"].splitlines():
        if "建议" in line and "处分或转让权的保证" in line:
            # A summary may display both covered facts and advice on one line;
            # only its recommendation portion is constrained to remaining gaps.
            assert "处分或转让权的保证" not in line.split("建议", 1)[1]
    infringement = next(row for row in result["items"] if row["item_id"] == "TC-12-01")
    assert "受让人或被许可人" in infringement["applicable_if"]
    assert "受让人、受让人" not in result["result"]


def test_current_file_upload_and_api_schema_preserve_final_element_truth():
    from fastapi import UploadFile
    from routers.review_v2 import ReviewV2Result, review_v2_file

    path = next((Path(__file__).resolve().parents[2] / "test_data").glob("44*.docx"))
    with path.open("rb") as source:
        response = asyncio.run(review_v2_file(
            file=UploadFile(filename=path.name, file=source), contract_type="technology",
            include_low_risk=True, use_llm=False,
            profile_json='{"user_identity":"甲方/受让方"}', current_user={},
        ))
    payload = ReviewV2Result(**response).model_dump()
    test_current_contract_registration_and_confidentiality_final_state(payload)
    test_current_contract_missing_lists_and_active_role_text_are_consistent(payload)
    test_current_contract_claim_context_does_not_invent_protective_duties(payload)
    test_current_authority_gap_and_recommendations_are_consistent_in_all_sections(payload)
