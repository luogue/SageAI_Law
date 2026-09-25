"""输出层：处理LLM输出，提取关键信息并格式化"""
import json
import time
import io
import re
import logging
import html
from typing import Dict, Optional, Tuple

logger = logging.getLogger(__name__)


class OutputLayer:
    SUPPORTED_FORMATS = ["md", "docx", "pdf"]

    def __init__(self):
        # 章节标题关键词映射（支持多种格式）
        # 新版提示词输出5个章节，这里只提取用户可见的3个模块：
        #   审查前提（来自"1. 前置步骤：关键前提确认"）
        #   风险评估（来自"3. 风险评估与商业可行性评估"）
        #   结论与行动建议（来自"5. 结论与行动建议"）
        # 排除的章节："2. 逻辑推理"和"4. 自检"（内部推理内容，不展示给用户）
        self.section_keywords = {
            '审查前提': ['前置步骤', '关键前提', '前提确认'],
            '风险评估': ['风险评估', '商业可行性评估', '风险分析'],
            '结论与行动建议': ['结论与行动建议', '应该怎么做', '总结']
        }

    def _extract_section(self, content: str, section_name: str) -> str:
        """
        从内容中提取指定章节
        使用更灵活的正则匹配，适配多种markdown格式
        """
        keywords = self.section_keywords.get(section_name, [])
        if not keywords:
            return ""

        for keyword in keywords:
            # 模式1：匹配带序号的加粗标题 **3. xxx** 或 **xxx**
            pattern1 = rf'\*\*(\d+\.\s+)?{keyword}[^\*]*\*\*([\s\S]*?)(?=\n\*\*|\n---|\Z)'
            # 模式2：匹配普通标题 3. xxx
            pattern2 = rf'(\d+\.\s+){keyword}([\s\S]*?)(?=\n\d+\.\s+|\n---|\Z)'
            # 模式3：匹配markdown标题 ## xxx 或 ### xxx
            pattern3 = rf'#{1,4}\s+{keyword}([\s\S]*?)(?=\n#|\n---|\Z)'
            # 模式4：匹配纯文本标题（无格式）
            pattern4 = rf'^[\s]*{keyword}([\s\S]*?)(?=\n[A-Za-z0-9一二三四五六七八九十]{{1,3}}\.\s+|\n---|\Z)'

            for pattern in [pattern1, pattern2, pattern3, pattern4]:
                match = re.search(pattern, content)
                if match:
                    result = match.group(0).strip()
                    # 移除标题标记，保留内容
                    result = re.sub(rf'^\*\*(\d+\.\s+)?{keyword}[^\*]*\*\*\s*', '', result)
                    result = re.sub(rf'^(\d+\.\s+){keyword}\s*', '', result)
                    result = re.sub(rf'^#{{1,4}}\s+{keyword}\s*', '', result)
                    result = re.sub(rf'^{keyword}\s*', '', result)
                    return result.strip()

        return ""

    def process(self, raw_output: str) -> str:
        """
        处理原始输出，提取并格式化3个用户可见章节。
        排除"逻辑推理"和"自检"两个内部章节。
        
        :param raw_output: LLM原始输出
        :return: 格式化后的用户可见报告内容
        """
        try:
            # 提取3个用户可见章节
            premise_section = self._extract_section(raw_output, '审查前提')
            risk_section = self._extract_section(raw_output, '风险评估')
            conclusion_section = self._extract_section(raw_output, '结论与行动建议')

            # 构建最终报告
            sections = []
            
            if premise_section:
                sections.append(f"## 1. 审查前提\n\n{premise_section}")
            
            if risk_section:
                sections.append(f"## 2. 风险评估\n\n{risk_section}")
            
            if conclusion_section:
                sections.append(f"## 3. 结论与行动建议\n\n{conclusion_section}")

            # 如果没有找到任何章节，返回原始内容（降级处理）
            if not sections:
                logger.warning("[OutputLayer] 未找到任何章节，返回原始内容")
                return self._clean_output(raw_output)

            return "\n\n---\n\n".join(sections)

        except Exception as e:
            logger.error(f"[OutputLayer] 处理失败: {str(e)}", exc_info=True)
            return self._clean_output(raw_output)

    def _filter_internal_content(self, raw_report: str) -> str:
        """
        过滤 LLM 原始输出中的 XML 标签，提取 <FINAL_REPORT> 内的纯内容。
        
        注意：此方法只做 XML 标签清理，不做章节过滤。
        章节级别的过滤（排除"逻辑推理""自检"）由 process() / package_output() 负责。
        
        Args:
            raw_report: LLM 原始输出（可能包含 XML 标签）
        
        Returns:
            去除 XML 标签后的报告内容（仍包含全部5个章节）
        """
        if not raw_report:
            return ""
        
        content = raw_report
        
        # 1. 优先提取 <FINAL_REPORT> 标签（防御性提取，即使上游已解析）
        final_pattern = r'<FINAL_REPORT>\s*([\s\S]*?)\s*</FINAL_REPORT>'
        final_match = re.search(final_pattern, content, re.IGNORECASE)
        if final_match:
            content = final_match.group(1).strip()
            logger.debug(f"[OutputLayer] 从 <FINAL_REPORT> 标签提取内容，长度: {len(content)}")
        
        # 2. 移除 <INTERNAL>...</INTERNAL> 标签及其内容（如果上游未清除）
        internal_pattern = r'<INTERNAL>\s*[\s\S]*?\s*</INTERNAL>'
        content = re.sub(internal_pattern, '', content, flags=re.IGNORECASE)
        
        # 3. 移除残留的独立 XML 标签
        content = re.sub(r'</?FINAL_REPORT>', '', content, flags=re.IGNORECASE)
        content = re.sub(r'</?INTERNAL>', '', content, flags=re.IGNORECASE)
        
        # 4. 清理多余空行
        content = re.sub(r'\n{4,}', '\n\n\n', content)
        
        # 5. 基础清理
        content = self._clean_output(content)
        
        logger.debug(f"[OutputLayer] _filter_internal_content: 输入 {len(raw_report)} → 输出 {len(content)} 字符")
        return content

    def _clean_output(self, text: str) -> str:
        """清理输出内容，移除多余格式"""
        text = re.sub(r'<\|.*?\|>', '', text)
        text = re.sub(r'---+', '', text)
        text = text.strip()
        return text

    # ============================================================
    # 以下方法从旧版 output_layer.py 恢复，适配新版 3 模块提取
    # ============================================================

    def _generate_metadata(self, validation_result: Dict) -> Dict:
        """生成元数据"""
        metadata = {
            "timestamp": time.strftime("%Y-%m-%d %H:%M:%S", time.localtime()),
            "model_version": "qwen-plus",
            "confidence_score": validation_result.get('confidence_check', {}).get('confidence_score', 0),
            "validation_status": validation_result.get('status', 'unknown'),
            "has_warnings": validation_result.get('has_warnings', False),
            "warnings": validation_result.get('warnings', []),
            "format_check_status": validation_result.get('format_check', {}).get('status', 'unknown'),
            "confidence_check_status": validation_result.get('confidence_check', {}).get('status', 'unknown'),
            "source_check_status": validation_result.get('source_check', {}).get('status', 'unknown')
        }
        return metadata

    def _format_markdown(self, raw_report: str) -> str:
        """
        格式化 Markdown 输出：先提取3个用户可见模块，再包装为完整报告。
        """
        # 使用 process() 提取3个用户可见模块（排除"逻辑推理""自检"）
        filtered_report = self.process(raw_report)
        
        # 兜底：如果过滤后内容为空或过短，使用原始报告全文
        if not filtered_report or not filtered_report.strip() or len(filtered_report.strip()) < 50:
            filtered_len = len(filtered_report.strip()) if filtered_report else 0
            raw_len = len(raw_report.strip()) if raw_report else 0
            logger.warning("[OUTPUT] process() 后内容过短(len=%d)，使用原始报告全文(%d chars)", filtered_len, raw_len)
            if raw_report and raw_report.strip():
                filtered_report = raw_report.strip()
            else:
                filtered_report = "报告生成遇到问题，请稍后重试或联系管理员。"
        
        markdown = f"""# 合同审查报告

---

{filtered_report}

---

*报告生成时间：{time.strftime("%Y-%m-%d %H:%M:%S", time.localtime())}*
*模型版本：qwen-plus*
"""
        return markdown

    def _extract_key_info(self, report: str) -> Dict:
        """从报告中提取关键信息"""
        risk_level_pattern = r'风险等级[：:]?\s*(低|中|高|极高)'
        risk_level_match = re.search(risk_level_pattern, report)
        risk_level = risk_level_match.group(1) if risk_level_match else '未知'
        
        risk_pattern = r'[一二三四五六七八九十]+\.\s*([^\n]+)'
        risk_points = re.findall(risk_pattern, report)[:10]
        
        suggestion_pattern = r'(建议|修改建议)[：:]?\s*([^\n]+)'
        suggestions = re.findall(suggestion_pattern, report)[:5]
        suggestions = [s[1] if len(s) > 1 else s[0] for s in suggestions]
        
        conclusion_pattern = r'(结论与建议)[：:]?\s*([\s\S]*?)(?=\n### |$)'
        conclusion_match = re.search(conclusion_pattern, report)
        summary = conclusion_match.group(2).strip() if conclusion_match else ''
        
        return {
            "risk_level": risk_level,
            "risk_points": risk_points,
            "suggestions": suggestions,
            "summary": summary
        }

    def generate_html_preview(self, report_markdown: str) -> str:
        """生成HTML预览版本"""
        try:
            import markdown
            return markdown.markdown(
                report_markdown,
                extensions=['tables', 'fenced_code', 'codehilite']
            )
        except ImportError:
            logger.warning("markdown package is not installed; using built-in HTML preview fallback")
            return self._generate_basic_html_preview(report_markdown)

    def _generate_basic_html_preview(self, report_markdown: str) -> str:
        """Generate a safe minimal HTML preview without optional markdown dependency."""
        html_lines = []
        in_list = False

        for raw_line in report_markdown.splitlines():
            line = raw_line.strip()

            if not line:
                if in_list:
                    html_lines.append("</ul>")
                    in_list = False
                continue

            if line.startswith("- ") or line.startswith("* "):
                if not in_list:
                    html_lines.append("<ul>")
                    in_list = True
                html_lines.append(f"<li>{self._format_inline_html(line[2:])}</li>")
                continue

            if in_list:
                html_lines.append("</ul>")
                in_list = False

            if line.startswith("# "):
                html_lines.append(f"<h1>{self._format_inline_html(line[2:])}</h1>")
            elif line.startswith("## "):
                html_lines.append(f"<h2>{self._format_inline_html(line[3:])}</h2>")
            elif line.startswith("### "):
                html_lines.append(f"<h3>{self._format_inline_html(line[4:])}</h3>")
            elif line.startswith("#### "):
                html_lines.append(f"<h4>{self._format_inline_html(line[5:])}</h4>")
            elif line.startswith("|"):
                html_lines.append(f"<pre>{html.escape(line)}</pre>")
            elif line.startswith("---"):
                html_lines.append("<hr>")
            else:
                html_lines.append(f"<p>{self._format_inline_html(line)}</p>")

        if in_list:
            html_lines.append("</ul>")

        return "\n".join(html_lines)

    def _format_inline_html(self, text: str) -> str:
        escaped = html.escape(text)
        escaped = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", escaped)
        escaped = re.sub(r"\*([^*]+)\*", r"<em>\1</em>", escaped)
        return escaped

    def package_output(self, raw_report: str, validation_result: Dict, input_data: Dict = None) -> Dict:
        """
        打包输出：3模块提取 → Markdown 格式化 → 结构化包装
        
        流程：
        1. 调用 _format_markdown → process() 提取3个用户可见模块
        2. 排除"逻辑推理""自检"等内部章节
        3. 包装为结构化输出供前端使用
        """
        metadata = self._generate_metadata(validation_result)
        markdown_content = self._format_markdown(raw_report)
        
        if input_data and 'contract_text' in input_data:
            metadata['contract_text'] = input_data['contract_text']
        
        key_info = self._extract_key_info(raw_report)
        
        output = {
            "status": "success",
            "metadata": metadata,
            "content": {
                "markdown": markdown_content,
                "html_preview": self.generate_html_preview(markdown_content),
                "summary": key_info.get('summary', ''),
                "risk_level": key_info.get('risk_level', '未知'),
                "risk_points": key_info.get('risk_points', []),
                "suggestions": key_info.get('suggestions', [])
            },
            "validation": {
                "has_warnings": validation_result.get('has_warnings', False),
                "warnings": validation_result.get('warnings', []),
                "confidence_score": validation_result.get('confidence_check', {}).get('confidence_score', 0)
            }
        }
        
        return output

    def format_for_frontend(self, output: Dict) -> Dict:
        """格式化输出供前端使用"""
        frontend_output = {
            "success": True,
            "message": "审查完成",
            "result": output['content']['markdown'],
            "html_preview": output['content'].get('html_preview', ''),
            "metadata": output['metadata'],
            "validation": output['validation'],
            "summary": {
                "risk_level": output['content']['risk_level'],
                "risk_count": len(output['content']['risk_points']),
                "suggestion_count": len(output['content']['suggestions']),
                "confidence": output['validation']['confidence_score']
            }
        }
        
        return frontend_output

    # ============================================================
    # 文件生成方法（用于报告下载）
    # ============================================================

    def generate_report_file(self, report_markdown: str, format: str = "md", filename: str = "合同审查报告") -> Tuple[bytes, str]:
        """
        根据格式生成报告文件
        
        :param report_markdown: Markdown格式的报告内容
        :param format: 输出格式 ("md", "docx", "pdf")
        :param filename: 文件名（不含扩展名）
        :return: (二进制文件内容, MIME类型)
        """
        if format == "md":
            return self._generate_markdown_file(report_markdown, filename)
        elif format == "docx":
            return self._generate_word(report_markdown, filename)
        elif format == "pdf":
            return self._generate_pdf(report_markdown, filename)
        else:
            raise ValueError(f"不支持的格式: {format}，支持的格式: {self.SUPPORTED_FORMATS}")

    def _generate_markdown_file(self, report_markdown: str, filename: str) -> Tuple[bytes, str]:
        """生成Markdown文件"""
        content = report_markdown.encode('utf-8')
        return content, "text/markdown"

    def _generate_word(self, report_markdown: str, filename: str) -> Tuple[bytes, str]:
        """将Markdown转换为Word文档"""
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        doc = Document()
        
        for line in report_markdown.split('\n'):
            line = line.rstrip()
            
            if not line:
                doc.add_paragraph()
                continue
            
            if line.startswith('# '):
                heading = doc.add_heading(line[2:], level=1)
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('#### '):
                doc.add_heading(line[5:], level=4)
            elif line.startswith('- ') or line.startswith('* '):
                paragraph = doc.add_paragraph(line[2:], style='List Bullet')
            elif line.startswith('| '):
                pass
            elif line.startswith('---'):
                pass
            elif line.startswith('*') and line.endswith('*'):
                paragraph = doc.add_paragraph(line[1:-1])
                for run in paragraph.runs:
                    run.italic = True
            else:
                formatted_line = self._parse_inline_formatting(line)
                paragraph = doc.add_paragraph()
                for text, bold in formatted_line:
                    run = paragraph.add_run(text)
                    run.bold = bold
        
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        content = buffer.getvalue()
        
        return content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    def _parse_inline_formatting(self, line: str) -> list:
        """解析行内格式（加粗等）"""
        result = []
        pattern = r'\*\*([^*]+)\*\*|\*([^*]+)\*'
        last_end = 0
        
        for match in re.finditer(pattern, line):
            if match.start() > last_end:
                result.append((line[last_end:match.start()], False))
            
            if match.group(1):
                result.append((match.group(1), True))
            elif match.group(2):
                result.append((match.group(2), False))
            
            last_end = match.end()
        
        if last_end < len(line):
            result.append((line[last_end:], False))
        
        if not result:
            result.append((line, False))
        
        return result

    def _generate_pdf(self, report_markdown: str, filename: str) -> Tuple[bytes, str]:
        """将Markdown转换为PDF"""
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=2*cm,
            leftMargin=2*cm,
            topMargin=2*cm,
            bottomMargin=2*cm
        )
        
        styles = getSampleStyleSheet()
        story = []
        
        for line in report_markdown.split('\n'):
            line = line.strip()
            if not line:
                story.append(Spacer(1, 6))
                continue
            
            if line.startswith('# '):
                story.append(Paragraph(line[2:], styles['Title']))
            elif line.startswith('## '):
                story.append(Paragraph(line[3:], styles['Heading2']))
            elif line.startswith('### '):
                story.append(Paragraph(line[4:], styles['Heading3']))
            elif line.startswith('---'):
                story.append(Spacer(1, 12))
            elif line.startswith('- '):
                story.append(Paragraph(f"• {line[2:]}", styles['Normal']))
            elif line.startswith('*') and not line.startswith('**'):
                story.append(Paragraph(f"<i>{line.strip('*')}</i>", styles['Normal']))
            else:
                story.append(Paragraph(line, styles['Normal']))
        
        doc.build(story)
        content = buffer.getvalue()
        return content, "application/pdf"
