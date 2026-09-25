"""Full pipeline acceptance for every new leaf, without network/model variance."""
from functools import lru_cache
import json
import pytest

from review_v2.engine import review_contract_v2
from review_v2.contract_registry import get_checklist, get_contract_meta, CONTRACT_TYPE_TREE
from review_v2.priority_contract_types import PRIORITY_TYPES
from review_v2.priority_checklist_helpers import PROFILES
from review_v2.clause_parser import parse_contract_clauses
from review_v2.domain_selection import resolve_domain_selection
from priority_contract_fixtures import NORMAL, MISSING, RISKY, CRITICAL_NODES, ADVERSE_CLAUSES


@lru_cache(maxsize=None)
def review(key, variant="normal"):
    return review_contract_v2({"normal": NORMAL, "missing": MISSING, "risky": RISKY}[variant][key], key, use_llm=False)


@pytest.mark.parametrize("key", PRIORITY_TYPES)
@pytest.mark.parametrize("variant", ("normal", "missing", "risky"))
def test_each_type_end_to_end(key, variant):
    result = review(key, variant)
    assert result["success"] is True
    assert result["contract_type"] == key
    assert result["contract_type_name"] == PRIORITY_TYPES[key].name
    assert result["report_context"]["contract_type_label"] == PRIORITY_TYPES[key].name
    assert result["result"].startswith("# " + PRIORITY_TYPES[key].name + "审查报告")
    assert result["release_quality_gate"]["status"] == "REPORT_READY", result["release_quality_gate"]["failures"]
    items = {row["node_id"]: row for row in result["items"]}
    assert items[CRITICAL_NODES[key]]["coverage_status"] == ("缺失" if variant == "missing" else "通过")
    if variant == "risky":
        adverse = items[ADVERSE_CLAUSES[key][1]]
        assert adverse["coverage_status"] == "风险条款"
        assert adverse["risk_level"] == "高"
        assert adverse["evidence_clauses"]
    for item in items.values():
        if item["node_id"] in PROFILES and item["coverage_status"] == "通过":
            assert item["evidence_clauses"]
            optional_ids = {e['element_id'] for e in PROFILES[item['node_id']]
                            if e.get('risk_relevant') is False}
            assert all(e["status"] == "COVERED" for e in item["coverage_elements"]
                       if e['element_id'] not in optional_ids)
        if item["coverage_status"] == "不适用":
            assert not item.get("needs_manual_review")
    findings = result.get("risk_findings", [])
    ids = [f["finding_id"] for f in findings]
    assert len(ids) == len(set(ids))
    themes = result["risk_themes"]
    fingerprints = [json.dumps(t["risk_fingerprint"], sort_keys=True) for t in themes]
    assert len(fingerprints) == len(set(fingerprints))
    assert sum(result["coverage_summary"].values()) == len(result["all_rule_statuses"])


@pytest.mark.parametrize("key", PRIORITY_TYPES)
def test_classification_overrides_incompatible_selection_and_tree(key):
    text = NORMAL[key]
    selection, domain, routed = resolve_domain_selection("sale", text, parse_contract_clauses(text))
    assert routed == key
    assert selection.resolution == "DOCUMENT_OVERRIDE"
    assert domain.contract_family == PRIORITY_TYPES[key].family
    assert key in CONTRACT_TYPE_TREE[domain.contract_family]
    assert get_contract_meta(PRIORITY_TYPES[key].name).response_key == key
    assert any(row.node_id in PROFILES for row in get_checklist(key))


@pytest.mark.parametrize("key,partial, node", [
    ("labor", "月工资为10000元，工资通过银行转账。", "employment_wage"),
    ("mortgage", "主合同编号：JK-1，债务人：远山有限公司，本金为100万元。", "security_main_debt"),
    ("movable_pledge", "签约后5日内将质物交付质权人占有。", "pledge_movable_delivery"),
    ("equity_transfer", "目标公司更新股东名册并签发出资证明。", "equity_closing_registration"),
    ("capital_increase", "新增注册资本20万元，计入资本公积80万元。", "equity_capital_increase"),
])
def test_partial_evidence_does_not_become_full_or_missing(key, partial, node):
    text = NORMAL[key].splitlines()[0] + "\n第一条 " + partial
    result = review_contract_v2(text, key, use_llm=False)
    item = next(i for i in result["items"] if i["node_id"] == node)
    assert item["coverage_status"] == "部分覆盖"
    states = {e["status"] for e in item["coverage_elements"]}
    assert {"COVERED", "MISSING"} <= states
    assert result["release_quality_gate"]["status"] == "REPORT_READY", result["release_quality_gate"]["failures"]


def test_ancillary_titles_do_not_override_existing_primary_transaction():
    for title, selected in [("借款合同", "loan"), ("买卖合同", "sale"), ("技术服务合同", "technology")]:
        text = title + "\n第一条 附件包括保密协议、股权质押合同和抵押合同，均用于辅助主合同履行。"
        _, _, routed = resolve_domain_selection(selected, text, parse_contract_clauses(text))
        assert routed == selected


@pytest.mark.parametrize("term, probation,expected", [("二个月", "一个月", True), ("六个月", "二个月", True), ("一年", "三个月", True), ("三年", "六个月", False), ("一年", "二个月", False)])
def test_probation_duration_limits(term, probation, expected):
    from review_v2.priority_contract_review import probation_risk
    assert probation_risk(f"劳动合同期限为{term}，试用期为{probation}。") is expected


def test_negative_and_protective_clauses_are_not_adverse_evidence():
    from review_v2.priority_contract_review import adverse_evidence
    cases = [
        ("security_forfeiture", "双方不得约定债务人不履行到期债务时质物直接归质权人所有。"),
        ("employment_illegal_waiver", "不得要求劳动者自愿放弃社会保险。"),
        ("equity_mandatory_rights_waiver", "不得约定股东永久放弃知情权。"),
    ]
    for node, clause in cases:
        assert not adverse_evidence(node, parse_contract_clauses("第一条 " + clause))


def test_optional_clauses_and_in_service_competition_are_not_missing():
    result = review("labor")
    for node in ("employment_training_service", "employment_noncompete_scope", "employment_noncompete_term"):
        assert next(i for i in result["items"] if i["node_id"] == node)["coverage_status"] == "不适用"
    result = review_contract_v2("竞业限制协议\n第一条 竞业限制仅适用于在职期间，劳动者实际接触商业秘密。", "non_compete", use_llm=False)
    assert next(i for i in result["items"] if i["node_id"] == "employment_noncompete_term")["coverage_status"] == "不适用"
    assert result["release_quality_gate"]["status"] == "REPORT_READY"
    result = review("non_compete", "missing")
    assert next(i for i in result["items"] if i["node_id"] == "employment_noncompete_term")["coverage_status"] == "缺失"


def test_repeated_contradictory_insurance_clause_is_one_theme():
    text = RISKY["labor"] + "\n第100条 " + ADVERSE_CLAUSES["labor"][0]
    result = review_contract_v2(text, "labor", use_llm=False)
    themes = [t for t in result["risk_themes"] if "EMP-13-01" in t["source_review_ids"]]
    assert len(themes) == 1
    assert result["release_quality_gate"]["status"] == "REPORT_READY"
    theme = themes[0]
    assert {"EMP-07-01", "EMP-13-01"} <= set(theme["source_review_ids"])
    assert "不缴纳社会保险" in result["result"]
    assert "依法按月缴纳社会保险" in result["result"]


def test_blank_values_and_bare_headings_do_not_pass():
    for source in ("工资标准、工资支付日期、支付方式。", "月工资为____元，每月__日支付，工资通过银行转账。"):
        result = review_contract_v2("劳动合同\n第一条 " + source, "labor", use_llm=False)
        wage = next(i for i in result["items"] if i["node_id"] == "employment_wage")
        assert wage["coverage_status"] in {"缺失", "部分覆盖"}
        assert result["release_quality_gate"]["status"] == "REPORT_READY"


@pytest.mark.parametrize("key", PRIORITY_TYPES)
def test_report_exports_use_the_same_production_narrative(key):
    import io
    from docx import Document
    from routers.review_v2 import ReviewV2PdfInput, _build_docx, _build_pdf, _strip_markdown_inline
    result = review(key, "risky")
    payload = ReviewV2PdfInput(**result)
    exported = Document(io.BytesIO(_build_docx(payload)))
    exported_text = "\n".join(p.text for p in exported.paragraphs)
    assert result["contract_type_name"] in exported_text
    assert ADVERSE_CLAUSES[key][0] in exported_text
    assert _build_pdf(payload).startswith(b"%PDF-")
    # Exporters must consume the production result rather than rebuild risks.
    for line in result["result"].splitlines():
        if line.startswith("- 风险说明："):
            assert _strip_markdown_inline(line[2:]) in exported_text


@pytest.mark.parametrize("key", ("labor", "mortgage", "movable_pledge", "equity_pledge", "rights_pledge", "equity_transfer", "capital_increase", "investment", "shareholders", "non_compete", "confidentiality", "employment_termination"))
def test_docx_upload_routes_new_types_through_the_existing_api(key):
    import asyncio
    import io
    from docx import Document
    from starlette.datastructures import UploadFile
    from routers.review_v2 import review_v2_file
    document = Document()
    for line in NORMAL[key].splitlines():
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    result = asyncio.run(review_v2_file(
        file=UploadFile(file=buffer, filename=key + ".docx"), contract_type=key,
        include_low_risk=True, use_llm=False, profile_json="{}", current_user={"id": 1}))
    assert result["success"]
    assert result["contract_type"] == key
    assert result["release_quality_gate"]["status"] == "REPORT_READY", result["release_quality_gate"]["failures"]


def test_specialized_evidence_cannot_be_overridden_by_model(monkeypatch):
    import review_v2.engine as engine
    def no_model_calls(*args, **kwargs):
        raise AssertionError("劳动专项清单的确定性结果不应发送给模型改写")
    monkeypatch.setattr(engine, "judge_checkpoints_with_llm", no_model_calls)
    result = engine.review_contract_v2(RISKY["labor"], "labor", use_llm=True)
    assert result["release_quality_gate"]["status"] == "REPORT_READY"
    assert next(i for i in result["items"] if i["node_id"] == "employment_illegal_waiver")["coverage_status"] == "风险条款"


@pytest.mark.parametrize("key,role,canonical", [
    ("labor", "劳动者", "WORKER"), ("labor", "用人单位", "LABOR_EMPLOYER"),
    ("mortgage", "抵押人", "SECURITY_PROVIDER"), ("equity_pledge", "质权人", "SECURED_CREDITOR"),
    ("equity_transfer", "受让方", "EQUITY_TRANSFEREE"), ("investment", "投资方", "INVESTOR"),
])
def test_new_perspectives_do_not_reuse_construction_or_patent_roles(key, role, canonical):
    from review_v2.perspective_resolution import resolve_perspective
    result = resolve_perspective(role, PRIORITY_TYPES[key].family, contract_subtype=key.upper())
    assert result.canonical_role == canonical


def test_specialized_nodes_do_not_change_legacy_checkpoints():
    from review_v2.contract_registry import CONTRACT_CHECKLISTS
    from review_v2.priority_checklist_helpers import ADVERSE, common_rows
    legacy = {row.node_id for key, rows in CONTRACT_CHECKLISTS.items() if key not in PRIORITY_TYPES for row in rows}
    assert not legacy.intersection(set(PROFILES) | set(ADVERSE))
    sale_rows = get_checklist("sale")
    assert common_rows()
    assert all(any(row is original for original in sale_rows) for row in common_rows())


def test_quality_gate_distinguishes_labor_contract_law_from_repealed_contract_law():
    from review_v2.quality4 import post_render_quality_gate
    for text, outdated in [("依据《劳动合同法》第47条。", False), ("依据《中华人民共和国合同法》第60条。", True), ("《劳动合同法》及《合同法》。", True)]:
        result = post_render_quality_gate(text, [], [], [])
        assert any(f["requirement"] == "Legal Freshness Finding" for f in result["failures"]) is outdated


def test_new_domain_gate_rejects_unrelated_checklist_rows():
    from review_v2.domain_selection import domain_selection_consistency_gate
    text = NORMAL["mortgage"]
    selection, _, _ = resolve_domain_selection("mortgage", text, parse_contract_clauses(text))
    valid = [row.item_id for row in get_checklist("mortgage")]
    assert domain_selection_consistency_gate(selection, valid)["status"] == "REPORT_READY"
    assert domain_selection_consistency_gate(selection, valid + ["GUA-01-01"])["status"] == "BLOCKED"


def test_arbitrary_dismissal_is_distinct_from_social_insurance_waiver():
    result = review_contract_v2(NORMAL["labor"] + "\n第99条 用人单位有权随时辞退劳动者，无需支付经济补偿。", "labor", use_llm=False)
    items = {i["node_id"]: i for i in result["items"]}
    assert items["employment_arbitrary_exit"]["coverage_status"] == "风险条款"
    assert items["employment_illegal_waiver"]["coverage_status"] == "通过"
    assert result["release_quality_gate"]["status"] == "REPORT_READY"
