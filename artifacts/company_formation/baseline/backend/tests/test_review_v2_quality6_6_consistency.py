from pathlib import Path

from docx import Document

from review_v2.engine import (
    _aggregate_review_provenance,
    _production_internal_token_gate,
    review_contract_v2,
)


ROOT = Path(__file__).resolve().parents[2]


def _base(extra: str) -> str:
    return f"""住宅房地产买卖合同
    出卖人（甲方）：甲某；买受人（乙方）：乙某。
    不动产权证号为苏（2025）宁字不动产权第999号，登记权利人为甲某，单独所有。
    房地产坐落南京市某区某路1号，建筑面积88平方米，土地用途为住宅用地。
    双方到房地产交易中心办理转移登记，交房时办理钥匙交接。
    {extra}
    """


def _item(output, review_id):
    return next(item for item in output["items"] if item["item_id"] == review_id)


def _finding_for(output, review_id):
    return [finding for finding in output["risk_findings"] if review_id in finding["source_review_ids"]]


def _docx_text(pattern: str) -> str:
    document = Document(next((ROOT / "test_data").glob(pattern)))
    rows = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            value = " | ".join(cell.text.strip() for cell in row.cells)
            if value.strip(" |"):
                rows.append(value)
    return "\n".join(rows)


def test_no_trigger_is_not_missing_clause_and_na_is_not_missing():
    output = review_contract_v2(_base("甲乙双方均本人签署。代理人：无。"), "sale", use_llm=False)
    item = _item(output, "RE-01-06")
    assert item["applicability"]["status"] == "NOT_APPLICABLE"
    assert item["coverage_status"] == "不适用"
    assert _finding_for(output, "RE-01-06") == []


def test_agency_trigger_mutations_are_fact_driven():
    incomplete = review_contract_v2(_base("甲方委托张某代为签署合同及办理房地产过户。"), "sale", use_llm=False)
    item = _item(incomplete, "RE-01-06")
    assert item["applicability"]["status"] == "APPLICABLE"
    assert item["coverage_status"] in {"部分覆盖", "缺失"}

    complete = review_contract_v2(_base("代理人张某，授权委托书作为附件，明确授权其代为签署合同、收款并办理房地产过户，授权期限至2027年1月1日。"), "sale", use_llm=False)
    assert _item(complete, "RE-01-06")["coverage_status"] == "通过"

    broker = review_contract_v2(_base("某房地产经纪公司仅提供居间服务。甲乙双方均本人签署，代理人：无。"), "sale", use_llm=False)
    assert _item(broker, "RE-01-06")["applicability"]["status"] == "NOT_APPLICABLE"


def test_composite_encumbrance_review_consumes_shared_claims():
    complete = review_contract_v2(_base("除已披露抵押外，卖方保证无司法查封、无未披露共有、无租赁、无产权争议及无其他影响正常转让的第三方权利。抵押权人及余额已披露并约定注销。"), "sale", use_llm=False)
    item = _item(complete, "RE-02-04")
    assert item["coverage_status"] == "通过"
    assert len(item["coverage_elements"]) == 6
    assert all(element["status"] == "COVERED" for element in item["coverage_elements"])
    claim_ids = complete["contract_facts"]["canonical_claims"]["claims"]
    assert "property.mortgage.disclosed" in claim_ids
    assert "seller.encumbrance_warranty.exists" in claim_ids

    mortgage_only = review_contract_v2(_base("房屋存在抵押，抵押权人为某银行。"), "sale", use_llm=False)
    assert _item(mortgage_only, "RE-02-04")["coverage_status"] == "部分覆盖"

    absent_text = _base("双方按约办理交易。").replace("，单独所有", "")
    absent = review_contract_v2(absent_text, "sale", use_llm=False)
    assert _item(absent, "RE-02-04")["coverage_status"] == "缺失"


def test_domain_classification_is_report_context_single_source_of_truth():
    real_estate = review_contract_v2(_base("代理人：无。无抵押、无查封、无租赁、无产权争议及无其他第三方权利。"), "sale", use_llm=False)
    assert real_estate["report_context"]["scenario_label"] == "二手住宅房地产买卖"
    assert real_estate["profile"]["contract_scenario"] == "二手住宅房地产买卖"
    assert real_estate["observability"]["report_context_debug"]["consistency"] == "PASS"

    goods = review_contract_v2("商品买卖合同。产品名称A，规格型号X，数量10件，交货地点南京。", "sale", use_llm=False)
    assert goods["report_context"]["scenario_label"] == "商品买卖"

    lease = review_contract_v2("房屋租赁合同。出租人甲方，承租人乙方，住宅用于居住。月租金3000元。", "house_lease", use_llm=False)
    assert lease["report_context"]["scenario_label"] == "住宅租赁"


def test_hb003_cross_layer_gates_and_false_finding_regression():
    path = next((ROOT / "test_data").glob("6*.docx"))
    document = Document(path)
    rows = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            value = " | ".join(cell.text.strip() for cell in row.cells)
            if value.strip(" |"):
                rows.append(value)
    output = review_contract_v2("\n".join(rows), "sale", use_llm=False)
    agency = _item(output, "RE-01-06")
    rights = _item(output, "RE-02-04")
    assert agency["applicability"]["status"] == "NOT_APPLICABLE"
    assert not _finding_for(output, "RE-01-06")
    assert rights["coverage_status"] != "缺失"
    gates = output["observability"]["cross_review_consistency_debug"]["gate_results"]
    assert all(value == "PASS" for value in gates.values())
    assert output["release_quality_gate"]["status"] == "REPORT_READY"


def test_shared_claim_pass_keeps_element_and_review_provenance_visible():
    output = review_contract_v2(
        _base("除已披露抵押外，卖方保证无司法查封、无未披露共有、无租赁、无产权争议及无其他影响正常转让的第三方权利。抵押权人及余额已披露并约定注销。"),
        "sale",
        use_llm=False,
    )
    item = _item(output, "RE-02-04")
    assert item["coverage_status"] == "通过"
    assert all(element["claim_id"] for element in item["coverage_elements"])
    assert all(element["source_clause_ids"] for element in item["coverage_elements"])
    assert item["review_evidence_provenance"]["evidence_ids"]
    assert item["review_evidence_provenance"]["primary_locations"]
    assert item["evidence_location"] != "未召回正式条款"
    assert output["release_quality_gate"]["gate_results"]["PASS_EVIDENCE_PROVENANCE_GATE"] == "PASS"


def test_generic_composite_provenance_union_is_review_agnostic():
    elements = [
        {
            "element_id": "generic_a", "status": "COVERED", "claim_id": "generic.a",
            "evidence_ids": ["EV-1"], "source_clause_ids": ["C1"],
            "source_spans": ["A条款"], "source_locations": ["第一条"],
        },
        {
            "element_id": "generic_b", "status": "COVERED", "claim_id": "generic.b",
            "evidence_ids": ["EV-2", "EV-1"], "source_clause_ids": ["C2", "C1"],
            "source_spans": ["B条款", "A条款"], "source_locations": ["第二条", "第一条"],
        },
    ]
    provenance = _aggregate_review_provenance(elements)
    assert provenance == {
        "claim_ids": ["generic.a", "generic.b"],
        "evidence_ids": ["EV-1", "EV-2"],
        "clause_ids": ["C1", "C2"],
        "source_spans": ["A条款", "B条款"],
        "primary_locations": ["第一条"],
        "supporting_locations": ["第二条"],
    }


def test_payment_partial_mutations_follow_element_truth():
    base = "商品买卖合同。甲方出卖A产品10件，总价10000元，乙方购买。"
    cases = {
        "account": ("收款账户：62220001，开户行：中国银行。", "部分覆盖", "收款账户或开户行信息", "货款支付方式"),
        "method": ("货款支付方式：银行转账。", "部分覆盖", "货款支付方式", "收款账户或开户行信息"),
        "both": ("货款支付方式：银行转账，收款账户：62220001，开户行：中国银行。", "通过", "", ""),
        "none": ("", "缺失", "", ""),
    }
    for extra, expected_status, covered_label, missing_label in cases.values():
        output = review_contract_v2(base + extra, "sale", use_llm=False)
        item = _item(output, "S-05-03")
        assert item["coverage_status"] == expected_status
        if expected_status != "部分覆盖":
            continue
        state = item["narrative_state"]
        assert covered_label in state["judgment"]
        assert missing_label in state["judgment"]
        assert covered_label not in state["issue"]
        assert covered_label not in state["consequence"]
        assert covered_label not in state["recommendation"]
        assert missing_label in state["issue"]
        assert missing_label in state["consequence"]
        assert state["recommendation"] == ""
        display_state = item["decision_trace"]["recommendation_display_state"]
        assert display_state["active_recommendation_ids"]
        assert display_state["has_active_recommendation"] is True
        assert output["release_quality_gate"]["gate_results"]["PARTIAL_ELEMENT_NARRATIVE_GATE"] == "PASS"


def test_partial_narrative_gate_is_generic_across_review_families():
    cases = [
        (_docx_text("2*.docx"), "sale", "S-03-03"),
        (_docx_text("2*.docx"), "sale", "S-06-01"),
        (_docx_text("1*.docx"), "house_lease", "HL-035"),
    ]
    checked = 0
    for text, contract_type, review_id in cases:
        output = review_contract_v2(text, contract_type, use_llm=False)
        item = _item(output, review_id)
        if item["coverage_status"] != "部分覆盖":
            continue
        checked += 1
        state = item["narrative_state"]
        assert state["covered_elements"] or state["partial_elements"]
        assert state["issue"] and state["consequence"]
        assert state["recommendation"] == ""
        assert "recommendation_display_state" in item["decision_trace"]
        for covered in state["covered_elements"]:
            label = covered.replace("合同明确", "").replace("合同或附件提供", "")
            assert label not in state["issue"] or "其余关键要素" in state["issue"]
        assert output["release_quality_gate"]["gate_results"]["PARTIAL_ELEMENT_NARRATIVE_GATE"] == "PASS"
    assert checked >= 2


def test_production_report_uses_user_reason_and_token_gate_blocks_debug_keys():
    output = review_contract_v2(_base("甲乙双方均本人签署。代理人：无。"), "sale", use_llm=False)
    item = _item(output, "RE-01-06")
    assert "agent.exists" in item["applicability"]["machine_reason"]
    assert "agent.exists" not in item["applicability"]["user_reason"]
    assert "agent.exists" not in output["result"]
    assert output["release_quality_gate"]["gate_results"]["PRODUCTION_INTERNAL_TOKEN_GATE"] == "PASS"
    blocked = _production_internal_token_gate("debug: agent.exists")
    assert blocked["status"] == "BLOCKED"
    assert blocked["count"] == 1
