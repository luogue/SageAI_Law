import io
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from auth.routes import get_current_user
from main import app
from review_v2.clause_parser import parse_contract_clauses
from review_v2.engine import (
    _analyze_liability_clause_facts,
    _bind_structured_risk_evidence_lineage,
    _new_technology_shared_item,
)
from review_v2.technology_service_analysis import analyze_payment_schedule, analyze_timeline


ROOT = Path(__file__).resolve().parents[2]
GOLD = ROOT / "test_data" / "18 委托开发.docx"


def _docx_text(path: Path) -> str:
    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    tables = [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    return "\n".join([*paragraphs, *tables])


def _docx_bytes(text: str) -> bytes:
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _review_file(file_bytes: bytes) -> dict:
    app.dependency_overrides[get_current_user] = lambda: {"id": 1, "phone": "test"}
    try:
        with TestClient(app) as client:
            response = client.post(
                "/review-v2/file",
                files={"file": ("entrusted-development.docx", file_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                data={
                    "contract_type": "technology",
                    "include_low_risk": "true",
                    "use_llm": "false",
                    "profile_json": '{"user_identity":"甲方/委托方"}',
                },
            )
    finally:
        app.dependency_overrides.pop(get_current_user, None)
    assert response.status_code == 200, response.text
    return response.json()


def _payment_text(percentages, amounts, total=1_000_000):
    rows = [f"开发费用总价为人民币（￥{total:,.2f}元）。"]
    for index, (percentage, amount) in enumerate(zip(percentages, amounts), start=1):
        rows.extend((f"第{index}阶段验收后", f"支付合同价款的{percentage}%，即人民币{amount:,}元。"))
    return "\n".join(rows)


@pytest.mark.parametrize(
    ("percentages", "amounts", "count", "percentage_total", "status"),
    [
        ([25, 25, 25, 25], [250_000] * 4, 4, 100, "CONSISTENT"),
        ([20, 20, 20, 20, 20], [200_000] * 5, 5, 100, "CONSISTENT"),
        ([30, 30, 20, 20], [300_000, 300_000, 200_000, 200_000], 4, 100, "CONSISTENT"),
        ([30, 30, 25, 20], [300_000, 300_000, 250_000, 200_000], 4, 105, "INCONSISTENT"),
        ([25, 25, 25, 25], [250_000, 250_000, 250_000, 300_000], 4, 100, "INCONSISTENT"),
    ],
)
def test_payment_fact_identity_and_arithmetic_mutations(percentages, amounts, count, percentage_total, status):
    result = analyze_payment_schedule(_payment_text(percentages, amounts))
    assert result["payment_fact_count"] == count
    assert len(result["facts_before_dedup"]) == len(result["facts_after_dedup"]) == count
    assert result["percentage_total"] == percentage_total
    assert result["status"] == status


def test_timeline_mutations():
    normal = analyze_timeline("本合同签名或盖章后生效。签订时间：2026年9月1日\n项目开发进度\n开发开始：2026年9月5日")
    inverted = analyze_timeline("项目开发进度\n开发开始：2026年9月1日\n本合同签名或盖章后生效。签订时间：2026年9月20日")
    assert normal["performance_before_effective_date"] is False
    assert normal["conflicts"] == []
    assert inverted["effective_date"] == "2026-09-20"
    assert inverted["earliest_performance_date"] == "2026-09-01"
    assert inverted["performance_before_effective_date"] is True


@pytest.fixture(scope="module")
def gold_review():
    return _review_file(GOLD.read_bytes())


def test_gold_payment_trace_and_generic_routing(gold_review):
    payment = gold_review["deterministic_analysis"]["payment_formula"]
    assert payment["payment_fact_count"] == 4
    assert [node["percentage"] for node in payment["nodes"]] == [30, 30, 25, 20]
    assert [node["amount"] for node in payment["nodes"]] == [180000, 180000, 150000, 120000]
    assert payment["percentage_total"] == 105
    assert payment["amount_total"] == 630000
    assert payment["contract_total"] == 600000
    assert payment["amount_delta"] == 30000
    items = {item["item_id"]: item for item in gold_review["items"]}
    assert items["TC-04-02"]["coverage_status"] == "通过"
    assert items["TC-09-04"]["coverage_status"] == "不适用"
    assert items["GEN-PAY-01"]["risk_group"] == "payment_formula_inconsistency"
    assert gold_review["deterministic_analysis"]["payment_risk_lifecycle"]["status"] == "FINAL_RISK"


def test_gold_timeline_shared_facts_and_relationship_risks(gold_review):
    timeline = gold_review["deterministic_analysis"]["timeline"]
    assert timeline["effective_date"] == "2026-09-20"
    assert timeline["earliest_performance_date"] == "2026-09-01"
    assert timeline["performance_before_effective_date"] is True
    items = {item["item_id"]: item for item in gold_review["items"]}
    assert items["TC-01-03"]["coverage_status"] == items["TC-04-01"]["coverage_status"] == "通过"
    groups = {finding["risk_group"] for finding in gold_review["risk_findings"]}
    assert {
        "payment_formula_inconsistency",
        "technical_timeline_conflict",
        "generic_demand_reduction_exit_cost",
        "generic_cumulative_liability_overlap",
    } <= groups
    assert "personal_information_processing_gap" not in groups


def test_hierarchical_clause_identity_does_not_collapse_repeated_item_numbers():
    clauses = parse_contract_clauses(
        "2. 项目范围\n2.1. 第一个章节中的第1项。\n"
        "14. 违约责任\n14.1. 第二个章节中的第1项。"
    )
    first = next(clause for clause in clauses if "第一个章节" in clause.text)
    second = next(clause for clause in clauses if "第二个章节" in clause.text)
    assert first.location == "第2.1项"
    assert second.location == "第14.1项"
    assert first.clause_id != second.clause_id
    assert first.start_offset != second.start_offset


def test_gold_personal_data_feature_scope_is_not_entrusted_processing(gold_review):
    assert not any(item["item_id"] == "TECH-SHARED-02" for item in gold_review["items"])
    assert not any(
        finding["risk_group"] == "personal_information_processing_gap"
        for finding in gold_review["risk_findings"]
    )


def test_gold_timeline_evidence_binds_the_two_conflicting_dates(gold_review):
    item = next(item for item in gold_review["items"] if item["item_id"] == "TECH-SHARED-03")
    evidence = item["evidence_clauses"]
    combined = " ".join(str(row.get("quote") or row.get("text") or "") for row in evidence)
    assert "第一阶段" in combined
    assert "2026年9月1日至2026年9月15日" in combined
    assert "本合同经各方签名或盖章后生效" in combined
    assert "签订时间：2026年9月20日" in combined
    assert "2026-09-01 < 2026-09-20" in combined
    assert "背景知识产权" not in combined
    assert "生效之日起5个工作日内" not in combined


def test_gold_demand_reduction_evidence_stays_on_clause_5_5(gold_review):
    item = next(item for item in gold_review["items"] if item["item_id"] == "GEN-COMM-01")
    assert item["decision_trace"]["evidence_consistency_gate"]["status"] == "PASSED"
    assert {row["location"] for row in item["evidence_clauses"]} == {"第5.5项"}
    combined = " ".join(row["quote"] for row in item["evidence_clauses"])
    assert all(term in combined for term in ("补贴下降数额的30%", "不得低于原总开发费用的80%", "重新协商费用", "有权解除"))
    assert "逾期付款" not in combined


def test_gold_liability_evidence_preserves_multiple_sources(gold_review):
    item = next(item for item in gold_review["items"] if item["item_id"] == "GEN-LIAB-01")
    assert item["decision_trace"]["evidence_consistency_gate"]["status"] == "PASSED"
    assert {row["fact_type"] for row in item["evidence_clauses"]} == {
        "daily_fixed_penalty", "daily_percentage_penalty", "termination_penalty", "all_loss_and_enforcement_costs",
    }
    assert len({row["clause_id"] for row in item["evidence_clauses"]}) >= 3
    combined = " ".join(row["quote"] for row in item["evidence_clauses"])
    assert all(term in combined for term in ("壹万元", "1%", "20%", "全部损失"))
    assert {row["location"] for row in item["evidence_clauses"]} == {
        "第9.4项", "第14.2项", "第14.6项", "第14.7项",
    }
    assert all(row["liable_party"] == "乙方" for row in item["evidence_clauses"])
    assert all(row["beneficiary_party"] == "甲方" for row in item["evidence_clauses"])
    assert "项目基本信息" not in combined
    assert item["risk_severity"] == "LOW"
    assert item["action_priority"] == "C_OPTIMIZATION"
    assert item["decision_trace"]["severity_calibration"]["result"] == "LOW_OPTIONAL_OPTIMIZATION"


def test_internal_uncertainty_metadata_never_enters_user_report(gold_review):
    forbidden = (
        "现有证据不足", "当前证据不足", "证据不足以确认", "现有信息不足以确定",
        "无法确认", "无法可靠判断", "未能确认", "检索结果不足", "模型无法判断",
        "系统无法确定", "证据质量不足", "evidence insufficient", "cannot confirm",
        "unable to determine", "not enough evidence",
    )
    report = gold_review["result"]
    assert not any(marker.lower() in report.lower() for marker in forbidden)

    items = {item["item_id"]: item for item in gold_review["items"]}
    assert items["TC-13-03"]["decision_trace"]["internal_uncertainty_metadata"][
        "excluded_from_user_risk_output"
    ] is True
    assert items["TC-05-03"]["decision_trace"]["internal_uncertainty_metadata"][
        "excluded_from_user_risk_output"
    ] is True


def _liability_analysis(text: str):
    return _analyze_liability_clause_facts(parse_contract_clauses(text))


def test_liability_case_a_builds_party_a_cluster():
    result = _liability_analysis(
        "1. 甲方逾期付款的，每逾期一天，应按逾期金额的5‰向乙方支付违约金。\n"
        "2. 因甲方违约导致乙方解除合同的，甲方应向乙方支付合同总价款20%作为违约金。\n"
        "3. 任何一方有其他违反合同情形的，应赔偿守约方全部损失，包括可得利益、第三方赔偿、律师费用和诉讼费用。"
    )
    cluster = result["selected_cluster"]
    assert (cluster["liable_party"], cluster["beneficiary_party"]) == ("甲方", "乙方")
    assert {row["remedy_type"] for row in cluster["facts"]} == {
        "daily_percentage_penalty", "termination_penalty", "all_loss_and_enforcement_costs",
    }


def test_liability_case_b_builds_party_b_cluster():
    result = _liability_analysis(
        "1. 乙方整改总耗时超过15个工作日的，每超过1个工作日，应向甲方支付10000元违约金。\n"
        "2. 乙方逾期交付的，每逾期一天，应按开发费用的1%向甲方支付违约金。\n"
        "3. 因乙方违约导致甲方解除合同的，乙方应向甲方支付合同总价款20%作为违约金。\n"
        "4. 任何一方有其他违反合同情形的，应赔偿守约方全部损失，包括可得利益、第三方赔偿、律师费用和诉讼费用。"
    )
    cluster = result["selected_cluster"]
    assert (cluster["liable_party"], cluster["beneficiary_party"]) == ("乙方", "甲方")
    assert len(cluster["facts"]) == 4


def test_liability_case_c_same_percentage_does_not_merge_opposite_directions():
    result = _liability_analysis(
        "1. 因甲方违约导致乙方解除合同的，甲方应向乙方支付合同总价款20%作为违约金。\n"
        "2. 因乙方违约导致甲方解除合同的，乙方应向甲方支付合同总价款20%作为违约金。\n"
        "3. 任何一方有其他违反合同情形的，应赔偿守约方全部损失，包括可得利益、第三方赔偿、律师费用和诉讼费用。"
    )
    termination = [row for row in result["facts"] if row["remedy_type"] == "termination_penalty"]
    assert {(row["liable_party"], row["beneficiary_party"]) for row in termination} == {
        ("甲方", "乙方"), ("乙方", "甲方"),
    }
    assert result["clusters"] == []


def test_liability_case_d_different_parties_and_triggers_do_not_form_chain():
    result = _liability_analysis(
        "1. 甲方逾期付款的，每逾期一天，应按逾期金额的1%向乙方支付违约金。\n"
        "2. 因乙方违约导致甲方解除合同的，乙方应向甲方支付合同总价款20%作为违约金。\n"
        "3. 任何一方有其他违反合同情形的，应赔偿守约方全部损失，包括可得利益、第三方赔偿、律师费用和诉讼费用。"
    )
    assert result["present"] is False
    assert result["selected_cluster"] is None


def test_unsupported_synthetic_candidate_is_not_given_fallback_evidence():
    item = _new_technology_shared_item("GEN-TEST-01", "generic_test", "无来源测试风险")
    item.risk_group = "generic_demand_reduction_exit_cost"
    item.coverage_elements = [{"element_id": "generic_demand_reduction_exit_cost.adverse_arrangement", "status": "COVERED"}]
    _bind_structured_risk_evidence_lineage(item, {
        "entailment_status": "FAILED",
        "items": [],
        "required_fact_types": ["demand_reduction_trigger"],
        "covered_fact_types": [],
    })
    assert item.evidence_clauses == []
    assert item.needs_manual_review is True
    assert item.decision_trace["evidence_consistency_gate"]["status"] == "EVIDENCE_INSUFFICIENT"
    assert item.decision_trace["evidence_consistency_gate"]["wrong_evidence_fallback_allowed"] is False


def test_tc_13_01_does_not_expand_generic_technical_documentation(gold_review):
    item = next(item for item in gold_review["items"] if item["item_id"] == "TC-13-01")
    elements = {row["element_id"]: row["status"] for row in item["coverage_elements"]}
    assert item["coverage_status"] == "部分覆盖"
    assert elements["source_and_executable_delivery"] == "COVERED"
    assert elements["generic_technical_documentation"] == "COVERED"
    assert elements["test_report_delivery"] == "COVERED"
    assert elements["interface_architecture_deployment_documents"] == "MISSING"
    assert elements["defect_list_delivery"] == "MISSING"


def test_renderer_uses_lineage_evidence_instead_of_keyword_fallback(gold_review):
    report = gold_review["result"]

    def item_section(item_id: str) -> str:
        marker = f"（{item_id}）"
        start = report.index(marker)
        end = report.find("\n---", start)
        return report[start:end if end >= 0 else None]

    timeline = item_section("TECH-SHARED-03")
    demand = item_section("GEN-COMM-01")
    liability = item_section("GEN-LIAB-01")
    assert all(term in timeline for term in ("第一阶段", "2026年9月1日", "签订时间：2026年9月20日"))
    assert "背景知识产权" not in timeline
    assert all(term in demand for term in ("第5.5项", "补贴下降数额的30%", "不得低于原总开发费用的80%", "有权解除"))
    assert "逾期付款的，每逾期一天" not in demand
    assert all(term in liability for term in ("第9.4项", "第14.2项", "第14.6项", "第14.7项", "全部损失"))
    assert "第14.4项" not in liability
    assert "项目基本信息" not in liability
