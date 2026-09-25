"""审查引擎第二版 API。

该路由独立于裁判引擎和旧版 LLM 合同审查主流程。
"""

from __future__ import annotations

import html
import json
import hashlib
import inspect
import io
import logging
import os
import re
from datetime import datetime
from typing import Any, Dict, List, Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from pydantic import BaseModel, Field
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from auth.routes import get_current_user
from layers.input_layer import InputLayer
from review_v2.engine import (
    COVERAGE_NOT_APPLICABLE,
    EVALUATOR_VERSION,
    ISSUE_NOT_APPLICABLE,
    PIPELINE_VERSION,
    REPORT_FORMAT_NOTICE,
    RULE_VERSION,
    TECHNICAL_CONSULTING_EVIDENCE_IMPLEMENTATION,
    review_contract_v2,
)
from review_v2.contract_registry import get_checklist_for_resolution
from review_v2 import domain_packs as review_domain_packs_module
from review_v2 import engine as review_engine_module
from review_v2 import quality4 as review_quality4_module
from review_v2.element_resolution import conditional_element
from review_v2.clause_parser import strip_non_contract_tail
from review_v2.input_integrity import (
    build_input_integrity_diagnostics,
    input_integrity_user_message,
    log_input_integrity,
)
from review_v2.technology_service_analysis import (
    TECHNOLOGY_CLASSIFIER_VERSION,
    classify_technology_transaction,
    extract_claim_response_facts,
)
from review_v2.llm_semantic_reviewer import get_llm_review_status
from review_v2.loan_contract_checklist import get_loan_checklist
from review_v2.user_profile import UserRiskProfile
from review_v2.semantic_reviewer import (
    COVERAGE_MISSING,
    COVERAGE_PARTIAL,
    COVERAGE_PASS,
    COVERAGE_RISKY,
    ISSUE_CLAUSE_DEFECT,
    ISSUE_MISSING,
    ISSUE_PASS,
    ISSUE_RISKY,
    ISSUE_TEMPLATE_BLANK,
    has_blank_value,
)


router = APIRouter(prefix="/review-v2", tags=["contract_review_v2"])
input_layer = InputLayer()
logger = logging.getLogger(__name__)


def _runtime_source_identity(target: Any) -> Dict[str, Any]:
    path = inspect.getsourcefile(target) or ""
    payload = b""
    mtime = None
    if path and os.path.exists(path):
        with open(path, "rb") as source_file:
            payload = source_file.read()
        mtime = os.path.getmtime(path)
    return {
        "module": getattr(target, "__module__", ""),
        "function": getattr(target, "__qualname__", getattr(target, "__name__", "")),
        "file": path,
        "first_line": getattr(getattr(target, "__code__", None), "co_firstlineno", None),
        "mtime": mtime,
        "sha256": hashlib.sha256(payload).hexdigest() if payload else "",
    }


from review_v2.lease_semantics import apply as apply_lease_semantics
from review_v2.commission_agency_evidence import apply as apply_commission_agency_evidence
from review_v2.utility_supply_evidence import apply as apply_utility_supply_evidence, consistency_gate as utility_supply_consistency_gate
from review_v2.parameter_validation import extract_parameters
from review_v2.prerequisite_resolution import resolve_prerequisites
from review_v2.evidence_compatibility import evidence_compatible
from review_v2.utility_supply_applicability import extract_applicability_facts
from review_v2.utility_supply_contract_checklist import get_utility_supply_checklist
from review_v2.recommendation_children import child_set_gate as lease_child_set_gate
from review_v2 import employment_contract_checklist, employment_facts, employment_exit, employment_rights, priority_contract_review
from review_v2 import noncompete_semantics, noncompete_reasonableness
from review_v2 import nda_semantics
from review_v2 import primary_purpose
from review_v2 import agency_facts
from review_v2 import semantic_relations, semantic_consistency
from review_v2 import penalty_base
from review_v2 import equity_investment_checklist, equity_transaction_facts
from review_v2 import company_formation_checklist, transaction_scope, priority_contract_types
from review_v2 import termination_checklist, template_parser
from review_v2 import brokerage_facts, brokerage_contract_checklist, brokerage_concepts
from review_v2 import gift_evidence, gift_contract_checklist
from review_v2 import service_scenarios, service_contract_checklist
from review_v2 import guarantee_contract_facts, guarantee_party_facts, guarantee_prerequisites, guarantee_mode, guarantee_scope, guarantee_contract_checklist
from review_v2 import guarantee_scenarios, guarantee_closure
from review_v2 import narrative_composition
from review_v2.final_snapshot import SnapshotStore, SnapshotError, FIELDS as SNAPSHOT_FIELDS

_RUNTIME_SOURCE_TARGETS = {
    'company_formation_checklist': company_formation_checklist,
    'transaction_scope': transaction_scope,
    'priority_contract_types': priority_contract_types,
    'narrative_composition': narrative_composition,
    'guarantee_scenarios': guarantee_scenarios,
    'guarantee_closure': guarantee_closure,
    'guarantee_facts': guarantee_contract_facts,
    'guarantee_parties': guarantee_party_facts,
    'guarantee_prerequisites': guarantee_prerequisites,
    'guarantee_mode': guarantee_mode,
    'guarantee_scope': guarantee_scope,
    'guarantee_checklist': guarantee_contract_checklist,
    'service_scenarios': service_scenarios,
    'service_checklist': service_contract_checklist,
    'gift_evidence': gift_evidence,
    'gift_checklist': gift_contract_checklist,
    'brokerage_facts': brokerage_facts,
    'brokerage_concepts': brokerage_concepts,
    'brokerage_checklist': brokerage_contract_checklist,
    'employment_checklist': employment_contract_checklist,
    'employment_facts': employment_facts,
    'employment_exit': employment_exit,
    'employment_rights': employment_rights,
    'noncompete_semantics': noncompete_semantics,
    'noncompete_reasonableness': noncompete_reasonableness,
    'nda_semantics': nda_semantics,
    'primary_purpose': primary_purpose,
    'agency_facts': agency_facts,
    'semantic_relations': semantic_relations,
    'semantic_consistency': semantic_consistency,
    'penalty_base': penalty_base,
    'equity_checklist': equity_investment_checklist,
    'equity_transaction_facts': equity_transaction_facts,
    'termination_checklist': termination_checklist,
    'template_parser': template_parser,
    'priority_contract_review': priority_contract_review,
    "commission_agency_evidence": apply_commission_agency_evidence,
    "utility_supply_evidence": apply_utility_supply_evidence,
    "utility_supply_consistency_gate": utility_supply_consistency_gate,
    "parameter_validation": extract_parameters,
    "prerequisite_resolution": resolve_prerequisites,
    "evidence_compatibility": evidence_compatible,
    "utility_supply_applicability": extract_applicability_facts,
    "utility_supply_checklist": get_utility_supply_checklist,
    "lease_semantics": apply_lease_semantics,
    "lease_recommendation_children": lease_child_set_gate,
    "lease_fact_graph": review_engine_module.extract_house_lease_facts,
    "clause_parser": review_engine_module.parse_contract_clauses,
    "domain_selection": review_engine_module.resolve_domain_selection,
    "source_span_evidence": review_engine_module.extract_sale_evidence,
    "work_evidence": review_engine_module.apply_work_evidence,
    "finance_lease_facts": review_engine_module.FinanceLeaseFacts.__init__,
    "loan_costs": review_engine_module.analyze_costs,
    "loan_facts": review_engine_module.LoanFacts.from_clauses,
    "loan_checklist": get_loan_checklist,
    "engine": review_contract_v2,
    "domain_packs": review_domain_packs_module,
    "quality4": review_quality4_module,
    "classifier": classify_technology_transaction,
    "registry": get_checklist_for_resolution,
    "document_parser": InputLayer.extract_text_with_diagnostics,
    "normalization": strip_non_contract_tail,
    "conditional_elements": conditional_element,
    "claim_evidence": extract_claim_response_facts,
    "claim_projection": review_engine_module._apply_exact_claim_response,
    "license_evaluation": review_engine_module._apply_technology_transaction_evidence_consistency,
    "secret_proof_evaluation": review_engine_module._technical_secret_transfer_proof_results,
    "element_search": review_engine_module._resolve_final_element_searches,
    "recommendations": review_engine_module._finalize_element_recommendation_sources,
    "liability_scope_consistency": review_engine_module.liability_scope_consistency_failures,
    "requirement_key_normalization": review_engine_module.normalize_requirement_supports,
    "development_exit_elements": review_engine_module.development_exit_elements,
    "entrusted_patent_application_elements": review_engine_module.entrusted_patent_application_elements,
    "development_failure_elements": review_engine_module.development_failure_elements,
    "confidentiality_detail_elements": review_engine_module.confidentiality_detail_elements,
    "ancillary_service_detail_elements": review_engine_module.ancillary_service_detail_elements,
    "report_state_consistency_failures": review_engine_module.report_state_consistency_failures,
}
from review_v2 import security_property_checklist, security_context, security_semantic_aliases, requirement_kinds, shared_fact_concepts, safekeeping_contract_checklist, penalty_base
_RUNTIME_SOURCE_TARGETS.update({
    'shared_fact_concepts': shared_fact_concepts,
    'safekeeping_contract_checklist': safekeeping_contract_checklist,
    'penalty_base': penalty_base,
    'security_property_checklist': security_property_checklist,
    'security_context': security_context,
    'security_semantic_aliases': security_semantic_aliases,
    'requirement_kinds': requirement_kinds,
})
from review_v2 import company_governance_facts, contract_references, equity_common_checklist, adverse_theme_projection, report_statistics, shareholder_transaction_facts, formation_rule_semantics
_RUNTIME_SOURCE_TARGETS.update({m.__name__: m for m in (company_governance_facts, contract_references, equity_common_checklist, adverse_theme_projection, report_statistics, shareholder_transaction_facts, formation_rule_semantics)})
from review_v2 import legal_basis_selection, legal_basis_trace, legal_enrichment_boundary, verified_source, adverse_evidence, construction_evidence, construction_coverage, construction_applicability, construction_subtype
from review_v2 import adverse_consistency, monetary_consequences, contract_effects, semantic_reviewer
_RUNTIME_SOURCE_TARGETS.update(adverse_consistency=adverse_consistency, monetary_consequences=monetary_consequences,
    contract_effects=contract_effects, semantic_reviewer=semantic_reviewer)
_RUNTIME_SOURCE_TARGETS.update(legal_basis_selection=legal_basis_selection, legal_basis_trace=legal_basis_trace,
                               legal_enrichment_boundary=legal_enrichment_boundary)
_RUNTIME_SOURCE_TARGETS.update(verified_source=verified_source, adverse_evidence=adverse_evidence,
    construction_evidence=construction_evidence, construction_coverage=construction_coverage)
_RUNTIME_SOURCE_TARGETS.update(construction_applicability=construction_applicability)
_RUNTIME_SOURCE_TARGETS.update(construction_subtype=construction_subtype)
from review_v2 import prohibition_protection
_RUNTIME_SOURCE_TARGETS.update(prohibition_protection=prohibition_protection)
from review_v2 import employment_fact_schema, employment_fact_store, employment_fact_rules, employment_fact_pipeline, deepseek_facts, fact_validation, fact_verification
from review_v2 import fact_routing, fact_module_registry, deepseek_transport, fact_number_normalization
_FACT_FRAMEWORK_SOURCES = (fact_routing, fact_module_registry, deepseek_transport, fact_number_normalization)
from review_v2.unified_facts import loan_pipeline, loan_rules, loan_bindings, loan_costs as verified_loan_costs, loan as loan_schema, pipeline as unified_pipeline, store as unified_store, normalization as unified_normalization, schema as unified_schema
_FACT_FRAMEWORK_SOURCES += (loan_pipeline, loan_rules, loan_bindings, verified_loan_costs, loan_schema,
    unified_pipeline, unified_store, unified_normalization, unified_schema)
_RUNTIME_SOURCE_TARGETS.update({m.__name__: m for m in _FACT_FRAMEWORK_SOURCES})
_RUNTIME_SOURCE_TARGETS.update({m.__name__: m for m in
    (employment_fact_schema, employment_fact_store, employment_fact_rules, employment_fact_pipeline, deepseek_facts, fact_validation, fact_verification)})
_LOADED_DEPENDENCY_SOURCES = {name: _runtime_source_identity(target)
                              for name, target in _RUNTIME_SOURCE_TARGETS.items()}

_RUNTIME_SOURCE_CORE_DEPENDENCIES = {
    'prohibition_protection',
    'adverse_consistency','monetary_consequences','contract_effects','semantic_reviewer',
    'legal_basis_selection', 'legal_basis_trace', 'legal_enrichment_boundary',
    'verified_source','adverse_evidence','construction_evidence','construction_coverage','construction_applicability','construction_subtype',
    'agency_facts',
    'semantic_relations', 'semantic_consistency',
    'security_property_checklist', 'security_context', 'security_semantic_aliases', 'requirement_kinds',
    'shared_fact_concepts', 'safekeeping_contract_checklist', 'penalty_base',
    'primary_purpose',
    'penalty_base',
    'equity_checklist', 'equity_transaction_facts',
    'termination_checklist', 'template_parser',
    'narrative_composition',
    'service_scenarios', 'service_checklist',
    'gift_evidence', 'gift_checklist',
    "clause_parser",
    "domain_selection",
    "domain_packs",
    "quality4",
    "engine",
    "classifier",
    "registry",
    "document_parser",
    "normalization",
    "conditional_elements",
    "element_search",
    "recommendations",
    "report_state_consistency_failures",
    "requirement_key_normalization",
}

_RUNTIME_SOURCE_DOMAIN_DEPENDENCIES = {
    'guarantee': {'guarantee_facts', 'guarantee_parties', 'guarantee_prerequisites', 'guarantee_mode', 'guarantee_scope', 'guarantee_checklist', 'lease_semantics', 'guarantee_scenarios', 'guarantee_closure'},
    'brokerage': {'brokerage_facts', 'brokerage_checklist', 'brokerage_concepts', 'lease_semantics'},
    "utility_supply": {"utility_supply_evidence", "utility_supply_consistency_gate", "parameter_validation",
                       "prerequisite_resolution", "evidence_compatibility", "utility_supply_applicability", "utility_supply_checklist"},
    "commission_agency": {"commission_agency_evidence"},
    "house_lease": {
        "lease_semantics",
        "lease_recommendation_children",
        "lease_fact_graph",
    },
    "lease": {
        "lease_semantics",
        "lease_recommendation_children",
        "lease_fact_graph",
    },
    "sale": {"source_span_evidence"},
    "work_contract": {"work_evidence"},
    "finance_lease": {"finance_lease_facts"},
    "loan": {"loan_costs", "loan_facts", "loan_checklist"},
    "technology": {
        "claim_evidence",
        "claim_projection",
        "license_evaluation",
        "secret_proof_evaluation",
        "liability_scope_consistency",
        "development_exit_elements",
        "entrusted_patent_application_elements",
        "development_failure_elements",
        "confidentiality_detail_elements",
        "ancillary_service_detail_elements",
    },
    "labor": {"lease_semantics", "employment_checklist", "employment_facts", "employment_exit", "employment_rights", "priority_contract_review", "noncompete_semantics", "noncompete_reasonableness"},
    "non_compete": {"lease_semantics", "employment_checklist", "employment_facts", "employment_rights", "priority_contract_review", "noncompete_semantics", "noncompete_reasonableness"},
    "confidentiality": {"lease_semantics", "nda_semantics", "employment_checklist", "employment_facts", "priority_contract_review", "noncompete_semantics", "noncompete_reasonableness"},
    "employment_termination": {"lease_semantics"},
}


def _runtime_source_dependency_scope(contract_type: str = "") -> set[str]:
    requested = (contract_type or "").strip()
    if not requested:
        return set(_RUNTIME_SOURCE_TARGETS)
    scope = set(_RUNTIME_SOURCE_CORE_DEPENDENCIES)
    scope.update(m.__name__ for m in _FACT_FRAMEWORK_SOURCES)
    from review_v2.contract_registry import get_contract_meta, UnsupportedContractTypeError
    try:
        requested = get_contract_meta(requested).response_key
    except UnsupportedContractTypeError:
        pass
    scope.update(_RUNTIME_SOURCE_DOMAIN_DEPENDENCIES.get(requested, set()))
    if requested in {'labor', 'non_compete', 'confidentiality'}:
        scope.update(m.__name__ for m in (employment_fact_schema, employment_fact_store,
                     employment_fact_rules, employment_fact_pipeline, deepseek_facts, fact_validation, fact_verification))
    return scope


def _runtime_source_drift_status(contract_type: str = "") -> Dict[str, Any]:
    """Fail closed when the running worker has older code than the workspace."""
    current = {name: _runtime_source_identity(target) for name, target in _RUNTIME_SOURCE_TARGETS.items()}
    scoped_dependencies = _runtime_source_dependency_scope(contract_type)
    changed = [
        name for name, identity in current.items()
        if name in scoped_dependencies
        if str(identity.get("sha256") or "")
        != str((_LOADED_DEPENDENCY_SOURCES.get(name) or {}).get("sha256") or "")
    ]
    return {
        "status": "STALE_RUNTIME" if changed else "CURRENT",
        "changed_dependencies": changed,
        "checked_dependencies": sorted(scoped_dependencies),
        "loaded": _LOADED_DEPENDENCY_SOURCES,
        "current": current,
    }


class ReviewV2Input(BaseModel):
    review_task_id: Optional[str] = None
    contract_text: str = Field(..., min_length=1, description="合同文本")
    contract_type: str = Field("house_lease", description="合同类型；沿用已有类型，并支持 labor、non_compete、confidentiality、employment_termination、mortgage、movable_pledge、equity_pledge、rights_pledge、equity_transfer、capital_increase、investment、shareholders。")
    include_low_risk: bool = Field(True, description="是否返回低风险/已覆盖检查点")
    use_llm: bool = Field(True, description="是否启用大模型语义复核")
    profile: Dict[str, str] = Field(default_factory=dict, description="用户风险画像：{user_identity, contract_scenario, risk_appetite, bargaining_position}")


class ReviewV2Result(BaseModel):
    review_task_id: str = ""
    review_run_id: str = ""
    fact_snapshot_id: str = ""
    rule_result_version: str = ""
    report_version: str = ""
    gate_version: str = ""
    snapshot_hash: str = ""
    export_version: Optional[str] = None
    created_at: str = ""
    updated_at: str = ""
    final_review_snapshot: Dict[str, Any] = Field(default_factory=dict)
    contract_type_flow: Dict[str, Any] = Field(default_factory=dict)
    fact_migration: Dict[str, Any] = Field(default_factory=dict)
    structured_fact_debug: Dict[str, Any] = Field(default_factory=dict)
    employment_rule_adaptation: Dict[str, Any] = Field(default_factory=dict)
    employment_classification: Dict[str, Any] = Field(default_factory=dict)
    success: bool
    contract_type: str
    selected_contract_type: str = ""
    contract_type_name: str
    contract_review_trace_id: str = ""
    message: str
    items: List[Dict[str, Any]] = Field(default_factory=list)
    all_items_count: int = 0
    summary: Dict[str, int] = Field(default_factory=dict)
    risk_finding_summary: Dict[str, int] = Field(default_factory=dict)
    final_overview_summary: Dict[str, int] = Field(default_factory=dict)
    final_disposition_summary: Dict[str, int] = Field(default_factory=dict)
    coverage_summary: Dict[str, int] = Field(default_factory=dict)
    all_rule_statuses: List[Dict[str, Any]] = Field(default_factory=list)
    equity_corpus_consistency_gate: Dict[str, Any] = Field(default_factory=dict)
    statistics_consistency_gate: Dict[str, Any] = Field(default_factory=dict)
    recommendation_children_gate: Dict[str, Any] = Field(default_factory=dict)
    technical_quality_gate: Dict[str, Any] = Field(default_factory=dict)
    release_hold: Dict[str, Any] = Field(default_factory=dict)
    issue_type_summary: Dict[str, int] = Field(default_factory=dict)
    judge_source_summary: Dict[str, int] = Field(default_factory=dict)
    llm_enabled: bool = True
    llm_status: Dict[str, Any] = Field(default_factory=dict)
    global_fallback_reason: str = ""
    report_notice: str = ""
    result: str = ""
    filename: str = ""
    extracted_text_length: int = 0
    clause_count: int = 0
    formal_clause_count: int = 0
    contract_facts: Dict[str, Any] = Field(default_factory=dict)
    loan_shared_fact_consistency_gate: Dict[str, Any] = Field(default_factory=dict)
    deterministic_analysis: Dict[str, Any] = Field(default_factory=dict)
    document_integrity: Dict[str, Any] = Field(default_factory=dict)
    input_integrity: Dict[str, Any] = Field(default_factory=dict)
    quality_gate: Dict[str, Any] = Field(default_factory=dict)
    review_results_4_0: List[Dict[str, Any]] = Field(default_factory=list)
    risk_findings: List[Dict[str, Any]] = Field(default_factory=list)
    risk_themes: List[Dict[str, Any]] = Field(default_factory=list)
    risk_theme_count: int = 0
    rule_risk_hit_count: Optional[int] = None
    risk_clause_count: Optional[int] = None
    aggregated_risk_count: int = 0
    finding_count: int = 0
    risk_theme_summary: Dict[str, int] = Field(default_factory=dict)
    risk_theme_severity_summary: Dict[str, int] = Field(default_factory=dict)
    recommendations: List[Dict[str, Any]] = Field(default_factory=list)
    final_dispositions: List[Dict[str, Any]] = Field(default_factory=list)
    final_disposition_invariant_gate: Dict[str, Any] = Field(default_factory=dict)
    manual_review_items: List[Dict[str, Any]] = Field(default_factory=list)
    manual_review_summary: Dict[str, Any] = Field(default_factory=dict)
    contract_context: Dict[str, Any] = Field(default_factory=dict)
    report_context: Dict[str, Any] = Field(default_factory=dict)
    perspective_resolution: Dict[str, Any] = Field(default_factory=dict)
    domain_selection_resolution: Dict[str, Any] = Field(default_factory=dict)
    domain_selection_consistency_gate: Dict[str, Any] = Field(default_factory=dict)
    overall_assessment: Dict[str, Any] = Field(default_factory=dict)
    legal_reference_issues: List[Dict[str, Any]] = Field(default_factory=list)
    legal_citation_claims: List[Dict[str, Any]] = Field(default_factory=list)
    legal_authority_validation: Dict[str, Any] = Field(default_factory=dict)
    benchmark_metrics: Dict[str, Any] = Field(default_factory=dict)
    release_quality_gate: Dict[str, Any] = Field(default_factory=dict)
    terminology_consistency_gate: Dict[str, Any] = Field(default_factory=dict)
    formal_recommendation_validation: Dict[str, Any] = Field(default_factory=dict)
    technology_transaction_consistency_gate: Dict[str, Any] = Field(default_factory=dict)
    fatal_report_generation_gate: Dict[str, Any] = Field(default_factory=dict)
    report_generation_decision: Dict[str, Any] = Field(default_factory=dict)
    quality4_gate: Dict[str, Any] = Field(default_factory=dict)
    architecture_status: List[Dict[str, Any]] = Field(default_factory=list)
    system_metadata: Dict[str, Any] = Field(default_factory=dict)
    observability: Dict[str, Any] = Field(default_factory=dict)
    production_path_trace: Dict[str, Any] = Field(default_factory=dict)
    profile: Dict[str, Any] = Field(default_factory=dict)
    litigation: Dict[str, Any] = Field(default_factory=dict)
    litigation_result: str = ""
    negotiation: Dict[str, Any] = Field(default_factory=dict)
    negotiation_result: str = ""
    jurisdiction: Dict[str, Any] = Field(default_factory=dict)
    jurisdiction_result: str = ""


class ReviewV2PdfInput(BaseModel):
    final_review_snapshot: Dict[str, Any] = Field(default_factory=dict)
    filename: str = "合同审查报告"
    contract_type_name: str = "房屋租赁合同"
    result: str = ""
    summary: Dict[str, int] = Field(default_factory=dict)
    risk_finding_summary: Dict[str, int] = Field(default_factory=dict)
    coverage_summary: Dict[str, int] = Field(default_factory=dict)
    issue_type_summary: Dict[str, int] = Field(default_factory=dict)
    judge_source_summary: Dict[str, int] = Field(default_factory=dict)
    llm_status: Dict[str, Any] = Field(default_factory=dict)
    global_fallback_reason: str = ""
    report_notice: str = ""
    risk_findings: List[Dict[str, Any]] = Field(default_factory=list)
    recommendations: List[Dict[str, Any]] = Field(default_factory=list)
    overall_assessment: Dict[str, Any] = Field(default_factory=dict)
    benchmark_metrics: Dict[str, Any] = Field(default_factory=dict)
    release_quality_gate: Dict[str, Any] = Field(default_factory=dict)
    quality4_gate: Dict[str, Any] = Field(default_factory=dict)
    system_metadata: Dict[str, Any] = Field(default_factory=dict)
    items: List[Dict[str, Any]] = Field(default_factory=list)


class ReviewV2LlmStatus(BaseModel):
    enabled_by_env: bool
    has_model_key: bool
    available: bool
    configured_models: List[str] = Field(default_factory=list)
    priority: str = ""
    message: str = ""


def _register_cjk_font() -> str:
    candidates = [
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
    ]
    for path in candidates:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont("SageCJK", path))
                return "SageCJK"
            except Exception:
                continue
    return "Helvetica"


def _p(text: Any) -> str:
    return html.escape(str(text or "")).replace("\n", "<br/>")


def _item_evidence_text(item: Dict[str, Any]) -> str:
    return "\n".join(
        str(evidence.get("text") or "")
        for evidence in (item.get("evidence_clauses") or [])
        if isinstance(evidence, dict)
    )


def _item_issue_type(item: Dict[str, Any]) -> str:
    issue_type = str(item.get("issue_type") or "").strip()
    if issue_type:
        return issue_type

    coverage_status = str(item.get("coverage_status") or "").strip()
    evidence_text = _item_evidence_text(item)
    if coverage_status in {COVERAGE_PASS, COVERAGE_PARTIAL} and has_blank_value(evidence_text):
        return ISSUE_TEMPLATE_BLANK
    if coverage_status == COVERAGE_PASS:
        return ISSUE_PASS
    if coverage_status == COVERAGE_RISKY:
        return ISSUE_RISKY
    if coverage_status == COVERAGE_MISSING:
        return ISSUE_MISSING
    return ISSUE_CLAUSE_DEFECT


def _item_legal_consequence(item: Dict[str, Any]) -> str:
    if item.get("coverage_status") == COVERAGE_NOT_APPLICABLE:
        applicability = item.get("applicability") if isinstance(item.get("applicability"), dict) else {}
        return str(applicability.get("user_reason") or applicability.get("reason") or "本项不适用于当前合同。")
    if item.get("coverage_status") == COVERAGE_PASS:
        related = item.get("related_risks") or []
        if related:
            return f"关联风险 {', '.join(map(str, related))}：同一条款或事实仍存在需要处理的经济或法律风险。"
        return "无明显残余风险"
    narrative = item.get("narrative_state") if isinstance(item.get("narrative_state"), dict) else {}
    return (
        str(narrative.get("consequence") or "").strip()
        if item.get("coverage_status") == COVERAGE_PARTIAL else ""
    ) or (
        str(item.get("legal_consequence_if_missing") or "").strip()
        or str(item.get("risk_point_description") or "").strip()
        or "未补全该审查点会增加合同履行和争议处理的不确定性，建议签署前补充明确约定。"
    )


def _item_production_narrative(item: Dict[str, Any], field: str, fallback: str = "") -> str:
    narrative = item.get("narrative_state") if isinstance(item.get("narrative_state"), dict) else {}
    if item.get("coverage_status") == COVERAGE_PARTIAL and narrative.get(field):
        return str(narrative[field])
    if item.get("coverage_status") == COVERAGE_NOT_APPLICABLE:
        applicability = item.get("applicability") if isinstance(item.get("applicability"), dict) else {}
        return str(applicability.get("user_reason") or applicability.get("reason") or fallback)
    return str(fallback or "")


def _item_legal_label(item: Dict[str, Any]) -> str:
    if item.get("coverage_status") == COVERAGE_NOT_APPLICABLE:
        return "适用说明"
    if item.get("coverage_status") == COVERAGE_PASS:
        if item.get("related_risks"):
            return "关联风险"
        return "残余风险"
    return "缺失的法律后果"


def _item_suggested_label(item: Dict[str, Any]) -> str:
    coverage = item.get("coverage_status")
    if coverage in {COVERAGE_NOT_APPLICABLE, COVERAGE_PASS}:
        return ""
    if coverage == COVERAGE_PARTIAL:
        return "建议补充"
    if coverage == COVERAGE_RISKY:
        return "建议修改原条款"
    return "建议新增"


def _item_citation(item: Dict[str, Any]) -> Dict[str, str]:
    citation = item.get("citation") if isinstance(item.get("citation"), dict) else {}
    return {
        "clause_number": str(citation.get("clause_number") or "未召回正式条款"),
        "original_excerpt": str(citation.get("original_excerpt") or "无正式条款证据"),
        "position_description": str(citation.get("position_description") or "合同正式条款及附件中未定位到可验证条款"),
    }


def _user_facing_text(value: Any) -> str:
    text = str(value or "")
    replacements = {
        "统一事实图谱提取": "根据合同原文可确认",
        "统一事实图谱识别": "根据合同原文可识别",
        "统一事实图谱显示": "根据合同原文显示",
        "统一事实图谱": "合同事实",
        "Fact Graph": "合同事实",
        "Scenario Engine": "金额测算",
        "Final Quality Gate": "最终校验",
        "Cross Review Consistency Verifier": "一致性校验",
        "rule_fallback": "自动审查",
        "LLM": "语义复核",
        "建议人工复核": "建议重点核对",
        "需要人工复核": "需要重点核对",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def _priority_label(priority: str) -> str:
    return {
        "A_CORE": "优先处理",
        "B_PRACTICAL": "建议处理",
        "C_OPTIMIZATION": "可选优化",
        "N/A": "不适用",
    }.get(priority or "", priority or "未标注")


def _finding_title_map(payload: "ReviewV2PdfInput") -> Dict[str, str]:
    return {finding.get("finding_id", ""): finding.get("title", "关联风险") for finding in (payload.risk_findings or [])}


def _related_risk_titles(item: Dict[str, Any], title_map: Dict[str, str]) -> str:
    ids = item.get("related_risks") or []
    if not ids:
        return "无"
    return "、".join(title_map.get(risk_id, "关联风险") for risk_id in ids)


def _issue_type_summary_from_items(items: List[Dict[str, Any]]) -> Dict[str, int]:
    summary = {
        ISSUE_PASS: 0,
        ISSUE_TEMPLATE_BLANK: 0,
        ISSUE_CLAUSE_DEFECT: 0,
        ISSUE_MISSING: 0,
        ISSUE_RISKY: 0,
        ISSUE_NOT_APPLICABLE: 0,
    }
    for item in items:
        issue_type = _item_issue_type(item)
        summary[issue_type] = summary.get(issue_type, 0) + 1
    return summary


def _production_markdown_or_raise(payload: ReviewV2PdfInput) -> str:
    markdown = (payload.result or "").strip()
    if not markdown:
        raise ValueError("正式导出必须使用生产 Markdown result，禁止由导出器重建业务报告。")
    return markdown


def _markdown_blocks(markdown: str) -> List[Dict[str, str]]:
    blocks: List[Dict[str, str]] = []
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line == "---":
            blocks.append({"type": "separator", "text": ""})
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            blocks.append({"type": f"h{len(heading.group(1))}", "text": heading.group(2)})
            continue
        if line.startswith("- "):
            blocks.append({"type": "bullet", "text": line[2:]})
            continue
        if line.startswith("> "):
            blocks.append({"type": "quote", "text": line[2:]})
            continue
        blocks.append({"type": "paragraph", "text": line})
    return blocks


def _strip_markdown_inline(text: str) -> str:
    cleaned = re.sub(r"\*\*(.*?)\*\*", r"\1", text or "")
    cleaned = re.sub(r"`([^`]*)`", r"\1", cleaned)
    return cleaned.strip()


def _build_pdf(payload: ReviewV2PdfInput) -> bytes:
    markdown = _production_markdown_or_raise(payload)
    font_name = _register_cjk_font()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=1.6 * cm,
        leftMargin=1.6 * cm,
        topMargin=1.6 * cm,
        bottomMargin=1.6 * cm,
    )
    styles = getSampleStyleSheet()
    base = ParagraphStyle("SageBase", parent=styles["Normal"], fontName=font_name, fontSize=9, leading=14, alignment=TA_LEFT)
    title = ParagraphStyle("SageTitle", parent=base, fontSize=18, leading=24, spaceAfter=12)
    heading = ParagraphStyle("SageHeading", parent=base, fontSize=12, leading=18, spaceBefore=10, spaceAfter=6)
    subheading = ParagraphStyle("SageSubHeading", parent=base, fontSize=10, leading=15, spaceBefore=6, spaceAfter=4)
    quote = ParagraphStyle("SageQuote", parent=base, textColor=colors.HexColor("#4b5563"), leftIndent=10, leading=14)
    story = []
    for block in _markdown_blocks(markdown):
        text = _p(_strip_markdown_inline(block["text"]))
        if block["type"] == "separator":
            story.append(Spacer(1, 6))
        elif block["type"] == "h1":
            story.append(Paragraph(text, title))
        elif block["type"] in {"h2", "h3"}:
            story.append(Paragraph(text, heading))
        elif block["type"].startswith("h"):
            story.append(Paragraph(text, subheading))
        elif block["type"] == "bullet":
            story.append(Paragraph(f"• {text}", base))
        elif block["type"] == "quote":
            story.append(Paragraph(text, quote))
        else:
            story.append(Paragraph(text, base))
    doc.build(story)
    return buffer.getvalue()


def _build_debug_pdf(payload: ReviewV2PdfInput) -> bytes:
    font_name = _register_cjk_font()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=1.6 * cm,
        leftMargin=1.6 * cm,
        topMargin=1.6 * cm,
        bottomMargin=1.6 * cm,
    )

    styles = getSampleStyleSheet()
    base = ParagraphStyle(
        "SageBase",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=9,
        leading=14,
        alignment=TA_LEFT,
    )
    title = ParagraphStyle(
        "SageTitle",
        parent=base,
        fontSize=18,
        leading=24,
        spaceAfter=12,
    )
    heading = ParagraphStyle(
        "SageHeading",
        parent=base,
        fontSize=12,
        leading=18,
        spaceBefore=10,
        spaceAfter=6,
    )

    summary = payload.summary or {}
    risk_finding_summary = payload.risk_finding_summary or {}
    coverage_summary = payload.coverage_summary or {}
    issue_type_summary = _issue_type_summary_from_items(payload.items)
    judge_source_summary = payload.judge_source_summary or {}
    finding_titles = _finding_title_map(payload)
    story = [
        Paragraph("合同审查报告（第二版）", title),
        Paragraph(f"合同类型：{_p(payload.contract_type_name)}", base),
        Paragraph(f"文件名称：{_p(payload.filename)}", base),
        Paragraph(f"总体可签署性：{_p((payload.overall_assessment or {}).get('signability') or '未记录')}", base),
        Paragraph(f"模型状态：{_p((payload.llm_status or {}).get('message') or '未记录')}", base),
        Paragraph(f"报告提示：{_p(payload.report_notice or REPORT_FORMAT_NOTICE)}", base),
        Paragraph(f"全局回退原因：{_p(payload.global_fallback_reason or '无')}", base),
        Paragraph(f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", base),
        Spacer(1, 10),
        Table(
            [
                ["重大风险", "中等风险", "低风险", "人工复核"],
                [
                    str(risk_finding_summary.get("重大风险", 0)),
                    str(risk_finding_summary.get("中等风险", 0)),
                    str(risk_finding_summary.get("低风险", 0)),
                    str(risk_finding_summary.get("人工复核", 0)),
                ],
            ],
            colWidths=[4 * cm, 4 * cm, 4 * cm, 4 * cm],
            style=TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), font_name),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f6fb")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d8dee9")),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ]
            ),
        ),
        Spacer(1, 10),
        Table(
            [
                ["高风险", "中风险", "低风险", "不适用", "审查项总数"],
                [
                    str(summary.get("高", 0)),
                    str(summary.get("中", 0)),
                    str(summary.get("低", 0)),
                    str(coverage_summary.get(COVERAGE_NOT_APPLICABLE, 0)),
                    str(len(payload.items)),
                ],
            ],
            colWidths=[3.2 * cm, 3.2 * cm, 3.2 * cm, 3.2 * cm, 3.2 * cm],
            style=TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), font_name),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f6fb")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d8dee9")),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ]
            ),
        ),
        Spacer(1, 12),
    ]

    if coverage_summary:
        story.extend(
            [
                Paragraph("覆盖状态统计", heading),
                Table(
                    [
                        [*coverage_summary.keys(), "总计"],
                        [*[str(value) for value in coverage_summary.values()], str(sum(coverage_summary.values()))],
                    ],
                    colWidths=[16 * cm / (len(coverage_summary) + 1)] * (len(coverage_summary) + 1),
                    style=TableStyle(
                        [
                            ("FONTNAME", (0, 0), (-1, -1), font_name),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f6fb")),
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d8dee9")),
                            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                            ("TOPPADDING", (0, 0), (-1, -1), 8),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                        ]
                    ),
                ),
                Spacer(1, 12),
            ]
        )

    if issue_type_summary:
        story.extend(
            [
                Paragraph("问题类型统计", heading),
                Table(
                    [
                        ["通过", "模板空白", "条款缺陷", "完全缺失", "风险条款", "不适用"],
                        [
                            str(issue_type_summary.get("通过", 0)),
                            str(issue_type_summary.get("模板空白", 0)),
                            str(issue_type_summary.get("条款缺陷", 0)),
                            str(issue_type_summary.get("完全缺失", 0)),
                            str(issue_type_summary.get("风险条款", 0)),
                            str(issue_type_summary.get(ISSUE_NOT_APPLICABLE, 0)),
                        ],
                    ],
                    colWidths=[2.65 * cm, 2.65 * cm, 2.65 * cm, 2.65 * cm, 2.65 * cm, 2.65 * cm],
                    style=TableStyle(
                        [
                            ("FONTNAME", (0, 0), (-1, -1), font_name),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f6fb")),
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d8dee9")),
                            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                            ("TOPPADDING", (0, 0), (-1, -1), 8),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                        ]
                    ),
                ),
                Spacer(1, 12),
            ]
        )

    for index, item in enumerate(payload.items, start=1):
        issue_type = _item_issue_type(item)
        legal_consequence = _item_legal_consequence(item)
        legal_label = _item_legal_label(item)
        suggested_label = _item_suggested_label(item)
        citation = _item_citation(item)
        reasoning = str(item.get("reasoning") or "").strip()
        story.extend(
            [
                Paragraph(f"{index}. {_p(item.get('title'))}（{_p(item.get('item_id'))}）", heading),
                Paragraph(
                    f"覆盖状态：{_p(item.get('coverage_status'))}　"
                    f"审查优先级：{_p(_priority_label(item.get('review_priority') or 'N/A'))}　"
                    f"关联风险发现：{_p(_related_risk_titles(item, finding_titles))}　"
                    f"问题类型：{_p(issue_type)}　"
                    f"复核状态：{_p(item.get('review_status') or '待处理')}",
                    base,
                ),
                Paragraph(
                    f"证据来源：{_p('语义复核' if item.get('judge_source') == 'llm' else '合同原文/结构化事实')}　"
                    f"模型置信度：{_p(item.get('llm_confidence') or '无')}",
                    base,
                ),
                Paragraph(f"原文位置：{_p(item.get('evidence_location') or '未召回正式条款')}", base),
                Paragraph(f"系统判断：{_p(_user_facing_text(_item_production_narrative(item, 'judgment', item.get('reasoning_summary'))))}", base),
                Paragraph(f"【风险点描述】：{_p(_user_facing_text(_item_production_narrative(item, 'issue', item.get('risk_point_description'))))}", base),
                Paragraph(f"【{_p(legal_label)}】：{_p(_user_facing_text(legal_consequence))}", base),
                Paragraph(f"【法律依据索引】：{_p(item.get('legal_basis_index'))}", base),
                Paragraph("命中原文：", base),
                *[
                    Paragraph(
                        f"{_p(evidence.get('location'))}：{_p(evidence.get('text'))}",
                        base,
                    )
                    for evidence in (item.get("evidence_clauses") or [])[:3]
                ],
                *(
                    [
                        Paragraph("【推理过程】", heading),
                        *[Paragraph(_p(line), base) for line in reasoning.splitlines()[1:] if line.strip()],
                        Paragraph("【证据定位】", heading),
                        Paragraph(f"条款编号：{_p(citation['clause_number'])}", base),
                        Paragraph(f"原文片段：{_p(citation['original_excerpt'])}", base),
                        Paragraph(f"位置描述：{_p(citation['position_description'])}", base),
                    ]
                    if reasoning
                    else []
                ),
                Paragraph(f"人工备注：{_p(item.get('reviewer_note'))}", base),
                Spacer(1, 8),
            ]
        )
        if suggested_label:
            suggested_value = _item_production_narrative(item, "recommendation", item.get("suggested_clause_example"))
            story.insert(-1, Paragraph(f"【{_p(suggested_label)}】：{_p(suggested_value)}", base))

    doc.build(story)
    return buffer.getvalue()


def _set_docx_run_font(run, *, bold: bool = False, size: int = 10) -> None:
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    try:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    except Exception:
        pass


def _add_docx_line(doc: DocxDocument, label: str, value: Any = "", *, bold_label: bool = True) -> None:
    paragraph = doc.add_paragraph()
    label_run = paragraph.add_run(label)
    _set_docx_run_font(label_run, bold=bold_label)
    if value not in (None, ""):
        value_run = paragraph.add_run(str(value))
        _set_docx_run_font(value_run)


def _add_docx_table(doc: DocxDocument, headers: List[str], values: List[Any]) -> None:
    table = doc.add_table(rows=2, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
        table.cell(1, index).text = str(values[index])


def _build_debug_docx(payload: ReviewV2PdfInput) -> bytes:
    doc = DocxDocument()
    finding_titles = _finding_title_map(payload)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10)

    title = doc.add_heading("合同审查报告（第二版）", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _add_docx_line(doc, "合同类型：", payload.contract_type_name)
    _add_docx_line(doc, "文件名称：", payload.filename)
    _add_docx_line(doc, "总体可签署性：", (payload.overall_assessment or {}).get("signability") or "未记录")
    _add_docx_line(doc, "模型状态：", (payload.llm_status or {}).get("message") or "未记录")
    _add_docx_line(doc, "报告提示：", payload.report_notice or REPORT_FORMAT_NOTICE)
    _add_docx_line(doc, "全局回退原因：", payload.global_fallback_reason or "无")
    _add_docx_line(doc, "导出时间：", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

    summary = payload.summary or {}
    risk_finding_summary = payload.risk_finding_summary or {}
    coverage_summary = payload.coverage_summary or {}
    doc.add_heading("风险发现统计", level=2)
    _add_docx_table(
        doc,
        ["重大风险", "中等风险", "低风险", "人工复核"],
        [
            risk_finding_summary.get("重大风险", 0),
            risk_finding_summary.get("中等风险", 0),
            risk_finding_summary.get("低风险", 0),
            risk_finding_summary.get("人工复核", 0),
        ],
    )
    if coverage_summary:
        doc.add_heading("覆盖状态统计", level=2)
        _add_docx_table(
            doc,
            [*coverage_summary.keys(), "总计"],
            [*coverage_summary.values(), sum(coverage_summary.values())],
        )

    issue_type_summary = _issue_type_summary_from_items(payload.items)
    if issue_type_summary:
        doc.add_heading("问题类型统计", level=2)
        _add_docx_table(
            doc,
            ["通过", "模板空白", "条款缺陷", "完全缺失", "风险条款", "不适用"],
            [
                issue_type_summary.get("通过", 0),
                issue_type_summary.get("模板空白", 0),
                issue_type_summary.get("条款缺陷", 0),
                issue_type_summary.get("完全缺失", 0),
                issue_type_summary.get("风险条款", 0),
                issue_type_summary.get(ISSUE_NOT_APPLICABLE, 0),
            ],
        )

    for index, item in enumerate(payload.items, start=1):
        doc.add_heading(f"{index}. {item.get('title')}（{item.get('item_id')}）", level=2)
        issue_type = _item_issue_type(item)
        legal_consequence = _item_legal_consequence(item)
        legal_label = _item_legal_label(item)
        suggested_label = _item_suggested_label(item)
        citation = _item_citation(item)
        reasoning = str(item.get("reasoning") or "").strip()
        source = "语义复核" if item.get("judge_source") == "llm" else "合同原文/结构化事实"

        _add_docx_line(
            doc,
            "基本信息：",
            f"覆盖状态：{item.get('coverage_status') or ''}；审查优先级：{_priority_label(item.get('review_priority') or 'N/A')}；"
            f"关联风险发现：{_related_risk_titles(item, finding_titles)}；"
            f"问题类型：{issue_type}；复核状态：{item.get('review_status') or '待处理'}",
        )
        _add_docx_line(doc, "证据来源：", f"{source}；模型置信度：{item.get('llm_confidence') or '无'}")
        _add_docx_line(doc, "原文位置：", item.get("evidence_location") or "未召回正式条款")
        _add_docx_line(doc, "系统判断：", _user_facing_text(_item_production_narrative(item, "judgment", item.get("reasoning_summary"))))
        _add_docx_line(doc, "【风险点描述】：", _user_facing_text(_item_production_narrative(item, "issue", item.get("risk_point_description"))))
        _add_docx_line(doc, f"【{legal_label}】：", _user_facing_text(legal_consequence))
        if suggested_label:
            _add_docx_line(doc, f"【{suggested_label}】：", _item_production_narrative(item, "recommendation", item.get("suggested_clause_example")))
        _add_docx_line(doc, "【法律依据索引】：", item.get("legal_basis_index") or "")

        evidences = item.get("evidence_clauses") or []
        _add_docx_line(doc, "命中原文：")
        if evidences:
            for evidence in evidences[:3]:
                _add_docx_line(doc, "", f"{evidence.get('location') or ''}：{evidence.get('text') or ''}", bold_label=False)
        else:
            _add_docx_line(doc, "", "无正式条款证据", bold_label=False)

        if reasoning:
            _add_docx_line(doc, "【推理过程】：")
            for line in reasoning.splitlines()[1:]:
                if line.strip():
                    _add_docx_line(doc, "", line, bold_label=False)
            _add_docx_line(doc, "【证据定位】：")
            _add_docx_line(doc, "- 条款编号：", citation["clause_number"])
            _add_docx_line(doc, "- 原文片段：", citation["original_excerpt"])
            _add_docx_line(doc, "- 位置描述：", citation["position_description"])
        _add_docx_line(doc, "人工备注：", item.get("reviewer_note") or "")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _build_docx(payload: ReviewV2PdfInput) -> bytes:
    markdown = _production_markdown_or_raise(payload)
    doc = DocxDocument()
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10)
    for block in _markdown_blocks(markdown):
        text = _strip_markdown_inline(block["text"])
        if block["type"] == "separator":
            doc.add_paragraph("")
            continue
        if block["type"] == "h1":
            paragraph = doc.add_heading(text, level=1)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if block["type"] == "h2":
            doc.add_heading(text, level=2)
            continue
        if block["type"] == "h3":
            doc.add_heading(text, level=3)
            continue
        if block["type"].startswith("h"):
            doc.add_heading(text, level=4)
            continue
        paragraph = doc.add_paragraph()
        if block["type"] == "bullet":
            paragraph.style = "List Bullet"
        run = paragraph.add_run(text)
        _set_docx_run_font(run)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@router.get("/llm-status", response_model=ReviewV2LlmStatus, summary="合同审查第二版：大模型语义复核状态")
async def review_v2_llm_status(
    contract_type: str = '',
    current_user: dict = Depends(get_current_user),
):
    """返回审查第二版当前是否可以调用大模型语义复核。"""

    from review_v2.contract_registry import get_contract_meta, UnsupportedContractTypeError
    try:
        canonical = get_contract_meta(contract_type).key.value if contract_type else ''
    except UnsupportedContractTypeError:
        canonical = ''
    if canonical in {'labor', 'non_compete', 'confidentiality', 'loan'}:
        status = deepseek_facts.status()
        return {**status, 'enabled_by_env':True, 'has_model_key':status['available'], 'priority':'deepseek'}
    return get_llm_review_status()


def _review_with_fact_errors(*args, **kwargs):
    try:
        return review_contract_v2(*args, **kwargs)
    except deepseek_facts.FactExtractionError as exc:
        logger.error('[FACT_EXTRACTION_BLOCKED] provider=deepseek reason=%s', str(exc))
        first = next((d for d in exc.diagnostics if d['disposition']=='BLOCK'),exc.diagnostics[0])
        detail = {'code':'FACT_EXTRACTION_BLOCKED',
            'release_quality_gate':'BLOCKED', 'provider':'deepseek',
            'message':f"事实校验失败：{first['failed_stage']} / {first['field']} / {first['validation_rule']}；详见诊断。",
            'reason':str(exc), **first, 'diagnostics':exc.diagnostics, 'attempts':exc.attempts}
        from uuid import uuid4
        detail['trace_id'] = 'FACT-' + uuid4().hex
        for diagnostic in detail['diagnostics']:
            diagnostic.update(stage=diagnostic.get('failed_stage'), code=diagnostic.get('validation_rule'),
                fact_key=diagnostic.get('fact_key') or diagnostic.get('concept'),
                raw_value=diagnostic.get('raw_value', diagnostic.get('model_raw_value')),
                normalized_value=diagnostic.get('normalized_value'),
                allowed_values=diagnostic.get('allowed_values', diagnostic.get('expected')),
                clause=diagnostic.get('source_clause'),
                message='抽取值未通过事实类型、证据或语义校验；未生成报告，不能以默认值替代。',
                trace_id=detail['trace_id'])
        audit_dir = os.getenv('REVIEW_V2_RESULT_AUDIT_DIR')
        if audit_dir:
            from pathlib import Path
            from uuid import uuid4
            directory = Path(audit_dir)
            directory.mkdir(parents=True,exist_ok=True)
            (directory/f'fact-error-{uuid4().hex}.json').write_text(json.dumps(detail,ensure_ascii=False,indent=2),encoding='utf8')
        raise HTTPException(503, detail=detail) from exc


@router.get("/profile-options", summary="按合同类型提供风险画像选项")
async def review_v2_profile_options(current_user: dict = Depends(get_current_user)):
    from review_v2.profile_catalog import profile_catalog
    return profile_catalog()


@router.post("/litigation-analysis", summary="合同审查第二版：诉讼策略分析")
async def review_v2_litigation_analysis(
    payload: ReviewV2Input,
    current_user: dict = Depends(get_current_user),
):
    """仅对已标记高/中风险的条款执行诉讼策略分析（证据缺口+法院裁量+成本收益）。"""

    from review_v2.litigation_analyzer import analyze_litigation_strategy, format_litigation_markdown

    profile = UserRiskProfile.from_dict(payload.profile) if payload.profile else UserRiskProfile.default()

    # 先运行审查获取 items
    review_result = review_contract_v2(
        payload.contract_text,
        payload.contract_type,
        include_low_risk=payload.include_low_risk,
        use_llm=payload.use_llm,
        profile=profile,
    )

    if not review_result.get("success"):
        return {"success": False, "message": review_result.get("message", "审查失败"), "items": [], "summary": {}}

    profile = UserRiskProfile.from_dict(review_result["profile"])
    report = analyze_litigation_strategy(review_result["items"], profile)
    return {
        "success": True,
        "disclaimer": report.disclaimer,
        "items": [item.to_public_dict() for item in report.items],
        "summary": report.summary,
        "result": format_litigation_markdown(report),
        "profile": profile.to_dict(),
    }


@router.post("/negotiation-analysis", summary="合同审查第二版：商业谈判分析")
async def review_v2_negotiation_analysis(
    payload: ReviewV2Input,
    current_user: dict = Depends(get_current_user),
):
    """对已标记高/中风险的条款生成谈判筹码转换方案（谈判属性+对方意图+替代方案+交换筹码+话术）。"""

    from review_v2.negotiation_analyzer import analyze_negotiation_strategy, format_negotiation_markdown

    profile = UserRiskProfile.from_dict(payload.profile) if payload.profile else UserRiskProfile.default()

    review_result = review_contract_v2(
        payload.contract_text,
        payload.contract_type,
        include_low_risk=payload.include_low_risk,
        use_llm=payload.use_llm,
        profile=profile,
    )

    if not review_result.get("success"):
        return {"success": False, "message": review_result.get("message", "审查失败"), "items": [], "overall_strategy": ""}

    profile = UserRiskProfile.from_dict(review_result["profile"])
    report = analyze_negotiation_strategy(review_result["items"], profile)
    return {
        "success": True,
        "disclaimer": report.disclaimer,
        "items": [item.to_public_dict() for item in report.items],
        "must_fight": report.must_fight,
        "core_negotiate": report.core_negotiate,
        "flexible": report.flexible,
        "concession_order": report.concession_order,
        "overall_strategy": report.overall_strategy,
        "result": format_negotiation_markdown(report),
        "profile": profile.to_dict(),
    }


@router.post("/jurisdiction-prediction", summary="合同审查第二版：跨地域裁判预判")
async def review_v2_jurisdiction_prediction(
    payload: ReviewV2Input,
    current_user: dict = Depends(get_current_user),
):
    """结合用户所在地法院的司法实践，对高/中风险条款进行地域裁判预判。"""

    from review_v2.jurisdiction_analyzer import analyze_jurisdiction_prediction, format_jurisdiction_markdown

    profile = UserRiskProfile.from_dict(payload.profile) if payload.profile else UserRiskProfile.default()

    review_result = review_contract_v2(
        payload.contract_text,
        payload.contract_type,
        include_low_risk=payload.include_low_risk,
        use_llm=payload.use_llm,
        profile=profile,
    )

    if not review_result.get("success"):
        return {"success": False, "message": review_result.get("message", "审查失败"), "items": [], "environment_assessment": ""}

    profile = UserRiskProfile.from_dict(review_result["profile"])
    report = analyze_jurisdiction_prediction(review_result["items"], profile, payload.contract_text)
    return {
        "success": True,
        "disclaimer": report.disclaimer,
        "jurisdiction": report.jurisdiction,
        "items": [item.to_public_dict() for item in report.items],
        "environment_assessment": report.environment_assessment,
        "jurisdiction_strategy_summary": report.jurisdiction_strategy_summary,
        "result": format_jurisdiction_markdown(report),
        "profile": profile.to_dict(),
    }


@router.post("/checklist", response_model=ReviewV2Result, summary="合同审查第二版：清单式审查")
async def review_v2_checklist(
    payload: ReviewV2Input,
    current_user: dict = Depends(get_current_user),
):
    """按文本运行独立审查引擎，供调试和后续批量迁移使用。"""

    profile = UserRiskProfile.from_dict(payload.profile) if payload.profile else UserRiskProfile.default()

    store = SnapshotStore()
    run = _begin_snapshot(store, current_user, payload.review_task_id)
    result = _review_with_fact_errors(
        payload.contract_text,
        payload.contract_type,
        include_low_risk=payload.include_low_risk,
        use_llm=payload.use_llm,
        profile=profile,
    )
    return store.complete(run, result)


@router.post("/diagnostics/brokerage-recall")
async def diagnose_brokerage_runtime(file: UploadFile = File(...), current_user: dict = Depends(get_current_user)):
    """Authenticated in-worker evidence-only diagnostic; no report or LLM call."""
    import sys
    import uuid
    import json
    import subprocess
    from dataclasses import asdict
    from types import SimpleNamespace
    from review_v2.lease_semantics import supports, corpus_facts
    drift = _runtime_source_drift_status('brokerage')
    if drift['status'] != 'CURRENT':
        raise HTTPException(503, detail={'release_quality_gate': 'BLOCKED', 'runtime': drift})
    text, diagnostics = input_layer.extract_text_with_diagnostics(file)
    corpus = strip_non_contract_tail(text)
    clauses = review_engine_module.parse_contract_clauses(corpus)
    profiles = brokerage_facts.requirements(corpus)
    targets = {'BR-04-02', 'BR-09-01'}
    sentences = ['甲方有权要求乙方如实告知与订立合同有关的信息及情况',
        '乙方应当尽到必要、审慎的核实义务，向甲方提供与订立合同有关的完整、准确、真实的信息及情况',
        '甲方应当对乙方提供必要的协助', '甲方应当向乙方提供自身与订立合同有关的真实信息',
        '乙方故意隐瞒与订立合同有关的重要事实或者提供虚假情况', '不得请求支付服务费用', '乙方应当承担损失赔偿责任']
    rules = []
    for rule in brokerage_contract_checklist.get_brokerage_checklist():
        if rule.item_id not in targets:
            continue
        specs = profiles[rule.node_id]
        item = SimpleNamespace(item_id=rule.item_id, node_id=rule.node_id, decision_trace={'semantic_requirements': specs})
        audits = [review_engine_module._search_final_element(item, {**s, 'status': 'MISSING'}, clauses) for s in specs]
        facts = corpus_facts(clauses, specs)
        rules.append({'rule': asdict(rule), 'runtime_requirements': specs, 'audits': audits,
            'accepted_evidence': facts, 'element_status': {k: 'COVERED' if v else 'MISSING' for k,v in facts.items()},
            'status': 'PASS' if all(facts.values()) else 'PARTIAL' if any(facts.values()) else 'MISSING',
            'threshold': None, 'role_check': 'embedded in patterns',
            'score_note': 'Boolean supports(); audit score is post-decision, not a numeric acceptance threshold',
            'adjudicator': 'not invoked; evidence-only diagnostic'})
    sources = {f.__name__: _runtime_source_identity(f) for f in [supports, corpus_facts,
        review_engine_module._search_final_element, brokerage_facts.requirements,
        brokerage_contract_checklist.get_brokerage_checklist, input_layer.extract_text_with_diagnostics]}
    version = hashlib.sha256(json.dumps(drift['loaded'], sort_keys=True).encode()).hexdigest()
    run_id = str(uuid.uuid4())
    result = {'run_id': run_id, 'task_id': run_id, 'worker_pid': os.getpid(), 'python': sys.executable,
        'cwd': os.getcwd(), 'sys_path': sys.path, 'code_version': version, 'rule_version': RULE_VERSION,
        'git_commit': subprocess.run(['git','rev-parse','HEAD'], capture_output=True,text=True).stdout.strip(),
        'runtime': drift, 'sources': sources, 'contract_hash': hashlib.sha256(corpus.encode()).hexdigest(),
        'corpus_length': len(corpus), 'corpus': corpus, 'extraction': diagnostics,
        'substrings': [{'text': s, 'exists': s in corpus, 'index': corpus.find(s),
            'chunks': [{'chunk_id': c.clause_id, 'text': c.text} for c in clauses if s in c.text]} for s in sentences],
        'rules': rules, 'cache': 'No cache access; direct functions; LLM and report generation not invoked'}
    logger.warning('[BROKERAGE_RUNTIME_DIAGNOSTIC] run_id=%s worker_pid=%s code_version=%s rule_version=%s', run_id, os.getpid(), version, RULE_VERSION)
    return result


@router.post("/diagnostics/input-trace", summary="显式生产输入诊断（不生成报告）")
async def review_v2_input_trace(
    file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user),
):
    """Only an explicit diagnostic URL may stop before report generation.

    The legacy environment setting controls optional trace persistence, never
    the execution mode of a normal upload (nor the next request in the worker).
    """
    import threading
    trace = await diagnose_brokerage_runtime(file, current_user)
    trace['entrypoint'] = _runtime_source_identity(review_v2_input_trace)
    trace['thread_id'] = threading.get_ident()
    trace['report_id'] = None
    trace['report_generated'] = False
    trace['run_mode'] = 'diagnostic'
    diagnostic_output = os.environ.get("REVIEW_V2_ONCE_INPUT_TRACE", "")
    if diagnostic_output:
        import json
        import threading
        from pathlib import Path
        trace['rules'] = [r for r in trace['rules'] if r['rule']['item_id'] in {'BR-04-02', 'BR-09-01'}]
        trace['assertions'] = []
        for phrase in ['甲方应当对乙方提供必要的协助',
                       '故意隐瞒与订立合同有关的重要事实或者提供虚假情况',
                       '不得请求支付服务费用', '承担损失赔偿责任']:
            try:
                assert phrase in trace['corpus'], 'PHRASE_NOT_IN_PRODUCTION_CORPUS'
                assertion = 'PASS'
            except AssertionError:
                assertion = 'FAIL'
            trace['assertions'].append({'phrase': phrase, 'result': assertion,
                                        'index': trace['corpus'].find(phrase)})
        Path(diagnostic_output).write_text(json.dumps(trace, ensure_ascii=False, indent=2), encoding='utf-8')
    return trace


@router.post("/file", response_model=ReviewV2Result, summary="合同审查第二版：文件上传审查")
async def review_v2_file(
    file: UploadFile = File(...),
    contract_type: str = Form("house_lease"),
    include_low_risk: bool = Form(True),
    use_llm: bool = Form(True),
    profile_json: str = Form("{}"),
    diagnostic_only: bool = Form(False),
    review_task_id: Optional[str] = Form(None),
    current_user: dict = Depends(get_current_user),
):
    """Normal uploads always run the full pipeline, independently of diagnostic settings."""
    if diagnostic_only is True:
        from fastapi.encoders import jsonable_encoder
        from fastapi.responses import JSONResponse
        trace = await review_v2_input_trace(file, current_user)
        trace['entrypoint'] = _runtime_source_identity(review_v2_file)
        return JSONResponse(content=jsonable_encoder(trace))

    store = SnapshotStore()
    run = _begin_snapshot(store, current_user, review_task_id if isinstance(review_task_id, str) else None)
    source_drift = _runtime_source_drift_status(contract_type)
    if source_drift["status"] != "CURRENT":
        logger.critical("[STALE_REVIEW_RUNTIME] dependencies=%s process_id=%s",
                        source_drift["changed_dependencies"], os.getpid())
        raise HTTPException(status_code=503,
                            detail="审查后端已加载旧代码，请重启本地服务后重新审查。")

    filename = file.filename or ""
    if not filename.lower().endswith((".doc", ".docx", ".pdf")):
        raise HTTPException(status_code=400, detail="审查第二版当前支持 .doc、.docx 和 .pdf 文件")

    import json as _json
    try:
        profile_dict = _json.loads(profile_json) if profile_json else {}
    except _json.JSONDecodeError:
        profile_dict = {}
    profile = UserRiskProfile.from_dict(profile_dict)

    try:
        contract_text, source_input_diagnostics = input_layer.extract_text_with_diagnostics(file)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"合同文本读取失败：{exc}") from exc

    normalized_text = strip_non_contract_tail(contract_text)
    preflight_integrity = build_input_integrity_diagnostics(
        extracted_text=contract_text,
        normalized_text=normalized_text,
        classifier_input=normalized_text,
        clause_extraction_input=normalized_text,
        rule_engine_input=normalized_text,
        search_audit_input=normalized_text,
        source_diagnostics=source_input_diagnostics,
        downstream_used=False,
    )
    log_input_integrity(logger, preflight_integrity, filename=filename)
    if preflight_integrity["status"] == "BLOCKED":
        logger.error(
            "[REVIEW_INPUT_BLOCKED] filename=%s failures=%s",
            filename,
            preflight_integrity.get("failures") or [],
        )
        raise HTTPException(
            status_code=422,
            detail={
                "code": "CONTRACT_INPUT_INTEGRITY_FAILED",
                "message": input_integrity_user_message(preflight_integrity),
                "input_integrity": preflight_integrity,
            },
        )

    runtime_audit = None
    if os.getenv('REVIEW_V2_RUNTIME_AUDIT') == '1' and contract_type == 'rights_pledge':
        from review_v2.runtime_security_audit import snapshot
        runtime_audit = snapshot(contract_text, contract_type)
        logger.warning('[LOADED_SECURITY_RULE_AUDIT] %s', _json.dumps(runtime_audit, ensure_ascii=False))
    result = _review_with_fact_errors(
        contract_text,
        contract_type,
        include_low_risk=include_low_risk,
        use_llm=use_llm,
        profile=profile,
        source_input_diagnostics=source_input_diagnostics,
    )
    log_input_integrity(
        logger,
        result.get("input_integrity") or preflight_integrity,
        trace_id=str(result.get("contract_review_trace_id") or ""),
        filename=filename,
    )
    actual_pipeline = str((result.get("system_metadata") or {}).get("pipeline_version") or "")
    if actual_pipeline != PIPELINE_VERSION:
        logger.critical(
            "[PIPELINE_ASSERTION_FAILED] expected=%s actual=%s entrypoint=review_v2_file",
            PIPELINE_VERSION,
            actual_pipeline,
        )
        raise HTTPException(status_code=500, detail="技术合同审查未进入 canonical review-v2 pipeline")
    if runtime_audit is not None:
        runtime_audit['final_items'] = [i for i in result.get('items', []) if i.get('item_id') in {'RPL-01-01','RPL-02-01','SEC-03-01','S-00-03'}]
        runtime_audit['report_text'] = result['result']
        runtime_audit['release_quality_gate'] = result.get('release_quality_gate')
        runtime_audit['report_sha256'] = hashlib.sha256(result['result'].encode('utf-8')).hexdigest()
        from pathlib import Path
        audit_path = Path(os.getenv('REVIEW_V2_RUNTIME_AUDIT_DIR', '../artifacts/rights-runtime-audit'))
        audit_path.mkdir(parents=True, exist_ok=True)
        (audit_path / ('run-' + str(os.getpid()) + '-' + __import__('uuid').uuid4().hex + '.json')).write_text(_json.dumps(runtime_audit, ensure_ascii=False, default=str, indent=2), encoding='utf-8')
        logger.warning('[FINAL_SECURITY_ELEMENT_AUDIT] %s', _json.dumps([{ 'rule_id':i['item_id'],'elements':i.get('coverage_elements'),'coverage_status':i.get('coverage_status')} for i in runtime_audit['final_items']], ensure_ascii=False))
    result["filename"] = filename
    result["extracted_text_length"] = len(contract_text)
    result["production_path_trace"] = {
        "run_mode": "review",
        "entrypoint": "routers.review_v2.review_v2_file",
        "entrypoint_file": inspect.getsourcefile(review_v2_file) or __file__,
        "entrypoint_line": review_v2_file.__code__.co_firstlineno,
        "engine_function": "review_v2.engine.review_contract_v2",
        "engine_file": inspect.getsourcefile(review_contract_v2) or "",
        "engine_line": review_contract_v2.__code__.co_firstlineno,
        "process_id": os.getpid(),
        "working_directory": os.getcwd(),
        "evidence_implementation": TECHNICAL_CONSULTING_EVIDENCE_IMPLEMENTATION,
        "pipeline_id": PIPELINE_VERSION,
        "engine_version": (result.get("system_metadata") or {}).get("system_version"),
        "classifier_version": TECHNOLOGY_CLASSIFIER_VERSION,
        "evaluator_version": EVALUATOR_VERSION,
        "ruleset_version": RULE_VERSION,
        "runtime_sources": {
            "entrypoint": _runtime_source_identity(review_v2_file),
            **source_drift["current"],
        },
        "runtime_source_drift": source_drift,
        "document_hash": hashlib.sha256(contract_text.encode("utf-8")).hexdigest(),
        "cache": {"status": "MISS", "reason": "review-v2/file has no rule, evidence, or report result cache"},
        "input_visibility": {
            "normalized_text_length": len(normalized_text),
            "confidentiality_terms_present": any(term in contract_text for term in (
                "保密", "商业秘密", "技术秘密", "不得泄露", "不得披露", "持续有效",
            )),
        },
        "input_integrity": result.get("input_integrity") or preflight_integrity,
        "classification_lifecycle": ((result.get("observability") or {}).get("classification_lifecycle") or {}),
        "classification_runtime": ((result.get("observability") or {}).get("classification_runtime") or {}),
        "registry_runtime": ((result.get("observability") or {}).get("review_registry_debug") or {}),
        "payment_risk_lifecycle": ((result.get("observability") or {}).get("payment_risk_lifecycle") or {}),
    }
    payment_events = result["production_path_trace"]["payment_risk_lifecycle"].get("events") or []
    for event in payment_events:
        fact_clause_ids = [row.get("clause_id") for row in event.get("source_clauses") or []]
        logger.warning(
            "%s review_id=%s filename=%s candidate_id=%s source_rule=%s source_clause_ids=%s "
            "scenario=%s applicability=%s status=%s filter_reason=%s invocations=%s arithmetic_status=%s "
            "percentage_total=%s amount_total=%s merged=%s finding_id=%s final_present=%s rendered=%s",
            event.get("event"),
            result.get("contract_review_trace_id"),
            filename,
            event.get("candidate_id", ""),
            event.get("source_rule", ""),
            event.get("source_clause_ids") or fact_clause_ids,
            event.get("scenario", ""),
            event.get("applicability", {}),
            event.get("status", ""),
            event.get("filter_reason") or event.get("reason", ""),
            event.get("invocation_count", ""),
            (event.get("output") or {}).get("status", ""),
            event.get("structured_percentages") or (event.get("output") or {}).get("percentage_total", ""),
            event.get("structured_amounts") or (event.get("output") or {}).get("amount_total", ""),
            event.get("merged", ""),
            event.get("finding_id", ""),
            event.get("present", ""),
            event.get("rendered", ""),
        )
    if os.getenv('REVIEW_V2_RESULT_AUDIT_DIR'):
        from pathlib import Path
        audit_dir = Path(os.environ['REVIEW_V2_RESULT_AUDIT_DIR'])
        audit_dir.mkdir(parents=True, exist_ok=True)
        (audit_dir / ('result-' + str(os.getpid()) + '-' + __import__('uuid').uuid4().hex + '.json')).write_text(
            _json.dumps(result, ensure_ascii=False, default=str, indent=2), encoding='utf-8')
    logger.warning(
        "[TRACE_REVIEW_V2] pid=%s pipeline=%s classifier=%s evaluator=%s ruleset=%s "
        "entry=%s:%s engine=%s:%s subtype=%s classifier_executions=%s registry_id=%s "
        "evidence_impl=%s doc_hash=%s confidentiality=%s",
        os.getpid(),
        PIPELINE_VERSION,
        TECHNOLOGY_CLASSIFIER_VERSION,
        EVALUATOR_VERSION,
        RULE_VERSION,
        result["production_path_trace"]["entrypoint_file"],
        result["production_path_trace"]["entrypoint_line"],
        result["production_path_trace"]["engine_file"],
        result["production_path_trace"]["engine_line"],
        ((result["production_path_trace"].get("classification_lifecycle") or {}).get("renderer") or {}).get("contract_subtype"),
        (result["production_path_trace"].get("classification_runtime") or {}).get("actual_execution_count"),
        (result["production_path_trace"].get("registry_runtime") or {}).get("registry_object_id"),
        TECHNICAL_CONSULTING_EVIDENCE_IMPLEMENTATION,
        result["production_path_trace"]["document_hash"],
        result["production_path_trace"]["input_visibility"]["confidentiality_terms_present"],
    )
    logger.info(
        "[REVIEW_RELEASE_GATE] review_id=%s pid=%s runtime=%s status=%s failures=%s",
        result.get("contract_review_trace_id"), os.getpid(), source_drift["status"],
        (result.get("release_quality_gate") or {}).get("status"),
        [row.get("requirement") for row in (result.get("release_quality_gate") or {}).get("failures", [])],
    )
    return store.complete(run, result)


def _begin_snapshot(store, owner, task_id):
    try:
        return store.begin(owner, task_id)
    except SnapshotError as exc:
        raise HTTPException(status_code=409, detail=exc.detail) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail='review_task_id 必须是 UUID。') from exc


def _resolve_export(payload, owner):
    supplied = payload.final_review_snapshot
    result, binding = SnapshotStore().resolve_export(owner, supplied)
    # A supplied browser body must agree, but the renderer always reads the
    # immutable server copy. Browser annotations cannot rewrite reviewed facts.
    if payload.result and payload.result != result.get('result'):
        raise SnapshotError('导出正文与已通过质量门的报告不一致。',
                            expected=binding, actual=supplied)
    if payload.release_quality_gate and payload.release_quality_gate != result.get('release_quality_gate'):
        raise SnapshotError('导出请求中的 Gate 诊断不属于该最终报告。', expected=binding, actual=supplied)
    return ReviewV2PdfInput(**result), binding


def _export_headers(binding):
    return {'X-Review-Run-Id': binding['review_run_id'],
            'X-Report-Version': binding['report_version'],
            'X-Gate-Version': binding['gate_version'],
            'X-Snapshot-Hash': binding['snapshot_hash'],
            'X-Export-Version': binding['export_version'],
            'X-Export-Artifact-Hash': binding['export_artifact_hash'],
            'Access-Control-Expose-Headers': 'X-Review-Run-Id,X-Report-Version,X-Gate-Version,X-Snapshot-Hash,X-Export-Version,X-Export-Artifact-Hash'}


@router.post("/export-pdf", summary="合同审查第二版：导出 PDF 报告")
async def review_v2_export_pdf(
    payload: ReviewV2PdfInput,
    current_user: dict = Depends(get_current_user),
):
    """将当前结构化审查报告导出为 PDF。"""

    try:
        payload, binding = _resolve_export(payload, current_user)
        content = _build_pdf(payload)
        exported = SnapshotStore().record_export(current_user, binding, 'pdf', content)
    except SnapshotError as exc:
        raise HTTPException(status_code=409, detail=exc.detail) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    base_name = (payload.filename or "合同审查报告").rsplit(".", 1)[0]
    download_name = f"{base_name}_审查第二版报告.pdf"
    encoded_name = quote(download_name)
    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_name}", **_export_headers(exported)},
    )


@router.post("/export-docx", summary="合同审查第二版：导出 DOCX 报告")
async def review_v2_export_docx(
    payload: ReviewV2PdfInput,
    current_user: dict = Depends(get_current_user),
):
    """将当前结构化审查报告导出为 Word 文档。"""

    try:
        payload, binding = _resolve_export(payload, current_user)
        content = _build_docx(payload)
        exported = SnapshotStore().record_export(current_user, binding, 'docx', content)
    except SnapshotError as exc:
        raise HTTPException(status_code=409, detail=exc.detail) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    base_name = (payload.filename or "合同审查报告").rsplit(".", 1)[0]
    download_name = f"{base_name}_审查第二版报告.docx"
    encoded_name = quote(download_name)
    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_name}", **_export_headers(exported)},
    )
