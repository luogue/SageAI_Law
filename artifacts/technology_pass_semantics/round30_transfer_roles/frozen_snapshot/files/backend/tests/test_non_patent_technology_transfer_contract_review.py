from pathlib import Path

import pytest
from docx import Document

from review_v2.engine import _partial_element_narrative_label, review_contract_v2
from review_v2.technology_service_analysis import (
    analyze_payment_schedule,
    analyze_timeline,
    classify_technology_transaction,
)
from review_v2.user_profile import UserRiskProfile


ROOT = Path(__file__).resolve().parents[2]
FIXTURE = ROOT / "test_data" / "29 非专利技术转让合同.docx"


def _fixture_text() -> str:
    document = Document(str(FIXTURE))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            text = " | ".join(cell.text.strip() for cell in row.cells)
            if text.strip(" |"):
                parts.append(text)
    return "\n".join(parts)


@pytest.fixture(scope="module")
def transfer_text() -> str:
    return _fixture_text()


@pytest.fixture(scope="module")
def transfer_review(transfer_text: str) -> dict:
    return review_contract_v2(
        transfer_text,
        "technology",
        include_low_risk=True,
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/技术秘密受让方"),
    )


def test_case_a_technical_secret_use_grant_remains_a_license():
    text = """
技术秘密许可合同
甲方（被许可方）与乙方（许可方）订立本合同。
乙方仅授予甲方在中国境内使用控制算法、源代码和参数的普通许可，核心权利仍归乙方。
技术秘密许可使用费为100万元，甲方不得转让或再许可。
"""
    result = classify_technology_transaction(text)
    assert result["primary_subtype"] == "TECHNOLOGY_LICENSE"
    assert result["tertiary_subtype"] == "TECHNICAL_SECRET_LICENSE"


def test_case_b_non_patent_complete_rights_transfer_routes_to_transfer(transfer_text: str):
    result = classify_technology_transaction(transfer_text)
    assert result["primary_subtype"] == "TECHNOLOGY_TRANSFER"
    assert result["tertiary_subtype"] == "TECHNICAL_SECRET_TRANSFER"
    assert result["primary_contract_type"] == "technology_transfer_technical_secret"
    assert result["operative_transaction_type"] == "TECHNOLOGY_TRANSFER"


def test_case_c_patent_ownership_transfer_is_not_a_license():
    text = """
专利权转让合同
甲方（受让方）与乙方（转让方）约定，乙方将ZL202410000001.1号发明专利权全部转让给甲方，
完成专利权转让登记后专利权归甲方所有，转让价款为100万元。
"""
    result = classify_technology_transaction(text)
    assert result["primary_subtype"] == "TECHNOLOGY_TRANSFER"
    assert result["tertiary_subtype"] == "PATENT_RIGHT_TRANSFER"


def test_case_d_patent_implementation_grant_is_a_license():
    text = """
专利实施许可合同
甲方（被许可方）与乙方（许可方）约定，乙方授予甲方在中国境内实施ZL202410000001.1号专利的普通许可，
许可期限五年，许可使用费为100万元，专利权仍归乙方。
"""
    result = classify_technology_transaction(text)
    assert result["primary_subtype"] == "TECHNOLOGY_LICENSE"
    assert result["tertiary_subtype"] == "PATENT_LICENSE"


def test_case_e_transfer_title_but_only_use_grant_routes_to_license_and_records_conflict():
    text = """
非专利技术转让合同
甲方（被许可方）与乙方（许可方）约定，乙方仅授予甲方使用技术秘密的普通许可，
乙方继续保有全部核心权利并可向其他主体许可。许可使用费为100万元，许可期限五年。
"""
    result = classify_technology_transaction(text)
    assert result["primary_subtype"] == "TECHNOLOGY_LICENSE"
    assert result["contract_type_conflict"]["exists"] is True


def test_case_f_license_title_but_all_rights_move_routes_to_transfer_and_records_conflict():
    text = """
技术秘密许可合同
甲方（受让方）与乙方（转让方）约定，乙方向甲方进行完整技术转让；
技术转让价款支付完成后，甲方取得该技术秘密的全部相关权利，乙方不再保留实施、许可或处分权。
"""
    result = classify_technology_transaction(text)
    assert result["primary_subtype"] == "TECHNOLOGY_TRANSFER"
    assert result["tertiary_subtype"] == "TECHNICAL_SECRET_TRANSFER"
    assert result["contract_type_conflict"]["exists"] is True


def test_case_g_retained_transferor_rights_do_not_reroute_and_tc_06_is_active(
    transfer_review: dict,
):
    by_id = {row["item_id"]: row for row in transfer_review["items"]}
    assert by_id["TC-06-01"]["coverage_status"] == "风险条款"
    assert by_id["TC-06-01"]["risk_level"] == "高"
    assert "转让方保留权利冲突" in by_id["TC-06-01"]["title"]
    assert all(by_id[item_id]["coverage_status"] != "不适用" for item_id in (
        "TC-06-01", "TC-06-02", "TC-06-03", "TC-06-04",
    ))
    assert all(by_id[item_id]["coverage_status"] == "不适用" for item_id in (
        "TC-07-01", "TC-07-02", "TC-07-03", "TC-07-04",
    ))


def test_case_h_transfer_payment_uses_fixed_total_and_excludes_sales_royalty(
    transfer_text: str,
):
    payment = analyze_payment_schedule(transfer_text)
    assert payment["status"] == "INCONSISTENT"
    assert payment["contract_total"] == 3_000_000
    assert payment["percentage_total"] == 120
    assert payment["amount_total"] == 3_600_000
    assert payment["amount_delta"] == 600_000
    assert len(payment["nodes"]) == 3


def test_case_i_confidentiality_periods_are_compared_by_party(transfer_review: dict):
    item = next(row for row in transfer_review["items"] if row["item_id"] == "TC-11-02")
    elements = {row["element_id"]: row for row in item["coverage_elements"]}
    assert item["coverage_status"] == "风险条款"
    assert elements["mutual_confidentiality_duration_balance"]["status"] == "ADVERSE"
    assert elements["licensor_post_termination_data_use"]["status"] == "ADVERSE"
    assert "明确且对等" not in item["reasoning_summary"]


def test_case_j_transfer_start_before_signature_has_direct_date_evidence(
    transfer_text: str,
    transfer_review: dict,
):
    timeline = analyze_timeline(transfer_text)
    item = next(row for row in transfer_review["items"] if row["item_id"] == "TECH-SHARED-03")
    assert timeline["transfer_start_date"] == "2026-10-10"
    assert timeline["contract_sign_date"] == "2026-10-20"
    assert timeline["primary_conflict_subject"] == "TRANSFER_START"
    assert item["coverage_status"] == "风险条款"
    assert "2026-10-10" in str(item["decision_trace"]["evidence_bundle"])
    assert "2026-10-20" in str(item["decision_trace"]["evidence_bundle"])
    assert item["evidence_location"]


def test_case_k_contractual_recordal_arrangement_activates_tc_06_04_without_patent_assumption(
    transfer_review: dict,
):
    item = next(row for row in transfer_review["items"] if row["item_id"] == "TC-06-04")
    elements = {row["element_id"]: row for row in item["coverage_elements"]}
    assert item["coverage_status"] == "风险条款"
    assert elements["responsible_party"]["status"] == "MISSING"
    assert elements["filing_deadline"]["status"] == "MISSING"
    assert "不当然要求专利式登记" in item["applicable_if"]


def test_case_k1_deferred_recordal_fact_is_not_negated_by_search_audit(
    transfer_review: dict,
):
    item = next(row for row in transfer_review["items"] if row["item_id"] == "TC-06-04")
    elements = {row["element_id"]: row for row in item["coverage_elements"]}
    finding = next(
        row for row in transfer_review["risk_findings"]
        if "TC-06-04" in row["source_review_ids"]
    )
    absence = (finding.get("evidence_bundle") or {}).get("absence_evidence") or {}
    missing = absence.get("missing_elements") or []
    direct_text = " ".join(
        str(row.get("quoted_text") or row.get("text") or "")
        for row in (finding.get("evidence_bundle") or {}).get("items") or []
        if row.get("source_type") != "SEARCH_AUDIT"
    )

    assert elements["responsible_party"]["observed_fact"] == "办理主体留待后续协商。"
    assert "后续协商确定" in direct_text
    assert "明确办理主体" in missing
    assert "明确办理材料" in missing
    assert "明确办理期限" in missing
    assert "明确未完成相关手续的责任" in missing
    assert all("留待后续协商" not in value for value in missing)


def test_case_k2_explicit_recordal_party_and_deadline_are_covered():
    text = """
非专利技术转让合同
甲方（受让方）与乙方（转让方）订立本合同。
乙方向甲方完整转让控制算法、参数及调试方法等技术秘密，转让价款支付后全部技术权利归甲方。
乙方应在合同生效后10日内负责办理登记备案并提交申请材料。
转让价款为100万元。
"""
    result = review_contract_v2(text, "technology", include_low_risk=True, use_llm=False)
    item = next(row for row in result["items"] if row["item_id"] == "TC-06-04")

    assert item["coverage_status"] == "通过"
    assert "留待后续协商" not in str(item.get("decision_trace") or {})


def test_case_k3_absent_recordal_arrangement_remains_missing():
    text = """
专利权转让合同
甲方（受让方）与乙方（转让方）订立本合同。
乙方将ZL202410000001.1号发明专利权全部转让给甲方。
转让价款为100万元。
"""
    result = review_contract_v2(text, "technology", include_low_risk=True, use_llm=False)
    item = next(row for row in result["items"] if row["item_id"] == "TC-06-04")

    assert item["coverage_status"] == "缺失"


def test_case_l_disclaimed_transferor_warranties_are_major_adverse_not_pass(
    transfer_review: dict,
):
    item = next(row for row in transfer_review["items"] if row["item_id"] == "TC-06-03")
    elements = {row["element_id"]: row for row in item["coverage_elements"]}
    assert item["coverage_status"] == "风险条款"
    assert item["risk_level"] == "高"
    assert elements["technology_origin_warranty"]["status"] == "ADVERSE"
    assert elements["third_party_rights_warranty"]["status"] == "ADVERSE"
    assert elements["non_infringement_warranty"]["status"] == "ADVERSE"


def test_current_full_review_is_report_ready_and_uses_non_patent_transfer_language(
    transfer_review: dict,
):
    report = transfer_review["result"]
    by_id = {row["item_id"]: row for row in transfer_review["items"]}
    assert transfer_review["release_quality_gate"]["status"] == "REPORT_READY"
    assert "甲方 / 技术秘密受让方" in report
    assert "技术合同—技术转让合同—技术秘密转让合同" in report
    assert "专利权转让还是专利实施许可" not in report
    assert by_id["GEN-PAY-01"]["action_suggestion"] == "优先处理"
    assert "合同正文及附件中未找到对应约定" not in str(by_id["TECH-SHARED-03"])


def test_current_transfer_report_respects_the_evidence_ceiling(transfer_review: dict):
    by_id = {row["item_id"]: row for row in transfer_review["items"]}
    competition = by_id["TC-02-01"]
    secret_loss = by_id["TC-12-03"]
    proof = {
        row["element_id"]: row
        for row in by_id["TC-03-04"]["coverage_elements"]
    }

    assert competition["coverage_status"] == "风险条款"
    assert "改进成果" in competition["risk_point_description"]
    assert "购买竞争技术" not in competition["risk_point_description"]
    assert "替代方案" not in competition["risk_point_description"]
    assert "禁止独立研发" not in competition["risk_point_description"]
    assert competition["decision_trace"]["evidence_claim_ceiling"]["status"] == "PASSED"

    assert secret_loss["coverage_status"] != "风险条款"
    assert "公开后仍须继续支付" not in secret_loss["risk_point_description"]
    assert "已付款不退" not in secret_loss["risk_point_description"]
    assert "不得解除" not in secret_loss["risk_point_description"]

    assert by_id["TC-01-04"]["coverage_status"] == "不适用"
    assert proof["third_party_confidentiality_obligations"]["status"] == "MISSING"
    assert proof["joint_holder_disclosure"]["status"] == "MISSING"
    assert proof["current_secret_status"]["status"] == "MISSING"
    assert proof["existing_confidentiality_measures"]["status"] == "MISSING"

    assert by_id["GEN-PAY-01"]["coverage_status"] == "风险条款"
    assert "120%" in by_id["GEN-PAY-01"]["risk_point_description"]
    assert "3,600,000元" in by_id["GEN-PAY-01"]["risk_point_description"]
    assert "3,000,000元" in by_id["GEN-PAY-01"]["risk_point_description"]


def test_partial_proof_narrative_does_not_claim_complete_disclosure(
    transfer_review: dict,
):
    report = transfer_review["result"]
    by_id = {row["item_id"]: row for row in transfer_review["items"]}

    for review_id in ("TC-03-04", "TC-13-02"):
        narrative = by_id[review_id]["decision_trace"]["narrative_elements"]
        judgment = narrative["judgment"]
        assert "仅部分明确：技术秘密的历史披露对象已经完整列明" not in judgment
        assert "技术秘密的历史披露对象已有部分约定，但仍未完整列明" in judgment
        assert "向第三方历史披露的具体技术内容和范围已有部分约定，但仍未完整列明" in judgment
        assert "既有许可、测试使用及其他第三方使用安排已有部分约定，但仍未完整披露" in judgment

    assert "仅部分明确：技术秘密的历史披露对象已经完整列明" not in report
    assert "建议补充明确技术秘密的历史披露对象已经完整列明" not in report


@pytest.mark.parametrize(
    ("proposition", "expected"),
    [
        (
            "测试数据已有交付安排，但具体样本和范围仍需固定",
            "测试数据已有交付安排，但具体样本和范围仍需固定",
        ),
        (
            "历史披露对象",
            "历史披露对象已有部分约定，但仍未完整列明",
        ),
        (
            "已有书面异议机制，但未固定最终验收确认文件",
            "已有书面异议机制，但未固定最终验收确认文件",
        ),
    ],
)
def test_partial_renderer_preserves_complete_gap_sentences(
    proposition: str,
    expected: str,
):
    assert _partial_element_narrative_label(proposition) == expected
