from pathlib import Path

import pytest
from docx import Document

from review_v2.engine import review_contract_v2
from review_v2.user_profile import UserRiskProfile


ROOT = Path(__file__).resolve().parents[2]
FIXTURE = ROOT / "test_data" / "31 纯合作开发合同.docx"


def _document_text(path: Path) -> str:
    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = [
        "\t".join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
    ]
    return "\n".join(paragraphs + tables)


@pytest.fixture(scope="module")
def latest_cooperative_review() -> dict:
    return review_contract_v2(
        _document_text(FIXTURE),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/合作开发方"),
    )


def _items(review: dict) -> dict:
    return {row["item_id"]: row for row in review["items"]}


def _elements(item: dict) -> dict:
    return {row["element_id"]: row for row in item["coverage_elements"]}


def test_latest_cooperative_route_and_final_quality_gate(latest_cooperative_review: dict):
    classification = latest_cooperative_review["deterministic_analysis"]["technology_classification"]
    assert classification["primary_subtype"] == "COOPERATIVE_DEVELOPMENT"
    assert latest_cooperative_review["report_context"]["contract_type_label"] == "技术合同—技术开发合同—合作开发合同"
    assert latest_cooperative_review["perspective_resolution"]["display_role"] == "甲方 / 合作开发方"
    assert latest_cooperative_review["release_quality_gate"]["status"] == "REPORT_READY"
    assert latest_cooperative_review["manual_review_summary"]["manual_review_count"] == 0


def test_competition_and_improvement_rules_reuse_direct_atomic_evidence(latest_cooperative_review: dict):
    by_id = _items(latest_cooperative_review)
    competition = by_id["TC-02-01"]
    assert competition["coverage_status"] == "风险条款"
    assert {row["status"] for row in competition["coverage_elements"]} == {"ADVERSE"}
    assert [row["location"] for row in competition["evidence_clauses"]] == ["第七条 禁止行为和保密"]
    assert all(row["entailment_status"] == "SUPPORTS" for row in competition["evidence_clauses"])

    for review_id in ("TC-10-01", "TC-10-02"):
        item = by_id[review_id]
        assert item["coverage_status"] == "风险条款"
        assert [row["location"] for row in item["evidence_clauses"]] == ["第五条 知识产权"]
        assert _elements(item)["compulsory_free_improvement_assignment"]["status"] == "ADVERSE"
    boundary = by_id["TC-10-04"]
    assert boundary["coverage_status"] == "风险条款"
    assert [row["location"] for row in boundary["evidence_clauses"]] == [
        "第七条 禁止行为和保密", "第五条 知识产权",
    ]
    assert by_id["TC-10-03"]["coverage_status"] == "缺失"

    improvement_finding = next(
        row for row in latest_cooperative_review["risk_findings"]
        if row["risk_group"] == "improvement_no_unlawful_restriction"
    )
    assert set(improvement_finding["source_review_ids"]) == {"TC-10-01", "TC-10-02", "TC-10-04"}


def test_tc_05_01_does_not_invent_covered_cooperation_facts(latest_cooperative_review: dict):
    item = _items(latest_cooperative_review)["TC-05-01"]
    elements = _elements(item)
    assert item["coverage_status"] == "部分覆盖"
    assert elements["budget_overrun_allocation"]["status"] == "COVERED"
    assert "新增或超预算研发费用" in elements["budget_overrun_allocation"]["proposition"]
    assert elements["major_decision_process"]["status"] == "PARTIALLY_COVERED"
    assert "尚未建立双方书面共同决策程序" in elements["major_decision_process"]["proposition"]
    assert [row["location"] for row in item["evidence_clauses"]] == ["第四条 合作方式"]
    rendered = " ".join((item["risk_point_description"], item["reasoning_summary"], item["suggested_clause_example"]))
    assert "超出预计投入的成本由各方自行承担" not in rendered
    assert "重大研发事项须双方书面一致同意" not in rendered


def test_acceptance_metrics_and_internal_reference_are_separately_resolved(latest_cooperative_review: dict):
    item = _items(latest_cooperative_review)["TC-13-03"]
    elements = _elements(item)
    assert item["coverage_status"] == "部分覆盖"
    assert item["issue_type"] == "CONFLICTING"
    assert elements["body_acceptance_metrics"]["status"] == "COVERED"
    assert elements["test_method"]["status"] == "MISSING"
    assert elements["test_report"]["status"] == "MISSING"
    assert elements["internal_cross_reference_semantics"]["status"] == "CONFLICTING"
    evidence_text = "\n".join(str(row.get("text") or row.get("quote") or "") for row in item["evidence_clauses"])
    assert all(term in evidence_text for term in ("800名", "1秒", "90%", "按照第七条", "保密"))


def test_claim_allocation_is_not_misreported_as_actual_infringement(latest_cooperative_review: dict):
    by_id = _items(latest_cooperative_review)
    infringement_screen = by_id["TC-02-02"]
    assert infringement_screen["coverage_status"] == "通过"
    assert infringement_screen["related_risks"] == []
    assert infringement_screen["decision_trace"]["evidence_judgment"] == {
        "evidence_semantics": "HYPOTHETICAL_CLAIM_OR_RISK_ALLOCATION",
        "affirmative_infringement": False,
        "affirmative_risk_evidence_ids": [],
        "judgment": "PASS",
    }

    claims = by_id["TC-12-02"]
    elements = _elements(claims)
    assert claims["coverage_status"] == "部分覆盖"
    assert claims["needs_manual_review"] is False
    assert claims["system_confidence"]["level"] == "high"
    assert [row["location"] for row in claims["evidence_clauses"]] == ["第十一条 第三方权利主张"]
    assert elements["claim_handling_responsibility"]["status"] == "COVERED"
    assert elements["claim_cost_allocation"]["status"] == "COVERED"
    assert elements["third_party_compensation_allocation"]["status"] == "COVERED"
    assert elements["claim_notice"]["status"] == "MISSING"
    assert elements["lawyer_selection"]["status"] == "MISSING"


def test_budget_investment_and_cooperation_timeline_use_shared_deterministic_checks(latest_cooperative_review: dict):
    by_id = _items(latest_cooperative_review)
    payment = by_id["GEN-PAY-01"]
    formula = payment["scenario_analysis"]["formula"]
    assert formula["type"] == "budget_investment_consistency"
    assert formula["contract_total"] == 1_500_000
    assert [row["amount"] for row in formula["nodes"]] == [1_200_000, 600_000]
    assert formula["amount_total"] == 1_800_000
    assert formula["amount_delta"] == 300_000
    assert formula["status"] == "INCONSISTENT"
    assert "项目研发预算总额" in payment["risk_point_description"]

    timeline = by_id["TECH-SHARED-03"]["scenario_analysis"]["timeline"]
    assert timeline["earliest_performance_date"] == "2026-10-01"
    assert timeline["effective_date"] == "2026-10-15"
    assert timeline["performance_before_effective_date"] is True
    timeline_finding = next(
        row for row in latest_cooperative_review["risk_findings"]
        if row["risk_group"] == "technical_timeline_conflict"
    )
    evidence_quotes = [row["quoted_text"] for row in timeline_finding["evidence_bundle"]["items"]]
    assert len(evidence_quotes) == len(set(evidence_quotes))


def test_cooperative_failure_allocation_is_an_evidence_bound_risk(latest_cooperative_review: dict):
    item = _items(latest_cooperative_review)["TECH-SHARED-DEV-FAIL-01"]
    elements = _elements(item)
    assert item["coverage_status"] == "风险条款"
    assert [row["location"] for row in item["evidence_clauses"]] == ["第八条 研发失败风险承担"]
    assert elements["incurred_cost_allocation"]["status"] == "ADVERSE"
    assert elements["future_unincurred_cost_allocation"]["status"] == "ADVERSE"
    assert elements["termination_decision"]["status"] == "ADVERSE"
    assert elements["failure_notice_and_mitigation"]["status"] == "MISSING"
    assert elements["stage_result_and_handover"]["status"] == "MISSING"
    assert any(
        row["risk_group"] == "cooperative_development_failure_allocation"
        for row in latest_cooperative_review["risk_findings"]
    )


def test_existing_result_license_and_revenue_terms_are_adverse_not_missing(latest_cooperative_review: dict):
    item = _items(latest_cooperative_review)["TC-05-03"]
    elements = _elements(item)
    assert item["coverage_status"] == "风险条款"
    assert elements["third_party_license_authority"]["status"] == "ADVERSE"
    assert elements["external_revenue_allocation"]["status"] == "ADVERSE"
    assert elements["counterparty_license_transfer_restriction"]["status"] == "ADVERSE"
    assert "许可第三方" in item["risk_point_description"]
    assert "甲方80%、乙方20%" in item["risk_point_description"]
    assert "未主动明确第三方普通许可" not in item["risk_point_description"]


def test_technical_material_attachment_preserves_body_scope_and_work_allocation(latest_cooperative_review: dict):
    item = _items(latest_cooperative_review)["TC-13-01"]
    elements = _elements(item)
    assert item["coverage_status"] == "部分覆盖"
    assert elements["body_development_scope"]["status"] == "COVERED"
    assert elements["body_research_work_allocation"]["status"] == "COVERED"
    assert elements["requirements_specification"]["status"] == "MISSING"
    assert elements["detailed_module_schedule"]["status"] == "MISSING"
    assert "正文已经明确基础研发范围" in item["risk_point_description"]
    assert "未明确研发分工" not in item["risk_point_description"]


def test_improvement_boundary_claim_is_limited_to_its_bound_clause(latest_cooperative_review: dict):
    item = _items(latest_cooperative_review)["TC-10-04"]
    assert [row["location"] for row in item["evidence_clauses"]] == [
        "第七条 禁止行为和保密", "第五条 知识产权",
    ]
    assert "无偿转归另一方" in item["risk_point_description"]
    assert "限制合法采购或使用替代技术" in item["risk_point_description"]
    assert "非互惠返授" not in item["risk_point_description"]


def test_cooperative_report_has_no_license_party_role_pollution(latest_cooperative_review: dict):
    assert "被许可方" not in latest_cooperative_review["result"]
    assert "许可方" not in latest_cooperative_review["result"]


def test_absence_search_uses_contract_writable_positive_facts(latest_cooperative_review: dict):
    by_id = _items(latest_cooperative_review)
    forbidden_judgment_markers = ("已约定，但", "仍需细化", "已有部分约定", "未覆盖要素")
    for review_id in ("TC-05-03", "TC-11-01", "TC-13-01"):
        audit = by_id[review_id]["decision_trace"].get("specialized_search_audit") or {}
        assert audit.get("query_semantics") == "POSITIVE_REQUIRED_FACT"
        targets = audit.get("required_positive_facts") or []
        assert targets
        assert not any(marker in target for marker in forbidden_judgment_markers for target in targets)


def test_major_decision_recommendation_is_direct_and_actionable(latest_cooperative_review: dict):
    recommendation = _items(latest_cooperative_review)["TC-05-01"]["suggested_clause_example"]
    assert recommendation == "建议约定重大需求变更、技术路线调整、预算增加、项目延期等重大事项须经双方书面确认。"
