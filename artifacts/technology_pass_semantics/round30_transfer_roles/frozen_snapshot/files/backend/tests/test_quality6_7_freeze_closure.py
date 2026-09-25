from pathlib import Path
from types import SimpleNamespace

from docx import Document

import review_v2.engine as engine


ROOT = Path(__file__).resolve().parents[2]


def _hb4_text():
    document = Document(next((ROOT / "test_data").glob("7*.docx")))
    rows = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    rows.extend(
        " | ".join(cell.text.strip() for cell in row.cells)
        for table in document.tables for row in table.rows
    )
    return "\n".join(row for row in rows if row.strip(" |"))


def _item(*, coverage="通过", applicability="APPLICABLE", polarity="POSITIVE_REQUIREMENT"):
    return SimpleNamespace(
        item_id="R-1", title="测试事项", coverage_status=coverage,
        applicability={"status": applicability, "reason": "未触发适用条件"},
        coverage_elements=[], decision_trace={"review_polarity": polarity},
        narrative_state={}, risk_point_description="", legal_consequence_if_missing="",
        reasoning_summary="", suggested_clause_example="",
    )


def test_not_applicable_stops_default_law_gap_finding_and_recommendation():
    item = _item(coverage=engine.COVERAGE_NOT_APPLICABLE, applicability="NOT_APPLICABLE")
    item.decision_trace.update({
        "default_law_assessment": {"executed": False, "status": "NOT_APPLICABLE"},
        "effective_gap": {
            "status": "NOT_APPLICABLE", "contractual_missing_elements": [],
        },
    })
    item.narrative_state = engine._build_narrative_state(item)
    gate = engine._applicability_polarity_gates([item], [], [])
    assert gate["gate_results"]["NOT_APPLICABLE_GAP_GATE"] == "PASS"
    assert gate["not_applicable_finding_count"] == 0
    assert gate["not_applicable_recommendation_count"] == 0
    assert item.narrative_state["missing_elements"] == []


def test_negative_pass_uses_absence_semantics_and_requires_sufficient_search():
    item = _item(polarity="NEGATIVE_CLAUSE_CHECK")
    item.decision_trace["negative_clause_check"] = {
        "adverse_clause_exists": False, "search_coverage": "SUFFICIENT", "confidence": "high",
    }
    item.narrative_state = engine._build_narrative_state(item)
    gate = engine._applicability_polarity_gates([item], [], [])
    assert item.narrative_state["judgment"].startswith("未发现")
    assert "合同已覆盖" not in item.narrative_state["judgment"]
    assert gate["gate_results"]["REVIEW_POLARITY_NARRATIVE_GATE"] == "PASS"
    assert gate["gate_results"]["NEGATIVE_CHECK_PASS_GATE"] == "PASS"

    item.decision_trace["negative_clause_check"]["search_coverage"] = "INSUFFICIENT"
    failed = engine._applicability_polarity_gates([item], [], [])
    assert failed["gate_results"]["NEGATIVE_CHECK_PASS_GATE"] == "FAIL"


def test_supported_default_capable_rule_cannot_diverge_without_explanation():
    item = _item(coverage="部分覆盖")
    item.item_id = "R-LEGAL"
    item.coverage_elements = [{
        "element_id": "foreseeability_limit", "status": "MISSING", "atomic": True,
        "legal_proposition_ids": ["legal.damages.foreseeability_limit"],
    }]
    item.legal_references = [{
        "authority_id": "CN-X", "support_type": "DIRECT_SUPPORT", "freshness_status": "CURRENT",
    }]
    item.decision_trace["default_law_assessment"] = {
        "status": "NO_DEFAULT",
        "rule": {
            "proposition_id": "legal.damages.foreseeability_limit",
            "proposition": "可预见性边界适用。", "authority_ids": ["CN-X"],
            "supplied_element_ids": ["foreseeability_limit"],
        },
    }
    gate = engine._legal_default_law_consistency_gate([item])
    assert gate["gate_results"]["LEGAL_DEFAULT_LAW_CONSISTENCY_GATE"] == "FAIL"
    assert gate["contradiction_count"] == 1


def test_non_atomic_element_without_children_fails_closed():
    item = _item(coverage="部分覆盖")
    item.legal_references = []
    item.coverage_elements = [{
        "element_id": "x_or_y", "proposition": "X 或 Y", "status": "MISSING",
        "atomic": False, "child_elements": [], "legal_proposition_ids": [],
    }]
    item.decision_trace["default_law_assessment"] = {"status": "NO_DEFAULT", "rule": None}
    gate = engine._legal_default_law_consistency_gate([item])
    assert gate["gate_results"]["ATOMIC_LEGAL_ELEMENT_GATE"] == "FAIL"
    assert gate["composite_error_count"] == 1


def test_hb4_freeze_closure_production_regressions():
    result = engine.review_contract_v2(_hb4_text(), "house_lease", use_llm=False)
    items = {item["item_id"]: item for item in result["items"]}
    gates = result["release_quality_gate"]["gate_results"]
    assert result["release_quality_gate"]["status"] == "REPORT_READY"
    for gate_name in (
        "LEGAL_DEFAULT_LAW_CONSISTENCY_GATE", "ATOMIC_LEGAL_ELEMENT_GATE",
        "NOT_APPLICABLE_NARRATIVE_GATE", "NOT_APPLICABLE_GAP_GATE",
        "REVIEW_POLARITY_NARRATIVE_GATE", "NEGATIVE_CHECK_PASS_GATE",
        "RECOMMENDATION_SINGLE_SOURCE_GATE", "RECOMMENDATION_CROSS_RENDERER_CONSISTENCY_GATE",
    ):
        assert gates[gate_name] == "PASS"

    s0803 = items["S-08-03"]
    element_ids = {element["element_id"] for element in s0803["coverage_elements"]}
    assert "foreseeability_limit" in element_ids
    assert "proof_or_evidentiary_requirement" in element_ids
    assert "foreseeability_or_provability" not in element_ids
    assert s0803["decision_trace"]["default_law_assessment"]["status"] == "PARTIAL_DEFAULT"

    for review_id in ("S-02-05", "S-08-04"):
        assert items[review_id]["decision_trace"]["review_polarity"] == "NEGATIVE_CLAUSE_CHECK"
        assert items[review_id]["narrative_state"]["judgment"].startswith("未发现")

    for review_id in ("S-04-03", "S-09-01", "S-09-02", "S-09-03", "S-09-04", "S-10-01", "S-10-02", "S-10-03", "S-10-04"):
        item = items[review_id]
        if item["coverage_status"] != engine.COVERAGE_NOT_APPLICABLE:
            continue
        assert item["decision_trace"]["default_law_assessment"]["executed"] is False
        assert item["decision_trace"]["effective_gap"]["status"] == "NOT_APPLICABLE"
        assert item["decision_trace"]["effective_gap_narrative"]["generated_narrative"] == ""
