"""Application-stage regression, plus controls for polarity and role routing."""
from pathlib import Path
import json
import re
from types import SimpleNamespace

import pytest
from docx import Document

from review_v2.engine import (
    _has_affirmative_improvement_restriction,
    _is_technology_checkpoint_applicable,
    _technology_ip_proof_model,
    review_contract_v2,
)
from review_v2.perspective_resolution import resolve_perspective
from review_v2.user_profile import UserRiskProfile


@pytest.fixture(scope="module")
def result():
    document = Document(next((Path(__file__).resolve().parents[2] / "test_data").glob("42*.docx")))
    text = "\n".join(p.text for p in document.paragraphs)
    text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for t in document.tables for row in t.rows)
    return review_contract_v2(text, "technology", use_llm=False, include_low_risk=True,
                              profile=UserRiskProfile(user_identity="甲方/受让方"))


def item(result, key):
    return next(row for row in result["items"] if row["item_id"] == key)


def test_application_subtype_survives_report_projection(result):
    assert result["deterministic_analysis"]["technology_classification"]["tertiary_subtype"] == "PATENT_APPLICATION_RIGHT_TRANSFER"
    assert result["report_context"]["scenario_label"] == "专利申请权转让"
    assert result["report_context"]["contract_type_label"] == "技术合同—技术转让合同—专利申请权转让合同"
    assert result["perspective_resolution"]["display_role"] == "甲方 / 专利申请权受让方"


def test_warranty_is_not_demoted_by_missing_proof_documents(result):
    warranty = item(result, "TC-06-03")
    assert warranty["coverage_status"] in {"通过", "部分覆盖"}
    assert warranty["risk_level"] != "高"
    assert "第六条" in warranty["evidence_location"]
    assert all(row["status"] == "COVERED" for row in warranty["coverage_elements"])


def test_duty_and_non_duty_have_separate_triggers(result):
    assert item(result, "TC-03-01")["coverage_status"] in {"通过", "部分覆盖"}
    assert item(result, "TC-03-02")["coverage_status"] == "不适用"


def test_claim_liability_reuses_seller_responsibility(result):
    liability = item(result, "TC-12-01")
    assert liability["coverage_status"] in {"通过", "部分覆盖"}
    assert liability["risk_level"] != "高"
    assert "第六条" in liability["evidence_location"]
    assert "874" in liability["legal_basis_index"]


def test_improvement_notice_does_not_use_confidentiality(result):
    notice = item(result, "TC-10-03")
    assert notice["coverage_status"] in {"缺失", "部分覆盖"}
    assert notice["risk_level"] == "低"
    assert "保密" not in notice["evidence_location"]
    assert not any("保密" in str(row.get("title", "")) for row in notice["evidence_clauses"])


def test_improvement_negative_screen_passes_without_restrictive_fact(result):
    boundary = item(result, "TC-10-04")
    assert boundary["coverage_status"] == "通过"
    assert boundary["decision_trace"]["negative_clause_check"]["adverse_clause_exists"] is False


def test_proof_model_uses_application_not_granted_fields(result):
    proof = item(result, "TC-13-02")
    text = json.dumps(proof["coverage_elements"], ensure_ascii=False) + proof["suggested_clause_example"]
    for forbidden in ("专利年费", "年费缴纳", "专利号", "专利权现时有效", "专利权转让登记"):
        assert forbidden not in text
    elements = {row["element_id"]: row for row in proof["coverage_elements"]}
    for key in ("application_number", "application_date", "current_applicant", "application_acceptance_document",
                "application_stage_fees", "application_procedural_documents", "application_pledge_disclosure",
                "application_seizure_disclosure", "application_existing_license_disclosure",
                "application_other_restrictions_disclosure",
                "application_transfer_registration_documents"):
        assert elements[key]["status"] == "COVERED"


def test_registration_effect_is_not_contract_formation(result):
    registration = item(result, "TC-06-04")
    assert registration["coverage_status"] == "通过"
    assert "专利法》第10条" in registration["legal_basis_index"]
    assert "自登记之日起" in registration["reasoning_summary"]
    assert "区别于合同成立" in registration["reasoning_summary"]
    assert "第十一条" in registration["evidence_location"]
    assert "第二条" not in registration["evidence_location"]


def test_report_release_and_positive_evidence_provenance(result):
    assert result["release_quality_gate"]["status"] != "BLOCKED", result["release_quality_gate"]["failures"]
    for review_id in ("TC-03-01", "TC-06-03", "TC-06-04", "TC-12-01"):
        row = item(result, review_id)
        assert row["review_evidence_provenance"]["evidence_ids"]
    notice_dispositions = [f for f in result["final_dispositions"] if "TC-10-03" in f["review_ids"]]
    assert notice_dispositions and all(f["final_action"] == "OPTIONAL_IMPROVEMENT" for f in notice_dispositions)


def test_actual_attachment_list_is_not_absent(result):
    material = item(result, "TC-13-01")
    elements = {row["element_id"]: row for row in material["coverage_elements"]}
    assert elements["closed_attachment_schedule"]["status"] == "COVERED"
    missing = material["decision_trace"].get("final_missing_element_ids", [])
    assert "closed_attachment_schedule" not in missing
    assert "未显示已经形成最终附件" not in material["risk_point_description"]


def test_named_inventors_are_not_missing_with_optional_reward_gap(result):
    creator = item(result, "TC-03-03")
    assert creator["coverage_status"] == "部分覆盖"
    assert creator["risk_level"] == "低"
    elements = {row["element_id"]: row for row in creator["coverage_elements"]}
    assert elements["creator_named_identity"]["status"] == "COVERED"
    assert "第一条" in creator["evidence_location"]
    assert creator["decision_trace"]["final_missing_element_ids"] == ["creator_reward_arrangement"]
    dispositions = [row for row in result["final_dispositions"] if "TC-03-03" in row["review_ids"]]
    assert all(row["final_action"] in {"OPTIONAL_IMPROVEMENT", "KEEP_CURRENT", "DO_NOT_ADD"} for row in dispositions)


def test_creator_reward_recommendation_is_consistent_across_report_sections(result):
    expected = "建议视需要核验让与方是否已依法落实完成人奖励、报酬或内部激励安排。"
    assert item(result, "TC-03-03")["suggested_clause_example"] == expected
    report = result["result"]
    for heading in ("核心风险摘要", "可选优化事项", "完整风险主题明细（聚合后）", "完整审查明细"):
        section = re.search(r"^##? [^\n]*" + re.escape(heading) + r"[^\n]*\n(.*?)(?=^##? [^\n]*|\Z)", report, re.M | re.S)
        assert section, heading
        # The detailed sections include nested per-rule/theme headings.
        if heading == "完整审查明细":
            text = report.split("完整审查明细", 1)[1].split("（TC-03-03）", 1)[1].split("\n## ", 1)[0]
        elif heading == "完整风险主题明细（聚合后）":
            text = report.split(heading, 1)[1].split("完整审查明细", 1)[0]
        else:
            text = section.group(1)
        assert expected in text, heading
    assert "建议补充明确完成人奖励" not in report
    dispositions = [row for row in result["final_dispositions"] if "TC-03-03" in row["review_ids"]]
    assert dispositions
    assert all(row["final_action"] == "OPTIONAL_IMPROVEMENT" for row in dispositions)
    assert all(row["recommendation"] == expected for row in dispositions)


def test_completed_core_becomes_pass_without_fabricated_gap_or_exposure(result):
    confidentiality = item(result, "TC-11-01")
    assert confidentiality["coverage_status"] == "通过"
    assert confidentiality["decision_trace"]["final_missing_element_ids"] == []
    assert not confidentiality["decision_trace"]["recommendation_display_state"]["has_active_recommendation"]
    narrative = confidentiality["narrative_state"]
    assert narrative["missing_elements"] == []
    assert narrative["issue"] == ""
    assert narrative["consequence"] in {"", "无明显残余风险。"}
    effective_narrative = confidentiality["decision_trace"]["effective_gap_narrative"]["generated_narrative"]
    assert effective_narrative in {"", "合同已明确约定该事项。"}
    detail = result["result"].split("完整审查明细", 1)[1].split("（TC-11-01）", 1)[1].split("\n## ", 1)[0]
    assert "合同已明确约定" in detail
    for contradiction in ("仍需补足", "缺失的法律后果", "实际法律暴露", "合同未明确约定"):
        assert contradiction not in detail


@pytest.mark.parametrize("text,adverse", [
    ("卖方无权要求无偿共有、转让或者分享买方独立形成的后续改进成果。", False),
    ("双方均有权自行进行后续改进。", False),
    ("买方不得自主研究开发或改进。", True),
    ("卖方不得要求无偿转让独立改进成果。买方不得自主改进。", True),
])
def test_improvement_polarity_is_sentence_scoped(text, adverse):
    assert _has_affirmative_improvement_restriction(text) is adverse


def test_granted_patent_model_and_both_party_roles_remain_distinct():
    model = _technology_ip_proof_model({"primary_subtype": "TECHNOLOGY_TRANSFER", "tertiary_subtype": "PATENT_RIGHT_TRANSFER"})
    assert "patent_legal_status_and_annual_fee" in {row[0] for row in model["specs"]}
    for role, expected in (("甲方/受让方", "甲方 / 专利申请权受让方"), ("乙方/让与方", "乙方 / 专利申请权让与方")):
        resolved = resolve_perspective(role, "TECHNOLOGY", contract_subtype="PATENT_APPLICATION_RIGHT_TRANSFER")
        assert resolved.display_role == expected
        assert resolved.transaction_type == "TECHNOLOGY_TRANSFER"


@pytest.mark.parametrize("fact,applicable", [
    ("申请人为企业，发明人为自然人。卖方已取得员工职务成果归属内部文件。", False),
    ("本项目不存在非职务发明。", False),
    ("本项目含非职务发明，发明人以个人名义投入。", True),
    ("员工主张其为成果所有权人，与单位存在争议。", True),
])
def test_non_duty_applicability_requires_its_own_positive_fact(fact, applicable):
    checkpoint = SimpleNamespace(item_id="TC-03-02", node_id="non_duty_achievement_ownership")
    actual, _ = _is_technology_checkpoint_applicable("专利申请权转让合同。" + fact, checkpoint)
    assert actual is applicable
