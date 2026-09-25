from pathlib import Path
from types import SimpleNamespace

from docx import Document

import review_v2.engine as engine
from review_v2.quality4 import validate_citation_entailment


ROOT = Path(__file__).resolve().parents[2]


def _hb4_text():
    document = Document(next((ROOT / "test_data").glob("7*.docx")))
    rows = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        rows.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
    return "\n".join(row for row in rows if row.strip(" |"))


def test_article_620_does_not_support_seller_certificate_obligation():
    result = validate_citation_entailment(
        proposition="出卖人必须随货提供产品合格证明和检验报告。",
        law_name="中华人民共和国民法典",
        article="第620条",
    )
    assert result["support_type"] == "NOT_SUPPORT"


def test_article_620_directly_supports_buyer_inspection_proposition():
    result = validate_citation_entailment(
        proposition="买受人收到标的物时应当及时检验。",
        law_name="中华人民共和国民法典",
        article="第620条",
    )
    assert result["support_type"] == "DIRECT_SUPPORT"


def test_articles_620_and_621_support_inspection_default_proposition():
    proposition = "提供合理检验期限和通知的法定补充。"
    for article in ("第620条", "第621条"):
        result = validate_citation_entailment(
            proposition=proposition, law_name="中华人民共和国民法典",
            article=article, claim_type="STATUTORY_DEFAULT",
        )
        assert result["support_type"] == "DIRECT_SUPPORT"


def test_legal_authority_selection_deduplicates_structured_authority_per_review():
    authority = {
        "authority_id": "CN-CC-605", "law_name": "中华人民共和国民法典",
        "article": "第605条", "effective_from": "2021-01-01",
        "text": "因买受人的原因致使标的物未按照约定期限交付的，买受人应当承担相应风险。",
    }
    item = SimpleNamespace(
        item_id="S-04-04", decision_trace={},
        legal_references=[
            {
                "claim_id": "A", "authority_id": "CN-CC-605", "law_name": "中华人民共和国民法典",
                "article": "第605条", "paragraph": "", "proposition": "迟延受领后风险依法转移。",
                "claim_type": "STATUTORY_DEFAULT", "support_type": "DIRECT_SUPPORT",
                "freshness_status": "CURRENT", "confidence": "high", "authority": authority,
            },
            {
                "claim_id": "B", "authority_id": "CN-CC-605", "law_name": "中华人民共和国民法典",
                "article": "第605条", "paragraph": "", "proposition": "迟延受领与风险承担有关。",
                "claim_type": "STATUTORY_RULE", "support_type": "PARTIAL_SUPPORT",
                "freshness_status": "CURRENT", "confidence": "medium", "authority": authority,
            },
        ],
    )
    gate = engine._select_legal_authorities([item])
    retained = [claim for claim in item.legal_references if claim.get("authority_id") == "CN-CC-605"]
    assert gate["gate_results"]["LEGAL_AUTHORITY_DEDUP_GATE"] == "PASS"
    assert gate["duplicate_count"] == 0
    assert len(retained) == 1
    assert retained[0]["claim_type"] == "STATUTORY_DEFAULT"


def test_hb4_final_recommendation_and_legal_truth_are_consistent():
    result = engine.review_contract_v2(_hb4_text(), "house_lease", use_llm=False)
    items = {item["item_id"]: item for item in result["items"]}
    report = result["result"]

    assert result["release_quality_gate"]["status"] == "REPORT_READY"
    assert result["recommendation_single_source_gate"]["gate_results"]["RECOMMENDATION_SINGLE_SOURCE_GATE"] == "PASS"
    assert result["recommendation_cross_renderer_consistency_gate"]["gate_results"]["RECOMMENDATION_CROSS_RENDERER_CONSISTENCY_GATE"] == "PASS"
    assert result["legal_authority_gate"]["gate_results"] == {
        "LEGAL_AUTHORITY_DEDUP_GATE": "PASS",
        "LEGAL_AUTHORITY_SELECTION_GATE": "PASS",
        "CITATION_ENTAILMENT_GATE": "PASS",
    }

    s0803_state = items["S-08-03"]["decision_trace"]["recommendation_display_state"]
    assert s0803_state["active_recommendation_ids"] == []
    assert s0803_state["blocked_recommendation_ids"]
    assert "建议仅补充明确：可预见性或可证明性限制。" not in report
    assert not set(s0803_state["blocked_recommendation_ids"]) & set(result["final_recommendation_set"]["recommendation_ids"])

    s0404 = items["S-04-04"]
    authority_keys = [
        (claim.get("authority_id"), claim.get("article"))
        for claim in s0404["legal_references"] if claim.get("authority_id")
    ]
    assert len(authority_keys) == len(set(authority_keys))
    default_narrative = s0404["decision_trace"]["effective_gap_narrative"]["generated_narrative"]
    assert "合同文本未完整约定" in default_narrative
    assert "现行法律已提供相应默认处理规则" in default_narrative
    assert "合同已明确约定该事项" not in default_narrative

    assert all(claim.get("article") != "第620条" for claim in items["S-12-02"]["legal_references"])
    assert result["default_law_narrative_truth_gate"]["contradiction_count"] == 0
