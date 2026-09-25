import json
from pathlib import Path
import pytest
from docx import Document
from review_v2.engine import review_contract_v2
from review_v2.clause_parser import parse_contract_clauses
from review_v2.template_parser import inactive_alternative_lines
from review_v2.priority_contract_types import classify_priority_contract
from review_v2.priority_contract_review import adverse_evidence, positive_rule_adverse_evidence
from review_v2.noncompete_semantics import temporal_roles


@pytest.mark.parametrize('selector,first,second', [
    ('按以下第（2）项执行：','(1)无需支付经济补偿。','(2)依法支付经济补偿。'),
    ('选择第二项：','（一）无需支付经济补偿。','（二）依法支付经济补偿。'),
    ('选择第1项：','(1)依法支付经济补偿。','(2)无需支付经济补偿。'),
    ('请勾选：','□无需支付经济补偿。','√依法支付经济补偿。'),
])
def test_selected_alternative_excludes_inactive_evidence(selector,first,second):
    text='离职结算协议\n3、'+selector+'\n'+first+'\n'+second+'\n4、办理工作交接。'
    inactive=inactive_alternative_lines(text)
    assert any('无需' in s for s in inactive)
    clauses=parse_contract_clauses(text)
    operative='\n'.join(c.text for c in clauses if c.section_type=='正式合同条款')
    assert '无需' not in operative
    assert '依法支付' in operative
    assert '无需' in '\n'.join(c.raw_text for c in clauses)


def test_ordinary_numbered_list_is_not_an_alternative():
    assert inactive_alternative_lines('1、工作内容\n（1）开发。\n（2）测试。') == set()


def test_checkbox_selection_is_scoped_to_its_own_article():
    text='1、选项\n√现金\n□转账\n2、另一个选项组\n□协商\n□仲裁'
    assert inactive_alternative_lines(text)=={'□转账'}


def test_release_of_future_covenant_does_not_mean_unpaid_competition():
    text='竞业限制协议\n本协议解除后，乙方无需再履行后续竞业限制义务，甲方无需再支付竞业限制经济补偿金。'
    assert not adverse_evidence('employment_noncompete_adverse', parse_contract_clauses(text))


@pytest.mark.parametrize('title', ['解除或终止劳动合同协议书','劳动关系终止协议','离职结算协议'])
def test_exit_subtype_heading(title):
    assert classify_priority_contract(title+'\n双方协商一致解除劳动关系。',[])[0]=='employment_termination'


def test_exit_purpose_is_not_ordinary_future_exit_clause():
    ordinary='劳动合同\n双方建立劳动关系，工作岗位为工程师。未来解除劳动合同按法律规定执行，并办理工作交接及经济补偿。'
    assert classify_priority_contract(ordinary,[])[0]=='labor'


def test_registry_scope_gate_rejects_ordinary_completeness_in_exit():
    from review_v2.priority_contract_review import applicability
    from review_v2.employment_contract_checklist import WORK, HOURS, SAFETY, WAGE
    for checkpoint in (WORK,HOURS,SAFETY,WAGE):
        assert applicability('双方办理离职结算。',checkpoint,'employment_termination')[0] is False


@pytest.mark.parametrize('days,adverse', [('15',False),('30',True),('二十',True),('10',False)])
def test_transfer_deadline_uses_numeric_value(days,adverse):
    cs=parse_contract_clauses('解除协议\n第一条 甲方在劳动合同解除后'+days+'日内办理档案及社会保险关系转移手续。')
    assert bool(positive_rule_adverse_evidence('employment_handover',cs)) == adverse


def test_noncompete_duration_and_compensation_waiver_are_distinct_facts():
    text='离职协议\n第一条 自劳动合同终止之日起两年内，不得进入存在竞争关系的企业任职。甲方无需就上述竞业限制另行支付经济补偿。'
    cs=parse_contract_clauses(text)
    assert any(r['role']=='NONCOMPETE_DURATION' and r['expression']=='两年' for r in temporal_roles(cs))
    assert positive_rule_adverse_evidence('employment_noncompete_term',cs)
    assert adverse_evidence('employment_noncompete_adverse',cs)


@pytest.mark.parametrize('node,text', [
    ('termination_payment_withholding','甲方有权暂缓支付劳动报酬，直到完成交接。'),
    ('termination_direct_deduction','甲方有权从应付工资中直接扣除培训费用及工作损失。'),
    ('termination_remedy_waiver','乙方放弃申请仲裁、诉讼、投诉的权利。'),
])
def test_direct_exit_risks_bind_adverse_text(node,text):
    rows=adverse_evidence(node,parse_contract_clauses('离职协议\n第一条 '+text))
    assert rows and all(row['text'] in text or text in row['text'] for row in rows)


def test_normal_compensation_handover_is_not_wage_withholding():
    cs=parse_contract_clauses('离职协议\n经济补偿在办结工作交接时支付。已到期工资当日付清。')
    assert not adverse_evidence('termination_payment_withholding',cs)
    assert not adverse_evidence('termination_remedy_waiver',parse_contract_clauses('双方确认已经结清款项，再无结算争议。'))


def test_current_exit_full_pipeline():
    path=Path(__file__).resolve().parents[2]/'test_data/74 劳动关系解除.docx'
    text='\n'.join(p.text for p in Document(path).paragraphs)
    result=review_contract_v2(text,'labor',use_llm=False)
    assert result['contract_type']=='employment_termination'
    items={i['item_id']:i for i in result['items']}
    assert not set(items)&{'EMP-02-01','EMP-03-01','EMP-04-01','EMP-05-01','EMP-06-01','EMP-08-01'}
    assert items['EMP-11-01']['coverage_status']=='通过'
    for key in ['ET-01-01','ET-04-01','ET-05-01','ET-06-01','NC-02-01','NC-03-01','EMP-14-01']:
        assert items[key]['coverage_status']=='风险条款', (key,items[key]['reasoning_summary'])
        assert items[key]['evidence_clauses']
    assert '无需支付乙方经济补偿金' not in json.dumps([i['evidence_clauses'] for i in items.values()],ensure_ascii=False)
    assert any(r['role']=='NONCOMPETE_DURATION' for r in items['NC-02-01']['decision_trace']['temporal_role_analysis'])
    assert result['release_quality_gate']['status']=='REPORT_READY',result['release_quality_gate']['failures']


@pytest.mark.parametrize('duty', [
    '乙方承担离职后竞业限制义务。',
    '乙方承担专项培训服务期义务。',
])
@pytest.mark.parametrize('amount', ['200000', '35000'])
def test_allowed_penalty_scope_is_not_illegal_because_of_amount(duty,amount):
    cs=parse_contract_clauses('解除劳动合同协议\n8、其他约定\n（1）'+duty+'\n（2）乙方违反本条约定，应支付违约金人民币'+amount+'元。')
    assert not adverse_evidence('employment_general_penalty',cs)


def test_penalty_scope_does_not_cross_parent_article():
    text='解除劳动合同协议\n7、保密\n（1）乙方承担保密义务。\n8、竞业\n（1）乙方承担竞业限制义务。\n（2）乙方违反本条约定，应支付违约金人民币70000元。'
    assert not adverse_evidence('employment_general_penalty',parse_contract_clauses(text))


def test_explicitly_limited_penalty_does_not_inherit_all_siblings():
    text='解除劳动合同协议\n8、其他约定\n（1）乙方承担保密义务。\n（2）乙方违反本条约定，仅对违反竞业限制义务支付违约金人民币70000元。'
    assert not adverse_evidence('employment_general_penalty',parse_contract_clauses(text))


def test_mixed_penalty_binds_both_source_roles():
    cs=parse_contract_clauses('解除劳动合同协议\n8、其他约定\n（1）乙方承担保密义务。\n（2）乙方承担竞业限制义务。\n（3）乙方违反本条约定，应支付违约金人民币70000元。')
    evidence=adverse_evidence('employment_general_penalty',cs)
    assert {e['binding_role'] for e in evidence}=={'PENALTY_CLAUSE','COVERED_ORDINARY_OBLIGATION'}
    assert len({tuple(e['scope_evidence_ids']) for e in evidence})==1


def test_exit_address_policy_does_not_mutate_ordinary_requirements():
    from review_v2.employment_facts import resolve_requirements
    from review_v2.priority_checklist_helpers import PROFILES
    for title,optional in [('解除劳动合同协议',True),('劳动合同',False)]:
        elements,_=resolve_requirements('employment_identity',PROFILES['employment_identity'],parse_contract_clauses(title))
        address=next(e for e in elements if e['element_id']=='employment_identity_1')
        assert (address.get('risk_relevant') is False)==optional
    assert PROFILES['employment_identity'][0].get('risk_relevant') is not False


def test_closeout_keeps_scope_evidence_in_final_report():
    path=Path(__file__).resolve().parents[2]/'test_data/74 劳动关系解除.docx'
    result=review_contract_v2('\n'.join(p.text for p in Document(path).paragraphs),'employment_termination',use_llm=False)
    items={i['item_id']:i for i in result['items']}
    assert items['EMP-01-01']['coverage_status']=='通过'
    assert '可选' in items['EMP-01-01']['reasoning_summary']
    penalty=items['EMP-14-01']
    assert penalty['coverage_status']=='风险条款'
    assert {e.get('binding_role') for e in penalty['evidence_clauses']} >= {'PENALTY_CLAUSE','COVERED_ORDINARY_OBLIGATION'}
    assert '保密义务' in json.dumps(penalty['presentation']['element_citations'],ensure_ascii=False)
    assert result['release_quality_gate']['status']=='REPORT_READY'
