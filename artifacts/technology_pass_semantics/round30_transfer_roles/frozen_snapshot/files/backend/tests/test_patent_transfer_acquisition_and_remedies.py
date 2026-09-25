"""Granted-patent title and remedies are independently source-bound facts."""
from pathlib import Path
import asyncio
from copy import deepcopy
from types import SimpleNamespace

import pytest
from docx import Document

from review_v2.clause_parser import parse_contract_clauses
from review_v2.technology_patent_transfer_analysis import (
    extract_patent_transfer_protection_facts,
    extract_third_party_claim_protection_facts,
)
from review_v2.engine import (
    review_contract_v2,
    _current_rights_disclosure_evidence,
    _PATENT_ENCUMBRANCE_DISCLOSURE_SPECS,
    _patent_transfer_quality_gate,
    _technology_infringement_remedy_context,
)
from review_v2.user_profile import UserRiskProfile


def facts(text):
    return extract_patent_transfer_protection_facts(parse_contract_clauses(text))


@pytest.mark.parametrize("event", ["双方签署本合同之日", "支付全部转让费之日", "技术资料交付之日"])
def test_non_registration_title_events_are_conflicting(event):
    actual = facts(f"第一条 标的\n专利权自{event}起归受让方所有。")
    assert actual["nonregistration_title_transfer"][0]["clause_id"]


@pytest.mark.parametrize("text", [
    "专利权转让合同自双方签署之日起生效。",
    "合同自双方签字盖章之日起生效。专利权自登记之日起转移给受让方。",
    "受让方可在签署合同后提前实施专利。专利权自登记之日起转移。",
    "专利权并非自签署合同之日起归受让方所有，而应以登记为准。",
    "合同签订后受让方支付价款，资料交付后办理登记。",
])
def test_contract_effectiveness_and_interim_license_are_not_title_transfer(text):
    assert not facts("第一条 生效\n" + text).get("nonregistration_title_transfer")


def test_later_registration_procedure_does_not_cure_a_signature_title_event():
    assert facts("第一条 标的\n专利权自签署之日起归买方所有，双方随后办理登记。")["nonregistration_title_transfer"]


@pytest.mark.parametrize("text,expected", [
    ("专利权全部被宣告无效后，受让方仍应支付尚未支付的全部转让费。", True),
    ("专利全部无效后，买方必须付清剩余价款。", True),
    ("专利全部被宣告无效后，已支付的转让费原则上不退。", False),
    ("专利全部无效后，受让方无需继续支付尚未支付的全部转让费。", False),
    ("专利部分被宣告无效，受让方应继续支付剩余价款。", False),
])
def test_unpaid_obligation_is_not_inferred_from_paid_fee_nonrefund(text, expected):
    assert bool(facts("第一条 专利无效\n" + text).get("invalidity_unpaid_fees_continue")) is expected


def test_cure_obligation_does_not_imply_deadline_termination_or_refund():
    rows = facts("第三条 资料交付\n资料验收不合格的，转让方负责补充，但未约定补充期限及次数。")
    assert rows["nonconforming_delivery_remedy"]
    for key in ("delivery_remedy_deadline", "delivery_failed_cure_termination", "delivery_failed_cure_refund"):
        assert not rows.get(key)


def test_deadline_does_not_imply_termination_or_refund():
    rows = facts("第三条 资料交付\n资料验收不合格的，转让方应在15日内补正。")
    assert rows["nonconforming_delivery_remedy"] and rows["delivery_remedy_deadline"]
    assert not rows.get("delivery_failed_cure_termination")
    assert not rows.get("delivery_failed_cure_refund")


def test_actual_failed_cure_rights_remain_covered_but_unrelated_refunds_do_not():
    rows = facts("第三条 资料交付\n资料验收不合格的，转让方应在15日内补正。逾期未完成补救，受让方有权解除合同，转让方应返还已收转让款。")
    assert rows["delivery_failed_cure_termination"] and rows["delivery_failed_cure_refund"]
    unrelated = facts("第三条 资料交付\n资料不合格时，转让方负责补充。\n第四条 不可抗力\n不可抗力超过90日可解除合同，转让方应返还预付款。")
    assert not unrelated.get("delivery_failed_cure_termination")
    assert not unrelated.get("delivery_failed_cure_refund")


@pytest.fixture(scope="module")
def result():
    doc = Document(next((Path(__file__).resolve().parents[2] / "test_data").glob("43*.docx")))
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n" + "\n".join("\t".join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
    return review_contract_v2(text, "technology", use_llm=False, include_low_risk=True,
                              profile=UserRiskProfile(user_identity="受让方"))


def item(result, key):
    return next(i for i in result["items"] if i["item_id"] == key)


def test_current_registration_conflict_is_not_pass(result):
    registration = item(result, "TC-06-04")
    assert registration["coverage_status"] == "风险条款"
    assert "专利法》第10条" in registration["legal_basis_index"]
    assert any(e["status"] == "CONFLICTING" for e in registration["coverage_elements"])
    assert not any(e["element_id"] == "registration_cost_and_effect" and e["status"] == "COVERED" for e in registration["coverage_elements"])
    assert any("TC-06-04" in r["review_ids"] and r["final_action"] in {"SHOULD_CHANGE", "MUST_CHANGE"} for r in result["final_dispositions"])
    payment = registration["decision_trace"]["payment_registration_alignment"]
    assert "70%" in str(payment["payment_nodes"])
    assert payment["priority"] == "SECONDARY_TO_TITLE_CONFLICT"


def test_current_invalidity_unpaid_fees_generate_actionable_risk(result):
    invalidity = item(result, "TC-06-03")
    assert invalidity["coverage_status"] == "风险条款"
    elements = {e["element_id"]: e for e in invalidity["coverage_elements"]}
    assert elements["invalidity_unpaid_fees_continue"]["status"] == "ADVERSE"
    assert elements["invalidity_costs_shifted"]["status"] == "ADVERSE"
    assert "专利法》第47条" in invalidity["legal_basis_index"]
    assert any("TC-06-03" in r["review_ids"] and r["final_action"] in {"SHOULD_CHANGE", "MUST_CHANGE"} for r in result["final_dispositions"])


def test_current_delivery_only_real_remedy_is_covered(result):
    delivery = item(result, "TC-06-02")
    assert delivery["coverage_status"] == "部分覆盖"
    elements = {e["element_id"]: e for e in delivery["coverage_elements"]}
    assert elements["nonconforming_delivery_remedy"]["status"] == "COVERED"
    for key in ("delivery_remedy_deadline", "delivery_failed_cure_termination", "delivery_failed_cure_refund"):
        assert elements[key]["status"] == "MISSING"
        assert not elements[key]["evidence_ids"]
    assert "nonconforming_delivery_remedy" not in delivery["decision_trace"]["final_missing_element_ids"]


def test_current_existing_and_retained_license_reduce_full_scope_coverage(result):
    scope = item(result, "TC-06-01")
    assert scope["coverage_status"] == "部分覆盖"
    keys = {e["element_id"] for e in scope["coverage_elements"]}
    assert {"transferor_retained_implementation", "existing_implementation_license"} <= keys


def test_existing_claim_and_daily_penalty_findings_remain_actionable(result):
    for key in ("TC-12-01", "GEN-LIAB-03"):
        assert any(key in r["review_ids"] and r["final_action"] == "SHOULD_CHANGE" for r in result["final_dispositions"])
    penalty = item(result, "GEN-LIAB-03")
    assert "0.1%" in str(penalty["evidence_clauses"])


def test_current_contract_passes_release_without_disabling_consistency_checks(result):
    assert result["release_quality_gate"]["status"] == "REPORT_READY"
    assert result["release_quality_gate"]["failures"] == []
    consistency = result["observability"]["cross_review_consistency_debug"]
    for key in ("applicability_misorders", "claim_conflicts", "element_identity_contradictions"):
        assert consistency[key] == []
    for key in ("APPLICABILITY_PRECONDITION_GATE", "PASS_EVIDENCE_PROPOSITION_MATCH_GATE",
                "PATENT_TRANSFER_QUALITY_GATE", "ELEMENT_IDENTITY_CONTRADICTION_GATE"):
        assert result["release_quality_gate"]["gate_results"][key] == "PASS"


@pytest.mark.parametrize("review_id", ["TC-10-04", "TC-12-04"])
def test_unmet_preconditions_are_not_later_projected_as_pass(result, review_id):
    review = item(result, review_id)
    assert review["applicability"]["status"].upper() == "NOT_APPLICABLE"
    assert review["coverage_presence"] == "NOT_APPLICABLE"
    assert review["coverage_status"] == "不适用"
    assert "不适用" in review["action_suggestion"]
    assert not any(review_id in f["source_review_ids"] for f in result["risk_findings"])


@pytest.mark.parametrize("review_id", ["TC-03-04", "TC-13-02"])
def test_current_proof_disclosures_are_atomic_and_never_independently_verified(result, review_id):
    review = item(result, review_id)
    elements = {e["element_id"]: e for e in review["coverage_elements"]}
    assert review["coverage_presence"] == "PARTIAL"
    assert "encumbrance_license_and_dispute_disclosure" not in elements
    for key in ("patent_pledge_disclosure", "patent_seizure_disclosure", "patent_ownership_dispute_disclosure"):
        assert elements[key]["status"] == "MISSING"
        assert elements[key]["evidence_ids"] == []
    for key in ("patent_existing_license_disclosure", "patent_invalidity_proceeding_disclosure"):
        assert elements[key]["status"] == "COVERED"
        assert elements[key]["evidence_level"] == "CONTRACTUAL_REPRESENTATION"
        assert key not in review["decision_trace"]["final_missing_element_ids"]
    assert elements["patent_invalidity_proceeding_disclosure"]["evidence_ids"] == ["C0010"]
    assert set(elements["patent_existing_license_disclosure"]["evidence_ids"]) == {"C0007", "C0010"}
    assert "质押、既有许可、权属争议和无效程序等权利负担已披露" not in result["result"]


@pytest.mark.parametrize("key,text", [
    ("patent_pledge_disclosure", "出让方声明该专利未设质押。"),
    ("patent_seizure_disclosure", "该专利目前处于司法保全状态。"),
    ("patent_ownership_dispute_disclosure", "双方确认目前不存在权属争议。"),
    ("patent_existing_license_disclosure", "让与人此前曾将专利许可给第三人使用。"),
    ("patent_invalidity_proceeding_disclosure", "该专利目前不存在已知无效宣告程序。"),
])
def test_each_disclosure_covers_only_its_own_element(key, text):
    rows = [{"clause_id": "right-status", "text": text}]
    matched = {
        element_id for element_id, _, terms in _PATENT_ENCUMBRANCE_DISCLOSURE_SPECS
        if _current_rights_disclosure_evidence(rows, terms)
    }
    assert matched == {key}


@pytest.mark.parametrize("text", [
    "提供专利证书及最近一期年费凭证。",
    "第九条 质押、权属争议和无效程序",
    "如专利被全部宣告无效，价款不退。专利无效程序及相关诉讼费用由受让方承担。",
    "受让方可以在过渡期内提前实施该专利，无需另行支付许可费用。",
    "转让方应披露质押及权属争议情况。",
])
def test_documents_headings_future_contingencies_and_promises_do_not_prove_disclosure(text):
    for _, _, terms in _PATENT_ENCUMBRANCE_DISCLOSURE_SPECS:
        assert not _current_rights_disclosure_evidence([{"clause_id": "unrelated", "text": text}], terms)


def coordination_gate_failures(result, context=None, coverage=None, add_link=False):
    reviews = [SimpleNamespace(**deepcopy(row)) for row in result["items"]]
    findings = [SimpleNamespace(**deepcopy(row)) for row in result["risk_findings"]]
    review = next(row for row in reviews if row.item_id == "TC-12-04")
    if context is not None:
        review.decision_trace["infringement_remedy_applicability"] = context
    if coverage is not None:
        review.coverage_status = coverage
    if add_link:
        review.review_risk_links.append({"risk_group": "transferee_infringement_liability_shift", "relation_type": "CONTRIBUTES_TO"})
    gate = _patent_transfer_quality_gate(
        {"tertiary_subtype": "PATENT_RIGHT_TRANSFER", "primary_subtype": "TECHNOLOGY_TRANSFER"},
        result["contract_context"], reviews, findings, None,
    )
    return [failure for failure in gate["failures"] if failure["requirement"] == "LINKED_RISK_PROPAGATION_GATE"]


def test_unrelated_infringement_risk_does_not_require_coordination_link(result):
    assert any(f["risk_group"] == "transferee_infringement_liability_shift" for f in result["risk_findings"])
    assert not coordination_gate_failures(result)


@pytest.mark.parametrize("kwargs", [{"context": {}}, {"coverage": "已覆盖"}, {"add_link": True}])
def test_gate_still_rejects_missing_precondition_trace_false_pass_and_fabricated_links(result, kwargs):
    assert coordination_gate_failures(result, **kwargs)


def test_actual_same_loss_remedies_cannot_be_hidden_as_not_applicable(result):
    text = ("第一条 侵权责任\n发生第三方侵权索赔，转让方承担全部侵权责任并赔偿损失。\n"
            "第二条 违约责任\n任何一方违反本合同任何义务，均应赔偿对方全部损失。")
    context = _technology_infringement_remedy_context(text)
    assert context["applicable"] and context["same_loss_general_remedy"]
    assert coordination_gate_failures(result, context=context)


def test_unrelated_payment_penalty_does_not_satisfy_same_loss_precondition():
    context = _technology_infringement_remedy_context(
        "第一条 侵权责任\n发生第三方侵权索赔，转让方承担赔偿责任。\n"
        "第二条 迟延付款\n受让方逾期付款，每日按应付款的0.1%支付违约金。"
    )
    assert context["specific_infringement_liability"]
    assert not context["applicable"] and not context["same_loss_general_remedy"]


def test_current_file_upload_and_response_schema_preserve_passed_gate():
    from fastapi import UploadFile
    from routers.review_v2 import ReviewV2Result, review_v2_file

    path = next((Path(__file__).resolve().parents[2] / "test_data").glob("43*.docx"))
    with path.open("rb") as source:
        response = asyncio.run(review_v2_file(
            file=UploadFile(filename=path.name, file=source), contract_type="technology",
            include_low_risk=True, use_llm=False,
            profile_json='{"user_identity":"受让方"}', current_user={},
        ))
    payload = ReviewV2Result(**response).model_dump()
    assert payload["release_quality_gate"]["status"] == "REPORT_READY"
    assert payload["release_quality_gate"]["failures"] == []
    for review_id in ("TC-10-04", "TC-12-04"):
        assert item(payload, review_id)["coverage_presence"] == "NOT_APPLICABLE"
    assert item(payload, "TC-03-04")["coverage_presence"] == "PARTIAL"
    test_current_registration_conflict_is_not_pass(payload)
    test_current_invalidity_unpaid_fees_generate_actionable_risk(payload)
    test_current_delivery_only_real_remedy_is_covered(payload)
    test_current_existing_and_retained_license_reduce_full_scope_coverage(payload)
    test_existing_claim_and_daily_penalty_findings_remain_actionable(payload)


def claim_facts(text):
    return extract_third_party_claim_protection_facts(parse_contract_clauses(text))


def test_current_claim_elements_reject_unrelated_remedies_and_share_real_evidence(result):
    for review_id in ("TC-12-01", "TC-12-02"):
        review = item(result, review_id)
        assert review["coverage_presence"] == "PARTIAL"
        for element in review["coverage_elements"]:
            if element["status"] in {"COVERED", "PARTIALLY_COVERED"}:
                assert set(element["evidence_ids"]) == {"C0010"}
            else:
                assert not element["evidence_ids"]
    first = {e["element_id"]: e for e in item(result, "TC-12-01")["coverage_elements"]}
    second = {e["element_id"]: e for e in item(result, "TC-12-02")["coverage_elements"]}
    assert first["transferor_claim_participation"]["status"] == "MISSING"
    assert first["liability_allocation_determinability"]["status"] == "MISSING"
    assert first["third_party_claim_trigger"]["evidence_ids"] == second["claim_trigger"]["evidence_ids"]
    for key in ("transferor_participation", "document_assistance", "damages_allocation", "refund_remedy", "termination_remedy", "rights_cure"):
        assert second[key]["status"] == "MISSING"
    assert "refund_or_termination_remedy" not in second
    for false_claim in ("让与方提供协助，但最终民事责任", "部分退款赔偿救济已有", "让与方必要协助、权利瑕疵责任及部分退款"):
        assert false_claim not in result["result"]


def test_current_material_types_do_not_depend_on_exhibit_completeness(result):
    review = item(result, "TC-13-01")
    elements = {e["element_id"]: e for e in review["coverage_elements"]}
    assert elements["major_document_types"]["status"] == "COVERED"
    assert elements["major_document_types"]["evidence_ids"] == ["C0004"]
    assert "major_document_types" not in review["decision_trace"]["final_missing_element_ids"]
    for key in ("delivery_quantity", "closed_attachment_schedule", "file_version", "file_identifier"):
        assert elements[key]["status"] == "MISSING"
    assert "未列举主要技术资料文件类型" not in result["result"]


def test_scope_narrative_does_not_invent_missing_territory_term_and_powers(result):
    scope = item(result, "TC-06-01")
    assert "商业排他性" in scope["risk_point_description"]
    assert "已有第三方实施许可" in scope["risk_point_description"]
    assert "地域、期限和权能不明" not in result["result"]


@pytest.mark.parametrize("duty,key", [
    ("受让方应在5个工作日内通知让与方。", "claim_notice_deadline"),
    ("让与方应参与处理并提供必要协助。", "claim_participation"),
    ("让与方应提供专利权属和技术来源文件。", "claim_document_assistance"),
    ("让与方负责抗辩。", "claim_defense_control"),
    ("让与方应承担合理律师费和诉讼费。", "claim_defense_cost"),
    ("让与方应承担全部侵权责任并赔偿损失。", "claim_damages"),
    ("让与方应退还全部转让价款。", "claim_refund"),
    ("受让方有权解除合同。", "claim_termination"),
    ("让与方应取得第三方授权，排除权利瑕疵。", "claim_rights_cure"),
])
def test_claim_protection_requires_its_own_affirmative_duty(duty, key):
    facts = claim_facts("第六条 第三方权利主张\n第三方对受让方实施专利提出侵权主张时，" + duty)
    assert facts[key][0]["clause_id"]
    for remedy in {"claim_refund", "claim_termination", "claim_damages"} - {key}:
        assert not facts.get(remedy)


@pytest.mark.parametrize("unrelated", [
    "第二条 登记\n转让方应提供必要协助，配合受让方办理专利权转让登记。",
    "第三条 保密\n泄密方应赔偿全部损失。",
    "第四条 不可抗力\n不可抗力持续超过90日，受让方有权解除合同，转让方应返还已收价款。",
    "第五条 无效\n专利全部被宣告无效后，转让方无需退还已付款项，未付款项仍应支付。",
    "第六条 交付\n资料验收不合格的，转让方应补救，无法补救时受让方有权解除合同并要求退款。",
])
def test_claim_response_cannot_borrow_another_event_remedies(unrelated):
    facts = claim_facts("第一条 第三方索赔\n转让方不保证不侵犯第三方知识产权。发生第三方侵权索赔时，双方另行协商责任承担。\n" + unrelated)
    assert set(facts) == {"claim_trigger", "claim_allocation_negotiation"}


@pytest.mark.parametrize("duty", [
    "转让方不负责处理，不承担律师费或赔偿责任。",
    "转让方无需退款，受让方不得解除合同。",
    "双方另行协商是否由转让方负责处理或赔偿损失。",
    "受让方自行负责抗辩并承担赔偿责任。",
])
def test_negative_or_unresolved_clauses_are_not_provider_protection(duty):
    facts = claim_facts("第六条 第三方索赔\n第三方提出侵权主张时，" + duty)
    for key in ("claim_participation", "claim_damages", "claim_defense_cost", "claim_refund", "claim_termination"):
        assert not facts.get(key)


def test_claim_scope_ends_at_a_different_event_even_inside_the_same_clause():
    facts = claim_facts(
        "第六条 协作事项\n第三方提出侵权主张时，双方另行协商责任。"
        "办理专利登记时，转让方应提供必要协助。"
        "专利全部被宣告无效后，转让方应返还未履行部分价款并赔偿损失。"
    )
    assert set(facts) == {"claim_trigger", "claim_allocation_negotiation"}


@pytest.mark.parametrize("documents,covered", [
    ("专利证书、专利说明书、权利要求书及附图", True),
    ("源代码", True),
    ("总体设计图、机械零件图", True),
    ("相关技术资料", False),
])
def test_identifiable_document_types_are_separate_from_versions_and_exhibits(documents, covered):
    text = ("专利权转让合同\n第一条 标的\n转让方将专利号ZL202010001234.5、名称为流量控制装置的专利权转让给受让方。\n"
            f"第二条 资料交付\n转让方应交付{documents}，通过电子邮件提供。\n"
            "第三条 价款及登记\n价款为100万元，专利权自登记完成时转移。")
    output = review_contract_v2(text, "technology", use_llm=False, include_low_risk=True)
    elements = {e["element_id"]: e for e in item(output, "TC-13-01")["coverage_elements"]}
    assert (elements["major_document_types"]["status"] == "COVERED") is covered
    assert elements["closed_attachment_schedule"]["status"] == "MISSING"
    assert elements["file_version"]["status"] == "MISSING"
