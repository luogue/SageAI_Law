from pathlib import Path

import pytest
from docx import Document

from review_v2.engine import review_contract_v2
from review_v2.user_profile import UserRiskProfile


ROOT = Path(__file__).resolve().parents[2]
FIXTURE = ROOT / "test_data" / "40 技术开发（合作）合同.docx"


def _document_text(path: Path) -> str:
    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = [
        "\t".join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
    ]
    return "\n".join(paragraphs + tables)


@pytest.fixture(scope="module")
def cooperative_review() -> dict:
    return review_contract_v2(
        _document_text(FIXTURE),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/合作开发方"),
    )


def _items(review: dict) -> dict:
    return {row["item_id"]: row for row in review["items"]}


def _elements(item: dict) -> dict:
    return {row["element_id"]: row for row in item["coverage_elements"]}


def test_cooperative_classification_uses_core_transaction_evidence(cooperative_review: dict):
    by_id = _items(cooperative_review)

    for review_id in ("TC-01-01", "TC-01-02"):
        evidence = by_id[review_id]["evidence_clauses"]
        locations = [str(row.get("location") or "") for row in evidence]
        rendered = "\n".join(
            str(row.get("quoted_text") or row.get("text") or "")
            for row in evidence
        )
        assert evidence
        assert not any("项目财产" in location for location in locations)
        assert any(
            term in rendered
            for term in ("项目要求", "各方分工", "货币资金", "阶段性成果", "最终成果")
        )

    assert "项目要求" in str(by_id["TC-01-01"]["evidence_clauses"][0].get("location") or "")
    assert "研究开发经费" in str(by_id["TC-01-02"]["evidence_clauses"][0].get("location") or "")


def test_investment_facts_override_absence_search(cooperative_review: dict):
    item = _items(cooperative_review)["TC-05-01"]
    elements = _elements(item)

    for element_id in (
        "research_task_allocation",
        "development_schedule_and_milestones",
        "funding_investment",
        "technical_and_personnel_investment",
        "data_equipment_environment_investment",
    ):
        assert elements[element_id]["status"] == "COVERED"
    assert elements["major_decision_process"]["status"] == "COVERED"
    assert elements["major_decision_process"]["direction"] == "FAVORABLE_TO_PARTY_A"
    assert elements["major_decision_process"]["risk_relevant"] is False

    missing = item["decision_trace"]["specialized_search_audit"]["missing_elements"]
    assert missing == [
        "新增或超预算研发费用的分担方式",
        "发生技术困难时的协作和帮助义务",
        "投入不足或未完成约定研发分工时的补足、接管和责任机制",
    ]
    assert "C0006" in {
        row["clause_id"] for row in item["evidence_clauses"]
        if "funding_investment" in (row.get("supporting_element_ids") or [])
    }


def test_confidentiality_damage_is_covered_without_duplicate_gap(cooperative_review: dict):
    by_id = _items(cooperative_review)
    obligation = _elements(by_id["TC-11-01"])
    liability_item = by_id["TC-11-04"]
    liability = _elements(liability_item)

    assert obligation["confidentiality_obligation"]["status"] == "COVERED"
    assert liability["confidentiality_obligation"]["status"] == "COVERED"
    assert liability["breach_damage_liability"]["status"] == "COVERED"
    assert "违约金" in liability["breach_damage_liability"]["proposition"]
    assert "不足部分" in liability["breach_damage_liability"]["proposition"]
    assert liability["specific_confidentiality_penalty"]["status"] == "COVERED"
    assert liability["specific_confidentiality_penalty"]["risk_relevant"] is False
    assert {
        element_id for element_id, row in liability.items()
        if row["status"] == "MISSING" and row.get("risk_relevant") is not False
    } == {"stop_disclosure", "stop_use", "mitigation_obligation", "incident_notification"}
    assert liability["return_or_destroy"]["status"] == "NOT_APPLICABLE"
    assert liability["disposition_confirmation"]["status"] == "NOT_APPLICABLE"
    assert "泄密损失赔偿责任" in liability_item["evidence_clauses"][0]["support_summary"]
    assert "赔偿责任缺失" not in liability_item["risk_point_description"]
    assert "已经覆盖泄密违约金及不足部分的全部损失赔偿" in liability_item["risk_point_description"]
    assert "合同虽未专项写明泄密损失赔偿" not in liability_item["legal_consequence_if_missing"]
    assert "合同虽未专项写明泄密损失赔偿" not in cooperative_review["result"]


def test_retention_does_not_imply_external_use_or_training(cooperative_review: dict):
    item = _items(cooperative_review)["TECH-SHARED-DATA-01"]
    elements = _elements(item)

    assert elements["post_termination_retention_and_deletion"]["status"] == "ADVERSE"
    for element_id in (
        "project_purpose_limitation",
        "research_training_or_other_project_use",
        "prior_written_consent",
        "anonymization_or_deidentification",
    ):
        assert elements[element_id]["status"] == "NOT_APPLICABLE"

    rendered = " ".join((
        item["risk_point_description"],
        item["legal_consequence_if_missing"],
        item["suggested_clause_example"],
    ))
    assert "可以用于合作项目之外" not in rendered
    assert "模型训练" not in rendered


def test_body_scope_and_allocation_are_not_reintroduced_as_attachment_gaps(cooperative_review: dict):
    item = _items(cooperative_review)["TC-13-01"]
    elements = _elements(item)

    assert elements["body_development_scope"]["status"] == "COVERED"
    assert elements["body_research_work_allocation"]["status"] == "COVERED"
    expected_missing = {
        "detailed_module_schedule",
        "interface_documentation",
        "source_code_version_list",
        "deployment_and_operation_documentation",
        "test_materials",
        "research_records",
        "stage_result_list",
    }
    assert {
        element_id for element_id, row in elements.items()
        if row["status"] == "MISSING"
    } == expected_missing
    assert "合同正文尚未明确基础研发范围" not in cooperative_review["result"]
    assert "合同正文尚未明确双方基本研发职责分工" not in cooperative_review["result"]


def test_explicit_improvement_and_high_salience_adverse_clauses_are_consumed(cooperative_review: dict):
    by_id = _items(cooperative_review)

    patent_rights = _elements(by_id["TC-05-02"])
    assert patent_rights["final_result_ownership"]["status"] == "COVERED"
    assert patent_rights["patent_application_right"]["status"] == "COVERED"
    assert "专利申请权归属已经明确" in patent_rights["patent_application_right"]["proposition"]
    assert "专利申请权未约定" not in cooperative_review["result"]
    assert by_id["TC-05-02"]["decision_trace"]["risk_to_current_party"]["current_party_directly_burdened"] is False

    result_rights = by_id["TC-05-03"]
    retention = _elements(result_rights)["material_copy_retention_and_disposal"]
    assert retention["status"] == "NOT_APPLICABLE"
    assert retention["owned_by_review_id"] == "TECH-SHARED-DATA-01"
    assert "C0005" not in {row["clause_id"] for row in result_rights["evidence_clauses"]}
    assert _elements(result_rights)["other_party_stage_result_use"]["status"] == "COVERED"
    result_absence = result_rights["decision_trace"]["specialized_search_audit"]["missing_elements"]
    assert not any("留存" in value or "返还" in value or "销毁" in value for value in result_absence)
    assert result_rights["decision_trace"]["risk_to_current_party"]["current_party_directly_burdened"] is False

    for review_id in ("TC-10-01", "TC-10-04"):
        item = by_id[review_id]
        elements = _elements(item)
        assert item["coverage_status"] == "风险条款"
        assert elements["continued_improvement_freedom"]["status"] == "COVERED"
        assert elements["universal_unilateral_improvement_assignment"]["status"] == "ADVERSE"
        assert [row["clause_id"] for row in item["evidence_clauses"]] == ["C0020"]

    expected = {"TECH-SHARED-ACCEPT-01": "C0014"}
    for review_id, clause_id in expected.items():
        item = by_id[review_id]
        assert item["coverage_status"] == "风险条款"
        assert [row["clause_id"] for row in item["evidence_clauses"]] == [clause_id]
        direction = item["decision_trace"]["risk_to_current_party"]
        assert direction["direct_contract_direction"] == "FAVORABLE_TO_PARTY_A"
        assert direction["current_party_directly_burdened"] is False


def test_sole_ownership_claim_cost_and_retention_branches_are_consistent(cooperative_review: dict):
    by_id = _items(cooperative_review)

    claim = by_id["TC-12-02"]
    claim_elements = _elements(claim)
    assert claim["coverage_status"] == "部分覆盖"
    for element_id in (
        "infringement_responsibility",
        "claim_cost_allocation",
        "third_party_compensation_allocation",
        "defense_cost_allocation",
    ):
        assert claim_elements[element_id]["status"] == "COVERED"
    assert "费用和赔偿责任承担已经覆盖" in claim["risk_point_description"]

    patent = by_id["TC-05-02"]
    patent_elements = _elements(patent)
    assert patent["coverage_status"] == "通过"
    assert patent["decision_trace"]["cooperative_final_rights_comparison"]["ownership_mode"] == "SOLE_OWNERSHIP"
    for element_id in (
        "patent_right_after_grant", "external_license_right",
        "result_transfer_right", "external_license_revenue",
    ):
        assert patent_elements[element_id]["status"] == "COVERED"
    assert "共同决策" not in patent["risk_point_description"]

    disposal = by_id["TC-11-03"]
    disposal_elements = _elements(disposal)
    assert disposal["coverage_status"] == "部分覆盖"
    assert disposal["needs_manual_review"] is False
    assert disposal_elements["original_material_return"]["status"] == "COVERED"
    assert {
        element_id for element_id, row in disposal_elements.items()
        if row["status"] == "MISSING"
    } == {
        "retained_copy_period", "retained_copy_access_control",
        "retained_copy_deletion_or_destroy", "written_disposition_confirmation",
    }


def test_duplicate_improvement_and_copy_retention_findings_are_aggregated(cooperative_review: dict):
    findings = cooperative_review["risk_findings"]

    improvement = [row for row in findings if row["risk_group"] == "followup_improvement_ownership"]
    assert len(improvement) == 1
    assert improvement[0]["source_review_ids"] == [
        "TC-02-01", "TC-10-01", "TC-10-02", "TC-10-04",
    ]
    assert "偏向完成方" not in improvement[0]["title"]
    assert "甲方无法继续使用" not in str(improvement[0])
    assert not any(
        row["risk_group"] in {"improvement_sharing_license", "improvement_no_unlawful_restriction"}
        for row in findings
    )

    retention = [row for row in findings if row["risk_group"] == "technology_data_secondary_use"]
    assert len(retention) == 1
    assert retention[0]["source_review_ids"] == ["TC-11-03", "TECH-SHARED-DATA-01"]


def test_party_a_favorable_rights_and_penalty_rendering_are_preserved(cooperative_review: dict):
    by_id = _items(cooperative_review)
    result_rights = by_id["TC-05-03"]
    result_finding = next(
        row for row in cooperative_review["risk_findings"]
        if row["risk_group"] == "cooperative_result_rights_and_commercialization"
    )

    assert result_rights["title"] == "最终成果归甲方的权利承接边界待完善"
    assert "对甲方有利" in result_rights["risk_point_description"]
    assert "不利安排" not in result_rights["risk_point_description"]
    assert "背景技术" in result_rights["legal_consequence_if_missing"]
    assert "独立成果" in result_rights["legal_consequence_if_missing"]
    assert "项目依赖型成果" in result_rights["legal_consequence_if_missing"]
    assert "保留甲方" in result_rights["suggested_clause_example"]
    assert "背景技术" in result_finding["issue"]
    assert "项目依赖型成果" in result_finding["issue"]
    assert "不利安排" not in result_finding["issue"]

    penalty = next(
        row for row in cooperative_review["risk_findings"]
        if row["risk_group"] == "generic_potential_excessive_liquidated_damages"
    )
    assert "乙方未按期完成研发任务时" in penalty["issue"]
    assert "乙方未时" not in penalty["issue"]
    assert "乙方未时" not in cooperative_review["result"]
    assert "当前审查方受益" in penalty["issue"]
    assert "保留当前审查方可主张的违约救济" in penalty["recommendation"]


def test_attachment_improvements_remain_optional(cooperative_review: dict):
    dispositions = {
        review_id: row
        for row in cooperative_review["final_dispositions"]
        for review_id in row["review_ids"]
    }
    findings = {
        row["risk_group"]: row for row in cooperative_review["risk_findings"]
    }

    for review_id in ("TC-13-01", "TC-13-03"):
        disposition = dispositions[review_id]
        assert disposition["severity"] == "LOW"
        assert disposition["final_action"] == "OPTIONAL_IMPROVEMENT"
        assert disposition["include_in_optional_improvements"] is True
        assert disposition["include_in_priority_changes"] is False

    assert findings["technical_material_attachment"]["severity"] == "LOW"
    assert findings["acceptance_test_attachment"]["severity"] == "LOW"


def test_report_layers_and_preserved_risks_remain_consistent(cooperative_review: dict):
    assert cooperative_review["release_quality_gate"]["status"] == "REPORT_READY"
    gates = cooperative_review["release_quality_gate"]["gate_results"]
    for gate in (
        "CROSS_LAYER_CONSISTENCY_GATE",
        "CROSS_RULE_EVIDENCE_CONSISTENCY",
        "REPORT_EVIDENCE_CONSISTENCY",
        "SUMMARY_DETAIL_RISK_ID_GATE",
        "PARTIAL_MISSING_ELEMENT_GATE",
    ):
        assert gates[gate] == "PASS"

    groups = {row["risk_group"] for row in cooperative_review["risk_findings"]}
    assert "cooperative_result_rights_and_commercialization" in groups
    assert "generic_potential_excessive_liquidated_damages" in groups
    assert "许可方" not in cooperative_review["result"]
    assert "被许可方" not in cooperative_review["result"]
    liquidated_damages = next(
        row for row in cooperative_review["risk_findings"]
        if row["risk_group"] == "generic_potential_excessive_liquidated_damages"
    )
    quotes = "\n".join(
        str(row.get("quoted_text") or "")
        for row in liquidated_damages["evidence_bundle"]["items"]
    )
    assert "每逾期一日按合同研发总金额的1%" in quotes


def test_party_a_failure_allocation_is_favorable_and_article_858_controls(
    cooperative_review: dict,
):
    item = _items(cooperative_review)["TECH-SHARED-DEV-FAIL-01"]
    elements = _elements(item)

    assert item["coverage_status"] == "部分覆盖"
    assert item["risk_level"] == "中"
    assert "乙方承担研发损失本身描述为甲方重大风险" in item["legal_consequence_if_missing"]
    assert "该分配对甲方有利" in item["risk_point_description"]
    assert elements["incurred_cost_allocation"]["status"] == "COVERED"
    assert elements["failure_notice"]["status"] == "COVERED"
    assert item["decision_trace"]["risk_to_current_party"]["direct_contract_direction"] == "FAVORABLE_TO_PARTY_A"
    assert {
        element_id for element_id, row in elements.items()
        if row["status"] == "MISSING"
    } == {
        "pause_further_investment",
        "loss_mitigation_measures",
        "future_unincurred_cost_allocation",
        "termination_decision",
        "stage_result_and_handover",
    }
    assert item["legal_basis_index"].startswith("《中华人民共和国民法典》第858条（核心依据）")
    assert item["legal_basis_index"].count("第858条") == 1
    assert item["legal_basis_index"].count("第509条") == 1
    article_858 = next(
        row for row in item["legal_references"]
        if row.get("article") == "第858条"
    )
    assert article_858["citation_existence_status"] == "FOUND"
    assert article_858["citation_entailment_status"] == "DIRECT_SUPPORT"
    assert item["narrative_state"]["consequence"] == item["legal_consequence_if_missing"]
    assert "付款期限、逾期责任和对账依据" not in item["narrative_state"]["consequence"]
    finding = next(
        row for row in cooperative_review["risk_findings"]
        if row["risk_group"] == "cooperative_development_failure_allocation"
    )
    assert finding["issue"] == item["risk_point_description"]
    assert finding["consequence"] == item["legal_consequence_if_missing"]
    assert finding["recommendation"] == item["suggested_clause_example"]


def test_party_a_ownership_and_final_decision_right_are_not_adverse(
    cooperative_review: dict,
):
    by_id = _items(cooperative_review)
    result_rights = by_id["TC-05-03"]
    result_elements = _elements(result_rights)
    governance = by_id["TECH-SHARED-GOV-01"]
    governance_elements = _elements(governance)

    assert result_rights["coverage_status"] == "部分覆盖"
    assert result_rights["risk_level"] == "低"
    assert result_elements["final_result_ownership_structure"]["status"] == "COVERED"
    assert result_elements["final_result_ownership_structure"]["direction"] == "FAVORABLE_TO_SOLE_OWNER"
    assert {
        element_id for element_id, row in result_elements.items()
        if row["status"] == "MISSING"
    } == {"background_technology_boundary", "project_dependent_result_boundary"}
    assert "对甲方有利" in result_rights["risk_point_description"]
    assert "保留甲方" in result_rights["suggested_clause_example"]
    result_finding = next(
        row for row in cooperative_review["risk_findings"]
        if row["risk_group"] == "cooperative_result_rights_and_commercialization"
    )
    assert result_finding["recommendation"] == result_rights["suggested_clause_example"]

    assert governance["coverage_status"] == "部分覆盖"
    assert governance["risk_level"] == "低"
    assert governance_elements["major_matter_voting"]["status"] == "COVERED"
    assert governance_elements["major_matter_voting"]["direction"] == "FAVORABLE_TO_DECISION_OWNER"
    assert governance["decision_trace"]["risk_to_current_party"]["direct_contract_direction"] == "FAVORABLE_TO_PARTY_A"
    assert not any(row["status"] == "ADVERSE" for row in governance_elements.values())
    assert {
        element_id for element_id, row in governance_elements.items()
        if row["status"] == "MISSING"
    } == {
        "major_decision_notice_and_record",
        "major_decision_objection_or_expert_review",
        "major_decision_authority_boundary",
    }
    assert "保留甲方的最终决定权" in governance["suggested_clause_example"]
    governance_finding = next(
        row for row in cooperative_review["risk_findings"]
        if row["risk_group"] == "cooperative_governance_control_conflict"
    )
    assert governance_finding["recommendation"] == governance["suggested_clause_example"]


def test_universal_independent_improvement_assignment_links_competition_and_deduplicates(
    cooperative_review: dict,
):
    by_id = _items(cooperative_review)
    related = [by_id[review_id] for review_id in (
        "TC-02-01", "TC-10-01", "TC-10-02", "TC-10-04",
    )]

    assert by_id["TC-02-01"]["coverage_status"] == "风险条款"
    scope = by_id["TC-02-01"]["decision_trace"]["improvement_assignment_scope"]
    assert scope == {
        "project_dependent_improvement_limited": False,
        "independent_improvement_excluded": False,
        "reasonable_consideration_present": False,
        "implicit_uncompensated_assignment": True,
    }
    assert len({row["title"] for row in related}) == 1
    assert len({row["legal_basis_index"] for row in related}) == 1

    finding = next(
        row for row in cooperative_review["risk_findings"]
        if row["risk_group"] == "followup_improvement_ownership"
    )
    assert finding["source_review_ids"] == [
        "TC-02-01", "TC-10-01", "TC-10-02", "TC-10-04",
    ]
    assert finding["severity"] == "MEDIUM"


def test_report_deduplicates_basis_and_evidence_and_uses_actual_acceptance_articles(
    cooperative_review: dict,
):
    by_id = _items(cooperative_review)
    acceptance = _elements(by_id["TC-13-03"])["acceptance_test_scene"]

    assert "第十三条、第十四条" in acceptance["proposition"]
    assert "第十二条" not in acceptance["proposition"]
    assert "甲乙双方" not in cooperative_review["result"]
    assert "合同正文已经明确各方的基本研发职责分工" in cooperative_review["result"]
    for line in cooperative_review["result"].splitlines():
        assert line.count("第509条") <= 1

    for finding in cooperative_review["risk_findings"]:
        quotes = [
            str(row.get("quoted_text") or "").strip()
            for row in finding.get("evidence_bundle", {}).get("items", [])
            if str(row.get("quoted_text") or "").strip()
        ]
        assert len(quotes) == len(set(quotes))


def test_report_distinguishes_aggregated_themes_from_rule_level_hits(
    cooperative_review: dict,
):
    rule_ids = list(dict.fromkeys(
        review_id
        for finding in cooperative_review["risk_findings"]
        for review_id in finding["source_review_ids"]
    ))

    assert cooperative_review["risk_theme_count"] == len(cooperative_review["risk_themes"])
    assert cooperative_review["aggregated_finding_count"] == len(cooperative_review["risk_findings"])
    assert cooperative_review["finding_count"] == cooperative_review["aggregated_finding_count"]
    assert cooperative_review["finding_count_semantics"] == "AGGREGATED_RISK_FINDINGS"
    assert cooperative_review["rule_level_finding_count"] == len(rule_ids)
    assert cooperative_review["rule_level_finding_ids"] == rule_ids
    assert cooperative_review["rule_level_finding_count"] > cooperative_review["risk_theme_count"]

    report = cooperative_review["result"]
    assert f"风险主题数（聚合后）：{cooperative_review['risk_theme_count']}" in report
    assert f"规则级审查命中数：{len(rule_ids)}" in report
    assert "完整风险主题明细（聚合后）" in report
    assert "底层审查发现" not in report
    assert "《民法典》、第" not in report


def test_result_rights_applicability_matches_actual_review_scope(cooperative_review: dict):
    applicable_if = _items(cooperative_review)["TC-05-03"]["applicable_if"]

    for term in ("最终成果", "背景技术", "既有知识产权", "独立成果"):
        assert term in applicable_if
    assert "仅适用于产生技术秘密成果" not in applicable_if
