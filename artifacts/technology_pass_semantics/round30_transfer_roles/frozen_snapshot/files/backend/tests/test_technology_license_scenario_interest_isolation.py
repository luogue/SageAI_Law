from pathlib import Path

from docx import Document

from review_v2.engine import (
    ReviewItemResult,
    _apply_technology_party_direction_to_items,
    review_contract_v2,
)
from review_v2.perspective_resolution import resolve_perspective
from review_v2.user_profile import UserRiskProfile


ROOT = Path(__file__).resolve().parents[2]
GOLD = ROOT / "test_data" / "19 技术许可.docx"
TRANSFER = ROOT / "test_data" / "9、技术转让.docx"


def _docx_text(path: Path) -> str:
    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = [
        " | ".join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
    ]
    return "\n".join([*paragraphs, *tables])


def _review(text: str, role: str = "甲方/受让方") -> dict:
    return review_contract_v2(
        text,
        "technology",
        include_low_risk=True,
        use_llm=False,
        profile=UserRiskProfile(user_identity=role),
    )


def _item(item_id: str, coverage: str, risk_level: str = "高") -> ReviewItemResult:
    return ReviewItemResult(
        item_id=item_id,
        node_id=item_id,
        title=item_id,
        risk_level=risk_level,
        status=coverage,
        risk_point_description="测试风险",
        suggested_clause_example="建议明确责任边界。",
        legal_basis_index="",
        legal_consequence_if_missing="测试后果",
        applicable_if="技术许可",
        matched_terms=[],
        coverage_status=coverage,
        issue_type="MISSING",
        evidence_clauses=[],
        evidence_location="",
        reasoning_summary="",
        risk_group="license_over_scope_liability" if item_id == "TC-07-03" else item_id,
        coverage_elements=[{
            "element_id": f"{item_id}.core",
            "status": "MISSING" if coverage != "通过" else "COVERED",
            "proposition": "责任边界",
        }],
    )


PATENT_LICENSE = """
专利实施许可合同
甲方（受让方）为被许可方，乙方（出让方）为许可方。
乙方系名称为工业控制方法、专利号CN202610000001.X的专利权人，授予甲方中国境内五年非独占实施许可。
乙方保证有权许可，并应披露共有权人、既有许可、独占或排他许可、质押、查封及其他权利负担。
如乙方并非登记专利权人，应提供上游授权和再许可权限文件。
甲方仅可在约定地域和产品范围内使用，不得擅自再许可。
甲方支付许可费100万元。双方的一般违约责任适用于违反许可范围的行为。
"""


def test_a_license_role_and_patent_proof_requirements_are_scenario_isolated():
    resolution = resolve_perspective(
        "甲方/受让方", "TECHNOLOGY", contract_subtype="PATENT_LICENSE",
    ).to_public_dict()
    output = _review(PATENT_LICENSE)
    by_id = {row["item_id"]: row for row in output["items"]}
    generated = "\n".join(
        str(value)
        for review_id in ("TC-03-04", "TC-13-02")
        for value in (
            by_id[review_id]["reasoning_summary"],
            by_id[review_id]["risk_point_description"],
            by_id[review_id]["legal_consequence_if_missing"],
            by_id[review_id]["suggested_clause_example"],
        )
    )

    assert resolution["transaction_type"] == "TECHNOLOGY_LICENSE"
    assert resolution["review_party"] == "PARTY_A"
    assert resolution["legal_role"] == "LICENSEE"
    assert resolution["contract_display_role"] == "受让方（被许可方）"
    assert resolution["canonical_role"] == "PATENT_LICENSEE"
    assert "专利权受让方" not in output["result"]
    assert "软件著作权" not in generated
    assert "开源组件" not in generated
    assert "让与方处分权" not in output["result"]
    assert "转让登记失败" not in output["result"]
    assert output["terminology_consistency_gate"]["status"] == "REPORT_READY"


def test_b_patent_transfer_keeps_transfer_legal_semantics():
    output = _review(_docx_text(TRANSFER), "甲方/专利权受让方")
    resolution = output["perspective_resolution"]

    assert resolution["transaction_type"] == "TECHNOLOGY_TRANSFER"
    assert resolution["legal_role"] == "TRANSFEREE"
    assert resolution["canonical_role"] == "PATENT_TRANSFEREE"
    assert "专利权受让方" in resolution["display_role"]
    assert any(term in output["result"] for term in ("转让登记", "权利转移", "让与方"))


def test_c_software_license_facts_enable_software_and_open_source_requirements():
    output = _review(
        """
        软件许可合同
        乙方许可甲方在中国境内使用软件及源代码，许可期限三年，许可费50万元。
        软件包含第三方组件和开源组件，乙方应提供软件著作权证明、第三方授权清单、开源许可证及SBOM。
        甲方不得超范围使用或擅自再许可，双方一般违约责任适用于本合同全部义务。
        """,
        "甲方/被许可方",
    )
    assert output["deterministic_analysis"]["technology_classification"]["involves_software_or_data"] is True
    text = "\n".join(
        str(value)
        for row in output["items"] if row["item_id"] in {"TC-03-04", "TC-13-02"}
        for value in (row["reasoning_summary"], row["suggested_clause_example"], row["decision_trace"])
    )
    assert "软件著作权" in text
    assert "开源" in text


def test_d_party_interest_check_distinguishes_uncapped_and_low_capped_licensor_liability():
    uncapped = _review(_docx_text(GOLD), "甲方/被许可方")
    uncapped_item = next(row for row in uncapped["items"] if row["item_id"] == "TC-12-04")

    assert uncapped_item["coverage_status"] == "不适用"
    assert uncapped_item["evidence_clauses"] == []
    assert all(term in uncapped_item["reasoning_summary"] for term in ("同一", "损失", "侵权", "一般救济"))
    assert "未限制许可方责任" not in uncapped["result"]

    low_cap = _review(
        PATENT_LICENSE
        + "\n许可方对侵权及权利瑕疵承担的全部赔偿责任最高不超过许可费的10%。"
        + "\n被许可方按约实施遭第三方权利主张时，由许可方负责处理并承担责任。",
        "甲方/被许可方",
    )
    low_item = next(row for row in low_cap["items"] if row["item_id"] == "TC-12-04")
    low_elements = {row["element_id"]: row for row in low_item["coverage_elements"]}
    assert low_item["decision_trace"]["liability_cap_facts"]["cap_ratio"] == 10
    assert low_elements["licensor_liability_cap"]["direction"] == "UNFAVORABLE"
    assert low_elements["licensor_liability_cap"]["risk_relevant"] is True
    assert low_elements["licensor_liability_cap"]["arrangement_status"] == "ADVERSE"


def test_e_missing_special_liability_is_not_mechanically_major_when_protected():
    protected = _review(_docx_text(GOLD), "甲方/被许可方")
    protected_item = next(row for row in protected["items"] if row["item_id"] == "TC-07-03")

    assert protected_item["coverage_status"] == "缺失"
    assert protected_item["risk_level"] == "低"
    assert protected_item["risk_severity"] == "LOW"
    assert protected_item["risk_direction"]["gate_result"] == "SUPPRESS_LICENSEE_ADVERSE_RISK"
    assert all("TC-07-03" not in row["source_review_ids"] for row in protected["risk_findings"])
    assert protected_item["decision_trace"]["severity_calibration"]["existing_protection"] == {
        "license_scope_clear": True,
        "sublicense_controlled": True,
        "general_remedy_fallback": True,
    }
    assert protected_item["decision_trace"]["severity_calibration"]["statutory_liability"]["authority"] == "《中华人民共和国民法典》第873条"
    assert "不主动新增或扩大被许可方赔偿责任" in protected_item["suggested_clause_example"]

    licensor = _review(_docx_text(GOLD), "乙方/许可方")
    licensor_item = next(row for row in licensor["items"] if row["item_id"] == "TC-07-03")
    licensor_finding = next(row for row in licensor["risk_findings"] if "TC-07-03" in row["source_review_ids"])
    assert licensor_item["risk_level"] == "高"
    assert licensor_item["risk_direction"]["gate_result"] == "ALLOW_LICENSOR_PROTECTION_RISK"
    assert licensor_finding["severity"] == "HIGH"
    assert "擅自再许可" in licensor_finding["recommendation"]

    exposed_items = [
        _item("TC-07-01", "缺失"),
        _item("TC-07-03", "缺失"),
        _item("TC-07-04", "风险条款"),
        _item("TC-12-04", "缺失"),
    ]
    _apply_technology_party_direction_to_items(
        exposed_items,
        {
            "display_role": "甲方 / 被许可方", "canonical_role": "PATENT_LICENSEE",
            "review_party_identity": "PARTY_A", "party_side": "PARTY_A",
            "transaction_type": "TECHNOLOGY_LICENSE", "review_party": "PARTY_A", "legal_role": "LICENSEE",
        },
        {"primary_subtype": "TECHNOLOGY_LICENSE"},
    )
    calibration = next(row for row in exposed_items if row.item_id == "TC-07-03").decision_trace["severity_calibration"]
    assert calibration["major_severity_reason"] == ""
    assert calibration["result"] == "LICENSEE_LOW_BOUNDARY_NOTICE"


def test_f_formal_recommendations_are_actionable_and_not_risk_copies():
    output = _review(_docx_text(GOLD), "甲方/被许可方")
    assert output["formal_recommendation_validation"]["status"] == "PASS"
    by_id = {row["finding_id"]: row for row in output["risk_findings"]}
    for disposition in output["final_dispositions"]:
        if disposition["final_action"] not in {"MUST_CHANGE", "SHOULD_CHANGE", "OPTIONAL_IMPROVEMENT"}:
            continue
        recommendation = disposition["recommendation"]
        finding = by_id[disposition["risk_id"]]
        assert recommendation
        assert recommendation != finding["issue"]
        assert any(action in recommendation for action in ("明确", "补充", "删除", "调整", "限制", "约定", "列明", "提供", "核验", "承担", "通知", "校正", "统一", "建立", "设置"))
