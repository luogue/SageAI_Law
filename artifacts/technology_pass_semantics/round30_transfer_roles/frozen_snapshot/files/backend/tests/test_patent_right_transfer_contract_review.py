from copy import deepcopy
import json
from pathlib import Path

from docx import Document

from review_v2.clause_parser import parse_contract_clauses
from review_v2.engine import (
    COVERAGE_NOT_APPLICABLE,
    FinalDisposition,
    ReviewItemResult,
    _apply_coverage_gap_consistency_guard,
    _build_narrative_state,
    _build_risk_findings,
    _contract_allocation_from_text,
    _final_disposition_invariant_gate,
    _patent_prepayment_economic_impact,
    _report_disposition_projection_gate,
    _report_view_model,
    review_contract_v2,
)
from review_v2.technology_service_analysis import classify_technology_transaction
from review_v2.technology_patent_transfer_analysis import analyze_patent_transfer
from review_v2.user_profile import UserRiskProfile


ROOT = Path(__file__).resolve().parents[2]
PATENT_FIXTURE = ROOT / "test_data" / "9、技术转让.docx"
CURRENT_PATENT_TRANSFER_FIXTURE = ROOT / "test_data" / "21 技术转让（专利权）.docx"
SERVICE_FIXTURE = ROOT / "test_data" / "8、技术服务.docx"


def _fixture_text(path: Path) -> str:
    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = ["\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    return "\n".join(paragraphs + tables)


def _patent_review() -> dict:
    return review_contract_v2(
        _fixture_text(PATENT_FIXTURE),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(
            user_identity="甲方/专利权受让方",
            contract_scenario="高校科技成果向企业协议转让",
        ),
    )


def _service_review() -> dict:
    return review_contract_v2(
        _fixture_text(SERVICE_FIXTURE),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(
            user_identity="乙方/技术服务提供方",
            contract_scenario="高校科研数据治理技术服务",
        ),
    )


def _current_patent_transfer_review() -> dict:
    return review_contract_v2(
        _fixture_text(CURRENT_PATENT_TRANSFER_FIXTURE),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(
            user_identity="甲方/专利权受让方",
            contract_scenario="专利权转让",
        ),
    )


def test_patent_transfer_has_single_three_level_classification_and_transferee_role():
    output = _patent_review()
    classification = output["deterministic_analysis"]["technology_classification"]
    assert classification["primary_category"] == "TECHNICAL_CONTRACT"
    assert classification["primary_subtype"] == "TECHNOLOGY_TRANSFER"
    assert classification["tertiary_subtype"] == "PATENT_RIGHT_TRANSFER"
    assert set(classification["excluded_subtypes"]) >= {
        "ENTRUSTED_DEVELOPMENT",
        "COOPERATIVE_DEVELOPMENT",
        "TECHNOLOGY_LICENSE",
        "TECHNICAL_CONSULTING",
        "TECHNICAL_SERVICE",
    }
    assert output["contract_context"]["user_role"] == "patent_transferee"
    assert output["contract_context"]["counterparty_role"] == "patent_transferor"
    assert output["report_context"]["contract_type_label"] == "技术合同—技术转让合同—专利权转让合同"
    assert output["report_context"]["scenario_label"] == "高校科技成果向企业协议转让"


def test_patent_transfer_golden_matrix_semantic_invariants():
    output = _patent_review()
    manifest = json.loads(
        (ROOT / "backend" / "tests" / "fixtures" / "technical_contracts" / "patent_transfer.golden.json")
        .read_text(encoding="utf-8")
    )
    classification = output["deterministic_analysis"]["technology_classification"]
    by_id = {row["item_id"]: row for row in output["items"]}
    applicable = {review_id for review_id, row in by_id.items() if row["coverage_status"] != COVERAGE_NOT_APPLICABLE}
    not_applicable = {review_id for review_id, row in by_id.items() if row["coverage_status"] == COVERAGE_NOT_APPLICABLE}

    assert manifest["sample_id"] == "TECH_TRANSFER_PATENT_GOLDEN_001"
    assert classification["primary_contract_type"] == manifest["primary_contract_type"]
    assert classification["primary_subtype"] == manifest["primary_subtype"]
    assert classification["tertiary_subtype"] == manifest["contract_subtype"]
    assert output["perspective_resolution"]["canonical_role"] == manifest["review_position"]
    assert output["coverage_summary"] == manifest["coverage_summary"]
    assert applicable == set(manifest["expected_applicable_rules"])
    assert not_applicable == set(manifest["expected_not_applicable_rules"])
    assert {row["risk_group"] for row in output["risk_findings"]} == set(manifest["expected_risk_groups"])
    assert len(output["risk_findings"]) == manifest["risk_count"]
    assert len(output["recommendations"]) == manifest["recommendation_count"]
    assert all(row["party_effect"] == "ADVERSE" for row in output["risk_findings"])
    assert all((row.get("risk_direction") or {}).get("direction") == "UNFAVORABLE" for row in output["risk_findings"])
    assert all(by_id[review_id]["coverage_status"] == COVERAGE_NOT_APPLICABLE for review_id in ("TC-07-01", "TC-07-02", "TC-07-03", "TC-07-04"))
    assert "适用于专利权转让中受让方按约实施受让专利" in by_id["TC-12-01"]["applicable_if"]
    assert "影响受让方行使和实施受让专利权" in by_id["TC-12-02"]["legal_consequence_if_missing"]
    assert output["release_quality_gate"]["status"] == "REPORT_READY"


def test_review_context_isolation_between_technical_service_and_patent_transfer():
    service = _service_review()
    patent = _patent_review()
    service_again = _service_review()

    assert service["contract_context"]["contract_subtype"] == "TECHNICAL_SERVICE"
    assert patent["contract_context"]["contract_subtype"] == "TECHNOLOGY_TRANSFER"
    assert patent["contract_context"]["user_role"] == "patent_transferee"
    assert service_again["contract_context"]["user_role"] == "service_provider"
    forbidden = ("服务提供方", "技术服务提供方", "高校科研数据治理技术服务", "乙方受托人")
    assert all(token not in patent["result"] for token in forbidden)

    patent_first = _patent_review()
    service_second = _service_review()
    assert patent_first["contract_context"]["user_role"] == "patent_transferee"
    assert service_second["contract_context"]["contract_subtype"] == "TECHNICAL_SERVICE"
    assert "专利权受让方" not in service_second["result"]


def test_historical_license_is_ancillary_and_subtype_packs_are_isolated():
    output = _patent_review()
    analysis = output["deterministic_analysis"]["patent_transfer"]
    assert analysis["ancillary_transactions"]["existing_patent_license"] is True
    assert analysis["ancillary_transactions"]["patent_pledge"] is True
    assert analysis["existing_license"]["is_current_transaction"] is False
    by_id = {item["item_id"]: item for item in output["items"]}
    for prefix in ("TC-04", "TC-05", "TC-07", "TC-08", "TC-09"):
        assert all(
            item["coverage_status"] == COVERAGE_NOT_APPLICABLE
            for item_id, item in by_id.items()
            if item_id.startswith(prefix)
        )
    assert all(by_id[item_id]["coverage_status"] != COVERAGE_NOT_APPLICABLE for item_id in by_id if item_id.startswith("TC-06"))


def test_patent_assets_pledge_existing_license_and_payment_are_structured():
    analysis = _patent_review()["deterministic_analysis"]["patent_transfer"]
    assets = analysis["patent_assets"]
    assert len(assets) == 2
    assert {row["patent_type"] for row in assets} == {"发明专利", "实用新型专利"}
    assert {row["transfer_price"] for row in assets} == {900000, 300000}
    assert analysis["total_transfer_price"] == 1200000
    assert analysis["patent_pledge"]["exists"] is True
    assert analysis["patent_pledge"]["pledgee"]
    assert analysis["patent_pledge"]["registration_no"]
    assert analysis["patent_pledge"]["release_after_first_payment"] is True
    assert analysis["existing_license"]["exists"] is True
    assert analysis["existing_license"]["license_type"] == "普通实施许可"
    assert analysis["existing_license"]["licensee"]
    assert analysis["existing_license"]["fee_retained_by_transferor"] is True
    assert analysis["existing_license"]["obligations_assumed_by_transferee"] is True
    assert analysis["payment"]["status"] == "CONSISTENT"


def test_patent_transfer_coverage_distinguishes_presence_from_adverse_terms():
    output = _patent_review()
    by_id = {item["item_id"]: item for item in output["items"]}
    for item_id in ("TC-01-03", "TC-06-02", "TC-06-03", "TC-06-04", "TC-10-01", "TC-11-01", "TC-11-02"):
        assert by_id[item_id]["coverage_presence"] == "PRESENT"
    assert by_id["TC-06-03"]["risk_status"] in {"ADVERSE", "LINKED_ADVERSE_RISK"}
    assert by_id["TC-10-01"]["risk_status"] in {"ADVERSE", "LINKED_ADVERSE_RISK"}


def test_patent_transfer_eight_gold_risks_have_claim_complete_evidence_bundles():
    output = _patent_review()
    findings = {row["risk_group"]: row for row in output["risk_findings"]}
    required = {
        "patent_pledge_prepayment",
        "existing_license_value_burden",
        "patent_invalidity_broad_disclaimer",
        "transferee_infringement_liability_shift",
        "transferor_retained_use_conflict",
        "asymmetric_breach_liability",
        "dispute_resolution_choice_conflict",
        "registration_completion_continuing_obligations_conflict",
    }
    assert required <= findings.keys()
    for group in required:
        bundle = findings[group]["evidence_bundle"]
        assert bundle["entailment_status"] == "PASSED"
        assert bundle["items"]
        assert set(bundle["required_claim_ids"]) <= set(bundle["supported_claim_ids"])
        assert len(bundle["items"]) == len({
            (row["clause_node_id"], row["start_offset"], row["end_offset"], row["quoted_text_hash"])
            for row in bundle["items"]
        })


def test_default_law_and_negative_assertions_for_patent_transfer():
    output = _patent_review()
    findings = {row["risk_group"]: row for row in output["risk_findings"]}
    infringement = findings["transferee_infringement_liability_shift"]
    assert infringement["issue_type"] == "DEFAULT_LAW_DEVIATION"
    assert infringement["default_law_status"] == "CONTRACT_DEVIATES_FROM_DEFAULT"
    assert infringement["agreement"]["affected_party"] == "PATENT_TRANSFEREE"

    assert output["deterministic_analysis"]["payment_formula"]["status"] == "CONSISTENT"
    assert not output["deterministic_analysis"]["patent_transfer"]["timeline"]["conflicts"]
    groups = set(findings)
    assert "payment_formula_inconsistency" not in groups
    assert "legal_reference" not in groups
    assert "personal_information_processing_gap" not in groups
    assert "technical_timeline_conflict" not in groups
    assert output["release_quality_gate"]["status"] == "REPORT_READY"


def test_classifier_generalizes_without_fixture_identity_or_values():
    classification = classify_technology_transaction(
        "技术转让（专利权）合同。受让方支付转让价款，让与方配合办理专利权变更登记。"
        "标的专利此前存在普通实施许可，该历史许可继续履行。"
    )
    assert classification["primary_subtype"] == "TECHNOLOGY_TRANSFER"
    assert classification["tertiary_subtype"] == "PATENT_RIGHT_TRANSFER"
    assert classification["current_transaction"] != "TECHNOLOGY_LICENSE"
    assert classification["referenced_contracts"] == [{
        "reference_type": "EXISTING_PATENT_LICENSE",
        "is_current_transaction": False,
    }]


def test_dispute_options_and_registration_authority_section_are_structured():
    clauses = parse_contract_clauses(_fixture_text(PATENT_FIXTURE))
    dispute = next(clause for clause in clauses if "仲裁委员会仲裁" in clause.text)
    group = dispute.option_groups[0]
    assert group["group_type"] == "DISPUTE_RESOLUTION"
    assert group["mutually_exclusive"] is True
    assert group["selected_count"] == 2

    authority = next(clause for clause in clauses if "以下内容由技术合同登记机构填写" in clause.text)
    assert authority.source_type == "REGISTRATION_AUTHORITY_SECTION"
    assert authority.party_obligation is False


def test_patent_facts_are_in_the_frozen_production_registry():
    output = _patent_review()
    facts = output["contract_facts"]["facts"]
    assert output["contract_facts"]["frozen"] is True
    assert facts["patent.count"][0]["normalized_value"] == 2
    assert facts["patent.pledge.exists"][0]["normalized_value"] is True
    assert facts["patent.existing_license.current_transaction"][0]["normalized_value"] is False
    assert facts["patent.registration.right_effective_event"][0]["normalized_value"] == "REGISTRATION_COMPLETED"


def test_patent_report_statistics_and_visibility_share_final_dispositions():
    output = _patent_review()
    finding_summary = output["risk_finding_summary"]
    disposition_summary = output["final_disposition_summary"]
    assert len(output["risk_findings"]) == sum(finding_summary[level] for level in ("重大风险", "中等风险", "低风险"))
    assert len(output["risk_findings"]) == len(output["final_dispositions"]) == disposition_summary["total"]
    assert disposition_summary["actionable"] == (
        disposition_summary["must_change"] + disposition_summary["should_change"]
    )
    optional_titles = [
        row["risk_title"] for row in output["final_dispositions"]
        if row["final_action"] == "OPTIONAL_IMPROVEMENT"
    ]
    assert len(optional_titles) == disposition_summary["optional_improvement"]
    if optional_titles:
        assert "## 四、可选优化事项" in output["result"]
        assert all(title in output["result"] for title in optional_titles)
    assert output["patent_transfer_quality_gate"]["status"] == "REPORT_READY"
    assert output["patent_transfer_quality_gate"]["gate_results"]["PATENT_TRANSFER_QUALITY_GATE"] == "PASS"


def test_patent_gold_closure_review_semantics_links_and_economic_impact():
    output = _patent_review()
    by_id = {item["item_id"]: item for item in output["items"]}

    assert by_id["TC-01-04"]["coverage_presence"] == "NOT_APPLICABLE"
    assert by_id["TC-01-04"]["coverage_status"] == COVERAGE_NOT_APPLICABLE
    for review_id in ("TC-10-02", "TC-10-03"):
        assert by_id[review_id]["coverage_presence"] == "PRESENT"
        assert by_id[review_id]["contract_allocation"] == "NO_OBLIGATION"

    tc1204 = by_id["TC-12-04"]
    linked_groups = {
        row["risk_group"]
        for row in tc1204["review_risk_links"]
    }
    assert "transferee_infringement_liability_shift" in linked_groups
    assert tc1204["risk_status"] == "LINKED_ADVERSE_RISK"
    assert tc1204["recommended_action"] == "MODIFY_WITH_LINKED_RISK"
    assert "需结合关联风险修改" in tc1204["action_suggestion"]

    tc0601 = by_id["TC-06-01"]
    assert tc0601["evidence_clauses"]
    assert any(
        any(term in (row.get("text") or "") for term in ("专利号", "普通实施许可", "不得继续", "转让登记"))
        for row in tc0601["evidence_clauses"]
    )
    assert not all("转让价款及支付方式" in (row.get("location") or "") for row in tc0601["evidence_clauses"])

    pledge = next(row for row in output["risk_findings"] if row["risk_group"] == "patent_pledge_prepayment")
    impact = pledge["economic_impact"]
    assert impact["has_direct_financial_exposure"] is True
    assert impact["direct_amount"] == 720000
    assert impact["percentage_of_contract"] == 60
    assert impact["formula_source_ids"]
    assert "720,000元" in output["result"]
    assert "60%" in output["result"]


def test_contract_allocation_recognizes_explicit_negative_arrangements():
    assert _contract_allocation_from_text("双方均无义务披露、许可或转让后续改进成果。") == "NO_OBLIGATION"
    assert _contract_allocation_from_text("双方无需通知另一方。") == "NO_OBLIGATION"
    assert _contract_allocation_from_text("乙方不得继续许可该专利。") == "PROHIBITION"
    assert _contract_allocation_from_text("甲方无需支付额外费用。") == "NO_OBLIGATION"


def test_patent_prepayment_economic_impact_is_calculated_from_facts():
    impact = _patent_prepayment_economic_impact({
        "total_transfer_price": 1000000,
        "payment": {
            "status": "CONSISTENT",
            "contract_total": 1000000,
            "nodes": [{"percentage": 50, "amount": 500000}],
        },
    }).to_public_dict()
    assert impact["direct_amount"] == 500000
    assert impact["percentage_of_contract"] == 50
    assert impact["has_direct_financial_exposure"] is True


def test_ptr_rules_do_not_run_for_technical_service_with_incidental_patent_terms():
    text = _fixture_text(SERVICE_FIXTURE) + "\n项目资料可能提及专利、质押和历史普通许可，但本合同主交易仍为技术服务。"
    output = review_contract_v2(
        text,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="乙方/技术服务提供方"),
    )
    assert output["deterministic_analysis"]["technology_classification"]["primary_subtype"] == "TECHNICAL_SERVICE"
    assert not any(item["item_id"].startswith("PTR-") for item in output["items"])
    assert not any(row["risk_group"].startswith("patent_") for row in output["risk_findings"])


def test_patent_report_projection_gate_matches_final_dispositions():
    output = _patent_review()
    view = output["report_view_model"]
    gate = output["report_disposition_projection_gate"]
    assert view["recommended_total"] == view["recommended_display_total"]
    assert view["optional_total"] == view["optional_display_total"]
    assert gate["status"] == "REPORT_READY"
    assert gate["recommended_count"] == gate["recommended_section_item_count"]
    assert gate["optional_improvement_count"] == gate["optional_section_item_count"]
    assert not gate["missing_priority_risk_ids"]
    assert not gate["missing_optional_risk_ids"]


def _disposition(index: int, action: str, recommendation: str = "建议补充具体条款。") -> FinalDisposition:
    return FinalDisposition(
        risk_id=f"SYN-{index:02d}",
        risk_title=f"合成风险{index}",
        review_ids=[f"SYN-R-{index:02d}"],
        severity="LOW" if action == "OPTIONAL_IMPROVEMENT" else "MEDIUM",
        client_effect="FAVORABLE",
        legal_default_protection="NONE",
        actionability="ACTIONABLE",
        final_action=action,
        final_reason="用于验证报告投影不变量。",
        recommendation=recommendation,
        recommendation_owner="REVIEW",
        formal_recommendation=recommendation or None,
        include_in_top_risk=action == "SHOULD_CHANGE",
        include_in_priority_changes=action == "SHOULD_CHANGE",
        include_in_optional_improvements=action == "OPTIONAL_IMPROVEMENT",
    )


def test_report_view_model_and_projection_preserve_nine_recommended_and_one_optional_ids():
    dispositions = [_disposition(index, "SHOULD_CHANGE") for index in range(1, 10)]
    dispositions.append(_disposition(10, "OPTIONAL_IMPROVEMENT"))
    view = _report_view_model([], dispositions)
    report = "\n".join([
        "## 三、建议修改事项",
        *[f"- **建议修改**：建议补充具体条款。（对应风险：合成风险{index}）" for index in range(1, 10)],
        "## 四、可选优化事项",
        "- **可选优化**：建议补充具体条款。（对应风险：合成风险10）",
    ])
    gate = _report_disposition_projection_gate(report, dispositions)

    assert view["recommended_risk_ids"] == [f"SYN-{index:02d}" for index in range(1, 10)]
    assert view["optional_risk_ids"] == ["SYN-10"]
    assert gate["recommended_section_item_count"] == 9
    assert gate["optional_section_item_count"] == 1
    assert gate["status"] == "REPORT_READY"


def test_actionable_final_disposition_without_recommendation_fails_gate():
    disposition = _disposition(1, "OPTIONAL_IMPROVEMENT", recommendation="")
    gate = _final_disposition_invariant_gate([], [disposition])
    assert gate["status"] == "BLOCKED"
    assert gate["gate_results"]["ACTIONABLE_RECOMMENDATION_COMPLETENESS_INVARIANT"] == "FAIL"


def _technology_review_with_confidentiality(extra_clause: str) -> dict:
    text = "\n".join([
        "技术转让合同（专利权转让）。让与方将专利权转让给受让方，双方办理专利权转让登记。",
        "保密条款：双方对履行中知悉的未公开技术信息承担保密义务，期限五年。",
        "违约责任：受让方逾期付款应支付违约金；让与方逾期交付资料应承担违约责任。",
        extra_clause,
    ])
    return review_contract_v2(
        text,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/专利权受让方"),
    )


def test_generic_payment_and_delivery_liability_do_not_satisfy_confidentiality_liability():
    output = _technology_review_with_confidentiality("")
    item = next(row for row in output["items"] if row["item_id"] == "TC-11-04")
    assert item["coverage_presence"] != "PRESENT"


def test_explicit_confidentiality_breach_liability_is_present_with_dedicated_fact_type():
    output = _technology_review_with_confidentiality(
        "任何一方违反本合同保密义务的，应立即停止披露并赔偿因此给另一方造成的损失。"
    )
    item = next(row for row in output["items"] if row["item_id"] == "TC-11-04")
    assert item["coverage_presence"] == "PRESENT"
    assert any(
        "CONFIDENTIALITY_BREACH_LIABILITY" in row.get("fact_types", [])
        for row in item["evidence_clauses"]
    )


def test_general_any_breach_liability_is_partial_without_confidentiality_link():
    output = _technology_review_with_confidentiality("任何一方违反合同均应承担违约责任。")
    item = next(row for row in output["items"] if row["item_id"] == "TC-11-04")
    assert item["coverage_presence"] == "PARTIAL"
    assert item["coverage_status"] == "部分覆盖"


def _synthetic_partial_item(item_id: str, title: str, elements: list[dict]) -> ReviewItemResult:
    return ReviewItemResult(
        item_id=item_id,
        node_id="synthetic_coverage_consistency",
        title=title,
        risk_level="中",
        status="部分覆盖",
        risk_point_description="现有信息不足以确定具体缺失字段。",
        suggested_clause_example="建议进一步确认。",
        legal_basis_index="",
        legal_consequence_if_missing="现有信息不足。",
        applicable_if="适用于当前合同。",
        matched_terms=[],
        coverage_status="部分覆盖",
        issue_type="PARTIAL",
        evidence_clauses=[{
            "clause_id": "C-TEST-01",
            "location": "第五条 技术资料及权利文件",
            "text": "乙方应向甲方提交专利证书、登记簿副本、说明书、附图及相关权利文件。",
        }],
        evidence_location="第五条 技术资料及权利文件",
        reasoning_summary="已覆盖适用范围、履行标准和责任安排。",
        applicability={"status": "APPLICABLE", "applicable": True},
        risk_severity="MEDIUM",
        residual_risk="现有信息不足。",
        risk_group="synthetic_coverage_consistency",
        coverage_elements=elements,
        coverage_presence="PARTIAL",
        risk_status="ADVERSE",
        recommended_action="RECOMMENDED",
        action_suggestion="建议处理",
    )


def test_partial_without_a_concrete_gap_is_promoted_to_pass_and_cannot_create_a_finding():
    item = _synthetic_partial_item("TC-13-01", "技术资料附件", [{
        "element_id": "core_contract_arrangement",
        "proposition": "合同存在可验证的技术资料交付清单及权利文件安排。",
        "status": "COVERED",
        "source_clause_ids": ["C-TEST-01"],
    }])

    _apply_coverage_gap_consistency_guard([item])
    state = _build_narrative_state(item)

    assert item.coverage_status == "通过"
    assert item.coverage_presence == "PRESENT"
    assert item.issue_type == "通过"
    assert item.action_suggestion == "已覆盖，无需修改。"
    assert item.recommended_action == "KEEP"
    assert item.risk_severity == "NONE"
    assert state["missing_elements"] == []
    assert state["partial_elements"] == []
    assert state["finding_exists"] is False
    assert "现有信息不足以确定具体缺失字段" not in " ".join(str(value) for value in state.values())
    assert _build_risk_findings([item]) == []
    assert item.decision_trace["coverage_gap_consistency_guard"]["resolution"] == "PROMOTED_TO_PASS"


def test_partial_with_a_concrete_missing_element_and_tc_11_04_evidence_remain_unchanged():
    item = _synthetic_partial_item("TC-11-04", "泄密责任", [
        {
            "element_id": "confidentiality_obligation",
            "proposition": "合同已明确保密义务。",
            "status": "COVERED",
            "source_clause_ids": ["C0009"],
        },
        {
            "element_id": "confidentiality_breach_liability",
            "proposition": "一般违约责任需明确延伸至违反保密义务，并补充停止披露、补救和损失赔偿机制。",
            "status": "PARTIALLY_COVERED",
            "source_clause_ids": ["C0013"],
        },
    ])
    item.evidence_clauses = [
        {"clause_id": "C0009", "evidence_role": "CONFIDENTIALITY_OBLIGATION", "quote": "合同约定保密义务。"},
        {"clause_id": "C0013", "evidence_role": "GENERAL_BREACH_LIABILITY", "quote": "合同约定一般违约责任。"},
    ]
    frozen_evidence = deepcopy(item.evidence_clauses)

    _apply_coverage_gap_consistency_guard([item])

    assert item.coverage_status == "部分覆盖"
    assert item.issue_type == "PARTIAL"
    assert item.evidence_clauses == frozen_evidence
    assert item.decision_trace["coverage_gap_consistency_guard"]["resolution"] == "PARTIAL_CONFIRMED"


def test_patent_gold_confidentiality_liability_uses_a_complete_two_clause_evidence_chain():
    output = _patent_review()
    item = next(row for row in output["items"] if row["item_id"] == "TC-11-04")

    assert item["action_suggestion"] == "建议处理"
    assert item["coverage_presence"] == "PARTIAL"
    assert item["coverage_status"] == "部分覆盖"
    assert item["issue_type"] == "PARTIAL"
    assert item["risk_level"] == "中"

    evidence_by_role = {
        row["evidence_role"]: row
        for row in item["evidence_clauses"]
        if row.get("evidence_role")
    }
    assert set(evidence_by_role) == {
        "CONFIDENTIALITY_OBLIGATION",
        "GENERAL_BREACH_LIABILITY",
    }
    confidentiality = evidence_by_role["CONFIDENTIALITY_OBLIGATION"]
    liability = evidence_by_role["GENERAL_BREACH_LIABILITY"]
    assert confidentiality["clause_id"] == "C0009"
    assert confidentiality["location"] == "第八条 技术秘密及保密"
    assert "保密信息" in confidentiality["quote"]
    assert "保密义务" in confidentiality["quote"]
    assert "技术秘密转让价款为0元" not in confidentiality["quote"]
    assert confidentiality["support_summary"] == "证明合同已约定保密信息范围、保密义务及保密期限。"
    assert liability["clause_id"] == "C0013"
    assert liability["location"] == "第十二条 违约责任"
    assert "资料" in liability["quote"]
    assert "违约金" in liability["quote"]
    assert liability["support_summary"] == "证明合同仅针对付款、资料交付等事项约定违约责任，未明确延伸至违反保密义务。"
    assert all(row["entailment_status"] == "PARTIAL_SUPPORT" for row in evidence_by_role.values())

    finding_section = output["result"].split("### 泄密责任", 1)[1].split("### 保密资料返还销毁", 1)[0]
    assert "第八条 技术秘密及保密" in finding_section
    assert "第十二条 违约责任" in finding_section
    assert "支持：存在保密义务" in finding_section
    assert "支持：仅约定一般/专项违约责任，未明确适用于保密义务" in finding_section
    assert "技术秘密转让价款为0元" not in finding_section

    review_section = output["result"].split("泄密责任（TC-11-04）", 1)[1].split("\n---\n", 1)[0]
    assert "证据1：第八条 技术秘密及保密" in review_section
    assert "证据2：第十二条 违约责任" in review_section
    assert "综合判断：合同存在保密义务" in review_section
    assert "技术秘密转让价款为0元" not in review_section

    expected_risks = {
        ("patent_pledge_prepayment", "质押解除和权利转移前存在高比例预付款", "HIGH", "ADVERSE"),
        ("asymmetric_breach_liability", "双方违约责任明显不对称", "HIGH", "ADVERSE"),
        ("patent_invalidity_broad_disclaimer", "专利无效免责范围过宽", "HIGH", "ADVERSE"),
        ("transferor_retained_use_conflict", "让与方转让后保留实施权内部冲突", "HIGH", "CONFLICTING"),
        ("transferee_infringement_liability_shift", "按约实施侵权责任全部转嫁给受让方", "HIGH", "DEFAULT_LAW_DEVIATION"),
        ("existing_license_value_burden", "既有许可负担影响专利商业价值", "HIGH", "ADVERSE"),
        ("dispute_resolution_choice_conflict", "仲裁与诉讼路径同时选中", "HIGH", "CONFLICTING"),
        ("registration_completion_continuing_obligations_conflict", "登记即全部履行与持续义务冲突", "HIGH", "CONFLICTING"),
        ("third_party_claim_defense", "第三方权利主张处理", "HIGH", "MISSING"),
        ("confidentiality_breach_liability", "泄密责任", "MEDIUM", "PARTIAL"),
        ("confidential_material_return_destroy", "保密资料返还销毁", "LOW", "MISSING"),
    }
    actual_risks = {
        (row["risk_group"], row["title"], row["severity"], row["issue_type"])
        for row in output["risk_findings"]
    }
    assert actual_risks == expected_risks

    technical_material = next(row for row in output["items"] if row["item_id"] == "TC-13-01")
    assert technical_material["coverage_status"] == "通过"
    assert technical_material["coverage_presence"] == "PRESENT"
    assert technical_material["action_suggestion"] == "已覆盖，无需修改。"
    assert technical_material["evidence_location"] == "第五条技术资料及权利文件"
    assert not any("TC-13-01" in row["source_review_ids"] for row in output["risk_findings"])
    assert output["coverage_summary"] == {
            "通过": 17,
        "部分覆盖": 1,
        "缺失": 2,
        "风险条款": 8,
            COVERAGE_NOT_APPLICABLE: 30,
    }
    assert "现有信息不足以确定具体缺失字段" not in output["result"]


def test_patent_gold_optional_recommendations_and_report_language_close_cleanly():
    output = _patent_review()
    by_id = {item["item_id"]: item for item in output["items"]}
    dispositions = {row["risk_id"]: row for row in output["final_dispositions"]}

    assert by_id["TC-11-03"]["coverage_presence"] == "MISSING"
    assert by_id["TC-11-04"]["coverage_presence"] in {"MISSING", "PARTIAL"}
    linked = [row for row in dispositions.values() if "TC-11-03" in row["review_ids"]]
    assert len(linked) == 1
    assert linked[0]["final_action"] == "OPTIONAL_IMPROVEMENT"
    assert linked[0]["recommendation"].strip()
    assert linked[0]["risk_id"] in output["report_view_model"]["optional_risk_ids"]
    assert linked[0]["risk_title"] in output["result"].split("## 五、", 1)[0]
    assert output["release_quality_gate"]["status"] == "REPORT_READY"
    assert all(
        not row["final_action"] in {"MUST_CHANGE", "SHOULD_CHANGE", "OPTIONAL_IMPROVEMENT"}
        or row["recommendation"].strip()
        for row in dispositions.values()
    )
    forbidden = ("服务提供方", "技术服务提供方", "乙方受托人", "技术提供方")
    duplicates = ("明确明确", "建议建议", "应应", "合同合同")
    assert all(token not in output["result"] for token in forbidden + duplicates)


def test_current_patent_transfer_registration_feasibility_blocker_is_core_risk():
    output = _current_patent_transfer_review()
    classification = output["deterministic_analysis"]["technology_classification"]
    facts = output["deterministic_analysis"]["patent_transfer"]["patent_pledge"]
    by_id = {row["item_id"]: row for row in output["items"]}
    by_group = {row["risk_group"]: row for row in output["risk_findings"]}

    assert classification["primary_subtype"] == "TECHNOLOGY_TRANSFER"
    assert classification["tertiary_subtype"] == "PATENT_RIGHT_TRANSFER"
    assert facts == {
        **facts,
        "exists": True,
        "pledgee_consent_required": True,
        "pledgee_consent_obtained": False,
        "pledge_released": False,
        "registration_feasibility_blocked": True,
    }
    registration = by_id["TC-06-04"]
    assert registration["coverage_status"] == "部分覆盖"
    assert registration["risk_level"] == "高"
    element_status = {row["element_id"]: row["status"] for row in registration["coverage_elements"]}
    assert element_status["registration_procedure"] == "COVERED"
    assert element_status["pledgee_consent_obtained"] == "MISSING"
    assert element_status["pledge_released"] == "MISSING"
    finding = by_group["registration_feasibility_blocker"]
    assert finding["severity"] == "HIGH"
    assert finding["priority"] == "A_CORE"
    assert "质权人书面同意" in finding["recommendation"]
    assert output["release_quality_gate"]["status"] == "REPORT_READY"


def test_current_patent_transfer_ancillary_service_materials_and_domestic_scope_are_fact_driven():
    output = _current_patent_transfer_review()
    classification = output["deterministic_analysis"]["technology_classification"]
    by_id = {row["item_id"]: row for row in output["items"]}

    assert set(classification["sub_transactions"]) >= {
        "technical_service", "acceptance", "technical_material_delivery",
    }
    assert by_id["TC-09-01"]["coverage_status"] == "通过"
    assert by_id["TC-09-02"]["coverage_status"] == "通过"
    acceptance_status = {
        row["element_id"]: row["status"] for row in by_id["TC-09-02"]["coverage_elements"]
    }
    assert acceptance_status == {
        "objective_acceptance_standard": "COVERED",
        "remedy_deadline": "COVERED",
        "reinspection_procedure": "COVERED",
        "reinspection_count": "COVERED",
        "final_failure_remedy": "COVERED",
        "refund_remedy": "COVERED",
        "damage_remedy": "COVERED",
    }
    assert by_id["TC-09-03"]["coverage_status"] == COVERAGE_NOT_APPLICABLE
    assert by_id["TC-09-04"]["coverage_status"] == "通过"
    assert "包含在主交易价款中" in by_id["TC-09-04"]["reasoning_summary"]
    assert all(by_id[item_id]["coverage_status"] == COVERAGE_NOT_APPLICABLE for item_id in (
        "TC-10-01", "TC-10-02", "TC-10-03", "TC-10-04",
    ))
    assert by_id["TC-13-04"]["coverage_status"] == COVERAGE_NOT_APPLICABLE
    assert "中国境内技术转让" in by_id["TC-13-04"]["reasoning_summary"]
    assert "技术进出口、技术转让场景" not in output["result"]

    proof_elements = {row["element_id"]: row["status"] for row in by_id["TC-03-04"]["coverage_elements"]}
    assert proof_elements["transaction_registration_documents"] == "COVERED"
    material_elements = {row["element_id"]: row["status"] for row in by_id["TC-13-01"]["coverage_elements"]}
    assert material_elements == {
        "technical_material_scope": "COVERED",
        "major_document_types": "COVERED",
        "delivery_obligation": "COVERED",
        "delivery_method": "COVERED",
        "delivery_quantity": "COVERED",
        "receipt_confirmation": "COVERED",
        "closed_attachment_schedule": "COVERED",
        "file_version": "COVERED",
        "file_identifier": "MISSING",
    }
    assert by_id["TC-13-01"]["risk_level"] == "低"
    assert "最终交付文件编号" in by_id["TC-13-01"]["suggested_clause_example"]

    registration_finding = next(
        row for row in output["risk_findings"] if row["risk_group"] == "registration_feasibility_blocker"
    )
    assert "专利权转让登记无法办理或被要求补正" in registration_finding["consequence"]
    assert "停用" not in registration_finding["consequence"]
    assert "替换成本" not in registration_finding["consequence"]
    assert "第三方" not in registration_finding["consequence"]
    assert "停用、替换成本" not in output["result"]


def test_patent_pledge_consent_or_release_controls_registration_feasibility_generically():
    blocked = analyze_patent_transfer(
        "标的专利存在专利权质押。截至签署日尚未取得质权人同意转让的书面文件。"
        "转让方应解除质押或取得质权人书面同意。"
    )["patent_pledge"]
    cleared = analyze_patent_transfer(
        "标的专利存在专利权质押。质押登记已注销，已经取得质权人同意转让的书面文件。"
    )["patent_pledge"]
    assert blocked["registration_feasibility_blocked"] is True
    assert cleared["pledgee_consent_obtained"] is True
    assert cleared["pledge_released"] is True
    assert cleared["registration_feasibility_blocked"] is False


def test_patent_transfer_review_context_does_not_leak_improvement_recommendations_between_contracts():
    contract_a = """
    专利权转让合同。乙方将标的专利权完整转让给甲方，专利权自转让登记完成时转移。
    转让价款为人民币100万元。任一方形成后续改进成果时，另一方在同等条件下享有优先许可权，具体条件另行协商。
    """
    first = review_contract_v2(contract_a, "technology", use_llm=False)
    second = _current_patent_transfer_review()
    second_by_id = {row["item_id"]: row for row in second["items"]}

    assert first["deterministic_analysis"]["technology_classification"]["tertiary_subtype"] == "PATENT_RIGHT_TRANSFER"
    assert "优先许可" in contract_a
    assert all(second_by_id[item_id]["coverage_status"] == COVERAGE_NOT_APPLICABLE for item_id in (
        "TC-10-01", "TC-10-02", "TC-10-03", "TC-10-04",
    ))
    assert "现有优先许可" not in second["result"]
    assert "另行协商机制基础上" not in second["result"]


def test_report_missing_notice_is_conditionally_rendered_from_existing_coverage_summary():
    current = _current_patent_transfer_review()
    with_missing = _patent_review()

    assert current["coverage_summary"].get("缺失", 0) == 0
    assert current["report_notice"] == ""
    assert "部分缺失项可能因合同格式简略所致" not in current["result"]

    assert with_missing["coverage_summary"].get("缺失", 0) > 0
    assert with_missing["report_notice"].startswith("本报告基于完整合同模板逐项审查")
    assert "部分缺失项可能因合同格式简略所致" in with_missing["result"]


def test_third_party_claim_advanced_only_gaps_are_optional_optimization():
    contract = """
    专利权转让合同。让与方将专利号CN100000001、名称为测试方法的发明专利权完整转让给受让方，
    双方应办理专利权转让登记，专利权自登记完成时转移。转让价款人民币100万元。
    如第三方对受让方实施受让专利提出权利主张，受让方应在收到主张后5个工作日内书面通知让与方。
    让与方应参与处理、提供专利权属和技术来源文件，并承担因权利瑕疵产生的第三方侵权责任。
    发生重大权利瑕疵时，受让方有权解除合同，让与方返还全部价款并赔偿损失。
    """
    output = review_contract_v2(contract, "technology", use_llm=False)
    item = next(row for row in output["items"] if row["item_id"] == "TC-12-02")
    finding = next(row for row in output["risk_findings"] if "TC-12-02" in row["source_review_ids"])
    tiers = item["decision_trace"]["third_party_claim_element_tiers"]
    element_status = {row["element_id"]: row["status"] for row in item["coverage_elements"]}

    assert tiers["necessary_protection_gaps"] == []
    assert set(tiers["advanced_transaction_optimization_gaps"]) == {
        "defense_control", "lawyer_selection", "settlement_authority", "defense_cost_allocation",
        "replacement_or_workaround", "rights_cure",
    }
    assert tiers["advanced_only_gap"] is True
    assert all(element_status[element_id] == "COVERED" for element_id in (
        "claim_trigger", "notice_requirement", "notice_deadline", "transferor_participation",
        "document_assistance", "damages_allocation", "refund_remedy", "termination_remedy",
    ))
    assert all(element_status[element_id] == "MISSING" for element_id in (
        "defense_control", "lawyer_selection", "settlement_authority", "defense_cost_allocation",
        "replacement_or_workaround", "rights_cure",
    ))
    assert item["risk_level"] == "低"
    assert finding["severity"] == "LOW"
    assert finding["priority"] == "C_OPTIMIZATION"
    assert "完全未约定第三方权利责任" not in finding["issue"]
    assert "现有证据不足以确认" not in finding["issue"]
    assert "通知" not in finding["consequence"] or "已经" in finding["consequence"]
    assert "合同已具备基本保护机制" in output["result"]
    assert "剩余事项主要属于争议处理程序的进一步完善" in output["result"]


def test_third_party_claim_missing_basic_protection_still_recalled():
    contract = """
    专利权转让合同。让与方将专利号CN100000002、名称为测试装置的发明专利权完整转让给受让方，
    双方应办理专利权转让登记，专利权自登记完成时转移。转让价款人民币100万元。
    """
    output = review_contract_v2(contract, "technology", use_llm=False)
    item = next(row for row in output["items"] if row["item_id"] == "TC-12-02")
    findings = [row for row in output["risk_findings"] if "TC-12-02" in row["source_review_ids"]]

    assert item["coverage_status"] == "缺失"
    assert item["risk_level"] in {"中", "高"}
    assert findings
    assert findings[0]["severity"] in {"MEDIUM", "HIGH"}
    assert findings[0]["priority"] != "C_OPTIMIZATION"
