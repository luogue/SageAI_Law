from pathlib import Path
import json
import re

import pytest
from docx import Document

from review_v2.engine import review_contract_v2
from review_v2.technology_service_analysis import (
    analyze_patent_license_consistency,
    analyze_payment_schedule,
    analyze_technical_secret_loss_arrangement,
    classify_technology_transaction,
)
from review_v2.user_profile import UserRiskProfile


ROOT = Path(__file__).resolve().parents[2]
FIXTURE = ROOT / "test_data" / "28 技术许可（技术秘密）合同.docx"
LATEST_FIXTURE = ROOT / "test_data" / "38 技术许可（技术秘密）合同.docx"


BASE_SECRET_LICENSE = """
技术许可（技术秘密）合同
甲方（被许可方）与乙方（许可方）订立本合同。
第一条 技术秘密内容：未公开的控制算法、源代码、控制参数和测试数据。
第二条 乙方将上述技术秘密以排他许可方式许可甲方实施，许可方应向被许可方交付完整技术秘密资料清单、源代码和接口说明。
第三条 许可期限五年，许可地域为中国境内，仅限生产合同产品。
第四条 技术秘密许可使用费为人民币100万元。
第五条 双方对在履行中取得的对方算法、源代码、控制参数、测试数据、生产数据、销售数据和客户资料承担保密义务，未经信息提供方书面同意不得向第三方披露或用于合同外目的；双方保密义务在合同终止后十年内持续有效。
"""


def _document_text(path: Path) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            text = " | ".join(cell.text.strip() for cell in row.cells)
            if text.strip(" |"):
                parts.append(text)
    return "\n".join(parts)


def _fixture_text() -> str:
    return _document_text(FIXTURE)


def _review(text: str) -> dict:
    return review_contract_v2(
        text,
        "technology",
        include_low_risk=True,
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/被许可方"),
    )


@pytest.fixture(scope="module")
def contract_text() -> str:
    return _fixture_text()


@pytest.fixture(scope="module")
def review(contract_text: str) -> dict:
    return _review(contract_text)


@pytest.fixture(scope="module")
def balanced_review() -> dict:
    return _review(BASE_SECRET_LICENSE)


@pytest.fixture(scope="module")
def latest_contract_text() -> str:
    return _document_text(LATEST_FIXTURE)


@pytest.fixture(scope="module")
def latest_review(latest_contract_text: str) -> dict:
    return _review(latest_contract_text)


def test_case_a_b_c_secret_loss_content_reasonableness():
    reasonable = analyze_technical_secret_loss_arrangement(
        "第七条 技术秘密公开后的处理。技术秘密因第三方合法公开或进入公共领域后，"
        "被许可方停止支付未来许可费并有权解除合同；已提供且具有独立价值的技术服务据实结算。"
    )
    adverse = analyze_technical_secret_loss_arrangement(
        "第七条 技术秘密公开后的处理。技术秘密因第三方公开、进入公共领域或许可方申请专利导致公开后，"
        "被许可方仍须继续支付尚未到期的技术秘密许可使用费、技术服务费和提成费，已付款不退，不得单方解除合同。"
    )

    assert classify_technology_transaction(BASE_SECRET_LICENSE)["tertiary_subtype"] == "TECHNICAL_SECRET_LICENSE"
    assert reasonable["reasonable"] is True
    assert reasonable["adverse"] is False
    assert adverse["adverse"] is True
    assert adverse["unconditional_future_payment"] is True
    assert adverse["termination_restricted"] is True


def test_case_d_e_f_license_type_consistency():
    exclusive_clean = analyze_patent_license_consistency(
        "技术秘密排他实施许可合同。许可方保留自行实施权，但不得再许可任何第三方，并确认不存在既有许可。"
    )
    exclusive_conflict = analyze_patent_license_consistency(
        "技术秘密排他实施许可合同。许可方此前已向第三方授予普通许可，该许可期限与本合同重叠；"
        "许可方仍可向科研合作企业提供相关技术。"
    )
    ordinary = analyze_patent_license_consistency(
        "技术秘密普通许可合同。许可方可同时向多个第三方授予普通许可。"
    )

    assert exclusive_clean["exclusive_license_conflict"] is False
    assert exclusive_conflict["existing_third_party_license"] is True
    assert exclusive_conflict["reserved_collaboration_permission"] is True
    assert exclusive_conflict["exclusive_license_conflict"] is True
    assert ordinary["licensor_third_party_license_permission"] is True
    assert ordinary["exclusive_license_conflict"] is False


def test_case_g_h_ancillary_service_and_acceptance_are_fact_routed():
    composite = BASE_SECRET_LICENSE + """
第六条 技术服务和技术指导。许可方负责软件部署、参数标定、系统调试、现场技术支持和人员培训。
第七条 验收标准。项目完成后按照测试环境、测试数据、性能指标和测试方法进行验收；未通过验收应整改并重新验收，验收合格后付款。
"""
    composite_classification = classify_technology_transaction(composite)
    plain_classification = classify_technology_transaction(BASE_SECRET_LICENSE)

    assert composite_classification["tertiary_subtype"] == "TECHNICAL_SECRET_LICENSE"
    assert {"technical_service", "technical_guidance", "acceptance"} <= set(
        composite_classification["sub_transactions"]
    )
    assert "technical_service" not in plain_classification["sub_transactions"]
    assert "acceptance" not in plain_classification["sub_transactions"]


def test_case_i_j_confidentiality_balance_is_directional(
    balanced_review: dict,
    review: dict,
):
    balanced = {row["item_id"]: row for row in balanced_review["items"]}["TC-11-02"]
    asymmetric = {row["item_id"]: row for row in review["items"]}["TC-11-02"]

    assert balanced["coverage_status"] == "通过"
    assert balanced["issue_type"] == "通过"
    assert asymmetric["coverage_status"] == "风险条款"
    assert "明显不对等" in asymmetric["title"]
    assert any(
        row.get("element_id") == "licensor_post_termination_data_use"
        and row.get("status") == "ADVERSE"
        for row in asymmetric["coverage_elements"]
    )


def test_case_k_l_infringement_presence_is_separate_from_direction(
    balanced_review: dict,
    review: dict,
):
    adverse = {row["item_id"]: row for row in review["items"]}["TC-12-01"]
    absent = {row["item_id"]: row for row in balanced_review["items"]}["TC-12-01"]

    assert adverse["coverage_status"] == "风险条款"
    assert adverse["issue_type"] == "DEFAULT_LAW_DEVIATION"
    assert any(
        row.get("element_id") == "compliant_use_infringement_allocation"
        and row.get("status") == "ADVERSE"
        for row in adverse["coverage_elements"]
    )
    assert absent["coverage_status"] == "缺失"


def test_current_secret_license_full_regression(contract_text: str, review: dict):
    classification = classify_technology_transaction(contract_text)
    payment = analyze_payment_schedule(contract_text)
    by_id = {row["item_id"]: row for row in review["items"]}

    assert classification["primary_subtype"] == "TECHNOLOGY_LICENSE"
    assert classification["tertiary_subtype"] == "TECHNICAL_SECRET_LICENSE"
    assert classification["primary_contract_type"] == "technology_license_technical_secret"
    assert classification["scenario"] == "TECHNICAL_SECRET_LICENSE"
    assert {"technical_service", "technical_guidance", "acceptance", "technical_secret"} <= set(
        classification["sub_transactions"]
    )

    assert payment["status"] == "INCONSISTENT"
    assert payment["contract_total"] == 2_400_000
    assert payment["percentage_total"] == 120
    assert payment["amount_total"] == 2_880_000
    assert payment["amount_delta"] == 480_000

    assert "技术秘密许可合同" in by_id["TC-01-01"]["risk_point_description"]
    assert "专利实施许可合同" not in by_id["TC-01-01"]["risk_point_description"]
    assert by_id["TC-07-01"]["coverage_status"] == "风险条款"
    assert "既有第三方普通许可" in by_id["TC-07-01"]["risk_point_description"]
    assert by_id["TC-09-01"]["coverage_status"] != "不适用"
    assert by_id["TC-09-02"]["coverage_status"] == "风险条款"
    assert "单方控制" in by_id["TC-09-02"]["title"]
    assert by_id["TC-09-04"]["coverage_status"] != "不适用"
    assert by_id["TC-12-03"]["coverage_status"] == "风险条款"
    assert "持续付款" in by_id["TC-12-03"]["title"]
    assert not any(
        row.get("review_id") == "TC-12-01"
        for row in review["review_evidence_element_gate"].get("mismatches", [])
    )
    assert by_id["TC-13-03"]["coverage_status"] == "部分覆盖"
    test_elements = {row["element_id"]: row for row in by_id["TC-13-03"]["coverage_elements"]}
    assert test_elements["test_plan_attachment_arrangement"]["status"] == "COVERED"
    assert test_elements["retest_mechanism"]["status"] == "MISSING"

    for item_id in ("TC-03-04", "TC-13-02"):
        item = by_id[item_id]
        assert item["coverage_status"] == "部分覆盖"
        recommendation = str(item.get("suggested_clause_example") or item.get("action_suggestion") or "")
        assert "专利证书" not in recommendation
        assert "专利登记簿" not in recommendation


def _item(result: dict, item_id: str) -> dict:
    return next(row for row in result["items"] if row["item_id"] == item_id)


def test_over_scope_liability_a_missing_when_no_responsibility_clause(contract_text: str):
    result = _review(contract_text.replace(
        "3．买方违反许可范围的，应支付人民币3,000,000元违约金，并赔偿卖方全部损失。",
        "",
    ))
    assert _item(result, "TC-07-03")["coverage_status"] == "缺失"


def test_over_scope_liability_b_existing_clause_is_not_missing(contract_text: str):
    result = _review(contract_text.replace(
        "3．买方违反许可范围的，应支付人民币3,000,000元违约金，并赔偿卖方全部损失。",
        "3．买方违反许可范围的，应承担违约责任并赔偿卖方实际直接损失。",
    ))
    item = _item(result, "TC-07-03")
    assert item["coverage_status"] != "缺失"
    assert any(
        row.get("element_id") == "over_scope_breach_remedy"
        and row.get("status") == "COVERED"
        for row in item["coverage_elements"]
    )


def test_over_scope_liability_c_fixed_penalty_and_all_losses_checks_cumulative_relation(
    contract_text: str,
):
    result = _review(contract_text)
    item = _item(result, "TC-07-03")
    elements = {row["element_id"]: row for row in item["coverage_elements"]}
    assert item["coverage_status"] == "风险条款"
    assert elements["over_scope_breach_remedy"]["status"] == "COVERED"
    assert elements["cumulative_remedy_boundary"]["status"] == "MISSING"
    assert "累计" in item["risk_point_description"]


def test_over_scope_liability_d_reasonable_clause_can_pass_or_be_partial(contract_text: str):
    result = _review(contract_text.replace(
        "3．买方违反许可范围的，应支付人民币3,000,000元违约金，并赔偿卖方全部损失。",
        "3．买方超出许可范围的，应立即停止并赔偿实际直接损失；"
        "一般违约赔偿总额不超过已付许可费的30%，故意或重大过失不适用该上限。",
    ))
    item = _item(result, "TC-07-03")
    assert item["coverage_status"] in {"通过", "部分覆盖"}


def test_positive_fact_evidence_wins_over_search_audit_absence(review: dict):
    item = _item(review, "TC-12-01")
    cap = next(
        row for row in item["coverage_elements"]
        if row.get("element_id") == "licensor_rights_defect_liability"
    )
    assert cap["status"] == "ADVERSE"
    assert cap.get("evidence_ids")
    rendered = json.dumps(review, ensure_ascii=False)
    assert "未找到许可方对其提供技术的权利瑕疵责任受到低额责任上限限制" not in rendered


def test_secret_label_without_direct_measures_does_not_cover_measures():
    result = _review("""
技术秘密许可合同
甲方（被许可方）与乙方（许可方）订立本合同。
乙方将未公开的控制参数作为技术秘密许可甲方使用。
""")
    item = _item(result, "TC-03-04")
    measure = next((
        row for row in item["coverage_elements"]
        if row.get("element_id") == "existing_confidentiality_measures"
    ), None)
    assert measure is None or measure["status"] != "COVERED"


def test_explicit_employee_nda_access_control_and_secret_list_cover_measures():
    result = _review(BASE_SECRET_LICENSE + """
第六条 许可方已建立技术秘密清单和分级管理制度，对源代码设置系统访问控制及加密存储；
所有研发人员已签署员工保密协议，向第三方披露前要求签署第三方保密承诺。
""")
    item = _item(result, "TC-03-04")
    measure = next(
        row for row in item["coverage_elements"]
        if row.get("element_id") == "existing_confidentiality_measures"
    )
    assert measure["status"] == "COVERED"
    assert measure.get("evidence_ids")


def _element_map(item: dict) -> dict:
    return {row["element_id"]: row for row in item.get("coverage_elements", [])}


def _evidence_ids(item: dict) -> set[str]:
    return {
        str(row.get("clause_id") or row.get("evidence_id") or "")
        for row in item.get("evidence_clauses", [])
        if row.get("clause_id") or row.get("evidence_id")
    }


def test_latest_secret_license_route_and_protective_polarity(
    latest_contract_text: str,
    latest_review: dict,
):
    classification = classify_technology_transaction(latest_contract_text)
    assert classification["primary_subtype"] == "TECHNOLOGY_LICENSE"
    assert classification["tertiary_subtype"] == "TECHNICAL_SECRET_LICENSE"
    assert classification["primary_subtype"] != "TECHNICAL_SERVICE"
    assert {"technical_service", "technical_guidance"} <= set(classification["sub_transactions"])

    competition = _item(latest_review, "TC-02-01")
    improvement_boundary = _item(latest_review, "TC-10-04")
    assert competition["coverage_status"] == "通过"
    assert improvement_boundary["coverage_status"] == "通过"
    assert all(row["status"] == "COVERED" for row in competition["coverage_elements"])
    assert all(row["status"] == "COVERED" for row in improvement_boundary["coverage_elements"])
    assert {"C0047", "C0048", "C0049"} <= _evidence_ids(competition)


def test_latest_secret_license_shared_fee_improvement_and_claim_facts(latest_review: dict):
    service_fee = _item(latest_review, "TC-09-04")
    improvement = _item(latest_review, "TC-10-01")
    claim = _item(latest_review, "TC-12-02")

    assert service_fee["coverage_status"] == "通过"
    assert {
        "service_fee_amount": "COVERED",
        "service_payment_schedule": "COVERED",
    }.items() <= {key: row["status"] for key, row in _element_map(service_fee).items()}.items()
    assert {"C0038", "C0039"} <= _evidence_ids(service_fee)

    improvement_elements = _element_map(improvement)
    assert improvement["coverage_status"] == "通过"
    assert improvement_elements["licensee_improvement_ownership"]["status"] == "COVERED"
    assert improvement_elements["licensor_improvement_ownership"]["status"] == "COVERED"
    assert {"C0051", "C0052"} <= _evidence_ids(improvement)

    claim_elements = _element_map(claim)
    assert claim["coverage_status"] == "部分覆盖"
    for element_id in (
        "notification_requirement",
        "notification_deadline",
        "licensor_participation",
        "defense_cost_allocation",
        "settlement_consultation",
        "damages_allocation",
    ):
        assert claim_elements[element_id]["status"] == "COVERED"
    assert str(claim.get("suggested_clause_example") or "").strip() not in {"", "建议。", "。"}


def test_latest_secret_license_confidentiality_consumes_cross_article_facts(latest_review: dict):
    confidentiality = _item(latest_review, "TC-11-01")
    liability = _item(latest_review, "TC-11-04")
    confidentiality_elements = _element_map(confidentiality)
    liability_elements = _element_map(liability)

    assert confidentiality["coverage_status"] == "通过"
    for element_id in (
        "confidential_information_scope",
        "confidential_personnel_control",
        "contract_purpose_use_restriction",
        "third_party_transfer_disclosure_restriction",
        "purpose_scoped_copying_boundary",
    ):
        assert confidentiality_elements[element_id]["status"] == "COVERED"
    assert {"C0020", "C0026"} <= _evidence_ids(confidentiality)

    for element_id in (
        "breach_damage_liability",
        "stop_disclosure",
        "stop_use",
    ):
        assert liability_elements[element_id]["status"] == "COVERED"
    assert liability_elements["mitigation_obligation"]["status"] == "PARTIALLY_COVERED"
    assert liability_elements["mitigation_obligation"]["residual_gap"].startswith("卖方")


def test_latest_secret_license_acceptance_evidence_is_subject_bound(latest_review: dict):
    acceptance = _item(latest_review, "TC-13-03")
    acceptance_ids = _evidence_ids(acceptance)
    assert acceptance["coverage_status"] == "部分覆盖"
    assert acceptance_ids == {"C0004", "C0042", "C0043", "C0044"}
    assert not acceptance_ids.intersection({
        "C0025", "C0026", "C0027", "C0028", "C0029", "C0030", "C0031", "C0032",
        "C0063", "C0064", "C0065", "C0066",
    })


def test_latest_secret_license_report_consistency_and_empty_template_gates(latest_review: dict):
    assert latest_review["technology_transaction_consistency_gate"]["status"] == "REPORT_READY"
    assert (
        latest_review["technology_transaction_consistency_gate"]["gate_results"]
        ["CROSS_RULE_EVIDENCE_CONSISTENCY"]
        == "PASS"
    )
    assert latest_review["covered_missing_consistency_gate"]["status"] == "PASS"
    assert latest_review["report_text_validation_gate"]["status"] == "REPORT_READY"
    assert latest_review["release_quality_gate"]["status"] == "REPORT_READY"

    forbidden_empty_outputs = {"建议。", "法律后果：。", "正式建议：", "。"}
    for item in latest_review["items"]:
        for field in (
            "risk_point_description",
            "legal_consequence_if_missing",
            "suggested_clause_example",
            "action_suggestion",
        ):
            assert str(item.get(field) or "").strip() not in forbidden_empty_outputs


def test_latest_secret_license_compliant_use_infringement_liability_is_covered(
    latest_review: dict,
):
    item = _item(latest_review, "TC-12-01")
    elements = _element_map(item)

    assert item["coverage_status"] == "通过"
    assert item["risk_severity"] == "NONE"
    assert _evidence_ids(item) == {"C0022"}
    assert elements["third_party_claim_trigger"]["status"] == "COVERED"
    assert elements["licensor_claim_handling"]["status"] == "COVERED"
    assert elements["licensor_legal_economic_liability"]["status"] == "COVERED"
    assert "第874条" in str(item.get("legal_basis_index") or "")


def test_latest_secret_license_service_acceptance_uses_shared_subject_evidence(
    latest_review: dict,
):
    item = _item(latest_review, "TC-09-02")
    elements = _element_map(item)
    evidence_ids = _evidence_ids(item)

    assert item["coverage_status"] == "部分覆盖"
    assert evidence_ids == {"C0004", "C0042", "C0043", "C0044"}
    assert elements["objective_acceptance_standard"]["status"] == "COVERED"
    assert "rectification_and_retest" not in elements
    for key in ("rectification_party", "rectification_measures", "rectification_procedure"):
        assert elements[key]["status"] == "MISSING"
    assert not evidence_ids.intersection({
        "C0025", "C0026", "C0027", "C0028", "C0029", "C0030", "C0031", "C0032",
        "C0063", "C0064", "C0065", "C0066",
    })
    rendered_gap = " ".join(str(item.get(field) or "") for field in (
        "risk_point_description",
        "legal_consequence_if_missing",
        "suggested_clause_example",
    ))
    assert "整改" in rendered_gap
    assert "再次验收" in rendered_gap


def test_latest_secret_license_confidentiality_partial_renderer_uses_final_missing_fact(
    latest_review: dict,
):
    item = _item(latest_review, "TC-11-04")
    elements = _element_map(item)
    finding = next(
        row for row in latest_review["risk_findings"]
        if "TC-11-04" in row.get("source_review_ids", [])
    )

    assert item["coverage_status"] == "部分覆盖"
    assert elements["incident_notification"]["status"] == "MISSING"
    for element_id in ("breach_damage_liability", "stop_disclosure", "stop_use"):
        assert elements[element_id]["status"] == "COVERED"
    assert elements["mitigation_obligation"]["status"] == "PARTIALLY_COVERED"

    rendered = " ".join(str(value or "") for value in (
        item.get("risk_point_description"),
        item.get("legal_consequence_if_missing"),
        item.get("suggested_clause_example"),
        finding.get("issue"),
        finding.get("consequence"),
        finding.get("recommendation"),
    ))
    assert "通知" in rendered
    for covered_fact in ("停止披露", "停止使用", "买方泄密后采取措施"):
        assert covered_fact not in finding["recommendation"]
    assert "卖方泄密后采取措施" in finding["recommendation"]
    assert set(finding.get("missing_element_ids")) == {"mitigation_obligation", "incident_notification"}


def test_secret_loss_explicit_article_consumes_trigger_notice_termination_and_damage(
    latest_contract_text: str,
):
    explicit_article = """
第七条 技术秘密公开
技术秘密被他人公开或非因买方原因丧失秘密性时，许可方应在10日内通知买方，买方有权解除合同。
许可方迟延通知造成买方损失的，应当赔偿。
"""
    # This is a semantic regression fixture.  The checked-in document's
    # Article 7 contains only a deadline and late-notice damages, so replace
    # that article with a complete version instead of inventing facts for the
    # stored source document.
    synthetic_text, replacements = re.subn(
        r"第七条\s*技术秘密公开.*?(?=第八条\s*保密义务)",
        explicit_article,
        latest_contract_text,
        count=1,
        flags=re.S,
    )
    assert replacements == 1

    item = _item(_review(synthetic_text), "TC-12-03")
    elements = _element_map(item)
    for element_id in (
        "secret_loss_trigger",
        "secret_loss_notification",
        "secret_loss_notification_deadline",
        "late_notification_damage",
        "termination_after_secret_loss",
    ):
        assert elements[element_id]["status"] == "COVERED"
    assert elements["future_fee_after_secret_loss"]["status"] == "MISSING"
    assert elements["separate_value_settlement"]["status"] == "MISSING"

    rendered_gap = " ".join(str(item.get(field) or "") for field in (
        "risk_point_description",
        "legal_consequence_if_missing",
        "suggested_clause_example",
    ))
    assert "许可费" in rendered_gap
    assert any(term in rendered_gap for term in ("停止", "调减"))
    assert "独立价值" in rendered_gap
    assert "公开触发" not in rendered_gap
    assert "解除权尚未明确" not in rendered_gap
    assert "转让方" not in rendered_gap


def test_latest_secret_loss_positive_notice_facts_never_reenter_missing(
    latest_review: dict,
):
    item = _item(latest_review, "TC-12-03")
    elements = _element_map(item)

    for element_id in (
        "secret_loss_notification",
        "secret_loss_notification_deadline",
        "late_notification_damage",
    ):
        assert elements[element_id]["status"] == "COVERED"
        assert elements[element_id].get("evidence_ids") == ["C0023"]

    final_missing = set(item.get("decision_trace", {}).get("final_missing_element_ids", []))
    assert not final_missing.intersection({
        "secret_loss_notification",
        "secret_loss_notification_deadline",
        "late_notification_damage",
    })
    rendered_gap = " ".join(str(item.get(field) or "") for field in (
        "risk_point_description",
        "legal_consequence_if_missing",
        "suggested_clause_example",
    ))
    assert "通知义务" not in rendered_gap
    assert "通知期限" not in rendered_gap
    assert "迟延通知" not in rendered_gap


def test_latest_improvement_notice_is_partial_with_only_residual_details(
    latest_review: dict,
):
    item = _item(latest_review, "TC-10-03")
    elements = _element_map(item)

    assert item["coverage_status"] == "部分覆盖"
    assert elements["improvement_notice_trigger"]["status"] == "COVERED"
    assert elements["improvement_notice_trigger"].get("evidence_ids") == ["C0052"]
    assert {
        element_id for element_id, element in elements.items()
        if element["status"] == "MISSING"
    } == {
        "improvement_notice_deadline",
        "improvement_notice_content",
        "improvement_material_scope",
    }
    assert item.get("decision_trace", {}).get("final_missing_element_ids") == [
        "improvement_notice_deadline",
        "improvement_notice_content",
        "improvement_material_scope",
    ]


def test_latest_acceptance_absence_search_excludes_covered_metrics(
    latest_review: dict,
):
    item = _item(latest_review, "TC-13-03")
    elements = _element_map(item)
    audit = item.get("decision_trace", {}).get("specialized_search_audit", {})

    assert elements["objective_acceptance_metrics"]["status"] == "COVERED"
    assert elements["objective_acceptance_metrics"].get("evidence_ids")
    assert "量化验收指标" not in audit.get("missing_elements", [])
    assert "量化验收指标" not in audit.get("required_positive_facts", [])
    assert "量化验收指标" in audit.get("removed_conflicting_absence_targets", [])
    assert "未找到量化验收指标" not in json.dumps(item, ensure_ascii=False)


def test_latest_confidentiality_period_recommendation_only_uses_final_exception_gap(
    latest_review: dict,
):
    item = _item(latest_review, "TC-11-02")
    elements = _element_map(item)
    finding = next(
        row for row in latest_review["risk_findings"]
        if "TC-11-02" in row.get("source_review_ids", [])
    )

    assert elements["confidential_information_scope"]["status"] == "COVERED"
    assert elements["confidentiality_period"]["status"] == "COVERED"
    assert elements["confidentiality_exceptions"]["status"] == "MISSING"
    assert item.get("decision_trace", {}).get("final_missing_element_ids") == [
        "confidentiality_exceptions"
    ]
    assert finding.get("missing_element_ids") == ["confidentiality_exceptions"]

    recommendations = " ".join(str(value or "") for value in (
        item.get("suggested_clause_example"),
        item.get("action_suggestion"),
        finding.get("recommendation"),
    ))
    assert "保密例外" in recommendations
    for stale_advice in ("明确保密对象", "约定保密期限", "终止后保密义务"):
        assert stale_advice not in recommendations


def test_confidentiality_liability_evidence_explanation_is_element_bound(
    latest_review: dict,
):
    item = _item(latest_review, "TC-11-04")
    direct_rows = [
        row for row in item.get("evidence_clauses", [])
        if row.get("clause_id") in {"C0028", "C0032"}
    ]
    assert direct_rows
    for row in direct_rows:
        supported = set(row.get("supporting_element_ids", []))
        assert {
            "breach_damage_liability",
            "stop_disclosure",
            "stop_use",
        } <= supported
        support_summary = str(row.get("support_summary") or "")
        assert "泄密损失赔偿责任" in support_summary
        assert "泄密后停止披露" in support_summary
        assert "泄密后停止使用" in support_summary
        if row["clause_id"] == "C0028":
            assert "mitigation_obligation" in supported
            assert "买方泄密后采取措施防止继续扩散" in support_summary
        else:
            assert "mitigation_obligation" not in supported
            assert "防止" not in support_summary
        assert "违约金" not in support_summary
        assert "保密信息范围" not in support_summary

    detail_match = re.search(
        r"## \d+\. 泄密责任（TC-11-04）(?P<section>.*?)(?=\n---\n|\Z)",
        latest_review["result"],
        flags=re.S,
    )
    assert detail_match
    detail = detail_match.group("section")
    assert "现有违约责任未形成覆盖全部合同义务的兜底" not in detail
    assert "泄密损失赔偿及专项救济机制仍不完整" not in detail
    assert "仅缺失：卖方泄密后采取措施防止继续扩散" in detail
    assert "泄密后及时通知受影响方" in detail


def test_partial_rule_text_is_generated_only_from_final_missing_elements(
    latest_review: dict,
):
    item = _item(latest_review, "TC-13-03")
    finding = next(
        row for row in latest_review["risk_findings"]
        if "TC-13-03" in row.get("source_review_ids", [])
    )
    expected_labels = [
        "独立验收测试方案附件",
        "测试方法和操作步骤",
        "固定测试数据集或样本",
        "测试设备和环境条件",
        "统计口径",
        "测试记录和签认要求",
        "争议复测或第三方检测机制",
    ]
    rendered = " ".join(str(value or "") for value in (
        item.get("risk_point_description"),
        item.get("reasoning_summary"),
        item.get("suggested_clause_example"),
        finding.get("issue"),
        finding.get("recommendation"),
    ))
    for label in expected_labels:
        assert label in rendered
    assert "量化验收指标尚未明确" not in rendered
    assert "客观技术指标尚未明确" not in rendered


def test_acceptance_report_blocks_share_one_final_missing_element_set(
    latest_review: dict,
):
    item = _item(latest_review, "TC-13-03")
    finding = next(
        row for row in latest_review["risk_findings"]
        if "TC-13-03" in row.get("source_review_ids", [])
    )
    expected_ids = [
        "test_plan_attachment_arrangement",
        "test_method_and_steps",
        "fixed_test_dataset",
        "test_equipment_environment",
        "statistical_criteria",
        "test_record_requirement",
        "retest_mechanism",
    ]
    expected_search_targets = [
        "独立验收测试方案附件",
        "测试方法和操作步骤",
        "固定测试数据集或样本",
        "测试设备和环境条件",
        "统计口径",
        "测试记录和签认要求",
        "争议复测或第三方检测机制",
    ]
    trace = item.get("decision_trace", {})
    audit = trace.get("specialized_search_audit", {})

    assert trace.get("final_missing_element_ids") == expected_ids
    assert finding.get("missing_element_ids") == expected_ids
    assert audit.get("final_missing_element_ids") == expected_ids
    assert audit.get("missing_elements") == expected_search_targets
    assert audit.get("required_positive_facts") == expected_search_targets
    assert "量化验收指标" not in audit.get("missing_elements", [])

    summary_line = next(
        line for line in latest_review["result"].splitlines()
        if line.startswith("- **验收测试附件**")
    )
    detail_match = re.search(
        r"## \d+\. 验收测试附件（TC-13-03）(?P<section>.*?)(?=\n---\n|\Z)",
        latest_review["result"],
        flags=re.S,
    )
    assert detail_match
    detail = detail_match.group("section")
    for label in expected_search_targets:
        assert label in summary_line
        assert label in detail
