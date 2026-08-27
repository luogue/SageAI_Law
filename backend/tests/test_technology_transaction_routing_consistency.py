import json
from pathlib import Path

from docx import Document

from review_v2.engine import (
    COVERAGE_NOT_APPLICABLE,
    extract_semantic_field_value,
    is_semantically_filled_field,
    review_contract_v2,
    validate_evidence_provenance,
)
from review_v2.technology_service_analysis import classify_technology_transaction
from review_v2.user_profile import UserRiskProfile


PATENT_LICENSE_TEXT = """技术许可（专利实施）合同
甲方（被许可人/买方）：测试公司
乙方（许可人/卖方）：专利权人
第一条 合同标的
乙方以排他许可方式许可甲方实施一种基于算法的工业检测方法发明专利，专利号为CN202410001234.5。
第二条 许可范围和期限
许可范围为中国境内工业检测设备，许可期限为五年，未经乙方书面同意不得再许可。
第三条 许可使用费
甲方向乙方支付专利许可使用费人民币100万元。
第四条 权利状态与文件
乙方提交专利证书、专利登记簿副本、法律状态及年费缴纳记录，并如实披露质押、既有第三方许可、权属争议和无效宣告请求。
乙方同时提交技术资料交付清单及许可备案所需文件。
"""


TRANSFER_LICENSE_CONFLICT_TEXT = """专利权转让合同
甲方（受让方）：测试公司
乙方（让与方/卖方）：专利权人
第一条 合同标的
乙方将专利号CN202410001234.5的发明专利权转让给甲方。
第二条 核心权利安排
卖方以排他许可方式许可买方实施本专利，许可期限五年，许可范围为中国境内。
卖方保证其有权向买方授予本专利实施许可，并已如实披露专利权质押、既有许可、权属争议及无效宣告请求。
第三条 价款
甲方支付许可使用费人民币100万元。
"""


def test_patent_license_has_one_primary_type_and_no_keyword_subtype_leakage():
    classification = classify_technology_transaction(PATENT_LICENSE_TEXT)

    assert classification["primary_contract_type"] == "technology_license_patent"
    assert classification["primary_subtype"] == "TECHNOLOGY_LICENSE"
    assert classification["secondary_contract_types"] == []
    assert classification["active_transaction_types"] == ["TECHNOLOGY_LICENSE"]
    assert classification["contract_type_conflict"]["exists"] is False
    assert classification["involves_software_or_data"] is False
    assert set(classification["classification_evidence_details"]) == {
        "contract_title",
        "subject_matter",
        "core_performance_obligations",
        "rights_transfer_method",
        "consideration_nature",
        "key_clauses",
    }

    output = review_contract_v2(PATENT_LICENSE_TEXT, "technology", use_llm=False)
    by_id = {row["item_id"]: row for row in output["items"]}

    assert by_id["TC-07-01"]["coverage_status"] != COVERAGE_NOT_APPLICABLE
    for review_id in ("TC-05-01", "TC-05-02", "TC-06-03", "TC-08-02", "TC-09-01"):
        assert by_id[review_id]["coverage_status"] == COVERAGE_NOT_APPLICABLE
    assert output["technology_transaction_consistency_gate"]["status"] == "REPORT_READY"
    assert all(
        by_id[review_id]["coverage_status"] != COVERAGE_NOT_APPLICABLE
        for finding in output["risk_findings"]
        for review_id in finding["source_review_ids"]
        if review_id in by_id
    )


def test_pure_patent_transaction_uses_fact_driven_proof_requirements():
    output = review_contract_v2(PATENT_LICENSE_TEXT, "technology", use_llm=False)
    by_id = {row["item_id"]: row for row in output["items"]}
    expected_elements = {
        "patent_identity_and_registration",
        "patent_legal_status_and_annual_fee",
        "encumbrance_license_and_dispute_disclosure",
        "transaction_registration_documents",
    }

    for review_id in ("TC-03-04", "TC-13-02"):
        item = by_id[review_id]
        assert {row["element_id"] for row in item["coverage_elements"]} == expected_elements
        assert set(item["decision_trace"]["excluded_irrelevant_elements"]) == {
            "software_copyright",
            "open_source_component",
            "software_background_ip",
        }
        assert "开源组件" not in item["reasoning_summary"]
        assert "软件著作权" not in item["reasoning_summary"]


def test_title_and_core_transaction_conflict_is_a_finding_not_a_second_rule_pack():
    classification = classify_technology_transaction(TRANSFER_LICENSE_CONFLICT_TEXT)

    assert classification["primary_contract_type"] == "technology_transfer_patent"
    assert classification["active_transaction_types"] == ["TECHNOLOGY_TRANSFER"]
    assert classification["contract_type_conflict"] == {
        "exists": True,
        "title_transaction_type": "TECHNOLOGY_TRANSFER",
        "operative_transaction_type": "TECHNOLOGY_LICENSE",
        "resolution": "PRIMARY_TITLE_WITH_CONFLICT_FINDING",
        "reason": "合同标题与核心权利处分条款指向不同交易性质，应单独提示合同性质冲突，但不得同时激活两套规则。",
    }

    output = review_contract_v2(TRANSFER_LICENSE_CONFLICT_TEXT, "technology", use_llm=False)
    by_id = {row["item_id"]: row for row in output["items"]}
    conflict_findings = [
        finding for finding in output["risk_findings"]
        if finding["risk_group"] == "contract_type_conflict"
    ]

    assert by_id["TC-01-02"]["issue_type"] == "CONFLICTING"
    assert len(conflict_findings) == 1
    assert conflict_findings[0]["source_review_ids"] == ["TC-01-02"]
    assert by_id["TC-06-01"]["coverage_status"] != COVERAGE_NOT_APPLICABLE
    for review_id in ("TC-05-01", "TC-07-01", "TC-08-02", "TC-09-01"):
        assert by_id[review_id]["coverage_status"] == COVERAGE_NOT_APPLICABLE
    assert output["technology_transaction_consistency_gate"]["status"] == "REPORT_READY"
    assert output["report_generation_decision"]["report_generated"] is True


def test_topic_words_without_transaction_evidence_do_not_activate_subtype_rules():
    text = """技术合作备忘录
双方可就开发、许可、咨询和技术服务机会进行沟通，本文件不设交易标的、给付义务或价款。
任何具体项目均须另行签署正式合同。
"""
    classification = classify_technology_transaction(text)
    output = review_contract_v2(text, "technology", use_llm=False)
    by_id = {row["item_id"]: row for row in output["items"]}

    assert classification["primary_contract_type"] == "general_technology"
    assert classification["active_transaction_types"] == ["GENERAL_TECHNOLOGY"]
    for review_id in ("TC-04-01", "TC-05-01", "TC-06-01", "TC-07-01", "TC-08-01", "TC-09-01"):
        assert by_id[review_id]["coverage_status"] == COVERAGE_NOT_APPLICABLE


def test_existing_third_party_rights_clause_is_partial_not_missing():
    text = PATENT_LICENSE_TEXT + """
第五条 第三方权利责任
如第三人指控甲方按照本合同约定实施专利侵犯第三方合法权利，甲方应在五个工作日内通知乙方。
卖方应配合买方提供专利权属、申请过程及技术来源资料。
因乙方无权许可或故意隐瞒已有权利负担造成的损失，由乙方承担。
因甲方擅自修改技术方案、超出许可范围或结合其他第三方技术导致的侵权，由甲方自行承担。
"""
    output = review_contract_v2(text, "technology", use_llm=False)
    item = next(row for row in output["items"] if row["item_id"] == "TC-12-01")
    finding = next(
        row for row in output["risk_findings"]
        if row["risk_group"] == "license_infringement_responsibility"
    )

    assert item["coverage_status"] == "部分覆盖"
    assert item["issue_type"] == "PARTIAL"
    assert "未找到对应约定" not in item["reasoning_summary"]
    assert {row["element_id"]: row["status"] for row in item["coverage_elements"]} == {
        "infringing_implementation_trigger": "COVERED",
        "fault_specific_loss_allocation": "COVERED",
        "compliant_use_residual_liability": "MISSING",
    }
    assert finding["issue_type"] == "PARTIAL"
    assert finding["evidence_bundle"]["entailment_status"] == "PASSED"
    assert {row["fact_type"] for row in finding["evidence_bundle"]["items"]} == {
        "INFRINGING_IMPLEMENTATION_TRIGGER",
        "LICENSEE_LOSS_ALLOCATION",
        "LICENSOR_ASSISTANCE_ONLY",
    }


EVIDENCE_JUDGMENT_LICENSE_TEXT = PATENT_LICENSE_TEXT + """
第五条 第三方权利主张处理
卖方保证其有权按照本合同约定向买方授予本专利实施许可，并保证如实披露专利权质押、既有许可、权属争议和无效宣告请求。
本条所称第三方权利主张，是指买方按照本合同约定实施专利侵犯第三方合法权利。
买方应在收到第三方正式权利主张后5个工作日内通知卖方。
卖方应配合买方提供专利权属、申请过程及技术来源资料。
卖方无权许可或者故意隐瞒已有权利负担造成的损失，由卖方承担。
买方擅自修改技术方案、超出许可范围实施或者结合其他第三方技术导致的侵权，由买方自行承担责任。
双方未明确约定第三方索赔处理中的律师选择、抗辩控制以及和解决策权限。
卖方无法合法许可时，应退还未履行期间许可费并支付合同总额10%的违约金。
第六条 不得限制技术竞争和发展 双方不得：
不得限制任何一方独立研究开发与本专利不同的新技术。
禁止任何一方通过合法渠道取得第三方技术。
不得要求另一方无偿转让与本合同专利无直接关联的后续技术成果。
第七条 后续改进成果
买方有权利用许可专利继续改进，买方独立完成的改进成果归买方所有。
卖方独立完成的改进成果归卖方所有，双方均无义务将独立改进成果无偿许可给另一方。
"""


def test_license_evidence_semantics_and_atomic_elements_are_consistent():
    output = review_contract_v2(EVIDENCE_JUDGMENT_LICENSE_TEXT, "technology", use_llm=False)
    by_id = {row["item_id"]: row for row in output["items"]}

    for review_id in ("TC-02-01", "TC-02-02"):
        item = by_id[review_id]
        assert item["coverage_status"] == "通过"
        assert item["decision_trace"]["evidence_judgment"]["evidence_semantics"] == "PROTECTIVE_CLAUSE"
        assert item["decision_trace"]["evidence_judgment"]["affirmative_risk_evidence_ids"] == []

    claim_states = {
        row["element_id"]: row["status"] for row in by_id["TC-12-02"]["coverage_elements"]
    }
    assert claim_states == {
        "notification_requirement": "COVERED",
        "notification_deadline": "COVERED",
        "licensor_participation": "COVERED",
        "defense_control": "MISSING",
        "lawyer_selection": "MISSING",
        "settlement_authority": "MISSING",
        "defense_cost_allocation": "MISSING",
        "damages_allocation": "PARTIALLY_COVERED",
        "replacement_or_workaround": "MISSING",
        "license_cure": "MISSING",
        "refund_or_termination_remedy": "COVERED",
    }
    assert by_id["TC-12-02"]["coverage_status"] == "部分覆盖"


def test_cross_rule_evidence_and_final_consistency_gate_use_classification_truth():
    output = review_contract_v2(EVIDENCE_JUDGMENT_LICENSE_TEXT, "technology", use_llm=False)
    by_id = {row["item_id"]: row for row in output["items"]}

    improvement = by_id["TC-10-04"]
    assert improvement["coverage_status"] == "通过"
    assert improvement["decision_trace"]["semantic_evidence_lookup"]["absence_search_skipped"] is True
    assert set(improvement["decision_trace"]["semantic_evidence_lookup"]["reused_clause_ids"]) == set(
        improvement["decision_trace"]["semantic_evidence_lookup"]["evaluated_clause_ids"]
    )

    attachment = by_id["TC-13-03"]
    assert attachment["coverage_status"] == COVERAGE_NOT_APPLICABLE
    assert "专利实施许可" in attachment["reasoning_summary"]
    assert "当前主交易为既有技术权利转让" not in attachment["reasoning_summary"]

    gate = output["technology_transaction_consistency_gate"]
    assert gate["status"] == "REPORT_READY"
    assert gate["gate_results"] == {
        "TECHNOLOGY_TRANSACTION_CONSISTENCY": "PASS",
        "CLASSIFICATION_NARRATIVE_CONSISTENCY": "PASS",
        "EVIDENCE_JUDGMENT_CONSISTENCY": "PASS",
        "ATOMIC_ELEMENT_CONSISTENCY": "PASS",
        "CROSS_RULE_EVIDENCE_CONSISTENCY": "PASS",
        "PARTY_DIRECTION_CONSISTENCY": "PASS",
        "SUBTYPE_RENDERING_CONSISTENCY": "PASS",
        "EVIDENCE_PROVENANCE_CONSISTENCY": "PASS",
        "EMPTY_FIELD_CONSISTENCY": "PASS",
        "REPORT_EVIDENCE_CONSISTENCY": "PASS",
        "EVIDENCE_POLARITY_CONSISTENCY": "PASS",
    }


def test_patent_transfer_asset_subtypes_and_transaction_layers_are_semantic():
    granted = classify_technology_transaction("""技术合同
标的专利当前状态：已取得授权，当前专利权人为乙方，授权日为2025年1月1日。
乙方将完整专利权转让给甲方，登记完成后专利权转移至甲方。
甲方支付专利权转让总价100万元。登记前乙方向甲方提供免费过渡期实施许可，并提供培训。
""")
    application = classify_technology_transaction("""技术合同
标的为尚未授权的发明专利申请，当前申请人为乙方。
乙方将专利申请权转让给甲方，双方办理申请人变更登记，甲方支付转让价款。
""")
    mixed = classify_technology_transaction("""技术合同
乙方将一项已取得授权的专利权及一项尚未授权的专利申请权一并转让给甲方。
登记完成后权利主体变更，甲方支付专利权及专利申请权转让价款。
""")

    assert granted["tertiary_subtype"] == "PATENT_RIGHT_TRANSFER"
    assert granted["primary_transaction"] == {
        "type": "technology_transfer", "subtype": "patent_right_transfer",
    }
    assert granted["active_transaction_types"] == ["TECHNOLOGY_TRANSFER"]
    assert "TRAINING" in granted["ancillary_obligations"]
    assert granted["transitional_arrangements"] == ["TEMPORARY_PATENT_LICENSE"]
    assert application["tertiary_subtype"] == "PATENT_APPLICATION_RIGHT_TRANSFER"
    assert application["primary_contract_type"] == "technology_transfer_patent_application_right"
    assert mixed["tertiary_subtype"] == "MIXED_PATENT_RIGHT_AND_APPLICATION_TRANSFER"


def test_new_patent_transfer_generalization_sample_routes_without_frozen_risks():
    fixture = Path(__file__).resolve().parents[2] / "test_data" / "14、专利转让.docx"
    document = Document(str(fixture))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    output = review_contract_v2(
        text, "technology", use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/专利权受让方"),
    )
    classification = output["deterministic_analysis"]["technology_classification"]
    by_id = {row["item_id"]: row for row in output["items"]}

    assert classification["primary_subtype"] == "TECHNOLOGY_TRANSFER"
    assert classification["tertiary_subtype"] == "PATENT_RIGHT_TRANSFER"
    assert classification["transitional_arrangements"] == ["TEMPORARY_PATENT_LICENSE"]
    assert output["perspective_resolution"]["canonical_role"] == "PATENT_TRANSFEREE"
    assert by_id["TC-01-01"]["coverage_status"] == "通过"
    assert by_id["TC-01-02"]["coverage_status"] == "通过"
    assert all(by_id[f"TC-06-0{index}"]["coverage_status"] != COVERAGE_NOT_APPLICABLE for index in range(1, 5))
    assert all(by_id[f"TC-07-0{index}"]["coverage_status"] == COVERAGE_NOT_APPLICABLE for index in range(1, 5))
    assert all(by_id[f"TC-09-0{index}"]["coverage_status"] == COVERAGE_NOT_APPLICABLE for index in range(1, 5))
    assert by_id["TC-11-01"]["coverage_status"] == "通过"
    assert by_id["TC-11-02"]["coverage_status"] == "通过"
    assert by_id["TC-11-03"]["coverage_status"] == "通过"
    assert by_id["TC-11-04"]["coverage_status"] == "通过"
    assert by_id["TC-12-01"]["coverage_status"] == "部分覆盖"
    assert {
        row["element_id"]: row["status"] for row in by_id["TC-12-01"]["coverage_elements"]
    } == {
        "third_party_claim_trigger": "COVERED",
        "transferor_claim_participation": "COVERED",
        "liability_allocation_exists": "PARTIALLY_COVERED",
        "liability_allocation_determinability": "MISSING",
        "defense_control": "MISSING",
        "replacement_or_workaround": "MISSING",
    }
    assert "尚未覆盖：受让方按约实施时的第三方侵权责任承担已明确" not in output["result"]
    assert "建议明确受让方按约实施时的第三方侵权责任承担已" not in output["result"]
    claim_response = by_id["TC-12-02"]
    assert claim_response["coverage_status"] == "部分覆盖"
    claim_states = {row["element_id"]: row["status"] for row in claim_response["coverage_elements"]}
    assert claim_states["claim_trigger"] == "COVERED"
    assert claim_states["transferor_participation"] == "COVERED"
    assert claim_states["document_assistance"] == "PARTIALLY_COVERED"
    assert claim_states["damages_allocation"] == "PARTIALLY_COVERED"
    assert claim_states["refund_or_termination_remedy"] == "COVERED"
    assert claim_states["defense_control"] == "MISSING"
    assert claim_response["decision_trace"]["semantic_evidence_lookup"]["reused_clause_ids"]
    assert by_id["TC-02-02"]["coverage_status"] == "通过"
    assert by_id["TC-02-02"]["decision_trace"]["evidence_judgment"]["evidence_semantics"] == "PROTECTIVE_CLAUSE"
    assert not by_id["TC-02-02"]["decision_trace"]["evidence_judgment"]["affirmative_risk_evidence_ids"]
    improvement_text = " ".join((
        by_id["TC-10-03"]["risk_point_description"],
        by_id["TC-10-03"]["legal_consequence_if_missing"],
        by_id["TC-10-03"]["suggested_clause_example"],
    ))
    assert "合作核心技术" not in improvement_text
    assert "推定改进成果共有" in improvement_text
    delivery_text = " ".join(
        str(row.get("quote") or row.get("text") or "") for row in by_id["TC-06-02"]["evidence_clauses"]
    )
    assert all(term in delivery_text for term in ("专利证书", "技术图纸", "交付时间", "交付方式", "验收不合格"))
    assert all(
        row["status"] == "COVERED"
        for review_id in ("TC-06-02", "TC-06-03", "TC-06-04")
        for row in by_id[review_id]["coverage_elements"]
    )
    assert by_id["TC-12-04"]["risk_direction"]["review_position_code"] == "PATENT_TRANSFEREE"
    assert output["release_quality_gate"]["status"] == "REPORT_READY"
    assert output["release_quality_gate"]["gate_results"]["COVERAGE_SENTENCE_CONSISTENCY_GATE"] == "PASS"
    assert "当前技术服务交易" not in output["result"]


def test_liability_cap_evidence_excludes_patent_maintenance_as_core_support():
    output = review_contract_v2(EVIDENCE_JUDGMENT_LICENSE_TEXT, "technology", use_llm=False)
    by_id = {row["item_id"]: row for row in output["items"]}
    item = by_id["TC-12-04"]
    finding = next(row for row in output["risk_findings"] if "TC-12-04" in row["source_review_ids"])

    assert item["coverage_status"] == "部分覆盖"
    assert item["decision_trace"]["excluded_irrelevant_evidence"] == ["patent_maintenance_clause"]
    evidence_text = " ".join(str(row.get("quote") or "") for row in item["evidence_clauses"])
    assert "维持专利有效并缴纳年费" not in evidence_text
    bundle = finding["evidence_bundle"]["items"]
    assert {row["fact_type"] for row in bundle} == {
        "FAULT_SPECIFIC_LOSS_ALLOCATION",
        "REFUND_AND_PENALTY_STRUCTURE",
        "ABSENCE_EVIDENCE",
    }
    assert any(row["source_type"] == "SEARCH_AUDIT" for row in bundle)


def _patent_license_golden_review() -> dict:
    fixture = Path(__file__).resolve().parents[2] / "test_data" / "13、技术许可.docx"
    document = Document(str(fixture))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    return review_contract_v2(
        text,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="被许可方/付款方"),
    )


def test_patent_license_direction_and_provenance_golden_sample():
    output = _patent_license_golden_review()
    by_id = {row["item_id"]: row for row in output["items"]}
    manifest_path = Path(__file__).parent / "fixtures" / "technical_contracts" / "patent_license.golden.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))

    assert output["technology_transaction_consistency_gate"]["primary_contract_type"] == "technology_license_patent"
    assert output["perspective_resolution"]["canonical_role"] == "PATENT_LICENSEE"
    assert output["perspective_resolution"]["display_role"] == "被许可方 / 付款方"
    assert output["coverage_summary"] == {"通过": 18, "不适用": 29, "缺失": 2, "部分覆盖": 4}
    assert len(output["risk_findings"]) == 6
    assert len(output["recommendations"]) == 6
    assert output["release_quality_gate"]["status"] == "REPORT_READY"
    assert output["report_generation_decision"]["report_generated"] is True
    assert manifest["sample_status"] == "GOLDEN_REGRESSION_SAMPLE"
    assert manifest["sample_id"] == "TECH_LICENSE_PATENT_GOLDEN_001"
    assert output["technology_transaction_consistency_gate"]["primary_contract_type"] == manifest["primary_contract_type"]
    assert output["perspective_resolution"]["canonical_role"] == manifest["review_position"]
    assert output["coverage_summary"] == manifest["coverage_summary"]
    assert len(output["risk_findings"]) == manifest["risk_count"]
    assert len(output["recommendations"]) == manifest["recommendation_count"]
    applicable = {row["item_id"] for row in output["items"] if row["coverage_status"] != COVERAGE_NOT_APPLICABLE}
    not_applicable = {row["item_id"] for row in output["items"] if row["coverage_status"] == COVERAGE_NOT_APPLICABLE}
    assert applicable == set(manifest["expected_applicable_rules"])
    assert not_applicable == set(manifest["expected_not_applicable_rules"])
    assert all(by_id[review_id]["coverage_status"] == status for review_id, status in manifest["key_statuses"].items())

    assert by_id["TC-02-01"]["coverage_status"] == "通过"
    assert by_id["TC-02-02"]["coverage_status"] == "通过"
    assert by_id["TC-10-04"]["coverage_status"] == "通过"

    liability = by_id["TC-12-04"]
    liability_elements = {row["element_id"]: row for row in liability["coverage_elements"]}
    direction_snapshot = manifest["expected_party_direction"]
    assert liability["risk_direction"]["review_position_code"] == "PATENT_LICENSEE"
    assert liability["risk_direction"]["direction"] == "UNFAVORABLE"
    assert liability_elements["licensor_liability_cap"]["status"] == "MISSING"
    assert liability_elements["licensor_liability_cap"]["direction"] == "FAVORABLE"
    assert liability_elements["licensor_liability_cap"]["risk_relevant"] is False
    assert liability_elements["licensee_liability_cap"]["direction"] == "UNFAVORABLE"
    assert liability["risk_direction"]["review_position_code"] == direction_snapshot["review_position_code"]
    assert liability["risk_direction"]["direction"] == direction_snapshot["direction"]
    assert liability_elements[direction_snapshot["favorable_missing_element"]]["direction"] == "FAVORABLE"
    assert liability_elements[direction_snapshot["unfavorable_missing_element"]]["direction"] == "UNFAVORABLE"
    finding = next(row for row in output["risk_findings"] if "TC-12-04" in row["source_review_ids"])
    final_missing_ids = {row["element_id"] for row in finding["final_missing_elements"]}
    assert "licensor_liability_cap" not in final_missing_ids
    assert "mutual_liability_cap" not in final_missing_ids
    assert "licensee_liability_cap" in final_missing_ids
    assert "不得因此限制许可方" in finding["recommendation"]

    claim_item = by_id["TC-12-02"]
    assert "持续实施许可专利" in claim_item["legal_consequence_if_missing"]
    assert "软件持续使用" not in claim_item["legal_consequence_if_missing"]
    improvement = by_id["TC-10-03"]
    assert all(term not in improvement["legal_consequence_if_missing"] for term in ("合作核心技术", "共有成果"))
    report_text = output["result"]
    assert "软件持续使用" not in report_text
    assert "合作核心技术" not in report_text
    assert "共有成果" not in report_text

    explicit_gap = next(row for row in claim_item["evidence_clauses"] if row["clause_id"] == "C0049")
    assert explicit_gap["evidence_type"] == "CONTRACT_TEXT"
    assert explicit_gap["derived_or_original"] == "ORIGINAL"
    assert explicit_gap["document_id"] == "contract_package/main"
    assert explicit_gap["source_span"]["raw_text"] == explicit_gap["raw_text"]
    assert "双方未明确约定第三方索赔诉讼中的律师选择、抗辩控制以及和解决策权限" in explicit_gap["raw_text"]
    claim_finding = next(row for row in output["risk_findings"] if "TC-12-02" in row["source_review_ids"])
    evidence_bundle = claim_finding["evidence_bundle"]["items"]
    provenance_snapshot = manifest["critical_evidence_provenance"]
    source_gap = next(row for row in evidence_bundle if row["clause_id"] == provenance_snapshot["clause_id"])
    assert source_gap["evidence_id"] == provenance_snapshot["evidence_id"]
    assert source_gap["evidence_type"] == "CONTRACT_TEXT"
    assert source_gap["original_document_reference"] == "contract_package/main"
    assert source_gap["source_start"] > 0
    assert source_gap["source_end"] > source_gap["source_start"]
    assert source_gap["source_span"]["raw_text"] == source_gap["quoted_text"]
    assert source_gap["evidence_type"] == provenance_snapshot["evidence_type"]
    assert source_gap["derived_or_original"] == provenance_snapshot["derived_or_original"]
    assert provenance_snapshot["raw_text_contains"] in source_gap["raw_text"]
    assert validate_evidence_provenance(source_gap)["status"] == "PASS"
    absence = next(row for row in evidence_bundle if row["evidence_type"] == "SEARCH_AUDIT")
    assert absence["result"] == "NOT_FOUND"
    assert absence["search_scope"]
    assert absence["search_targets"]
    assert absence["searched_sections"]
    assert validate_evidence_provenance(absence)["status"] == "PASS"

    confidentiality = by_id["TC-11-02"]
    assert confidentiality["coverage_status"] == "通过"
    assert confidentiality["decision_trace"]["empty_field_detection"] == {
        "confidentiality_scope_filled": True,
        "confidentiality_period_filled": True,
        "placeholder_only": False,
    }
    raw_confidentiality = " ".join(row["raw_text"] for row in confidentiality["evidence_clauses"])
    assert "双方员工" in raw_confidentiality
    assert "自技术秘密披露之日起5年" in raw_confidentiality
    scope_field = extract_semantic_field_value(raw_confidentiality, "保密义务适用于")
    period_field = extract_semantic_field_value(raw_confidentiality, "保密期限")
    assert scope_field["is_filled"] is True
    assert period_field["is_filled"] is True
    assert "双方员工" in scope_field["field_value"]
    assert "5年" in period_field["field_value"]
    confidentiality_report = output["result"].split("## 45. 泄密责任（TC-11-04）", 1)[1].split("\n---", 1)[0]
    assert "保密义务适用于：双方员工" in confidentiality_report
    assert "保密期限：自技术秘密披露之日起5年" in confidentiality_report
    assert '原文片段："保密义务适用于： 保密期限："' not in confidentiality_report
    assert output["technology_transaction_consistency_gate"]["gate_results"]["REPORT_EVIDENCE_CONSISTENCY"] == "PASS"


def test_semantically_empty_template_fields_do_not_count_as_filled():
    assert is_semantically_filled_field("保密期限：\n_____", "保密期限") is False
    assert is_semantically_filled_field("保密期限：\n待填写", "保密期限") is False
    assert is_semantically_filled_field("保密义务适用于：\n□", "保密义务适用于") is False
    assert is_semantically_filled_field("保密期限：\n自披露之日起5年", "保密期限") is True
