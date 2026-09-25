import asyncio
import io
import json
from pathlib import Path

import pytest
from docx import Document

from review_v2.engine import (
    COVERAGE_NOT_APPLICABLE,
    _normalize_missing_element_for_recommendation,
    _technology_ip_proof_context,
    review_contract_v2,
)
from review_v2.technology_service_analysis import analyze_technology_relationship_risks
from review_v2.user_profile import UserRiskProfile
from routers.review_v2 import ReviewV2Result, review_v2_file


ROOT = Path(__file__).resolve().parents[2]
# A numeric prefix also selects 121/122/... and silently changes the contract
# under test. Keep this suite on its original entrusted-development fixture.
FIXTURE = ROOT / "test_data" / "12、委托开发.docx"


class _UploadFile:
    def __init__(self, filename: str, content: bytes):
        self.filename = filename
        self.file = io.BytesIO(content)


def _document_text(path: Path) -> str:
    document = Document(path)
    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        blocks.extend(
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in table.rows
        )
    return "\n".join(blocks)


def _fixture_text() -> str:
    return _document_text(FIXTURE)


@pytest.fixture(scope="module")
def party_a_review():
    return review_contract_v2(
        _fixture_text(),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方"),
    )


def _by_id(review):
    return {item["item_id"]: item for item in review["items"]}


def test_entrusted_classification_is_canonical_across_report_and_execution(party_a_review):
    canonical = party_a_review["contract_context"]["canonical_classification"]
    report = party_a_review["report_context"]

    assert canonical["contract_family"] == "TECHNOLOGY"
    assert canonical["contract_category"] == "TECHNOLOGY_DEVELOPMENT"
    assert canonical["contract_subtype"] == "ENTRUSTED_DEVELOPMENT"
    assert canonical["transaction_structure"] == "B2B_ENTRUSTED_DEVELOPMENT"
    assert party_a_review["contract_context"]["contract_subtype"] == canonical["contract_subtype"]
    assert party_a_review["contract_context"]["transaction_structure"] == canonical["transaction_structure"]
    assert report["contract_subtype"] == canonical["contract_subtype"]
    assert report["transaction_structure"] == canonical["transaction_structure"]
    assert report["scenario_label"] == "委托开发"
    assert report["contract_type_label"] == "技术合同—技术开发合同—委托开发合同"


def test_entrusted_party_roles_follow_identity_and_subtype(party_a_review):
    party_a = party_a_review["perspective_resolution"]
    assert party_a["review_party_identity"] == "PARTY_A"
    assert party_a["canonical_role"] == "COMMISSIONING_PARTY"
    assert party_a["display_role"] == "甲方 / 委托方"
    assert "专利权受让方" not in party_a["display_role"]

    party_b_review = review_contract_v2(
        _fixture_text(),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="乙方"),
    )
    party_b = party_b_review["perspective_resolution"]
    assert party_b["review_party_identity"] == "PARTY_B"
    assert party_b["canonical_role"] == "DEVELOPMENT_CONTRACTOR"
    assert party_b["display_role"] == "乙方 / 开发方"


def test_entrusted_rule_states_and_partial_elements_are_consistent(party_a_review):
    items = _by_id(party_a_review)
    expected = {
        "TC-01-01": "通过",
        "TC-01-04": "不适用",
        "TC-04-01": "通过",
        "TC-04-02": "通过",
        "TC-04-03": "部分覆盖",
        "TC-04-04": "部分覆盖",
        "TC-05-04": "部分覆盖",
        "TC-12-02": "部分覆盖",
        "TC-12-04": "部分覆盖",
        "TC-13-01": "通过",
        "TC-13-03": "部分覆盖",
    }
    assert {review_id: items[review_id]["coverage_status"] for review_id in expected} == expected

    for review_id in ("TC-04-03", "TC-04-04", "TC-05-04", "TC-12-02", "TC-12-04"):
        statuses = {element["status"] for element in items[review_id]["coverage_elements"]}
        assert "COVERED" in statuses
        assert "MISSING" in statuses
        audit = items[review_id]["decision_trace"]["specialized_search_audit"]
        assert audit["searched_full_document"] is True
        assert audit["searched_aliases"]
        assert audit["missing_elements"]
        narrative = " ".join(
            str(items[review_id].get(key) or "")
            for key in ("reasoning_summary", "risk_point_description")
        )
        assert "适用范围、履行标准和责任安排" not in narrative


def test_entrusted_rule_evidence_covers_body_and_attachment_semantics(party_a_review):
    items = _by_id(party_a_review)

    expected_terms = {
        "TC-04-01": ("主要功能", "里程碑"),
        "TC-04-02": ("开发费用", "付款"),
        "TC-04-03": ("前景知识产权", "申请"),
        "TC-05-04": ("解除", "开发费用"),
        "TC-12-02": ("第三方", "许可"),
        "TC-13-01": ("源代码", "测试报告"),
        "TC-13-03": ("验收", "测试"),
    }
    for review_id, terms in expected_terms.items():
        evidence_text = " ".join(
            str(row.get("text") or row.get("quote") or "")
            for row in items[review_id]["evidence_clauses"]
        )
        assert all(term in evidence_text for term in terms)


def test_entrusted_report_is_generated_without_fatal_quality_gate(party_a_review):
    assert party_a_review["report_generation_decision"]["status"] == "GENERATED"
    assert party_a_review["report_generation_decision"]["report_generated"] is True
    assert party_a_review["report_generation_decision"]["fatal_rule_ids"] == []


def test_unsubstantiated_generic_missing_is_local_manual_review(party_a_review):
    generic_phrase = "适用范围、履行标准和责任安排"
    manual_items = []
    for item in party_a_review["items"]:
        if item["coverage_status"] == "缺失" and generic_phrase in item["risk_point_description"]:
            audit = item["decision_trace"].get("specialized_search_audit")
            assert audit and audit["searched_full_document"] is True
            assert audit["missing_elements"]
        if item["coverage_status"] == "INSUFFICIENT_CONTEXT":
            manual_items.append(item["item_id"])
            assert item["needs_manual_review"] is True
            assert item["decision_trace"]["missing_fallback_resolution"]["original_status"] == "MISSING"
            assert generic_phrase not in item["risk_point_description"]
    assert all(review_id for review_id in manual_items)


def test_acceptance_evidence_does_not_invent_third_party_testing(party_a_review):
    item = _by_id(party_a_review)["TC-13-03"]
    elements = {row["element_id"]: row for row in item["coverage_elements"]}

    assert item["coverage_status"] == "部分覆盖"
    assert elements["acceptance_metrics"]["status"] == "COVERED"
    assert elements["test_method_and_report"]["status"] == "COVERED"
    assert elements["acceptance_dispute_verification"]["status"] == "MISSING"
    assert elements["acceptance_dispute_verification"]["supporting_evidence_ids"] == []
    evidence_text = " ".join(str(row.get("quote") or row.get("text") or "") for row in item["evidence_clauses"])
    assert "第三方检测机构" not in evidence_text


def test_explicit_third_party_testing_clause_is_covered():
    text = _fixture_text() + "\n如双方对测试结果存在争议，共同委托有资质的第三方检测机构进行检测，检测结果作为验收依据。"
    review = review_contract_v2(
        text,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方"),
    )
    item = _by_id(review)["TC-13-03"]
    element = next(row for row in item["coverage_elements"] if row["element_id"] == "acceptance_dispute_verification")

    assert item["coverage_status"] == "通过"
    assert element["status"] == "COVERED"
    assert element["supporting_evidence_ids"]


def test_post_termination_confidentiality_is_directly_covered(party_a_review):
    item = _by_id(party_a_review)["TC-11-02"]
    elements = {row["element_id"]: row for row in item["coverage_elements"]}

    assert item["coverage_status"] == "通过"
    assert elements["during_contract_confidentiality"]["status"] == "COVERED"
    assert elements["during_contract_confidentiality"]["supporting_evidence_ids"]
    assert elements["post_termination_confidentiality"]["status"] == "COVERED"
    assert elements["post_termination_confidentiality"]["supporting_evidence_ids"]
    assert elements["confidentiality_duration"]["status"] == "COVERED"
    assert elements["confidentiality_duration"]["supporting_evidence_ids"]
    evidence_text = " ".join(str(row.get("quote") or row.get("text") or "") for row in item["evidence_clauses"])
    assert "终止或解除之后仍需履行" in evidence_text


def test_recommendation_normalization_removes_negative_prefixes(party_a_review):
    assert _normalize_missing_element_for_recommendation("未专项明确专利申请权归属") == ("专项明确", "专利申请权归属")
    report = party_a_review["result"]
    assert "建议明确未专项明确" not in report
    assert "建议明确未明确" not in report
    assert "建议明确未约定" not in report


def test_gap_consequences_are_bound_to_actual_missing_elements(party_a_review):
    items = _by_id(party_a_review)
    patent = items["TC-04-03"]["narrative_state"]["consequence"]
    failure = items["TC-04-04"]["narrative_state"]["consequence"]
    exit_consequence = items["TC-05-04"]["narrative_state"]["consequence"]

    assert all(term in patent for term in ("申请主体", "文件签署", "权利登记"))
    assert "第三方提出权利主张" not in patent
    assert all(term in failure for term in ("研发费用", "阶段成果", "资料交接"))
    assert "风险转移、质量异议" not in failure
    exit_elements = items["TC-05-04"]["coverage_elements"]
    unresolved = [row for row in exit_elements if row["status"] in {"MISSING", "PARTIALLY_COVERED"}]
    assert unresolved
    assert all(row["required_fact"] in exit_consequence for row in unresolved)
    assert all(row["required_fact"] not in exit_consequence
               for row in exit_elements if row["status"] == "COVERED")
    assert "付款期限、逾期责任" not in exit_consequence


def test_manual_review_and_insufficient_context_states_are_explained(party_a_review):
    missing_count = party_a_review["coverage_summary"].get("缺失", 0)
    insufficient = [item for item in party_a_review["items"] if item["coverage_status"] == "INSUFFICIENT_CONTEXT"]
    assert missing_count == sum(item["coverage_status"] == "缺失" for item in party_a_review["items"])
    for item in party_a_review["items"]:
        if item["coverage_status"] == "通过" and item["needs_manual_review"]:
            assert item["evidence_clauses"]
            assert item["quality_gate_failure_reason"]
        if item["coverage_status"] == "缺失" and item["decision_trace"].get("specialized_search_audit"):
            assert item["needs_manual_review"] is False
    for item in insufficient:
        assert item["issue_type"] == "人工复核"
        assert item["quality_gate_failure_reason"]


def test_entrusted_semantic_coverage_and_conditional_applicability(party_a_review):
    items = _by_id(party_a_review)
    patent = items["TC-04-03"]
    patent_elements = {row["element_id"]: row for row in patent["coverage_elements"]}

    assert patent["decision_trace"]["open_ended_enumeration"] == {
        "detected": True,
        "scope": "OPEN",
        "semantic_result": "SEMANTICALLY_COVERED",
        "applies_to": ["patent_application_right_ownership", "patent_right_after_grant"],
    }
    assert patent_elements["patent_application_right_ownership"]["status"] == "COVERED"
    assert patent_elements["patent_application_right_ownership"]["reason"] == "SEMANTICALLY_COVERED"
    assert patent_elements["document_signature_cooperation"]["status"] == "MISSING"
    assert "专利申请权归属及申请协助义务未专项明确" not in patent["risk_point_description"]

    secret = items["TC-05-03"]
    assert secret["applicability"]["status"] == "CONDITIONALLY_APPLICABLE"
    assert secret["applicability"]["applicable"] is True
    assert secret["needs_manual_review"] is False
    assert secret["coverage_status"] == "部分覆盖"
    secret_elements = {row["element_id"]: row for row in secret["coverage_elements"]}
    assert secret_elements["technical_secret_result_ownership"]["status"] == "COVERED"
    assert secret_elements["source_algorithm_document_delivery"]["status"] == "COVERED"
    assert secret_elements["developer_copy_retention"]["status"] == "MISSING"

    improvement = items["TC-10-03"]
    assert improvement["applicability"]["status"] == "NOT_APPLICABLE"
    assert improvement["coverage_status"] == "不适用"
    assert improvement["needs_manual_review"] is False


def test_confidentiality_liability_reuses_semantically_valid_evidence(party_a_review):
    item = _by_id(party_a_review)["TC-11-04"]
    evidence_by_role = {
        row.get("evidence_role"): str(row.get("quote") or row.get("text") or "")
        for row in item["evidence_clauses"]
    }

    assert item["applicability"]["status"] == "APPLICABLE"
    assert item["coverage_status"] == "部分覆盖"
    assert item["needs_manual_review"] is False
    assert "无法自公开渠道" in evidence_by_role["CONFIDENTIALITY_OBLIGATION"]
    assert "不得向任何第三方泄露" in evidence_by_role["CONFIDENTIALITY_OBLIGATION"]
    assert "其他违反本合同" in evidence_by_role["GENERAL_BREACH_LIABILITY"]
    assert "技术支持服务的对象" not in " ".join(evidence_by_role.values())


def test_general_breach_catch_all_propagates_only_damage_remedies(party_a_review):
    item = _by_id(party_a_review)["TC-11-04"]
    elements = {row["element_id"]: row["status"] for row in item["coverage_elements"]}
    propagation = item["decision_trace"]["general_breach_semantic_propagation"]
    finding = next(
        row for row in party_a_review["risk_findings"]
        if row["risk_group"] == "confidentiality_breach_liability"
    )

    assert propagation["classification"] == "GENERAL_BREACH_CATCH_ALL"
    assert all(elements[element_id] == "COVERED" for element_id in (
        "confidentiality_obligation", "breach_damage_liability", "direct_loss",
        "lost_profit", "attorney_fee", "litigation_cost",
    ))
    assert all(elements[element_id] == "MISSING" for element_id in (
        "stop_disclosure", "stop_use", "mitigation_obligation",
        "return_or_destroy", "disposition_confirmation",
    ))
    assert elements["specific_confidentiality_penalty"] == "NOT_APPLICABLE"
    assert "一般违约责任需明确延伸" not in item["risk_point_description"]
    assert "损失赔偿" in item["risk_point_description"]
    assert all(term in finding["recommendation"] for term in ("停止使用", "防止继续扩散", "通知受影响方"))
    assert "无需重复创设同一赔偿义务" in finding["recommendation"]
    assert "专项违约金" not in " ".join((
        item["risk_point_description"],
        finding["issue"],
        finding["recommendation"],
    ))


def test_secret_copy_retention_uses_matching_legal_consequence(party_a_review):
    item = _by_id(party_a_review)["TC-05-03"]
    consequence = item["narrative_state"]["consequence"]

    assert item["decision_trace"]["legal_consequence_mapping"]["semantic_domain"] == "SECRET_COPY_RETENTION"
    assert all(term in consequence for term in ("源代码", "副本", "未经授权", "披露", "删除"))
    assert all(term not in consequence for term in ("付款期限", "验收节点", "质量异议", "价格"))
    assert party_a_review["report_semantic_consistency_gate"]["gate_results"][
        "FINDING_CONSEQUENCE_CONSISTENCY_GATE"
    ] == "PASS"


def test_improvement_notification_rendering_respects_development_subtype(party_a_review):
    commissioned = _by_id(party_a_review)["TC-10-03"]
    commissioned_rendering = commissioned["decision_trace"]["subtype_rendering"]
    commissioned_consequence = commissioned["narrative_state"]["consequence"]

    assert commissioned_rendering["primary_subtype"] == "ENTRUSTED_DEVELOPMENT"
    assert commissioned_rendering["has_follow_on_entitlement"] is False
    assert commissioned["coverage_status"] == "不适用"
    assert commissioned["applicability"]["status"] == "NOT_APPLICABLE"
    assert commissioned_consequence == ""

    cooperative_fixture = next((ROOT / "test_data").glob("11*.docx"))
    cooperative_review = review_contract_v2(
        _document_text(cooperative_fixture),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方"),
    )
    cooperative = _by_id(cooperative_review)["TC-10-03"]
    cooperative_rendering = cooperative["decision_trace"]["subtype_rendering"]
    cooperative_consequence = cooperative["narrative_state"]["consequence"]

    assert cooperative_rendering["primary_subtype"] == "COOPERATIVE_DEVELOPMENT"
    assert cooperative_rendering["has_follow_on_entitlement"] is True
    assert "合作开发成果" in cooperative_consequence
    assert "共同权利" in cooperative_consequence


def test_entrusted_improvement_boundary_requires_affirmative_restriction():
    text = _fixture_text() + "\n乙方不得继续开展后续改进；形成的改进成果必须无偿转让给甲方。"
    review = review_contract_v2(
        text,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方"),
    )
    item = _by_id(review)["TC-10-04"]

    assert review["contract_context"]["canonical_classification"]["contract_subtype"] == "ENTRUSTED_DEVELOPMENT"
    assert item["applicability"]["status"] == "APPLICABLE"
    assert item["coverage_status"] == "风险条款"
    assert item["risk_group"] == "improvement_no_unlawful_restriction"
    assert item["decision_trace"]["entrusted_improvement_applicability"]["matched_clause_ids"]


def test_commissioning_party_liability_direction_excludes_favorable_no_cap_gap(party_a_review):
    item = _by_id(party_a_review)["TC-12-04"]
    elements = {row["element_id"]: row for row in item["coverage_elements"]}
    direction = item["risk_direction"]

    assert item["coverage_status"] == "部分覆盖"
    assert elements["developer_liability_cap"]["status"] == "NOT_APPLICABLE"
    assert elements["developer_liability_cap"]["direction"] == "FAVORABLE"
    assert elements["developer_liability_cap"]["risk_relevant"] is False
    assert elements["special_general_liability_coordination"]["direction"] == "UNFAVORABLE"
    assert elements["special_general_liability_coordination"]["risk_relevant"] is True
    assert direction["favorable_missing_elements"] == []
    assert direction["favorable_existing_conditions"] == ["developer_liability_cap"]
    assert direction["unfavorable_missing_elements"] == ["special_general_liability_coordination"]
    assert item["title"] == "侵权责任与一般救济衔接"
    assert all(term in item["narrative_state"]["consequence"] for term in ("适用顺序", "累计关系", "重复计算"))
    assert all(term not in item["narrative_state"]["consequence"] for term in ("缺少责任上限", "责任总额难以预估"))
    assert "无需主动为开发方新增责任上限" in item["suggested_clause_example"]
    finding = next(
        row for row in party_a_review["risk_findings"]
        if row["risk_group"] == "indemnity_and_liability_cap"
    )
    assert "无需主动为开发方新增责任上限" in finding["recommendation"]


def test_entrusted_partial_findings_have_positive_and_absence_evidence(party_a_review):
    findings = {row["risk_group"]: row for row in party_a_review["risk_findings"]}
    for group in (
        "entrusted_patent_application_right",
        "development_failure_risk",
        "technical_secret_result_rights",
        "cooperation_breach_and_exit",
        "confidentiality_breach_liability",
        "indemnity_and_liability_cap",
    ):
        evidence = findings[group]["evidence_bundle"]["items"]
        assert any(row["source_type"] != "SEARCH_AUDIT" for row in evidence)
        assert any(row["source_type"] == "SEARCH_AUDIT" for row in evidence)

    indemnity_absence = next(
        row for row in findings["indemnity_and_liability_cap"]["evidence_bundle"]["items"]
        if row["source_type"] == "SEARCH_AUDIT"
    )
    assert "适用和累计衔接" in indemnity_absence["quoted_text"]
    assert "责任总额上限" not in indemnity_absence["quoted_text"]


def test_commissioned_development_golden_manifest(party_a_review):
    manifest_path = Path(__file__).parent / "fixtures" / "technical_contracts" / "commissioned_development.golden.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    items = _by_id(party_a_review)
    classification = party_a_review["contract_context"]["canonical_classification"]

    assert manifest["sample_id"] == "TECH_DEV_COMMISSIONED_GOLDEN_001"
    assert manifest["sample_status"] == "GOLDEN_CANDIDATE"
    assert party_a_review["technology_transaction_consistency_gate"]["primary_contract_type"] == manifest["primary_contract_type"]
    assert classification["contract_family"] == manifest["contract_family"]
    assert classification["contract_category"] == manifest["contract_category"]
    assert classification["contract_subtype"] == manifest["contract_subtype"]
    assert party_a_review["perspective_resolution"]["canonical_role"] == manifest["review_position"]
    assert all(items[review_id]["coverage_status"] == status for review_id, status in manifest["key_statuses"].items())
    assert all(
        items[review_id]["applicability"]["status"] == "CONDITIONALLY_APPLICABLE"
        for review_id in manifest["conditional_applicability"]
    )
    assert all(items[review_id]["coverage_status"] == "不适用" for review_id in manifest["not_applicable"])
    assert not any(
        row["risk_group"] in manifest["suppressed_risk_groups"]
        for row in party_a_review["risk_findings"]
    )

    confidentiality = items[manifest["confidentiality_catch_all"]["review_id"]]
    confidentiality_elements = {row["element_id"]: row["status"] for row in confidentiality["coverage_elements"]}
    assert confidentiality["decision_trace"]["general_breach_semantic_propagation"]["classification"] == manifest[
        "confidentiality_catch_all"
    ]["classification"]
    assert all(
        confidentiality_elements[element_id] == "COVERED"
        for element_id in manifest["confidentiality_catch_all"]["covered_elements"]
    )
    assert all(
        confidentiality_elements[element_id] == "MISSING"
        for element_id in manifest["confidentiality_catch_all"]["missing_specific_remedies"]
    )

    for review_id, semantic_domain in manifest["legal_consequence_domains"].items():
        assert items[review_id]["decision_trace"]["legal_consequence_mapping"]["semantic_domain"] == semantic_domain

    improvement = items[manifest["improvement_notification"]["review_id"]]
    rendering = improvement["decision_trace"]["subtype_rendering"]
    assert rendering["primary_subtype"] == manifest["improvement_notification"]["primary_subtype"]
    assert rendering["has_follow_on_entitlement"] is manifest["improvement_notification"]["has_follow_on_entitlement"]
    assert improvement["coverage_status"] == manifest["improvement_notification"]["coverage_status"]
    assert all(
        term not in improvement["narrative_state"]["consequence"]
        for term in manifest["improvement_notification"]["forbidden_default_terms"]
    )


def test_production_file_api_renders_entrusted_development_snapshot():
    engine_output = asyncio.run(
        review_v2_file(
            file=_UploadFile(FIXTURE.name, FIXTURE.read_bytes()),
            contract_type="technology",
            include_low_risk=True,
            use_llm=False,
            profile_json=json.dumps({"user_identity": "甲方"}, ensure_ascii=False),
            current_user={"id": 1},
        )
    )
    output = ReviewV2Result.model_validate(engine_output).model_dump()
    items = _by_id(output)
    markdown = output["result"]
    ownership = {
        row["final_output"]: row
        for row in output["observability"]["final_output_ownership_matrix"]
    }

    assert output["filename"] == FIXTURE.name
    assert output["extracted_text_length"] > 1000
    assert output["report_generation_decision"]["report_generated"] is True
    assert output["release_quality_gate"]["status"] == "REPORT_READY"
    assert output["perspective_resolution"]["display_role"] == "甲方 / 委托方"
    assert "审查立场：甲方 / 委托方" in markdown

    expected_statuses = {
        "TC-04-01": "通过",
        "TC-04-02": "通过",
        "TC-04-03": "部分覆盖",
        "TC-04-04": "部分覆盖",
        "TC-05-04": "部分覆盖",
        "TC-12-02": "部分覆盖",
        "TC-12-04": "部分覆盖",
        "TC-13-01": "通过",
        "TC-13-03": "部分覆盖",
    }
    assert {review_id: items[review_id]["coverage_status"] for review_id in expected_statuses} == expected_statuses

    third_party = next(
        element
        for element in items["TC-13-03"]["coverage_elements"]
        if element["element_id"] == "acceptance_dispute_verification"
    )
    assert third_party["status"] == "MISSING"
    assert third_party["supporting_evidence_ids"] == []
    assert "第三方检测机制明确" not in markdown

    confidentiality = items["TC-11-02"]
    assert confidentiality["coverage_status"] == "通过"
    assert "终止或解除之后仍需履行" in " ".join(
        str(row.get("quote") or row.get("text") or "")
        for row in confidentiality["evidence_clauses"]
    )
    assert "建议明确未明确" not in markdown
    assert "建议明确未约定" not in markdown
    assert "建议明确未专项明确" not in markdown
    assert "建议补充明确合同尚未" not in markdown
    assert all(term in markdown for term in ("申请主体", "文件签署", "权利登记"))
    assert all(term in markdown for term in ("研发费用", "阶段成果", "资料交接"))
    assert all(term in markdown for term in ("接管项目", "继续开发", "既有投入"))

    assert ownership["正式建议"]["current_producer"].endswith("FinalDisposition.recommendation")
    assert ownership["缺失的法律后果"]["current_producer"].startswith("RiskFinding.consequence")
    assert output["manual_review_summary"]["manual_review_count"] == 0


def test_explicit_review_role_is_not_overridden_by_generic_subtype_identity():
    engine_output = asyncio.run(
        review_v2_file(
            file=_UploadFile(FIXTURE.name, FIXTURE.read_bytes()),
            contract_type="technology",
            include_low_risk=True,
            use_llm=False,
            profile_json=json.dumps(
                {
                    "review_role": "甲方 / 委托方",
                    "user_identity": "委托开发合同当事人",
                },
                ensure_ascii=False,
            ),
            current_user={"id": 1},
        )
    )
    output = ReviewV2Result.model_validate(engine_output).model_dump()
    items = _by_id(output)

    assert output["profile"]["user_identity"] == "甲方 / 委托方"
    assert output["perspective_resolution"]["canonical_role"] == "COMMISSIONING_PARTY"
    assert output["perspective_resolution"]["display_role"] == "甲方 / 委托方"
    assert output["contract_context"]["review_party_identity"] == "PARTY_A"
    assert "审查立场：甲方 / 委托方" in output["result"]
    assert "审查立场：委托开发合同当事人" not in output["result"]
    assert len(output["risk_findings"]) == 11
    assert engine_output["final_disposition_summary"]["actionable"] == 8
    assert engine_output["final_disposition_summary"]["optional_improvement"] == 3
    assert output["coverage_summary"]["不适用"] == 34
    assert all(items[review_id]["coverage_status"] == "不适用" for review_id in (
        "TC-10-01", "TC-10-02", "TC-10-03", "TC-10-04",
    ))


def test_future_development_warranty_does_not_trigger_existing_ip_proof_attachments():
    text = """
    软件委托开发合同。甲方委托乙方开发新的业务系统。
    开发成果和源代码交付甲方，乙方保证开发成果不侵犯第三方知识产权。
    """
    output = review_contract_v2(
        text,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方", contract_scenario="委托开发"),
    )
    items = _by_id(output)
    assert _technology_ip_proof_context(text)["proof_asset_transaction"] is False
    assert items["TC-03-04"]["coverage_status"] == COVERAGE_NOT_APPLICABLE
    assert items["TC-13-02"]["coverage_status"] == COVERAGE_NOT_APPLICABLE
    assert "技术合同均适用" not in str(items["TC-03-04"]["applicability"])
    assert not any(row["risk_group"] == "technology_ip_rights_provenance" for row in output["risk_findings"])


def test_existing_patent_transfer_still_triggers_ip_proof_review():
    text = """
    专利权转让合同。让与方将其现有发明专利权转让给受让方。
    标的专利名称为测试专利，专利号为CN123456，双方应共同办理专利权转让登记。
    """
    output = review_contract_v2(
        text,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="受让方", contract_scenario="专利权转让"),
    )
    items = _by_id(output)
    assert _technology_ip_proof_context(text)["existing_ip_asset_transaction"] is True
    assert items["TC-03-04"]["coverage_status"] != COVERAGE_NOT_APPLICABLE
    assert items["TC-13-02"]["coverage_status"] != COVERAGE_NOT_APPLICABLE


def test_shared_personal_information_trigger_requires_business_processing_fact():
    generic = analyze_technology_relationship_risks(
        "软件委托开发涉及业务数据、数据库、接口、用户模块和账号系统。"
    )
    explicit = analyze_technology_relationship_risks(
        "乙方受托处理客户姓名、手机号和联系地址，用于完成软件委托开发。"
    )
    contact_only = analyze_technology_relationship_risks(
        "项目联系人张三，手机号13800000000；双方通过指定电子邮箱提供项目文档。"
    )
    feature_only = analyze_technology_relationship_risks(
        "系统需处理的数据包括员工姓名、手机号、企业邮箱、系统账号、操作日志和消费者收货地址；"
        "软件部署在甲方本地，乙方不接触真实生产数据。"
    )
    production_access = analyze_technology_relationship_risks(
        "系统处理客户姓名和手机号；乙方负责生产环境运维并可访问生产数据库。"
    )
    assert generic["personal_information_processing_gap"]["present"] is False
    assert contact_only["personal_information_processing_gap"]["present"] is False
    assert feature_only["personal_information_processing_gap"]["present"] is False
    assert feature_only["personal_information_processing_gap"]["analysis"]["trigger_strength"] == "WEAK"
    assert explicit["personal_information_processing_gap"]["present"] is True
    assert production_access["personal_information_processing_gap"]["present"] is True
    assert production_access["personal_information_processing_gap"]["analysis"]["trigger_strength"] == "STRONG"
    assert {"姓名", "手机号", "联系地址"} <= set(
        explicit["personal_information_processing_gap"]["analysis"]["matched_categories"]
    )


def test_current_commissioning_party_liability_overlap_is_an_enforceability_risk(party_a_review):
    item = _by_id(party_a_review)["GEN-LIAB-01"]
    trace = item["decision_trace"]
    cluster = trace["risk_recall_layer"]["analysis"]["selected_cluster"]
    assert cluster["liable_party"] == "乙方"
    assert cluster["beneficiary_party"] == "甲方"
    assert trace["party_interest_check"]["result"] == "ENFORCEABILITY_UNCERTAINTY"
    assert trace["risk_to_current_party"]["multiple_remedies_are_adverse"] is False
    assert "执行确定性" in item["legal_consequence_if_missing"] or "实际可执行" in item["legal_consequence_if_missing"]
    assert "最大限度保留甲方" in item["suggested_clause_example"]
