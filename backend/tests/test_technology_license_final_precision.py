from pathlib import Path

from docx import Document

from review_v2.engine import review_contract_v2
from review_v2.technology_service_analysis import analyze_license_lifecycle, analyze_timeline
from review_v2.user_profile import UserRiskProfile


ROOT = Path(__file__).resolve().parents[2]
GOLD = ROOT / "test_data" / "19 技术许可.docx"

LICENSE_BASE = """
技术许可合同
乙方（许可方）授予甲方（被许可方）在中国境内五年非独占实施专利技术的许可，专利号CN202610000001.X。
甲方支付许可费100万元。甲方不得超出许可范围实施或擅自再许可。
"""


def _review(text: str) -> dict:
    return review_contract_v2(
        text,
        "technology",
        include_low_risk=True,
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/被许可方"),
    )


def _item(output: dict, item_id: str) -> dict:
    return next(row for row in output["items"] if row["item_id"] == item_id)


def _finding(output: dict, risk_group: str) -> dict:
    return next(row for row in output["risk_findings"] if row["risk_group"] == risk_group)


def _gold_text() -> str:
    document = Document(str(GOLD))
    return "\n".join([
        *(paragraph.text for paragraph in document.paragraphs),
        *(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows),
    ])


def test_technical_material_types_are_positive_evidence_not_a_false_absence():
    output = _review(
        LICENSE_BASE
        + "\n许可方应交付总体设计图、机械零件图、电气原理图、PLC源程序、安装调试手册和维修手册。"
    )
    item = _item(output, "TC-13-01")
    elements = {row["element_id"]: row for row in item["coverage_elements"]}

    assert elements["major_document_types"]["status"] == "COVERED"
    assert "未明确技术资料具体清单及文件类型" not in item["risk_point_description"]
    finding = _finding(output, "technical_material_attachment")
    audit = finding["evidence_bundle"]["absence_evidence"]
    assert audit["query_semantics"] == "POSITIVE_REQUIRED_FACT"
    assert all(not target.startswith(("未明确", "未约定", "未覆盖", "缺失", "不存在", "没有")) for target in audit["required_elements"])
    assert "未找到：未明确" not in output["result"]


def test_vague_material_obligation_can_report_missing_file_types():
    output = _review(LICENSE_BASE + "\n许可方应提供相关技术资料。")
    item = _item(output, "TC-13-01")
    elements = {row["element_id"]: row for row in item["coverage_elements"]}

    assert elements["major_document_types"]["status"] == "MISSING"
    assert "未列举主要技术资料文件类型" in item["risk_point_description"]


def test_material_categories_without_version_or_attachment_are_not_overstated():
    output = _review(
        LICENSE_BASE
        + "\n许可方应向被许可方交付总体设计图、机械图、电气图、源程序及安装手册。"
    )
    item = _item(output, "TC-13-01")
    elements = {row["element_id"]: row for row in item["coverage_elements"]}

    assert item["coverage_status"] == "部分覆盖"
    assert item["risk_level"] == "低"
    assert elements["major_document_types"]["status"] == "COVERED"
    assert elements["closed_attachment_schedule"]["status"] == "MISSING"
    assert elements["file_version"]["status"] == "MISSING"
    assert elements["file_identifier"]["status"] == "MISSING"
    assert "未明确技术资料具体清单及文件类型" not in item["risk_point_description"]


def test_permanent_license_rights_are_favorable_to_licensee():
    facts = analyze_license_lifecycle(
        "合同期限5年。合同终止后被许可方仍有权继续使用技术、使用技术资料、制造合同产品并销售合同产品；"
        "上述权利永久、不可撤销。终止后是否继续支付提成费未约定。"
    )

    assert "终止后继续使用技术" in facts["favorable_elements"]
    assert "终止后继续销售合同产品" in facts["favorable_elements"]
    assert "上述许可权利永久存续" in facts["favorable_elements"]
    assert "上述许可权利不可撤销" in facts["favorable_elements"]
    assert facts["adverse_elements"] == []
    assert facts["party_interest"]["overall"] == "AMBIGUOUS"


def test_undefined_post_termination_royalty_is_the_only_lifecycle_risk():
    output = _review(_gold_text())
    finding = _finding(output, "license_lifecycle_boundary_gap")
    item = _item(output, "TECH-SHARED-LICENSE-LIFECYCLE-01")
    interest = item["decision_trace"]["party_interest_check"]

    assert finding["title"] == "合同终止后继续销售的提成费义务不明确"
    assert finding["severity"] == "MEDIUM"
    assert finding["party_effect"] == "AMBIGUOUS"
    assert interest["adverse_elements"] == []
    assert interest["favorable_elements"]
    assert interest["ambiguous_elements"]
    assert "权利本身对被许可方有利" in finding["issue"]
    assert finding["severity_reason"]["major_severity_reason"] == ""


def test_license_timeline_uses_material_delivery_not_service_start():
    text = (
        LICENSE_BASE
        + "\n合同生效日期为2026年12月20日。"
        + "\n第一批技术资料应于2026年11月20日前交付。"
    )
    timeline = analyze_timeline(text)
    output = _review(text)
    finding = _finding(output, "technical_timeline_conflict")

    assert timeline["primary_conflict_subject"] == "TECHNICAL_MATERIAL_DELIVERY"
    assert "技术资料交付义务早于合同生效日" in finding["issue"]
    assert "服务开始" not in finding["issue"]
    assert "服务开始" not in finding["recommendation"]


def test_cross_border_explicit_gaps_survive_fallback_and_keep_medium_severity():
    output = _review(
        LICENSE_BASE
        + "\n许可方为德国企业，被许可方为中国企业，本合同属于跨境技术交易。"
        + "\n双方应遵守技术进出口、出口管制、制裁、外汇、税务、知识产权和数据跨境规则，"
        + "但具体审批、登记、备案、许可文件、办理材料、期限及责任分工未约定，双方未形成专项附件。"
    )
    item = _item(output, "TC-13-04")
    finding = _finding(output, "technology_import_export_control")
    missing = {row["element_id"] for row in finding["final_missing_elements"]}

    assert item["coverage_status"] == "部分覆盖"
    assert finding["severity"] == "MEDIUM"
    assert item["decision_trace"]["missing_element_resolution"]["selected_source"] == "structured_missing_elements"
    assert {
        "technology_import_license_responsible_party",
        "technology_import_registration_responsible_party",
    } <= missing
    assert item["decision_trace"]["regulatory_reasoning_chain"]["control_category"] == "UNDETERMINED_REQUIRES_CATALOG_LOOKUP"
    assert "现有信息不足以确定具体缺失字段" not in output["result"]
    assert "查询现行技术进出口目录" in finding["issue"]
    assert any(term in finding["recommendation"] for term in ("禁止履行", "不得履行"))
    assert all(term in finding["recommendation"] for term in ("进口许可", "合同登记", "责任方"))


def test_generic_compliance_sentence_without_cross_border_fact_does_not_invent_specific_gap():
    output = _review(LICENSE_BASE + "\n双方应遵守相关法律法规。")
    item = _item(output, "TC-13-04")

    assert item["coverage_status"] in {"不适用", "信息不足"}
    assert not any(row["risk_group"] == "technology_import_export_control" for row in output["risk_findings"])


def test_open_ended_material_types_do_not_prove_closed_attachment_schedule():
    output = _review(
        LICENSE_BASE
        + "\n技术资料包括但不限于总体设计图、机械零件图、电气原理图、PLC源程序和安装手册。"
    )
    elements = {row["element_id"]: row for row in _item(output, "TC-13-01")["coverage_elements"]}

    assert elements["major_document_types"]["status"] == "COVERED"
    assert elements["major_document_types"]["evidence_strength"] == "EXPLICIT"
    assert elements["closed_attachment_schedule"]["status"] == "MISSING"
    assert elements["closed_attachment_schedule"]["evidence_strength"] == "WEAK_INFERENCE"


def test_explicit_attachment_schedule_with_identifiers_and_versions_is_covered():
    output = _review(
        LICENSE_BASE
        + "\n附件一：《技术资料交付清单》\n"
        + "文件名称：总体设计图；文件编号：DWG-001；版本：V1.0；数量：1套。\n"
        + "文件名称：PLC源程序；文件编号：PLC-002；版本：V2.1；数量：1套。"
    )
    elements = {row["element_id"]: row for row in _item(output, "TC-13-01")["coverage_elements"]}

    assert elements["closed_attachment_schedule"]["status"] == "COVERED"
    assert elements["closed_attachment_schedule"]["evidence_strength"] in {"EXPLICIT", "STRONG_INFERENCE"}
    assert elements["file_version"]["status"] == "COVERED"
    assert elements["file_identifier"]["status"] == "COVERED"


def test_material_attachment_gaps_are_aggregated_into_one_low_risk_finding():
    output = _review(
        LICENSE_BASE
        + "\n技术资料包括但不限于总体设计图、机械零件图、电气原理图、PLC源程序和安装手册。"
    )
    findings = [row for row in output["risk_findings"] if row["risk_group"] == "technical_material_attachment"]

    assert len(findings) == 1
    assert findings[0]["severity"] == "LOW"
    assert "主要类别已经明确" in findings[0]["issue"]


def test_high_risks_default_to_core_priority_without_unexplained_downgrade():
    output = _review(_gold_text())
    payment = _finding(output, "payment_formula_inconsistency")
    improvement = _finding(output, "improvement_no_unlawful_restriction")

    assert payment["severity"] == "HIGH"
    assert payment["priority"] == "A_CORE"
    assert payment["aggregation_signals"]["priority_projection_source"] == "final_review_item_priority"
    assert improvement["severity"] == "HIGH"
    assert improvement["priority"] == "A_CORE"


def test_confidentiality_disposal_gap_is_owned_only_by_tc_11_03():
    output = _review(_gold_text())
    disposal = _item(output, "TC-11-03")
    liability = _item(output, "TC-11-04")
    liability_elements = {row["element_id"]: row for row in liability["coverage_elements"]}

    assert disposal["coverage_status"] != "通过"
    for element_id in ("return_or_destroy", "return_delete_destroy", "disposition_confirmation"):
        if element_id in liability_elements:
            assert liability_elements[element_id].get("risk_relevant") is False
    assert all(
        term not in liability["risk_point_description"]
        for term in ("返还", "删除", "销毁", "书面处置确认")
    )
    finding = _finding(output, "confidentiality_breach_liability")
    assert all(term not in finding["recommendation"] for term in ("返还", "删除", "销毁", "书面处置确认"))


def test_license_third_party_claim_advanced_only_gap_is_not_a_required_risk():
    output = _review(
        LICENSE_BASE
        + "\n被许可方按本合同约定实施技术时，如收到第三方知识产权权利主张，应在5个工作日内书面通知许可方。"
        + "\n许可方应负责与第三方交涉、处理和应诉，提供权属证明及技术来源文件，并承担赔偿责任和被许可方的直接经济损失。"
        + "\n未经被许可方书面同意，许可方不得与第三方达成损害被许可方权益的和解。"
        + "\n因许可方权利瑕疵导致被许可方无法继续合法实施时，被许可方有权解除合同，许可方应返还全部许可费并赔偿损失。"
    )
    item = _item(output, "TC-12-02")
    tiers = item["decision_trace"]["third_party_claim_element_tiers"]

    assert tiers["necessary_protection_gaps"] == []
    assert tiers["advanced_only_gap"] is True
    assert item["risk_level"] == "低"
    assert "基础责任和救济机制" in item["risk_point_description"]
    assert "缺少必要保护要素" not in item["risk_point_description"]
    assert not any(row["risk_group"] == "third_party_claim_defense" for row in output["risk_findings"])


def test_tc_12_04_requires_same_loss_general_remedy_before_activation():
    output = _review(_gold_text())
    item = _item(output, "TC-12-04")

    assert item["coverage_status"] == "不适用"
    assert item["evidence_clauses"] == []
    assert "同一侵权损失" in item["reasoning_summary"]
    assert not any(row["risk_group"] == "indemnity_and_liability_cap" for row in output["risk_findings"])


def test_tc_13_04_uses_cross_border_consequence_mapping_not_ip_claim_fallback():
    output = _review(
        LICENSE_BASE
        + "\n许可方为德国企业，被许可方为中国企业，本合同属于跨境技术交易。"
        + "\n双方遵守技术进出口、出口管制及数据跨境规则，但具体审批、登记、备案、许可和责任分工未约定。"
    )
    item = _item(output, "TC-13-04")
    finding = _finding(output, "technology_import_export_control")

    assert item["decision_trace"]["field_provenance"]["legal_consequence_source"] == "rule_family:technology_import_export_control"
    assert all(term in finding["consequence"] for term in ("跨境交付", "许可履行", "付款结算"))
    assert "第三方提出权利主张" not in finding["consequence"]
