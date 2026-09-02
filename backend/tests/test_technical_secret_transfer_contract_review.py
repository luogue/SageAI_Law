from pathlib import Path

import pytest
from docx import Document

from review_v2.engine import review_contract_v2
from review_v2.technology_service_analysis import (
    analyze_improvement_license_facts,
    analyze_payment_schedule,
    build_technology_fact_registry,
    classify_technology_transaction,
)
from review_v2.user_profile import UserRiskProfile


ROOT = Path(__file__).resolve().parents[2]
FIXTURE = ROOT / "test_data" / "26 技术转让（技术秘密）合同.docx"


def _fixture_text() -> str:
    document = Document(str(FIXTURE))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            text = " | ".join(cell.text.strip() for cell in row.cells)
            if text.strip(" |"):
                parts.append(text)
    return "\n".join(parts)


@pytest.fixture(scope="module")
def secret_transfer_text() -> str:
    return _fixture_text()


@pytest.fixture(scope="module")
def secret_transfer_review(secret_transfer_text: str) -> dict:
    return review_contract_v2(
        secret_transfer_text,
        "technology",
        include_low_risk=True,
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/技术秘密受让方"),
    )


def test_classification_prefers_the_core_technical_secret_transfer_transaction(
    secret_transfer_text: str,
):
    classification = classify_technology_transaction(secret_transfer_text)

    assert classification["primary_subtype"] == "TECHNOLOGY_TRANSFER"
    assert classification["tertiary_subtype"] == "TECHNICAL_SECRET_TRANSFER"
    assert classification["primary_contract_type"] == "technology_transfer_technical_secret"
    assert classification["transaction_structure"] == "B2B_TECHNICAL_SECRET_TRANSFER"
    assert classification["scenario"] == "TECHNICAL_SECRET_TRANSFER"
    assert "technical_secret" in classification["sub_transactions"]


def test_payment_consistency_is_shared_by_technical_secret_transfer(
    secret_transfer_text: str,
):
    payment = analyze_payment_schedule(secret_transfer_text)

    assert payment["status"] == "INCONSISTENT"
    assert payment["contract_total"] == 3_000_000
    assert payment["percentage_total"] == 110
    assert payment["amount_total"] == 3_300_000
    assert payment["amount_delta"] == 300_000
    assert payment["currency"] == "CNY"


def test_existing_adverse_third_party_allocation_is_not_treated_as_missing(
    secret_transfer_text: str,
    secret_transfer_review: dict,
):
    claim = build_technology_fact_registry(secret_transfer_text)["third_party_claim"]
    by_id = {row["item_id"]: row for row in secret_transfer_review["items"]}
    claim_elements = {
        row["element_id"]: row for row in by_id["TC-12-02"]["coverage_elements"]
    }

    assert claim["authority_warranty"] is True
    assert claim["claim_trigger"] is True
    assert claim["claim_assistance"] is True
    assert claim["claim_cost_allocation"] is True
    assert claim["claim_cost_bearer"] == "LICENSEE"
    assert claim["licensor_liability_cap"] is True
    assert claim_elements["defense_cost_allocation"]["status"] == "COVERED"
    assert claim_elements["defense_cost_allocation"]["direction"] == "ADVERSE"
    assert claim_elements["damages_allocation"]["status"] == "COVERED"
    assert claim_elements["damages_allocation"]["direction"] == "ADVERSE"
    assert claim_elements["defense_control"]["status"] == "MISSING"
    assert claim_elements["lawyer_selection"]["status"] == "MISSING"


def test_forced_improvement_transfer_separates_presence_from_direction(
    secret_transfer_text: str,
    secret_transfer_review: dict,
):
    improvement = analyze_improvement_license_facts(secret_transfer_text)
    by_id = {row["item_id"]: row for row in secret_transfer_review["items"]}

    assert improvement["forced_free_improvement_transfer"] is True
    assert improvement["unreasonable_improvement_restriction"] is True
    for review_id in ("TC-10-01", "TC-10-02", "TC-10-04"):
        item = by_id[review_id]
        elements = {row["element_id"]: row for row in item["coverage_elements"]}
        assert item["coverage_status"] == "风险条款"
        assert elements["improvement_ownership_written"]["status"] == "COVERED"
        assert elements["forced_free_improvement_transfer"]["status"] == "ADVERSE"
        assert elements["post_termination_improvement_use"]["status"] == "ADVERSE"


def test_secret_loss_and_competition_rules_are_active_and_adverse(
    secret_transfer_review: dict,
):
    by_id = {row["item_id"]: row for row in secret_transfer_review["items"]}
    secret_loss = {
        row["element_id"]: row for row in by_id["TC-12-03"]["coverage_elements"]
    }

    assert by_id["TC-02-01"]["coverage_status"] == "风险条款"
    assert by_id["TC-12-03"]["coverage_status"] == "风险条款"
    assert secret_loss["secret_loss_trigger"]["status"] == "COVERED"
    assert secret_loss["future_transfer_fee_after_disclosure"]["status"] == "ADVERSE"
    assert secret_loss["refund_and_termination_after_disclosure"]["status"] == "ADVERSE"
    assert secret_loss["transferor_caused_disclosure"]["status"] == "ADVERSE"


def test_full_report_uses_secret_transfer_role_and_non_patent_proof_semantics(
    secret_transfer_review: dict,
):
    by_id = {row["item_id"]: row for row in secret_transfer_review["items"]}
    report = secret_transfer_review["result"]

    assert secret_transfer_review["release_quality_gate"]["status"] == "REPORT_READY"
    assert secret_transfer_review["contract_context"]["subject_matter"] == "TECHNICAL_SECRET"
    assert secret_transfer_review["contract_context"]["transaction_structure"] == "B2B_TECHNICAL_SECRET_TRANSFER"
    assert "甲方 / 技术秘密受让方" in report
    assert "技术合同—技术转让合同—技术秘密转让合同" in report
    assert "专利权受让方" not in report
    assert "专利申请权转让" not in report
    assert "专利转让登记" not in report
    assert by_id["TC-06-04"]["coverage_status"] == "不适用"
    assert by_id["GEN-PAY-01"]["risk_level"] == "高"
    assert "110%" in by_id["GEN-PAY-01"]["risk_point_description"]
    assert "3,300,000元" in by_id["GEN-PAY-01"]["risk_point_description"]
    assert "3,000,000元" in by_id["GEN-PAY-01"]["risk_point_description"]


def test_improvement_checklist_rows_aggregate_without_losing_auditability(
    secret_transfer_review: dict,
):
    finding = next(
        row for row in secret_transfer_review["risk_findings"]
        if row["risk_group"] == "improvement_no_unlawful_restriction"
    )

    assert set(finding["source_review_ids"]) == {"TC-10-01", "TC-10-02", "TC-10-04"}
    assert sum(
        row["risk_group"] == "improvement_no_unlawful_restriction"
        for row in secret_transfer_review["risk_findings"]
    ) == 1


def test_ancillary_service_fee_keeps_covered_base_and_only_reports_excess_fee_gaps(
    secret_transfer_review: dict,
):
    item = next(row for row in secret_transfer_review["items"] if row["item_id"] == "TC-09-04")
    elements = {row["element_id"]: row for row in item["coverage_elements"]}

    assert item["coverage_status"] == "部分覆盖"
    assert elements["service_fee_amount"]["status"] == "COVERED"
    assert elements["service_payment_schedule"]["status"] == "COVERED"
    assert elements["extra_service_pricing_standard"]["status"] == "MISSING"
    assert elements["extra_service_billing_method"]["status"] == "MISSING"
    assert elements["extra_service_payment_node"]["status"] == "MISSING"
    assert "仍需补足：技术服务费用已明确安排" not in secret_transfer_review["result"]
    assert "建议补充明确技术服务费用已明确安排" not in secret_transfer_review["result"]


def test_positive_adverse_evidence_overrides_stale_absence_search(
    secret_transfer_review: dict,
):
    item = next(row for row in secret_transfer_review["items"] if row["item_id"] == "TC-02-01")
    elements = {row["element_id"]: row for row in item["coverage_elements"]}
    finding = next(
        row for row in secret_transfer_review["risk_findings"]
        if "TC-02-01" in row["source_review_ids"]
    )

    access = elements["lawful_competing_technology_access"]
    assert access["status"] == "ADVERSE"
    assert access["evidence_ids"]
    assert access["element_id"] not in finding["missing_element_ids"]
    absence_text = " ".join(
        " ".join(row.get("missing_elements") or [])
        for row in (finding.get("evidence_bundle") or {}).get("items") or []
        if row.get("source_type") == "SEARCH_AUDIT"
    )
    assert "购买竞争技术或形成替代方案受到限制" not in absence_text


def test_technical_secret_proof_uses_partial_atomic_facts_not_principle_warranty_pass(
    secret_transfer_review: dict,
):
    by_id = {row["item_id"]: row for row in secret_transfer_review["items"]}
    for review_id in ("TC-03-04", "TC-13-02"):
        item = by_id[review_id]
        elements = {row["element_id"]: row for row in item["coverage_elements"]}

        assert item["coverage_status"] == "部分覆盖"
        assert elements["technical_secret_identity_and_delivery_scope"]["status"] == "COVERED"
        assert elements["technical_source"]["status"] == "PARTIALLY_COVERED"
        assert elements["lawful_possession_basis"]["status"] == "MISSING"
        assert elements["disposition_authority"]["status"] == "PARTIALLY_COVERED"
        assert elements["historical_disclosure_objects"]["status"] == "PARTIALLY_COVERED"
        assert elements["historical_disclosure_scope"]["status"] == "PARTIALLY_COVERED"
        assert elements["third_party_confidentiality_obligations"]["status"] == "MISSING"
        assert elements["existing_license_and_use_arrangements"]["status"] == "PARTIALLY_COVERED"
        assert elements["joint_holder_disclosure"]["status"] == "MISSING"
        assert elements["current_secret_status"]["status"] == "MISSING"
        assert elements["existing_confidentiality_measures"]["status"] == "MISSING"


def test_no_actual_infringement_fact_does_not_create_tc_02_02_finding(
    secret_transfer_review: dict,
):
    by_id = {row["item_id"]: row for row in secret_transfer_review["items"]}

    assert by_id["TC-02-02"]["coverage_status"] == "通过"
    assert all(
        "TC-02-02" not in finding["source_review_ids"]
        for finding in secret_transfer_review["risk_findings"]
    )
    assert "未陈述" in by_id["TC-02-02"]["reasoning_summary"]
    assert "侵害他人技术成果" in by_id["TC-02-02"]["reasoning_summary"]


def test_secret_proof_consequence_never_uses_patent_registration_or_pledge_language(
    secret_transfer_review: dict,
):
    by_id = {row["item_id"]: row for row in secret_transfer_review["items"]}
    forbidden = ("转让登记补正", "登记失败", "质押")

    for review_id in ("TC-03-04", "TC-13-02"):
        consequence = by_id[review_id]["legal_consequence_if_missing"]
        assert not any(term in consequence for term in forbidden)
        assert "合法持有" in consequence
        assert "秘密性" in consequence


def test_competition_and_forced_improvement_consequences_include_validity_risk_without_overreach(
    secret_transfer_review: dict,
):
    by_id = {row["item_id"]: row for row in secret_transfer_review["items"]}

    for review_id in ("TC-02-01", "TC-10-01", "TC-10-02", "TC-10-04"):
        consequence = by_id[review_id]["legal_consequence_if_missing"]
        assert "技术竞争" in consequence
        assert "技术发展" in consequence
        assert "效力" in consequence
        assert "风险" in consequence
        assert "合同无效" not in consequence


def test_acceptance_test_attachment_is_fact_activated_and_theme_deduplicated(
    secret_transfer_review: dict,
):
    by_id = {row["item_id"]: row for row in secret_transfer_review["items"]}
    attachment = by_id["TC-13-03"]
    theme = next(
        row for row in secret_transfer_review["risk_themes"]
        if "TC-13-03" in row["source_rule_ids"]
    )

    assert attachment["coverage_status"] == "部分覆盖"
    assert attachment["risk_level"] == "低"
    assert attachment["action_priority"] == "C_OPTIMIZATION"
    assert set(theme["source_rule_ids"]) == {"TC-09-02", "TC-13-03"}
    # The shared liability-cap validity and excessive-liquidated-damages
    # screens are independently evidenced in this fixture as well.
    assert secret_transfer_review["risk_theme_count"] == 13
    assert len(secret_transfer_review["risk_findings"]) == 15


def test_acceptance_test_attachment_narrative_has_complete_sentences(
    secret_transfer_review: dict,
):
    attachment = next(
        row for row in secret_transfer_review["items"]
        if row["item_id"] == "TC-13-03"
    )
    report = secret_transfer_review["result"]
    narrative = attachment["decision_trace"]["narrative_elements"]

    assert "测试争议的复测机制未明确。" in narrative["judgment"]
    assert "测试争议的复测机制未。" not in report
    assert "尚未。" not in report
    assert "建议补充明确测试争议的复测机制未明确" not in report
    assert "建议在测试方案中固定具体样本和范围" in report
    assert "明确测试争议的复测机制" in report


def test_service_fee_and_confidentiality_use_only_direct_operational_clauses(
    secret_transfer_review: dict,
):
    by_id = {row["item_id"]: row for row in secret_transfer_review["items"]}
    service_fee = by_id["TC-09-04"]
    confidentiality = by_id["TC-11-01"]
    service_finding = next(
        row for row in secret_transfer_review["risk_findings"]
        if row["risk_group"] == "service_payment_and_expense"
    )

    assert {row["clause_id"] for row in service_fee["evidence_clauses"]} == {"C0011", "C0012"}
    assert {
        row["clause_id"]
        for row in service_finding["evidence_bundle"]["items"]
        if row.get("source_type") != "SEARCH_AUDIT"
    } == {"C0011", "C0012"}
    assert [row["clause_id"] for row in confidentiality["evidence_clauses"]] == ["C0009"]
