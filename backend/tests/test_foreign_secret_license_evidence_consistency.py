from pathlib import Path
import re

import pytest
from docx import Document

from review_v2.engine import review_contract_v2
from review_v2.user_profile import UserRiskProfile


ROOT = Path(__file__).resolve().parents[2]
FIXTURE = next((ROOT / "test_data").glob("41*.docx"))


def _fixture_text() -> str:
    document = Document(str(FIXTURE))
    rows = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    rows.extend(
        " | ".join(cell.text.strip() for cell in row.cells)
        for table in document.tables
        for row in table.rows
    )
    return "\n".join(rows)


@pytest.fixture(scope="module")
def result() -> dict:
    return review_contract_v2(
        _fixture_text(),
        "technology",
        use_llm=False,
        include_low_risk=True,
        profile=UserRiskProfile(user_identity="甲方/受让方"),
    )


def _item(result: dict, review_id: str) -> dict:
    return next(row for row in result["items"] if row["item_id"] == review_id)


def _elements(item: dict) -> dict:
    return {row["element_id"]: row for row in item.get("coverage_elements") or []}


def test_secret_license_scope_reuses_subject_territory_term_and_survival_facts(result: dict):
    classification = result["deterministic_analysis"]["technology_classification"]
    assert classification["primary_subtype"] == "TECHNOLOGY_LICENSE"
    assert classification["tertiary_subtype"] == "TECHNICAL_SECRET_LICENSE"

    item = _item(result, "TC-07-01")
    elements = _elements(item)
    assert elements["licensed_secret_subject"]["status"] == "COVERED"
    assert elements["license_territory"]["status"] == "COVERED"
    assert "C0014" in elements["license_territory"]["evidence_ids"]
    assert elements["license_term"]["status"] == "COVERED"
    assert "C0083" in elements["license_term"]["evidence_ids"]
    assert elements["post_termination_use_boundary"]["status"] == "PARTIALLY_COVERED"
    assert "C0063" in elements["post_termination_use_boundary"]["evidence_ids"]
    assert item["decision_trace"]["final_missing_element_ids"] == ["post_termination_use_boundary"]
    assert "许可期限缺失" not in item["risk_point_description"]
    assert "终止后继续使用权" in item["risk_point_description"]


def test_money_refund_is_not_confidential_material_disposal_evidence(result: dict):
    item = _item(result, "TC-11-03")
    assert item["coverage_status"] == "缺失"
    elements = _elements(item)
    assert {
        "confidential_material_return",
        "confidential_copy_deletion",
        "confidential_backup_deletion",
        "confidential_material_destruction",
        "confidential_disposal_confirmation",
    } == set(elements)
    assert all(row["status"] == "MISSING" for row in elements.values())
    evidence_ids = {
        row.get("clause_id") or row.get("evidence_id")
        for row in item.get("evidence_clauses") or []
    }
    assert "C0056" not in evidence_ids
    assert "C0057" not in evidence_ids
    assert "退款或返还已付款项" in item["decision_trace"]["material_disposal_direct_evidence"]["excluded_semantic_matches"]


def test_improvement_provision_is_shared_without_inventing_a_restriction(result: dict):
    ownership = _item(result, "TC-10-01")
    sharing = _item(result, "TC-10-02")
    notice = _item(result, "TC-10-03")
    restriction = _item(result, "TC-10-04")

    assert ownership["coverage_status"] == "缺失"
    assert sharing["coverage_status"] == "部分覆盖"
    assert _elements(sharing)["improvement_material_sharing"]["status"] == "COVERED"
    assert "C0050" in _elements(sharing)["improvement_material_sharing"]["evidence_ids"]
    assert notice["coverage_status"] == "部分覆盖"
    assert _elements(notice)["improvement_notice_trigger"]["status"] == "COVERED"
    assert "C0050" in _elements(notice)["improvement_notice_trigger"]["evidence_ids"]
    assert notice["decision_trace"]["final_missing_element_ids"] == [
        "improvement_notice_deadline",
        "improvement_notice_method",
        "improvement_material_scope",
    ]
    assert restriction["coverage_status"] == "通过"
    assert restriction["risk_status"] == "ACCEPTABLE"
    assert restriction["decision_trace"]["evidence_judgment"]["affirmative_risk_evidence_ids"] == []


def test_public_information_exception_is_covered_and_removed_from_residual_gaps(result: dict):
    item = _item(result, "TC-12-03")
    elements = _elements(item)
    assert item["coverage_status"] == "部分覆盖"
    assert elements["technical_secret_scope"]["status"] == "COVERED"
    assert elements["technical_secret_consideration"]["status"] == "COVERED"
    assert elements["secret_loss_trigger"]["status"] == "COVERED"
    assert "C0061" in elements["secret_loss_trigger"]["evidence_ids"]
    assert "secret_loss_trigger" not in item["decision_trace"]["final_missing_element_ids"]
    assert "依法公开" in item["risk_point_description"]
    assert "公开信息处理条款" in item["suggested_clause_example"]
    assert "触发情形" not in item["suggested_clause_example"]


def test_secret_license_uses_precise_technology_import_partial_route(result: dict):
    item = _item(result, "TC-13-04")
    elements = _elements(item)
    assert item["coverage_status"] == "部分覆盖"
    assert item["needs_manual_review"] is True
    assert item["decision_trace"]["regulatory_reasoning_chain"]["transaction_type"] == "TECHNICAL_SECRET_LICENSE"
    assert elements["cross_border_transaction_direction"]["status"] == "COVERED"
    assert elements["technology_import_characterization"]["status"] == "COVERED"
    assert elements["cross_border_transaction_type"]["status"] == "COVERED"
    assert elements["technology_import_formality_cooperation"]["status"] == "COVERED"
    assert elements["technology_import_formality_deadline"]["status"] == "COVERED"
    assert elements["foreign_export_license_document"]["status"] == "COVERED"
    assert elements["technology_category_and_control_status"]["status"] == "UNCERTAIN"
    assert elements["applicable_import_management_path"]["status"] == "UNCERTAIN"
    assert elements["technology_import_license_responsible_party"]["status"] == "MISSING"
    assert elements["technology_import_registration_responsible_party"]["status"] == "MISSING"
    assert elements["technology_import_change_formality"]["status"] == "MISSING"
    assert elements["royalty_payment_method"]["status"] == "NOT_APPLICABLE"
    assert "办理协作义务" not in item["suggested_clause_example"]
    assert "办理期限或节点" not in item["suggested_clause_example"]
    assert "目录" in item["quality_gate_failure_reason"]


def test_cross_rule_positive_facts_never_reappear_as_missing(result: dict):
    pairs = (
        ("TC-10-02", "improvement_material_sharing", "TC-10-03", "improvement_notice_trigger"),
        ("TC-11-02", "confidentiality_exceptions", "TC-12-03", "secret_loss_trigger"),
    )
    for source_id, source_element, consumer_id, consumer_element in pairs:
        source = _elements(_item(result, source_id)).get(source_element)
        consumer = _elements(_item(result, consumer_id)).get(consumer_element)
        if source and source["status"] == "COVERED":
            assert consumer is not None
            assert consumer["status"] == "COVERED"

    infringement = _item(result, "TC-12-01")
    assert infringement["applicability"]["status"] == "APPLICABLE"
    assert result["release_quality_gate"]["status"] == "REPORT_READY"


def test_report_statistics_are_clean_decimal_integers(result: dict):
    assert result["rule_level_finding_count"] == 15
    assert "规则级审查命中数：15" in result["result"]
    assert "1a5" not in result["result"]
    statistic_lines = [
        line for line in result["result"].splitlines()
        if any(label in line for label in (
            "风险主题数（聚合后）", "规则级审查命中数", "重大风险数量", "中等风险数量",
            "低风险数量", "需要真正修改的风险数", "优先处理风险", "建议处理风险",
            "可选优化主题数（聚合后）", "人工复核项",
        ))
    ]
    assert statistic_lines
    assert all(re.search(r"：\d+$", line) for line in statistic_lines)


def test_scope_recommendation_and_evidence_use_precise_term_survival_clauses(result: dict):
    item = _item(result, "TC-07-01")
    evidence_ids = [row.get("clause_id") for row in item.get("evidence_clauses") or []]
    assert evidence_ids[:3] == ["C0014", "C0083", "C0063"]
    assert all(row.get("location") != "全文" for row in item.get("evidence_clauses") or [])
    assert "合同期限届满或提前终止后继续使用权的存续期间" in item["suggested_clause_example"]
    assert "是否继续付费" in item["suggested_clause_example"]
    assert "与保密义务的衔接" in item["suggested_clause_example"]
    assert "建议补充明确二者" not in result["result"]


def test_confidentiality_breach_core_gaps_exclude_optional_damage_terms(result: dict):
    item = _item(result, "TC-11-04")
    elements = _elements(item)
    assert item["decision_trace"]["final_missing_element_ids"] == [
        "stop_disclosure", "stop_use", "mitigation_obligation", "incident_notification",
    ]
    assert elements["breach_damage_liability"]["risk_relevant"] is False
    assert "一般违约赔偿责任" not in item["risk_point_description"]
    assert "律师费" not in item["suggested_clause_example"]
    assert "诉讼费" not in item["suggested_clause_example"]


def test_improvement_restriction_pass_uses_negative_clause_check(result: dict):
    item = _item(result, "TC-10-04")
    assert item["coverage_status"] == "通过"
    assert item.get("evidence_clauses") == []
    assert item["decision_trace"]["review_mode"] == "ABSENCE_REQUIRED"
    assert item["decision_trace"]["review_polarity"] == "NEGATIVE_CLAUSE_CHECK"
    assert item["decision_trace"]["negative_clause_check"]["adverse_clause_exists"] is False
    assert "未发现限制一方自行改进" in item["risk_point_description"]
    assert "C0050" not in item["decision_trace"].get("final_scope_evidence_ids", [])


def test_import_review_wording_matches_secret_license_and_fixed_price(result: dict):
    item = _item(result, "TC-13-04")
    chain = item["decision_trace"]["regulatory_reasoning_chain"]
    assert chain["transaction_type"] == "TECHNICAL_SECRET_LICENSE"
    assert chain["payment_scheme"] == "FIXED_PRICE"
    assert item["decision_trace"]["manual_review_materiality"] == "MEDIUM"
    assert "专利许可履行" not in result["result"]
    assert "提成费付款结算受阻" not in result["result"]
    assert "许可费/合同价款结算受阻" in result["result"]
    assert "人工核验优先级：MEDIUM" in result["result"]


def test_licensee_over_scope_wording_stays_optional_and_party_aligned(result: dict):
    item = _item(result, "TC-07-03")
    assert item["risk_level"] == "低"
    assert item["action_priority"] == "C_OPTIMIZATION"
    assert "是否构成超范围实施、停止使用、解除合同或赔偿责任发生争议" in item["legal_consequence_if_missing"]
    assert "被许可人扩大使用" not in item["risk_point_description"]
