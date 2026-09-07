"""Exercise the authenticated file-report pipeline, not isolated evaluators."""
import io
import json
import threading
from copy import deepcopy
from dataclasses import fields
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from auth.routes import get_current_user
from main import app
from review_v2 import engine


ROOT = Path(__file__).resolve().parents[2]
CURRENT = ROOT / "test_data" / "38 技术许可（技术秘密）合同.docx"


def review_upload(payload, filename="secret-license.docx"):
    previous = app.dependency_overrides.get(get_current_user)
    app.dependency_overrides[get_current_user] = lambda: {"id": 1, "phone": "test"}
    try:
        with TestClient(app) as client:
            response = client.post("/review-v2/file", files={
                "file": (filename, payload,
                         "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                data={"contract_type": "technology", "include_low_risk": "true",
                      "use_llm": "false", "profile_json": '{"user_identity":"甲方/被许可方"}'})
    finally:
        if previous is None:
            app.dependency_overrides.pop(get_current_user, None)
        else:
            app.dependency_overrides[get_current_user] = previous
    assert response.status_code == 200, response.text
    return response.json()


@pytest.fixture(scope="module")
def current_report():
    called = set()
    def observe(frame, event, arg):
        if event == "call" and "Sage" in frame.f_code.co_filename:
            called.add(frame.f_code.co_name)
    previous = threading.getprofile()
    threading.setprofile(observe)
    try:
        report = review_upload(CURRENT.read_bytes(), CURRENT.name)
    finally:
        threading.setprofile(previous)
    report["test_observed_functions"] = called
    return report


def elements(report, rule):
    item = next(i for i in report["items"] if i["item_id"] == rule)
    return item, {e["element_id"]: e for e in item["coverage_elements"]}


def test_scope_liability_survives_file_pipeline(current_report):
    item, rows = elements(current_report, "TC-07-03")
    assert item["coverage_status"] == "通过"
    assert rows["over_scope_breach_remedy"]["status"] == "COVERED"
    assert not item["decision_trace"]["final_missing_element_ids"]
    assert not item["legal_consequence_if_missing"]
    assert not any("TC-07-03" in f["source_review_ids"] for f in current_report["risk_findings"])


def test_responsible_party_synonyms_cannot_disagree(current_report):
    item, rows = elements(current_report, "TC-12-02")
    assert rows["liability_responsible_party"]["status"] == "COVERED"
    assert rows["liability_responsible_party"].get("supporting_evidence_ids")
    assert any(row["status"] == "COVERED" and "赔偿" in row["proposition"] for row in rows.values())


def test_each_final_notice_gap_has_its_own_search(current_report):
    item, rows = elements(current_report, "TC-10-03")
    assert rows["improvement_notice_trigger"]["status"] == "COVERED"
    missing = {key for key, row in rows.items() if row["status"] == "MISSING"}
    audits = item["decision_trace"].get("element_absence_searches", [])
    absent = {row["element_id"] for row in audits if row["result"] == "NOT_FOUND"}
    assert missing <= absent
    assert all(row["searched_clause_ids"] and row["query_patterns"] for row in audits)
    finding = next(f for f in current_report["risk_findings"] if "TC-10-03" in f["source_review_ids"])
    package = finding["evidence_bundle"]
    bound = {eid for e in package["items"] if e.get("source_type") == "SEARCH_AUDIT"
             for eid in e.get("searched_absent_element_ids", [])}
    assert missing <= bound


def test_unknown_proof_conditions_are_verification_not_drafting(current_report):
    pending = {rule: row for row in current_report["manual_review_items"] for rule in row["source_rules"]}
    assert current_report["final_overview_summary"]["manual_review"] == len(current_report["manual_review_items"])
    assert pending["TC-03-04"] is pending["TC-13-02"]
    assert set(pending["TC-03-04"]["source_rules"]) == {"TC-03-04", "TC-13-02"}
    for rule in ("TC-03-04", "TC-13-02"):
        assert rule in pending
        assert "核验" in pending[rule]["review_instruction"]
        item, rows = elements(current_report, rule)
        for key in ("joint_holder_disclosure", "staff_assignment", "open_source_compliance",
                    "software_copyright_proof", "third_party_code_authorization"):
            row = rows[key]
            assert row["status"] == "CONDITION_UNRESOLVED", (key, row)
            assert row["prerequisite_status"] == "UNKNOWN"
            assert row["risk_relevant"] is False
        finding = next(f for f in current_report["risk_findings"] if rule in f["source_review_ids"])
        excluded = {key for key, row in rows.items() if row.get("prerequisite_status") == "UNKNOWN"}
        assert not excluded.intersection(finding.get("recommendation_element_ids", []))
        assert not excluded.intersection(e["element_id"] for e in finding["final_missing_elements"])


def test_existing_license_atomic_facts_survive_absence_and_all_report_sections(current_report):
    positive = {"existing_license_disclosure", "existing_license_counterparty",
                "existing_license_term", "existing_license_territory"}
    for rule in ("TC-03-04", "TC-13-02"):
        item, rows = elements(current_report, rule)
        assert "existing_license_and_use_arrangements" not in rows
        for key in positive:
            assert rows[key]["status"] == "COVERED", rows[key]
            assert rows[key]["supporting_evidence_ids"]
        for key in ("test_use_arrangement", "other_third_party_use_arrangement"):
            assert rows[key]["status"] == "MISSING", rows[key]
        assert positive.isdisjoint(item["decision_trace"]["final_missing_element_ids"])
        finding = next(f for f in current_report["risk_findings"] if rule in f["source_review_ids"])
        assert positive.isdisjoint(e["element_id"] for e in finding["final_missing_elements"])
        assert positive.isdisjoint(finding.get("recommendation_element_ids", []))
        absent = {key for e in finding["evidence_bundle"]["items"] if e.get("source_type") == "SEARCH_AUDIT"
                  for key in e.get("searched_absent_element_ids", [])}
        assert positive.isdisjoint(absent)
    assert "既有许可、测试使用及其他第三方使用安排" not in current_report["result"]
    manual_section = current_report["result"].split("人工复核事项", 1)[1].split("完整风险主题明细", 1)[0]
    assert manual_section.count("核验交付技术是否包含开源组件") == 1


def test_production_report_gate_and_previous_correct_states(current_report):
    gate = current_report["release_quality_gate"]
    assert gate["status"] in {"PASSED", "REPORT_READY"}, json.dumps(gate.get("failures"), ensure_ascii=False)
    assert not gate.get("failures"), gate
    for rule in ("TC-10-01", "TC-10-04", "TC-12-01"):
        assert elements(current_report, rule)[0]["coverage_status"] == "通过"
    trace = current_report["production_path_trace"]
    assert trace["runtime_source_drift"]["status"] == "CURRENT"
    assert trace["document_hash"]
    assert {
        "review_v2_file", "extract_text_with_diagnostics", "review_contract_v2",
        "extract_claim_response_facts", "_set_multi_element_state",
        "_apply_technology_transaction_evidence_consistency", "_technical_secret_transfer_proof_results",
        "_resolve_final_element_searches", "_build_risk_findings", "_finalize_element_recommendation_sources",
    } <= current_report["test_observed_functions"]


def test_confidentiality_remedies_remain_atomic_and_party_bound_in_report(current_report):
    item, rows = elements(current_report, "TC-11-04")
    assert item["coverage_status"] == "部分覆盖"
    assert rows["breach_damage_liability"]["status"] == "COVERED"
    assert rows["specific_confidentiality_penalty"]["status"] == "NOT_APPLICABLE"
    assert rows["excess_loss_compensation"]["status"] == "NOT_APPLICABLE"
    mitigation = rows["mitigation_obligation"]
    assert mitigation["status"] == "PARTIALLY_COVERED"
    by_party = {r["obligated_party"]: r for r in mitigation["party_coverage"]}
    assert by_party["买方"]["status"] == "COVERED"
    assert by_party["卖方"]["status"] == "MISSING"
    assert all(f["obligated_party"] == "买方" for f in mitigation["supporting_evidence"])
    finding = next(f for f in current_report["risk_findings"] if "TC-11-04" in f["source_review_ids"])
    statements = [item["risk_point_description"], item["reasoning_summary"], item["legal_consequence_if_missing"],
                  finding["issue"], finding["recommendation"],
                  *[e["proposition"] for e in rows.values() if e["status"] in {"COVERED", "PARTIALLY_COVERED"}],
                  *[e.get("support_summary", "") for e in item["evidence_clauses"]]]
    assert all("违约金" not in text for text in statements), statements
    for evidence in finding["evidence_bundle"]["items"]:
        if evidence.get("source_type") == "SEARCH_AUDIT":
            continue
        if "mitigation_obligation" in evidence.get("supporting_element_ids", []):
            assert "防止" in evidence["quoted_text"]
            assert "卖方泄密后采取措施" not in evidence.get("proposition", "")
    assert "卖方泄密后采取措施" in finding["recommendation"]


def test_partial_proof_requirements_survive_summary_recommendation_and_evidence(current_report):
    finding = next(f for f in current_report["risk_findings"] if "TC-03-04" in f["source_review_ids"])
    aggregate = {e["element_id"] for e in finding["final_missing_elements"]}
    package = {e["element_id"] for e in finding["evidence_bundle"]["unresolved_requirements"]}
    for rule in ("TC-03-04", "TC-13-02"):
        item, rows = elements(current_report, rule)
        expected = {e["element_id"] for e in engine.unresolved_requirements(item["coverage_elements"])}
        assert expected == aggregate == package
        assert expected == set(item["decision_trace"]["final_missing_element_ids"])
        recs = [r for r in current_report["recommendations"] if finding["finding_id"] in r["source_finding_ids"]]
        recommendation_ids = {key for rec in recs for fragment in rec["fragments"] for key in fragment["source_element_ids"]}
        assert expected <= recommendation_ids
        for row in rows.values():
            if row["status"] == "PARTIALLY_COVERED":
                gap = row["residual_gap"]
                assert gap in finding["issue"]
                assert gap in finding["recommendation"]
                assert current_report["result"].count(gap) >= 3


@pytest.mark.parametrize("mutation,reason", [
    ("penalty", "PENALTY_NARRATIVE_WITHOUT_PENALTY_EVIDENCE"),
    ("party", "OBLIGATED_PARTY_EVIDENCE_MISMATCH"),
    ("aggregate", "AGGREGATE_UNRESOLVED_ELEMENTS_MISMATCH"),
    ("evidence", "EVIDENCE_UNRESOLVED_ELEMENTS_MISMATCH"),
    ("narrative", "UNRESOLVED_NARRATIVE_OMISSION"),
])
def test_release_gate_rejects_party_penalty_or_unresolved_projection_drift(current_report, mutation, reason):
    def load(cls, key):
        return [cls(**{f.name: row[f.name] for f in fields(cls) if f.name in row}) for row in deepcopy(current_report[key])]
    items = load(engine.ReviewItemResult, "items")
    findings = load(engine.RiskFinding, "risk_findings")
    liability = next(i for i in items if i.item_id == "TC-11-04")
    proof = next(f for f in findings if "TC-03-04" in f.source_review_ids)
    if mutation == "penalty":
        liability.risk_point_description = "合同已覆盖泄密违约金。"
    elif mutation == "party":
        row = next(e for e in liability.coverage_elements if e["element_id"] == "mitigation_obligation")
        row["party_coverage"][1].update(status="COVERED", supporting_evidence=row["supporting_evidence"])
    elif mutation == "aggregate":
        proof.final_missing_elements = [e for e in proof.final_missing_elements if e["status"] != "PARTIALLY_COVERED"]
    elif mutation == "evidence":
        proof.evidence_bundle["unresolved_requirements"] = []
    else:
        proof.issue = "仅需补充完全缺失的文件。"
    failures = engine._exact_element_report_invariants(items, findings)
    assert reason in {f["reason"] for f in failures}
    assert engine._release_quality_gate(items, findings, current_report["result"])["status"] == "BLOCKED"


def amended_upload(clauses):
    document = Document(CURRENT)
    document.add_paragraph("附件：合同技术来源及程序说明")
    for clause in clauses:
        document.add_paragraph(clause)
    stream = io.BytesIO()
    document.save(stream)
    return review_upload(stream.getvalue())


@pytest.mark.parametrize("statement,status", [
    ("本合同技术不含开源组件，不涉及第三方代码；不存在共同持有人，不涉及职务成果；未登记软件著作权。", "NOT_APPLICABLE"),
    ("本合同技术使用开源组件和第三方代码；共同持有人为南桥技术有限公司；员工研发形成的技术属于职务成果；已办理软件著作权。", "MISSING"),
    ("是否包含开源组件、第三方代码及职务成果，以及是否存在共同持有人、软件著作权登记，需另行核验。", "CONDITION_UNRESOLVED"),
])
def test_conditional_states_through_uploaded_document(statement, status):
    report = amended_upload([statement])
    for rule in ("TC-03-04", "TC-13-02"):
        item, rows = elements(report, rule)
        for key in ("open_source_compliance", "joint_holder_disclosure", "staff_assignment",
                    "third_party_code_authorization", "software_copyright_proof"):
            assert rows[key]["status"] == status, (key, rows[key])
            if status != "MISSING":
                finding = next(f for f in report["risk_findings"] if rule in f["source_review_ids"])
                assert key not in {e["element_id"] for e in finding["final_missing_elements"]}
    assert report["release_quality_gate"]["status"] == "REPORT_READY", report["release_quality_gate"]["failures"]


def test_notice_positive_parameters_defeat_old_absence_through_upload():
    report = amended_upload([
        "改进成果形成后应于七个工作日内通知对方，通知应包含改进内容和基本说明，并提交研发资料。",
        "保密事件发生后应在三日内通知。",
    ])
    _, rows = elements(report, "TC-10-03")
    for key in ("improvement_notice_deadline", "improvement_notice_content", "improvement_material_scope"):
        assert rows[key]["status"] == "COVERED", rows[key]


def test_removing_one_absence_binding_cannot_pass_release_invariant(current_report):
    def load(cls, name):
        names = {f.name for f in fields(cls)}
        return [cls(**{k: v for k, v in row.items() if k in names})
                for row in deepcopy(current_report[name])]
    items = load(engine.ReviewItemResult, "items")
    findings = load(engine.RiskFinding, "risk_findings")
    notice = next(f for f in findings if "TC-10-03" in f.source_review_ids)
    notice.evidence_bundle["items"] = [row for row in notice.evidence_bundle["items"]
        if "improvement_notice_deadline" not in row.get("searched_absent_element_ids", [])]
    failures = engine._exact_element_report_invariants(items, findings)
    assert any(row.get("element_id") == "improvement_notice_deadline"
               and row["reason"] == "FINAL_MISSING_WITHOUT_ELEMENT_ABSENCE_SEARCH" for row in failures)
