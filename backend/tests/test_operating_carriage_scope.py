import asyncio
import io
from types import SimpleNamespace
from dataclasses import replace
import pytest
from review_v2.engine import review_contract_v2
from review_v2.clause_parser import parse_contract_clauses
from review_v2.domain_selection import resolve_domain_selection, domain_selection_consistency_gate
from review_v2.lease_semantics import consistency_failures

OP='甲方将营业部交由乙方承包经营。甲方授权乙方经营并使用标识。经营费用由乙方承担。乙方向甲方上交管理费。具体运输业务另行签署合同。'

def rows(text):
    return {r['item_id']:r for r in review_contract_v2(text,'transportation',use_llm=False)['items']}

def test_body_operation_over_transport_title():
    r=rows('运输合同\n'+OP)
    assert 'OP-01-01' in r and not any(k.startswith('T-') for k in r)

def test_cooperation_title_carriage_body():
    text='合作协议\n乙方为承运人，货物为十吨钢材，起运地为东城，目的地为西城。运费为800元。'
    assert resolve_domain_selection('service',text,parse_contract_clauses(text))[2]=='transportation'

def test_future_shipment_framework():
    r=rows('运输合作框架协议。每票另行签署运输单。')
    assert all(r[k]['coverage_status']=='不适用' for k in ['T-01-02','T-02-02','T-08-01'])

def test_insurance_synonym_and_safe_operation():
    r=rows('乙方必须按规定办理各种保险，并确保安全规范操作。')
    assert r['T-12-03']['coverage_status']=='部分覆盖'
    assert r['T-04-01']['coverage_status']=='部分覆盖'

def test_one_of_five_not_pass_and_wrong_quote_blocked():
    r=rows('承运人应将货物交付收货人。')['T-08-01']
    assert r['coverage_status']=='部分覆盖'
    r['evidence_clauses'][0]['quote']='甲方全力支持乙方业务。'
    assert consistency_failures([SimpleNamespace(**r)])

def test_freight_settlement_not_amount():
    r=rows('甲方依据运输合同和派车单与乙方结算。')['T-09-01']
    assert r['coverage_status']=='部分覆盖'
    assert any(e['element_id']=='transport_freight_amount' and e['status']=='MISSING' for e in r['coverage_elements'])

def test_notice_direct_binding():
    r=rows('甲方必须全力支持乙方业务。双方认为需要终止，应提前一个月通知对方。')['T-11-02']
    assert r['coverage_status']=='部分覆盖'
    assert all('支持' not in e.get('quote','') for e in r['evidence_clauses'])

def test_dangling_obligation_and_route_gate():
    text='运输合同\n'+OP+'\n乙方必须因自身原因导致'
    r=rows(text)
    assert r['COMMON-22-01']['coverage_status']=='风险条款'
    selection,_,_=resolve_domain_selection('transportation',text,parse_contract_clauses(text))
    assert domain_selection_consistency_gate(selection,[k for k in r if k!='COMMON-22-01'])['status']=='BLOCKED'
    wrong=replace(selection,selected_domain='TRANSPORTATION')
    assert domain_selection_consistency_gate(wrong,list(r))['status']=='BLOCKED'

def test_temporal_restriction_not_negated_by_prior_distinct_ban():
    r=rows(OP+'\n乙方不得私自接单，本协议提前终止后乙方在两年内不得从事同类竞争业务。')
    assert r['OP-08-01']['coverage_status']=='风险条款'

def test_file_production_entry():
    from docx import Document
    from routers.review_v2 import review_v2_file
    document=Document()
    document.add_paragraph('运输合同')
    document.add_paragraph(OP)
    document.add_paragraph('乙方必须因自身原因导致')
    stream=io.BytesIO(); document.save(stream); stream.seek(0)
    result=asyncio.run(review_v2_file(file=SimpleNamespace(filename='scope.docx',file=stream),contract_type='transportation',include_low_risk=True,use_llm=False,profile_json='{}',current_user={'id':'test'}))
    assert result['contract_type']=='operating_cooperation'
    assert result['result'] and result['release_quality_gate']['status']=='REPORT_READY'
