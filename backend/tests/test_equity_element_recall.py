import json
from pathlib import Path
import pytest
from docx import Document
from review_v2.engine import review_contract_v2
from review_v2.clause_parser import parse_contract_clauses
from review_v2.priority_contract_review import judge
from review_v2.equity_investment_checklist import PRICE, CONTRIBUTION, WARRANTY, RESTRICTIONS, CLOSING
from review_v2.priority_checklist_helpers import PROFILES
from review_v2.lease_semantics import corpus_facts


@pytest.mark.parametrize('amount', ['480万元','三百六十万元','壹佰万元','850000元'])
@pytest.mark.parametrize('label', ['转让金','股权转让款'])
def test_price_synonyms_amounts_and_date_before_payment(amount,label):
    cs=parse_contract_clauses('股权转让协议\n第一条 '+label+amount+'。乙方于2028年4月12日前向甲方支付上述股权转让款。')
    assert judge(PRICE,cs).coverage_status=='通过'


def test_subscribed_amount_is_independent_of_paid_amount_and_deadline():
    cs=parse_contract_clauses('股权转让协议\n第一条 原认缴出资250万元。')
    facts=corpus_facts(cs,PROFILES[CONTRIBUTION.node_id])
    assert facts['equity_contribution_liability_1']
    assert not facts['equity_contribution_liability_2']
    assert not facts['equity_contribution_liability_3']
    assert judge(CONTRIBUTION,cs).coverage_status=='部分覆盖'


def test_missing_all_contribution_facts_still_missing():
    assert judge(CONTRIBUTION,parse_contract_clauses('股权转让协议\n双方协商交易。')).coverage_status=='缺失'


def test_warranty_and_otherwise_remedy_are_not_financial_disclosure():
    cs=parse_contract_clauses('股权转让协议\n第二条 保证\n甲方保证股权未设定任何质押。否则一切损失由转让方赔偿。')
    facts=corpus_facts(cs,PROFILES[WARRANTY.node_id])
    assert facts['equity_warranties_remedies_1']
    assert not facts['equity_warranties_remedies_2']
    assert facts['equity_warranties_remedies_3']
    assert judge(WARRANTY,cs).coverage_status=='部分覆盖'


def test_other_shareholders_signature_is_not_preemption_waiver():
    cs=parse_contract_clauses('股权转让协议\n第八条\n其他股东签署：\n赵磊')
    facts=corpus_facts(cs,PROFILES[RESTRICTIONS.node_id])
    assert facts['equity_transfer_restrictions_4']
    assert not facts['equity_transfer_restrictions_2']
    assert not facts['equity_transfer_restrictions_3']
    assert judge(RESTRICTIONS,cs).coverage_status=='部分覆盖'


def test_shareholder_effective_date_does_not_prove_registration():
    cs=parse_contract_clauses('股权转让协议\n第三条 从2028年5月1日起，乙方即成为乙公司股东。')
    assert judge(CLOSING,cs).coverage_status=='部分覆盖'
    facts=corpus_facts(cs,PROFILES[CLOSING.node_id])
    assert facts['equity_closing_registration_1']
    assert not facts['equity_closing_registration_3']


def test_actual_equity_transfer_full_pipeline_and_shared_clause_binding():
    source=Path(__file__).resolve().parents[2]/'test_data/79 股权转让.docx'
    text='\n'.join(p.text for p in Document(source).paragraphs)
    result=review_contract_v2(text,'equity_transfer',use_llm=False)
    items={i['item_id']:i for i in result['items']}
    for key in ('EQT-01-01','EQT-04-01'):
        assert items[key]['coverage_status']=='通过'
    for key in ('EQT-02-01','EQT-03-01','EQT-05-01','EQT-11-01'):
        assert items[key]['coverage_status']=='部分覆盖'
    price=json.dumps(items['EQT-04-01']['evidence_clauses'],ensure_ascii=False)
    assert '360万元' in price and '2026年9月30日' in price
    for i in items.values():
        covered={e['element_id'] for e in i['coverage_elements'] if e['status']=='COVERED'}
        assert not any(a['element_id'] in covered and a['result']=='NOT_FOUND' for a in i['decision_trace'].get('element_absence_searches',[]))
    warranty={e['element_id']:e for e in items['EQT-11-01']['coverage_elements']}
    assert all('否则' not in e['quote'] for e in warranty['equity_warranties_remedies_1']['supporting_evidence'])
    assert all('否则' in e['quote'] for e in warranty['equity_warranties_remedies_3']['supporting_evidence'])
    assert result['release_quality_gate']['status']=='REPORT_READY',result['release_quality_gate']['failures']


def test_closing_internal_date_does_not_assert_statutory_registration_completed():
    source=Path(__file__).resolve().parents[2]/'test_data/79 股权转让.docx'
    text='\n'.join(p.text for p in Document(source).paragraphs)
    result=review_contract_v2(text,'equity_transfer',use_llm=False)
    item=next(i for i in result['items'] if i['item_id']=='EQT-05-01')
    assert item['coverage_status']=='部分覆盖'
    evidence=json.dumps(item['evidence_clauses'],ensure_ascii=False)
    assert '2026年10月1日起' in evidence and '乙方即成为' in evidence
    report=result['result']
    assert '内部' in report and '权益承接' in report
    assert '不能替代股东名册更新、出资证明更新及相关变更登记手续' in report
    assert '仍需明确办理责任人、办理期限及未及时办理的责任' in report
    for unsupported in ('已完成股东名册变更','已完成工商变更'):
        assert unsupported not in report
    elements={e['element_id']:e for e in item['coverage_elements']}
    assert elements['equity_closing_registration_1']['status']=='COVERED'
    assert all(elements['equity_closing_registration_'+str(n)]['status']=='MISSING' for n in (2,3,4,5))
    assert result['release_quality_gate']['status']=='REPORT_READY'
