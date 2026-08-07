"""房屋租赁合同审查第二版测试。"""

import json
from pathlib import Path

import pytest
from docx import Document

from review_v2.clause_parser import ParsedClause, SECTION_FORMAL, SECTION_SPECIAL_NOTICE, SECTION_USAGE, parse_contract_clauses
from review_v2.engine import review_contract_v2
from review_v2.house_lease_checklist import HOUSE_LEASE_REVIEW_CHECKLIST
from review_v2.house_lease_checklist import get_house_lease_checklist
from review_v2.llm_semantic_reviewer import _judgment_from_model_result
from review_v2.llm_semantic_reviewer import get_llm_review_status
from review_v2.user_profile import UserRiskProfile
from routers.review_v2 import ReviewV2PdfInput, _build_docx, _build_pdf


@pytest.fixture(autouse=True)
def disable_llm_by_default(monkeypatch):
    monkeypatch.setenv("REVIEW_V2_LLM_ENABLED", "false")


def test_house_lease_checklist_has_50_core_points():
    checklist = get_house_lease_checklist()
    high_risk_ids = {checkpoint.item_id for checkpoint in checklist if checkpoint.risk_level == "高"}

    assert len(checklist) == 50
    assert checklist[0].item_id == "HL-001"
    assert checklist[-1].item_id == "HL-050"
    assert checklist[37].node_id == "overbroad_damage_clause"
    assert high_risk_ids == {
        "HL-010",
        "HL-039",
        "HL-043",
    }
    natural_wear = next(checkpoint for checkpoint in checklist if checkpoint.item_id == "HL-034")
    assert natural_wear.risk_level == "低"


def test_llm_status_reports_env_and_model_key(monkeypatch):
    monkeypatch.setenv("REVIEW_V2_LLM_ENABLED", "true")
    monkeypatch.delenv("DASHSCOPE_API_KEY", raising=False)
    monkeypatch.delenv("ARK_API_KEY", raising=False)
    monkeypatch.delenv("DEEPSEEK_API_KEY", raising=False)

    status = get_llm_review_status()
    assert status["enabled_by_env"] is True
    assert status["has_model_key"] is False
    assert status["available"] is False

    monkeypatch.setenv("DASHSCOPE_API_KEY", "test-key")
    status = get_llm_review_status()
    assert status["has_model_key"] is True
    assert status["available"] is True
    assert "qwen" in status["configured_models"]

    monkeypatch.setenv("REVIEW_V2_LLM_ENABLED", "false")
    status = get_llm_review_status()
    assert status["enabled_by_env"] is False
    assert status["available"] is False


def test_review_v2_returns_fixed_review_format_without_verdict_language():
    output = review_contract_v2(
        "出租人将房屋出租给承租人，约定租金每月3000元，押金6000元。",
        "house_lease",
    )

    assert output["success"] is True
    assert output["contract_type"] == "house_lease"
    assert output["all_items_count"] == 50
    assert len(output["items"]) == 50
    first_item = output["items"][0]
    assert "risk_point_description" in first_item
    assert "suggested_clause_example" in first_item
    assert "legal_basis_index" in first_item
    assert "coverage_status" in first_item
    assert "issue_type" in first_item
    assert "legal_consequence_if_missing" in first_item
    assert "evidence_clauses" in first_item
    assert "evidence_location" in first_item
    assert "reasoning_summary" in first_item
    assert "judge_source" in first_item
    assert "llm_confidence" in first_item
    assert "fallback_reason" in first_item
    assert "judge_source_summary" in output
    assert output["judge_source_summary"]["rule_fallback"] == 50
    assert "issue_type_summary" in output
    assert sum(output["issue_type_summary"].values()) == output["all_items_count"]
    assert sum(output["summary"].values()) == output["all_items_count"]
    assert all(item["issue_type"] for item in output["items"])
    assert all(item["legal_consequence_if_missing"] for item in output["items"])
    assert "llm_status" in output
    assert "global_fallback_reason" in output
    assert "【风险点描述】" in output["result"]
    assert "【建议条款示例】" in output["result"]
    assert "【法律依据索引】" in output["result"]
    assert "本报告基于完整合同模板逐项审查" in output["result"]
    assert output["report_notice"].startswith("本报告基于完整合同模板逐项审查")
    forbidden_words = ("判决结果", "裁判结论", "驳回请求", "押金全额返还", "扣除维修费后返还余额", "承租方应赔偿")
    assert not any(word in output["result"] for word in forbidden_words)


def test_review_v2_can_filter_low_risk_items():
    output = review_contract_v2(
        "出租人和承租人约定房屋所在地法院管辖，押金为6000元，退租交接后七日内退还押金。",
        "house",
        include_low_risk=False,
    )

    assert output["success"] is True
    assert len(output["items"]) < 50
    assert output["summary"]["高"] > 0
    assert all(item["risk_level"] in {"高", "中"} for item in output["items"])


def test_review_v2_detects_overbroad_damage_clause_as_risky_clause():
    output = review_contract_v2(
        "合同约定：租赁期间任何损坏均由租客赔偿，出租人可直接扣除押金。",
        "房屋租赁合同",
    )

    item = next(item for item in output["items"] if item["item_id"] == "HL-038")
    assert item["risk_level"] == "中"
    assert item["status"] == "发现风险条款"
    assert item["coverage_status"] == "风险条款"
    assert "明显加重承租人责任" in item["risk_point_description"]
    assert item["evidence_clauses"]


def test_review_v2_marks_covered_item_low_risk():
    output = review_contract_v2(
        "双方约定租赁期满后承租人继续居住且出租人继续收租、未续签书面合同的，转为不定期租赁，任何一方提前一个月通知。",
        "house_lease",
    )

    item = next(item for item in output["items"] if item["item_id"] == "HL-013")
    assert item["risk_level"] == "低"
    assert item["coverage_status"] == "通过"
    assert item["status"] == "已覆盖"


def test_conservative_profile_does_not_bulk_promote_medium_house_lease_risks():
    output = review_contract_v2(
        "第十一条 争议解决：依法向房屋所在地浦东新区人民法院起诉。"
        "第九条 乙方逾期不支付租金累计超过1个月的，甲方可以解除本合同。",
        "house_lease",
        include_low_risk=True,
        use_llm=False,
        profile=UserRiskProfile.conservative_individual_tenant(),
    )
    by_id = {item["item_id"]: item for item in output["items"]}

    assert output["summary"]["高"] <= 12
    assert by_id["HL-001"]["risk_level"] == "低"
    assert by_id["HL-001"]["coverage_status"] == "通过"
    assert by_id["HL-015"]["risk_level"] == "低"
    assert by_id["HL-015"]["coverage_status"] == "通过"
    for item_id in ("HL-007", "HL-008", "HL-009", "HL-013", "HL-019", "HL-026", "HL-041", "HL-042", "HL-047", "HL-049"):
        assert by_id[item_id]["risk_level"] != "高"
    assert "对你而言建议修改" in by_id["HL-007"]["reasoning_summary"]


def test_real_house_lease_jurisdiction_renewal_and_termination_corrections():
    output = review_contract_v2(
        "\n".join(
            [
                "第八条 租赁期满，乙方需继续承租的，应提前一个月向甲方提出续租要求，双方另行签订租赁合同。",
                "第九条 解除条件。甲方未按时交付房屋、乙方擅自改变用途、乙方逾期不支付租金累计超过1个月或者擅自转租的，守约方可以书面通知解除本合同。",
                "第十一条 争议解决。双方协商不成的，依法向房屋所在地浦东新区人民法院起诉。",
            ]
        ),
        "house_lease",
        include_low_risk=True,
        use_llm=False,
    )
    by_id = {item["item_id"]: item for item in output["items"]}

    assert by_id["HL-001"]["risk_level"] == "低"
    assert by_id["HL-001"]["coverage_status"] == "通过"
    assert "浦东新区人民法院" in by_id["HL-001"]["evidence_clauses"][0]["text"]

    assert by_id["HL-007"]["risk_level"] == "中"
    assert by_id["HL-007"]["coverage_status"] == "部分覆盖"
    assert "仅约定了租赁期满后的续租" in by_id["HL-007"]["reasoning_summary"]

    assert by_id["HL-017"]["risk_level"] == "低"
    assert by_id["HL-017"]["coverage_status"] == "通过"
    assert "逾期不支付租金累计超过1个月" in by_id["HL-017"]["evidence_clauses"][0]["text"]


def test_indefinite_lease_missing_reason_explains_article_734_and_suggested_risk():
    output = review_contract_v2(
        "租赁期限为一年，租期届满后乙方应当交还房屋。",
        "house_lease",
        use_llm=False,
    )

    item = next(item for item in output["items"] if item["item_id"] == "HL-013")
    assert item["risk_level"] == "中"
    assert item["coverage_status"] == "缺失"
    assert "《民法典》第734条" in item["reasoning_summary"]
    assert "不定期" in item["reasoning_summary"]
    assert "建议性补充条款" in item["reasoning_summary"]


def test_review_v2_rejects_unsupported_type_without_crashing():
    output = review_contract_v2("测试合同文本", "unknown_contract")

    assert output["success"] is False
    assert output["items"] == []
    assert "当前支持" in output["message"]
    assert "房屋租赁合同" in output["message"]
    assert "买卖合同" in output["message"]
    assert "赠与合同" in output["message"]


def test_clause_parser_separates_notice_and_formal_terms():
    text = """
使用说明
本说明仅供参考。
特别提示
出租人应注意提示承租人不得违法使用。
甲方（出租人）：
乙方（承租人）：
第一条 房屋基本情况
（一）该房屋坐落于北京市朝阳区。
第二条 租赁用途
该房屋的租赁用途为居住。未经甲方同意，乙方不得变更租赁用途。
附件一 房屋交割单
水电读数：____。
"""
    clauses = parse_contract_clauses(text)

    assert any(clause.section_type == SECTION_USAGE for clause in clauses)
    assert any(clause.section_type == SECTION_SPECIAL_NOTICE for clause in clauses)
    assert any(clause.section_type == SECTION_FORMAL for clause in clauses)
    assert any(clause.location.startswith("第二条") for clause in clauses)


def test_clause_parser_splits_unstructured_contract_lines():
    text = (
        "2023年3月1日，原告（出租人）与被告（承租人）签订《房屋租赁合同》，约定：\n"
        "月租金为人民币12,000元，按季度支付，每季度首月5日前支付当季租金；\n"
        "押金为人民币36,000元，合同终止时结清各项费用后无息返还；"
    )
    clauses = parse_contract_clauses(text)

    assert len(clauses) >= 3
    assert clauses[0].location == "第1段"
    assert not any(clause.location == "全文" for clause in clauses)


def test_negated_phrase_does_not_pollute_later_matches():
    output = review_contract_v2(
        "第六条 押金。押金不计息。出租人收到押金后应向承租人出具押金收款凭证。"
        "第二条 租赁用途。该房屋的租赁用途为居住，乙方不得变更租赁用途。",
        "house_lease",
    )

    deposit = next(item for item in output["items"] if item["item_id"] == "HL-028")
    use_restriction = next(item for item in output["items"] if item["item_id"] == "HL-048")
    assert deposit["coverage_status"] == "通过"
    assert use_restriction["coverage_status"] == "通过"


def test_blank_values_are_partial_coverage():
    output = review_contract_v2(
        "第五条 租金。该房屋月租金为             元。"
        "第六条 押金。乙方应向甲方支付押金            元，甲方收到押金后出具收款凭证。",
        "house_lease",
    )

    item = next(item for item in output["items"] if item["item_id"] == "HL-028")
    assert item["coverage_status"] == "部分覆盖"
    assert item["issue_type"] == "模板空白"
    assert item["risk_level"] == "中"
    assert "空白" in item["reasoning_summary"]


def test_rule_review_does_not_pass_generic_party_or_deposit_deadline():
    output = review_contract_v2(
        "2023年3月1日，原告（出租人）与被告（承租人）签订《房屋租赁合同》，约定：\n"
        "押金为人民币36,000元，合同终止时结清各项费用后无息返还；\n"
        "租赁期满或合同解除后，承租人应在7日内搬离并交还房屋。",
        "house_lease",
        use_llm=False,
    )

    party = next(item for item in output["items"] if item["item_id"] == "HL-002")
    deadline = next(item for item in output["items"] if item["item_id"] == "HL-043")
    sublease = next(item for item in output["items"] if item["item_id"] == "HL-016")
    demolition = next(item for item in output["items"] if item["item_id"] == "HL-024")
    repair = next(item for item in output["items"] if item["item_id"] == "HL-026")
    indefinite = next(item for item in output["items"] if item["item_id"] == "HL-013")
    offset = next(item for item in output["items"] if item["item_id"] == "HL-042")
    assert party["coverage_status"] == "通过"
    assert party["issue_type"] == "通过"
    assert party["risk_level"] == "低"
    assert party["fallback_reason"] == "案件文书摘录来源豁免"
    assert deadline["coverage_status"] == "部分覆盖"
    assert deadline["risk_level"] == "高"
    assert deadline["issue_type"] == "条款缺陷"
    assert sublease["coverage_status"] != "通过"
    assert sublease["coverage_status"] == "缺失"
    assert sublease["evidence_clauses"] == []
    assert demolition["coverage_status"] != "通过"
    assert repair["coverage_status"] != "通过"
    assert indefinite["coverage_status"] != "通过"
    assert offset["coverage_status"] != "通过"
    assert any("押金" in clause["text"] for clause in offset["evidence_clauses"])
    assert party["evidence_location"] != "全文"


def test_case_document_excerpt_exempts_party_identity_details():
    output = review_contract_v2(
        "2023年3月1日，原告（出租人）与被告（承租人）签订《房屋租赁合同》，约定月租金12000元、押金36000元。"
        "原告诉称被告退租后未结清费用，被告辩称押金应予返还。",
        "house_lease",
        use_llm=False,
    )

    party = next(item for item in output["items"] if item["item_id"] == "HL-002")
    assert party["coverage_status"] == "通过"
    assert party["issue_type"] == "通过"
    assert party["risk_level"] == "低"
    assert party["fallback_reason"] == "案件文书摘录来源豁免"


def test_third_party_tort_does_not_recall_fee_or_decoration_terms():
    output = review_contract_v2(
        "第七条 费用。甲方承担物业费、水费、电费；乙方承担网络使用费、燃气费。\n"
        "第九条 装修。乙方装修过程中不得破坏房屋主体结构，并遵守物业管理规定。",
        "house_lease",
        use_llm=False,
    )

    item = next(item for item in output["items"] if item["item_id"] == "HL-006")
    assert item["coverage_status"] == "缺失"
    assert item["issue_type"] == "完全缺失"
    assert item["evidence_clauses"] == []

    defect = next(item for item in output["items"] if item["item_id"] == "HL-023")
    assert defect["risk_level"] == "低"
    assert defect["coverage_status"] == "缺失"
    assert defect["issue_type"] == "完全缺失"
    assert defect["evidence_clauses"] == []


def test_notice_delivery_does_not_recall_advance_notice_only():
    output = review_contract_v2(
        "第十五条 提前解除。乙方提前退租的，应提前30日书面通知甲方，并按实际使用期间结算费用。",
        "house_lease",
        use_llm=False,
    )

    item = next(item for item in output["items"] if item["item_id"] == "HL-005")
    assert item["risk_level"] == "中"
    assert item["coverage_status"] == "缺失"
    assert item["issue_type"] == "完全缺失"
    assert item["evidence_clauses"] == []


def test_expropriation_demolition_requires_direct_clause_not_generic_termination():
    output = review_contract_v2(
        "第五条 押金。合同终止时，乙方结清费用后，甲方无息返还押金。\n"
        "第十五条 合同解除。任何一方提前解除合同的，应提前30日书面通知对方。",
        "house_lease",
        use_llm=False,
    )

    item = next(item for item in output["items"] if item["item_id"] == "HL-024")
    assert item["risk_level"] == "中"
    assert item["coverage_status"] == "缺失"
    assert item["issue_type"] == "完全缺失"
    assert item["evidence_clauses"] == []


def test_llm_selected_surface_keyword_evidence_is_downgraded_to_missing():
    by_id = {checkpoint.item_id: checkpoint for checkpoint in HOUSE_LEASE_REVIEW_CHECKLIST}
    cases = [
        ("HL-005", "第十五条 提前解除。乙方提前退租的，应提前30日书面通知甲方。"),
        ("HL-023", "第七条 费用。水费、电费、燃气费、物业管理费由乙方承担。"),
        ("HL-024", "第五条 押金。合同终止时，乙方结清费用后，甲方无息返还押金。"),
    ]

    for item_id, text in cases:
        clause = ParsedClause(
            clause_id=f"{item_id}-C1",
            location="测试条款",
            title=text[:20],
            text=text,
            section_type=SECTION_FORMAL,
            level=1,
        )
        judgment = _judgment_from_model_result(
            by_id[item_id],
            [clause],
            [],
            {
                "coverage_status": "部分覆盖",
                "risk_level": by_id[item_id].risk_level,
                "reasoning_summary": "模型认为表面关键词可部分覆盖。",
                "evidence_clause_ids": [clause.clause_id],
                "confidence": "high",
                "needs_human_review": False,
            },
        )

        assert judgment.coverage_status == "缺失"
        assert judgment.issue_type == "完全缺失"
        assert judgment.evidence_clauses == []
        assert judgment.evidence_location == ""
        assert judgment.reasoning_summary == "未在正式合同条款中找到能覆盖该审查点的有效约定。"


def test_substantive_notice_and_expropriation_clauses_remain_valid_evidence():
    output = review_contract_v2(
        "第十八条 通知、送达。双方保证在本合同中记载的联系电话、联系地址、电子邮箱等信息均真实有效，"
        "作为本合同履行以及争议解决时的有效联系方式和送达地址。任何一方变更联系电话、联系地址、电子邮箱的，"
        "应当书面通知对方；未及时通知的，对方按照本合同约定的联系方式和送达地址进行送达的，视为有效送达。\n"
        "第十六条 合同解除。该房屋被依法征收征用的，本合同解除，双方按实际使用期间结算费用。",
        "house_lease",
        use_llm=False,
    )
    by_id = {item["item_id"]: item for item in output["items"]}

    assert by_id["HL-005"]["coverage_status"] == "通过"
    assert by_id["HL-005"]["evidence_clauses"]
    assert by_id["HL-024"]["coverage_status"] == "通过"
    assert by_id["HL-024"]["evidence_clauses"]


def test_term_expiry_settlement_prefers_deposit_evidence_over_breach_clause():
    output = review_contract_v2(
        "第五条 押金。合同终止时，乙方结清各项费用后，甲方无息返还押金。\n"
        "第十二条 违约责任。任何一方违约的，应按实际损失承担违约责任。",
        "house_lease",
        use_llm=False,
    )

    item = next(item for item in output["items"] if item["item_id"] == "HL-031")
    assert item["risk_level"] == "中"
    assert item["coverage_status"] == "部分覆盖"
    assert item["issue_type"] == "条款缺陷"
    assert any("押金" in clause["text"] for clause in item["evidence_clauses"])
    assert not any("违约责任" in clause["text"] for clause in item["evidence_clauses"])


def test_blank_early_surrender_and_penalty_terms_are_not_passed():
    output = review_contract_v2(
        "第十五条 提前解除。甲方提前收回该房屋的，或者乙方提前退租的，应当提前   日书面通知对方，"
        "并经对方同意后，应当按照           的标准向对方支付违约金。",
        "house_lease",
        use_llm=False,
    )

    early = next(item for item in output["items"] if item["item_id"] == "HL-029")
    deduction = next(item for item in output["items"] if item["item_id"] == "HL-039")
    assert early["coverage_status"] == "部分覆盖"
    assert early["issue_type"] == "模板空白"
    assert early["risk_level"] == "中"
    assert "违约金比例" in early["risk_point_description"]
    assert "实际损失" in early["risk_point_description"]
    assert "模板空白" in early["reasoning_summary"]
    assert "调减" in early["reasoning_summary"]
    assert deduction["coverage_status"] == "部分覆盖"
    assert deduction["issue_type"] == "模板空白"
    assert deduction["risk_level"] == "高"


def test_termination_notice_distinguishes_early_surrender_notice():
    output = review_contract_v2(
        "第十五条 提前解除。乙方提前退租的，应提前30日书面通知甲方，并按实际使用期间结算费用。",
        "house_lease",
        use_llm=False,
    )

    item = next(item for item in output["items"] if item["item_id"] == "HL-019")
    assert item["risk_level"] == "中"
    assert item["coverage_status"] == "部分覆盖"
    assert "提前退租的书面通知" in item["reasoning_summary"]
    assert "解除通知的送达方式" in item["reasoning_summary"]
    assert "前者启动解除程序，后者使解除行为生效" in item["reasoning_summary"]


def test_loss_based_penalty_adjustment_can_pass_when_complete():
    output = review_contract_v2(
        "押金扣款、违约金或赔偿金额应以实际损失为基础；约定金额明显高于实际损失的，双方可协商按实际损失及合理费用调整。"
        "承租人提前退租的，应提前30日书面通知出租人，并按实际使用期间、未结费用和合理空置损失结算。",
        "house_lease",
        use_llm=False,
    )

    early = next(item for item in output["items"] if item["item_id"] == "HL-029")
    deduction = next(item for item in output["items"] if item["item_id"] == "HL-039")
    assert early["coverage_status"] == "通过"
    assert deduction["coverage_status"] == "通过"


def test_formal_contract_still_requires_party_identity_details():
    output = review_contract_v2(
        "出租人与承租人签订《房屋租赁合同》，约定月租金12000元、押金36000元。",
        "house_lease",
        use_llm=False,
    )

    party = next(item for item in output["items"] if item["item_id"] == "HL-002")
    assert party["coverage_status"] != "通过"
    assert party["risk_level"] == "中"


def test_missing_item_reports_clause_absence_and_legal_consequence():
    output = review_contract_v2(
        "第六条 押金。乙方应向甲方支付押金6000元，甲方收到押金后出具收款凭证。",
        "house_lease",
    )

    item = next(item for item in output["items"] if item["item_id"] == "HL-020")
    assert item["coverage_status"] == "缺失"
    assert item["issue_type"] == "完全缺失"
    assert item["legal_consequence_if_missing"]
    assert "新房主" in item["legal_consequence_if_missing"] or "接续义务" in item["legal_consequence_if_missing"]


def test_demo_house_lease_contract_has_evidence_and_no_verdict_words():
    source = Path("test_data")
    files = list(source.glob("*2614*.docx"))
    if not files:
        return

    doc = Document(str(files[0]))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())
    output = review_contract_v2(text, "house_lease")

    assert output["all_items_count"] == 50
    assert output["clause_count"] > 20
    assert output["formal_clause_count"] > 10
    assert any(item["evidence_clauses"] for item in output["items"])
    assert any(item["coverage_status"] == "部分覆盖" for item in output["items"])
    assert not any(word in output["result"] for word in ("判决结果", "裁判结论", "驳回请求", "应赔偿"))


def test_llm_judgment_is_used_when_model_returns_valid_json(monkeypatch):
    monkeypatch.setenv("REVIEW_V2_LLM_ENABLED", "true")
    monkeypatch.setenv("DASHSCOPE_API_KEY", "test-key")

    def fake_call_llm(prompt, response_format="text", max_retries=3):
        assert response_format == "json"
        return json.dumps(
            {
                "reviews": [
                    {
                        "item_id": "HL-028",
                        "coverage_status": "通过",
                        "risk_level": "低",
                        "reasoning_summary": "正式条款已经明确约定押金金额和收款凭证，能够覆盖该审查点。",
                        "evidence_clause_ids": ["C0001"],
                        "confidence": "high",
                        "needs_human_review": False,
                    }
                ]
            },
            ensure_ascii=False,
        )

    monkeypatch.setattr("reasoning.llm_client.call_llm", fake_call_llm)
    output = review_contract_v2(
        "第六条 押金。乙方应向甲方支付押金6000元，甲方收到押金后出具收款凭证。",
        "house_lease",
    )

    item = next(item for item in output["items"] if item["item_id"] == "HL-028")
    assert item["judge_source"] == "llm"
    assert item["llm_confidence"] == "high"
    assert item["fallback_reason"] == ""
    assert item["coverage_status"] == "通过"


def test_llm_pass_is_downgraded_when_evidence_has_blank_values(monkeypatch):
    monkeypatch.setenv("REVIEW_V2_LLM_ENABLED", "true")
    monkeypatch.setenv("DASHSCOPE_API_KEY", "test-key")
    monkeypatch.setenv("REVIEW_V2_LLM_BATCH_SIZE", "50")

    def fake_call_llm(prompt, response_format="text", max_retries=3):
        payload_text = prompt.split("审查点和候选条款 JSON：", 1)[1].split("请输出严格 JSON", 1)[0].strip()
        payload = json.loads(payload_text)
        reviews = []
        for item in payload:
            candidate_ids = [clause["clause_id"] for clause in item["candidate_clauses"]]
            reviews.append(
                {
                    "item_id": item["item_id"],
                    "coverage_status": "通过" if item["item_id"] == "HL-028" else item["rule_judgment"]["coverage_status"],
                    "risk_level": "低" if item["item_id"] == "HL-028" else item["rule_judgment"]["risk_level"],
                    "reasoning_summary": "正式条款已经覆盖该审查点。",
                    "evidence_clause_ids": candidate_ids[:1],
                    "confidence": "high",
                    "needs_human_review": False,
                }
            )
        return json.dumps({"reviews": reviews}, ensure_ascii=False)

    monkeypatch.setattr("reasoning.llm_client.call_llm", fake_call_llm)
    output = review_contract_v2(
        "第六条 押金。乙方应向甲方支付押金            元，甲方收到押金后出具收款凭证。",
        "house_lease",
    )

    item = next(item for item in output["items"] if item["item_id"] == "HL-028")
    assert item["judge_source"] == "llm"
    assert item["coverage_status"] == "部分覆盖"
    assert item["issue_type"] == "模板空白"
    assert item["risk_level"] == "中"
    assert item["llm_confidence"] == "low"


def test_llm_cannot_override_risk_whitelist_or_pass_generic_deposit_deadline(monkeypatch):
    monkeypatch.setenv("REVIEW_V2_LLM_ENABLED", "true")
    monkeypatch.setenv("DASHSCOPE_API_KEY", "test-key")
    monkeypatch.setenv("REVIEW_V2_LLM_BATCH_SIZE", "50")

    def fake_call_llm(prompt, response_format="text", max_retries=3):
        payload_text = prompt.split("审查点和候选条款 JSON：", 1)[1].split("请输出严格 JSON", 1)[0].strip()
        payload = json.loads(payload_text)
        reviews = []
        for item in payload:
            candidate_ids = [clause["clause_id"] for clause in item["candidate_clauses"]]
            if item["item_id"] == "HL-042":
                reviews.append(
                    {
                        "item_id": "HL-042",
                        "coverage_status": "部分覆盖",
                        "risk_level": "高",
                        "reasoning_summary": "模型认为押金抵扣顺序存在高风险。",
                        "evidence_clause_ids": candidate_ids[:1],
                        "confidence": "high",
                        "needs_human_review": False,
                    }
                )
            elif item["item_id"] == "HL-013":
                reviews.append(
                    {
                        "item_id": "HL-013",
                        "coverage_status": "通过",
                        "risk_level": "低",
                        "reasoning_summary": "模型认为期满搬离交房条款已经覆盖不定期租赁安排。",
                        "evidence_clause_ids": candidate_ids[:1],
                        "confidence": "high",
                        "needs_human_review": False,
                    }
                )
            elif item["item_id"] == "HL-043":
                reviews.append(
                    {
                        "item_id": "HL-043",
                        "coverage_status": "通过",
                        "risk_level": "低",
                        "reasoning_summary": "模型认为押金返还已有约定。",
                        "evidence_clause_ids": candidate_ids[:1],
                        "confidence": "high",
                        "needs_human_review": False,
                    }
                )
            else:
                reviews.append(
                    {
                        "item_id": item["item_id"],
                        "coverage_status": item["rule_judgment"]["coverage_status"],
                        "risk_level": item["rule_judgment"]["risk_level"],
                        "reasoning_summary": "批量语义复核已完成该审查点判断。",
                        "evidence_clause_ids": candidate_ids[:1],
                        "confidence": "medium",
                        "needs_human_review": False,
                    }
                )
        return json.dumps({"reviews": reviews}, ensure_ascii=False)

    monkeypatch.setattr("reasoning.llm_client.call_llm", fake_call_llm)
    output = review_contract_v2(
        "第六条 押金。押金为36000元，合同终止时结清各项费用后无息返还。"
        "押金可用于抵扣租金、水电费、维修费。"
        "租赁期满或合同解除后，承租人应在7日内搬离并交还房屋。",
        "house_lease",
    )

    offset_order = next(item for item in output["items"] if item["item_id"] == "HL-042")
    indefinite = next(item for item in output["items"] if item["item_id"] == "HL-013")
    deadline = next(item for item in output["items"] if item["item_id"] == "HL-043")
    assert offset_order["coverage_status"] == "部分覆盖"
    assert offset_order["risk_level"] == "中"
    assert indefinite["coverage_status"] != "通过"
    assert indefinite["risk_level"] == "中"
    assert indefinite["llm_confidence"] == "low"
    assert deadline["coverage_status"] == "部分覆盖"
    assert deadline["risk_level"] == "高"
    assert deadline["llm_confidence"] == "low"
    assert "具体天数" in deadline["reasoning_summary"] or "一般性表述" in deadline["reasoning_summary"]


def test_llm_review_uses_batches_instead_of_one_call_per_checkpoint(monkeypatch):
    monkeypatch.setenv("REVIEW_V2_LLM_ENABLED", "true")
    monkeypatch.setenv("DASHSCOPE_API_KEY", "test-key")
    monkeypatch.setenv("REVIEW_V2_LLM_BATCH_SIZE", "20")
    calls = []

    def fake_call_llm(prompt, response_format="text", max_retries=3):
        calls.append(prompt)
        payload_text = prompt.split("审查点和候选条款 JSON：", 1)[1].split("请输出严格 JSON", 1)[0].strip()
        payload = json.loads(payload_text)
        return json.dumps(
            {
                "reviews": [
                    {
                        "item_id": item["item_id"],
                        "coverage_status": item["rule_judgment"]["coverage_status"],
                        "risk_level": item["rule_judgment"]["risk_level"],
                        "reasoning_summary": "批量语义复核已完成该审查点判断。",
                        "evidence_clause_ids": [
                            item["candidate_clauses"][0]["clause_id"]
                        ] if item["candidate_clauses"] else [],
                        "confidence": "medium",
                        "needs_human_review": False,
                    }
                    for item in payload
                ]
            },
            ensure_ascii=False,
        )

    monkeypatch.setattr("reasoning.llm_client.call_llm", fake_call_llm)
    output = review_contract_v2(
        "第一条 房屋基本情况。该房屋坐落于北京市朝阳区。"
        "第五条 租金。该房屋月租金为3000元。"
        "第六条 押金。乙方应向甲方支付押金6000元，甲方收到押金后出具收款凭证。"
        "第二条 租赁用途。该房屋的租赁用途为居住，乙方不得变更租赁用途。",
        "house_lease",
    )

    assert output["judge_source_summary"]["llm"] > 0
    assert 0 < len(calls) < 50


def test_llm_generic_breach_clause_cannot_pass_specialized_successor_review(monkeypatch):
    monkeypatch.setenv("REVIEW_V2_LLM_ENABLED", "true")
    monkeypatch.setenv("DASHSCOPE_API_KEY", "test-key")
    monkeypatch.setenv("REVIEW_V2_LLM_BATCH_SIZE", "50")

    def fake_call_llm(prompt, response_format="text", max_retries=3):
        payload_text = prompt.split("审查点和候选条款 JSON：", 1)[1].split("请输出严格 JSON", 1)[0].strip()
        payload = json.loads(payload_text)
        reviews = []
        for item in payload:
            candidate_ids = [clause["clause_id"] for clause in item["candidate_clauses"]]
            if item["item_id"] == "HL-012":
                reviews.append(
                    {
                        "item_id": "HL-012",
                        "coverage_status": "通过",
                        "risk_level": "低",
                        "reasoning_summary": "条款约定继续履行责任，可以覆盖出租人死亡或权利继受安排。",
                        "evidence_clause_ids": candidate_ids[:1],
                        "confidence": "high",
                        "needs_human_review": False,
                    }
                )
            else:
                reviews.append(
                    {
                        "item_id": item["item_id"],
                        "coverage_status": item["rule_judgment"]["coverage_status"],
                        "risk_level": item["rule_judgment"]["risk_level"],
                        "reasoning_summary": "批量语义复核已完成该审查点判断。",
                        "evidence_clause_ids": candidate_ids[:1],
                        "confidence": "medium",
                        "needs_human_review": False,
                    }
                )
        return json.dumps({"reviews": reviews}, ensure_ascii=False)

    monkeypatch.setattr("reasoning.llm_client.call_llm", fake_call_llm)
    output = review_contract_v2(
        "第十五条 违约责任。任何一方有其他违约行为的，应当承担继续履行、采取补救措施或者赔偿损失等违约责任。",
        "house_lease",
    )

    item = next(item for item in output["items"] if item["item_id"] == "HL-012")
    assert item["judge_source"] == "llm"
    assert item["coverage_status"] == "部分覆盖"
    assert item["risk_level"] == "低"
    assert item["llm_confidence"] == "low"
    assert "一般性表述" in item["reasoning_summary"]


def test_use_llm_false_skips_model_call_even_when_key_exists(monkeypatch):
    monkeypatch.setenv("REVIEW_V2_LLM_ENABLED", "true")
    monkeypatch.setenv("DASHSCOPE_API_KEY", "test-key")

    def fake_call_llm(prompt, response_format="text", max_retries=3):
        raise AssertionError("LLM should not be called when use_llm is false")

    monkeypatch.setattr("reasoning.llm_client.call_llm", fake_call_llm)
    output = review_contract_v2(
        "第六条 押金。乙方应向甲方支付押金6000元，甲方收到押金后出具收款凭证。",
        "house_lease",
        use_llm=False,
    )

    item = next(item for item in output["items"] if item["item_id"] == "HL-028")
    assert item["judge_source"] == "rule_fallback"
    assert "本次审查未启用大模型语义复核" in item["fallback_reason"]
    assert output["global_fallback_reason"] == item["fallback_reason"]
    assert output["llm_enabled"] is False


def test_llm_invalid_json_falls_back_to_rule(monkeypatch):
    monkeypatch.setenv("REVIEW_V2_LLM_ENABLED", "true")
    monkeypatch.setenv("DASHSCOPE_API_KEY", "test-key")

    def fake_call_llm(prompt, response_format="text", max_retries=3):
        return "不是 JSON"

    monkeypatch.setattr("reasoning.llm_client.call_llm", fake_call_llm)
    output = review_contract_v2(
        "第六条 押金。乙方应向甲方支付押金6000元，甲方收到押金后出具收款凭证。",
        "house_lease",
    )

    item = next(item for item in output["items"] if item["item_id"] == "HL-028")
    assert item["judge_source"] == "rule_fallback"
    assert "模型判定不可用" in item["fallback_reason"]
    assert item["coverage_status"] == "通过"


def test_llm_verdict_language_falls_back_to_rule(monkeypatch):
    monkeypatch.setenv("REVIEW_V2_LLM_ENABLED", "true")
    monkeypatch.setenv("DASHSCOPE_API_KEY", "test-key")

    def fake_call_llm(prompt, response_format="text", max_retries=3):
        return json.dumps(
            {
                "reviews": [
                    {
                        "item_id": "HL-028",
                        "coverage_status": "通过",
                        "risk_level": "低",
                        "reasoning_summary": "承租方应赔偿，押金应按判决结果处理。",
                        "evidence_clause_ids": ["C0001"],
                        "confidence": "high",
                        "needs_human_review": False,
                    }
                ]
            },
            ensure_ascii=False,
        )

    monkeypatch.setattr("reasoning.llm_client.call_llm", fake_call_llm)
    output = review_contract_v2(
        "第六条 押金。乙方应向甲方支付押金6000元，甲方收到押金后出具收款凭证。",
        "house_lease",
    )

    item = next(item for item in output["items"] if item["item_id"] == "HL-028")
    assert item["judge_source"] == "rule_fallback"
    assert "裁判式表述" in item["fallback_reason"] or "未返回该审查点" in item["fallback_reason"]
    assert "应赔偿" not in output["result"]


def test_review_v2_exports_pdf_and_docx_bytes(tmp_path):
    output = review_contract_v2(
        "2023年3月1日，原告（出租人）与被告（承租人）签订《房屋租赁合同》，约定押金36000元。",
        "house_lease",
        use_llm=False,
    )
    payload = ReviewV2PdfInput(
        filename="房屋租赁合同测试.docx",
        contract_type_name=output["contract_type_name"],
        summary=output["summary"],
        coverage_summary=output["coverage_summary"],
        issue_type_summary=output["issue_type_summary"],
        judge_source_summary=output["judge_source_summary"],
        llm_status=output["llm_status"],
        global_fallback_reason=output["global_fallback_reason"],
        report_notice=output["report_notice"],
        items=output["items"],
    )

    pdf_content = _build_pdf(payload)
    assert pdf_content.startswith(b"%PDF")
    assert len(pdf_content) > 1000

    docx_content = _build_docx(payload)
    docx_path = tmp_path / "review_v2.docx"
    docx_path.write_bytes(docx_content)
    exported = Document(str(docx_path))
    full_text = "\n".join(paragraph.text for paragraph in exported.paragraphs)
    assert "合同审查报告（第二版）" in full_text
    assert "出租人与承租人身份信息" in full_text
