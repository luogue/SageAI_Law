import json
from pathlib import Path

from docx import Document

from review_v2.clause_parser import parse_contract_clauses
from review_v2.domain_selection import resolve_domain_selection
from review_v2.engine import review_contract_v2
from review_v2.user_profile import UserRiskProfile


ROOT = Path(__file__).resolve().parents[2]
COMMISSIONED_FIXTURE = next((ROOT / "test_data").glob("12*.docx"))
CLASSIFICATION_GOLDEN = (
    Path(__file__).parent
    / "fixtures"
    / "technical_contracts"
    / "commissioned_development.classification.golden.json"
)


def _fixture_text() -> str:
    document = Document(COMMISSIONED_FIXTURE)
    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        blocks.extend(
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in table.rows
        )
    return "\n".join(blocks)


def _resolve(text: str, selected: str = "service"):
    return resolve_domain_selection(selected, text, parse_contract_clauses(text))


def test_software_development_with_training_routes_to_commissioned_development():
    text = """
    软件委托开发合同
    甲方作为委托方委托乙方作为开发方开发企业管理软件。甲方支付开发费用，乙方负责开发工作。
    乙方按需求规格和项目里程碑交付源代码、API接口文档及架构文档，开发成果经最终验收。
    验收后乙方培训甲方人员使用软件。
    """
    resolution, _, routed_type = _resolve(text)

    assert routed_type == "technology"
    assert resolution.resolution == "DOCUMENT_OVERRIDE"
    assert resolution.selected_domain == "TECHNOLOGY"
    assert resolution.document_domain["primary_contract_type"] == "commissioned_development"
    assert "TRAINING" in resolution.classification_context["ancillary_obligations"]
    assert resolution.classification_context["consumer_service_context"]["gate_status"] == "NOT_APPLICABLE"


def test_software_development_with_support_remains_commissioned_development():
    text = """
    软件委托开发合同
    甲方委托乙方研究开发数据平台，甲方支付开发费，乙方负责开发工作。
    乙方交付源代码和开发成果，项目验收后提供技术支持和维护。
    """
    resolution, _, routed_type = _resolve(text)

    assert routed_type == "technology"
    assert resolution.document_domain["primary_subtype"] == "ENTRUSTED_DEVELOPMENT"
    assert "TECHNICAL_SUPPORT" in resolution.classification_context["ancillary_obligations"]


def test_pure_training_without_development_remains_service_transaction():
    text = """
    软件使用培训服务合同
    乙方向甲方学员提供软件操作培训课程，共20课时。培训费用为20000元。
    本合同不涉及软件研发、新成果形成或源代码交付。
    """
    resolution, _, routed_type = _resolve(text)

    assert routed_type == "service"
    assert resolution.selected_domain == "SERVICE"
    consumer = resolution.classification_context["consumer_service_context"]
    assert consumer["is_training_primary_transaction"] is True
    assert consumer["gate_status"] == "APPLICABLE"


def test_development_with_deployment_maintenance_training_keeps_primary_hierarchy():
    text = """
    信息系统委托开发合同
    委托方甲方委托开发方乙方开发新的业务系统，甲方支付研发费用，乙方负责开发工作。
    按需求规格、开发里程碑交付源代码、API接口和开发成果并验收。
    乙方另提供部署、维护、培训和售后技术支持。
    """
    resolution, _, routed_type = _resolve(text)

    assert routed_type == "technology"
    assert resolution.document_domain["primary_contract_type"] == "commissioned_development"
    assert set(resolution.classification_context["ancillary_obligations"]) >= {
        "TRAINING", "TECHNICAL_SUPPORT",
    }


def test_real_commissioned_fixture_overrides_service_and_matches_classification_golden():
    manifest = json.loads(CLASSIFICATION_GOLDEN.read_text(encoding="utf-8"))
    result = review_contract_v2(
        _fixture_text(),
        "service",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方"),
    )
    resolution = result["domain_selection_resolution"]
    review_ids = [item["item_id"] for item in result["items"]]
    context = resolution["classification_context"]

    assert manifest["sample_id"] == "COMMISSIONED_DEV_CLASSIFICATION_REGRESSION"
    assert resolution["resolution"] == "DOCUMENT_OVERRIDE"
    assert resolution["document_domain"]["family"] == manifest["contract_family"]
    assert resolution["document_domain"]["primary_contract_type"] == manifest["primary_contract_type"]
    assert resolution["document_domain"]["primary_subtype"] == manifest["contract_subtype"]
    assert result["perspective_resolution"]["canonical_role"] == manifest["review_position"]
    assert any(item_id.startswith(manifest["rule_pack_prefix"]) for item_id in review_ids)
    assert not any(
        item_id.startswith(tuple(manifest["forbidden_rule_pack_prefixes"]))
        for item_id in review_ids
    )
    assert {
        key: context["consumer_service_context"][key]
        for key in manifest["consumer_service_context"]
    } == manifest["consumer_service_context"]
    assert result["domain_selection_consistency_gate"]["gate_status"] == "PASS"
    assert result["release_quality_gate"]["status"] == "REPORT_READY"
