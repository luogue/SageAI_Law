from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from auth.routes import get_current_user
from main import app
from review_v2.clause_parser import parse_contract_clauses
from review_v2.contract_registry import get_checklist_for_resolution
from review_v2.domain_selection import resolve_domain_selection
from review_v2.technology_service_analysis import (
    analyze_improvement_license_facts,
    analyze_license_lifecycle,
    analyze_primary_transaction,
    analyze_proposition_polarity,
    analyze_third_party_infringement_allocation,
    classify_technology_transaction,
)


ROOT = Path(__file__).resolve().parents[2]
GOLD = ROOT / "test_data" / "19 技术许可.docx"


def _docx_text(path: Path) -> str:
    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    tables = [
        " | ".join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
    ]
    return "\n".join([*paragraphs, *tables])


def _review_file(path: Path) -> dict:
    app.dependency_overrides[get_current_user] = lambda: {"id": 1, "phone": "test"}
    try:
        with TestClient(app) as client:
            response = client.post(
                "/review-v2/file",
                files={
                    "file": (
                        path.name,
                        path.read_bytes(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                # Deliberately select sale: high-confidence document facts must override it.
                data={
                    "contract_type": "sale",
                    "include_low_risk": "true",
                    "use_llm": "false",
                    "profile_json": '{"user_identity":"甲方/被许可方"}',
                },
            )
    finally:
        app.dependency_overrides.pop(get_current_user, None)
    assert response.status_code == 200, response.text
    return response.json()


@pytest.fixture(scope="module")
def gold_review() -> dict:
    return _review_file(GOLD)


def test_gold_routes_raw_docx_to_technology_license_not_sale(gold_review: dict):
    classification = gold_review["deterministic_analysis"]["technology_classification"]
    model = classification["primary_transaction_model"]
    review_ids = [item["item_id"] for item in gold_review["items"]]

    assert gold_review["contract_type"] == "technology"
    assert gold_review["selected_contract_type"] == "sale"
    assert gold_review["report_context"]["contract_type_label"] == "技术合同—技术许可合同"
    assert gold_review["report_context"]["scenario_label"] == "技术许可"
    assert classification["primary_subtype"] == "TECHNOLOGY_LICENSE"
    assert classification["license_mode"] == "NON_EXCLUSIVE"
    assert model["type"] == "TECHNOLOGY_LICENSE"
    assert model["dominant_family"] == "TECHNOLOGY"
    assert model["rights_change"] == ["license to use technology"]
    assert model["license_semantics"]["boundaries"]["territory"] is True
    assert model["sale_semantics"]["independent_economic_value"] is False
    assert model["secondary_transactions"] == ["ANCILLARY_FUTURE_GOODS_SUPPLY"]
    assert not any(item_id.startswith("S-") for item_id in review_ids)
    assert all(
        next(item for item in gold_review["items"] if item["item_id"] == item_id)["coverage_status"] != "不适用"
        for item_id in ("TC-07-01", "TC-07-02", "TC-07-03", "TC-07-04", "TC-13-04")
    )


def test_gold_generic_calculations_and_relationship_risks_remain_active(gold_review: dict):
    deterministic = gold_review["deterministic_analysis"]
    payment = deterministic["payment_formula"]
    royalty = deterministic["royalty_calculation"]
    timeline = deterministic["timeline"]
    findings = {row["risk_group"]: row for row in gold_review["risk_findings"]}
    items = {row["item_id"]: row for row in gold_review["items"]}

    assert payment["status"] == "INCONSISTENT"
    assert payment["contract_total"] == 300_000
    assert payment["percentage_total"] == 105
    assert payment["amount_total"] == 315_000
    assert payment["amount_delta"] == 15_000
    assert {row["currency"] for row in payment["nodes"]} == {"USD"}
    assert findings["payment_formula_inconsistency"]["evidence_bundle"]["entailment_status"] == "PASSED"

    assert royalty["status"] == "CONSISTENT"
    assert royalty["sales"] == 10_000_000
    assert royalty["deduction_total"] == 2_000_000
    assert royalty["net_sales"] == 8_000_000
    assert royalty["royalty_rate"] == 4
    assert royalty["royalty_fee"] == 320_000
    assert "royalty_calculation_inconsistency" not in findings

    assert timeline["effective_date"] == "2026-12-20"
    assert timeline["earliest_performance_date"] == "2026-11-20"
    assert timeline["performance_before_effective_date"] is True
    assert "technical_timeline_conflict" in findings
    assert "improvement_no_unlawful_restriction" in findings
    assert items["TC-12-02"]["coverage_status"] == "部分覆盖"
    assert gold_review["release_quality_gate"]["status"] == "REPORT_READY"
    assert gold_review["report_generation_decision"]["status"] == "GENERATED"


def test_gold_semantic_facts_and_rendering_are_consistent(gold_review: dict):
    deterministic = gold_review["deterministic_analysis"]
    items = {row["item_id"]: row for row in gold_review["items"]}
    findings = {row["risk_group"]: row for row in gold_review["risk_findings"]}

    improvement = deterministic["improvement_license"]
    assert improvement["asymmetric_improvement_license"] is True
    by_direction = {(row["grantor"], row["grantee"]): row for row in improvement["facts"]}
    assert by_direction[("LICENSEE", "LICENSOR")]["license_fee"] == "ZERO"
    assert by_direction[("LICENSEE", "LICENSOR")]["license_duration"] == "PERPETUAL"
    assert by_direction[("LICENSEE", "LICENSOR")]["territory"] == "WORLDWIDE"
    assert by_direction[("LICENSEE", "LICENSOR")]["revocability"] == "IRREVOCABLE"
    assert by_direction[("LICENSOR", "LICENSEE")]["license_fee"] == "NEGOTIATED"
    assert items["TC-10-02"]["coverage_status"] == "风险条款"
    assert items["TC-10-04"]["coverage_status"] == "风险条款"

    assert items["TC-11-03"]["coverage_status"] != "通过"
    disposal_evidence = items["TC-11-03"]["evidence_clauses"][0]
    assert disposal_evidence["polarity_analysis"]["polarity"] == "NEGATED"
    assert disposal_evidence["polarity_analysis"]["fact_value"] is False
    assert items["TC-12-01"]["coverage_status"] == "部分覆盖"
    assert items["TC-12-02"]["coverage_status"] == "部分覆盖"
    assert deterministic["third_party_infringement_allocation"]["basic_liability_allocation"] is True

    lifecycle = deterministic["license_lifecycle"]
    assert lifecycle["contract_duration_years"] == 8
    assert lifecycle["technology_license_duration_years"] == 8
    assert lifecycle["technology_rights_survive_contract"] is True
    assert lifecycle["technology_rights_survival_duration"] == "PERPETUAL"
    assert lifecycle["trademark_rights_survive_contract"] is False
    assert lifecycle["post_termination_royalty_defined"] is False
    assert lifecycle["license_lifecycle_risk"] is True
    assert "license_lifecycle_boundary_gap" in findings

    payment_issue = findings["payment_formula_inconsistency"]["issue"]
    assert "315,000美元" in payment_issue
    assert "300,000美元" in payment_issue
    assert "15,000美元" in payment_issue
    assert findings["payment_formula_inconsistency"]["economic_impact"]["currency"] == "USD"
    assert gold_review["perspective_resolution"]["display_role"] == "甲方 / 被许可方"
    assert gold_review["contract_context"]["transaction_role_object"] == {
        "transaction_type": "TECHNOLOGY_LICENSE",
        "review_party": "PARTY_A",
        "legal_role": "LICENSEE",
        "contract_display_role": "被许可方",
    }
    assert gold_review["terminology_consistency_gate"]["status"] == "REPORT_READY"
    assert gold_review["formal_recommendation_validation"]["status"] == "PASS"
    assert gold_review["technology_transaction_consistency_gate"]["status"] == "REPORT_READY"
    forbidden_terms = (
        "专利权受让方", "专利权转移", "让与方处分权", "转让登记失败",
        "软件著作权及开源组件合规证明", "未限制许可方责任",
    )
    assert all(term not in gold_review["result"] for term in forbidden_terms)


def test_improvement_license_mutations_distinguish_asymmetry_symmetry_and_absence():
    asymmetric = analyze_improvement_license_facts(
        "A方形成的改进成果应无偿许可B方。B方形成的改进成果许可A方使用时应另行支付许可费。"
    )
    symmetric = analyze_improvement_license_facts(
        "双方形成的改进成果均应相互无偿许可，许可期限均为5年，许可地域均为中国。"
    )
    missing = analyze_improvement_license_facts("本合同未设置改进成果许可条款。")

    assert asymmetric["arrangement_status"] == "ASYMMETRIC"
    assert asymmetric["asymmetric_improvement_license"] is True
    assert len(asymmetric["facts"]) == 2
    assert symmetric["arrangement_status"] == "SYMMETRIC"
    assert symmetric["asymmetric_improvement_license"] is False
    assert all(row["reciprocal_obligation"] for row in symmetric["facts"])
    assert missing["arrangement_status"] == "MISSING"
    assert missing["facts"] == []


def test_polarity_mutations_distinguish_absence_from_prohibitive_duty():
    absence = analyze_proposition_polarity(
        "双方未约定资料返还删除机制。", ("资料返还删除机制",)
    )
    prohibition = analyze_proposition_polarity(
        "未经许可不得向第三方披露保密信息。", ("向第三方披露",)
    )

    assert absence["polarity"] == "NEGATED"
    assert absence["fact_value"] is False
    assert absence["semantic_role"] == "EXPLICIT_GAP"
    assert prohibition["polarity"] == "PROHIBITIVE_OBLIGATION"
    assert prohibition["fact_value"] is True
    assert prohibition["semantic_role"] == "AFFIRMATIVE_DUTY"


def test_infringement_and_lifecycle_mutations():
    liability = analyze_third_party_infringement_allocation(
        "许可方保证有权许可。被许可方按约实施遭第三方权利主张时，"
        "由许可方负责处理并承担法律和经济责任。"
    )
    aligned_lifecycle = analyze_license_lifecycle(
        "合同期限5年，技术许可期限5年，合同终止后被许可方不得继续使用技术。"
    )
    priced_survival = analyze_license_lifecycle(
        "合同期限5年，技术许可为永久许可，合同终止后继续使用技术仍支付净销售额2%的提成费。"
    )

    assert liability["basic_liability_allocation"] is True
    assert aligned_lifecycle["contract_duration_years"] == 5
    assert aligned_lifecycle["technology_license_duration_years"] == 5
    assert aligned_lifecycle["technology_rights_survive_contract"] is False
    assert aligned_lifecycle["license_lifecycle_risk"] is False
    assert priced_survival["technology_rights_survive_contract"] is True
    assert priced_survival["post_termination_royalty_defined"] is True
    assert priced_survival["post_termination_royalty_obligation"] == "CONTINUES"
    assert priced_survival["license_lifecycle_risk"] is False


LICENSE_CORE = """
{title}
许可方授予被许可方在中华人民共和国境内非独占、不可转让的制造技术使用许可。
被许可方有权使用该技术设计、制造、使用和销售合同产品，不得再许可。
许可方交付技术资料和制造工艺，并提供技术指导和人员培训。
被许可方支付入门费100000美元，并按净销售额5%支付销售提成费。
"""


@pytest.mark.parametrize(
    "title",
    ("技术合作协议", "技术使用协议", "技术引进协议", "许可与服务协议"),
)
def test_title_variants_are_classified_from_transaction_facts(title: str):
    classification = classify_technology_transaction(LICENSE_CORE.format(title=title))
    assert classification["primary_subtype"] == "TECHNOLOGY_LICENSE"
    assert classification["primary_transaction_model"]["type"] == "TECHNOLOGY_LICENSE"


@pytest.mark.parametrize(
    ("text", "expected_type", "expected_family"),
    [
        (
            "设备买卖合同\n德国卖方向中国买方出售10台设备，每台100万美元，交货后设备所有权转移给买方，无技术许可。",
            "GOODS_SALE",
            "SALE",
        ),
        (
            LICENSE_CORE.format(title="技术使用协议"),
            "TECHNOLOGY_LICENSE",
            "TECHNOLOGY",
        ),
        (
            LICENSE_CORE.format(title="技术合作协议") + "\n被许可方可另行购买少量备件，具体供货另行订立合同。",
            "TECHNOLOGY_LICENSE",
            "TECHNOLOGY",
        ),
        (
            "设备买卖合同\n卖方向买方出售10台控制设备，每台100万元，交付后设备所有权转移给买方。"
            "设备配套软件仅为设备运行所需，随设备许可使用，无独立许可费。",
            "GOODS_SALE",
            "SALE",
        ),
        (
            "设备与技术许可合同\n卖方向买方出售10台设备，每台100万元，交付后设备所有权转移给买方。"
            "许可方另行授予被许可方非独占、不可转让的制造技术使用许可，许可地域为中国。"
            "许可方交付技术资料并培训，被许可方另行支付入门费50万元及净销售额5%的提成费。",
            "MIXED_SALE_TECHNOLOGY_LICENSE",
            "TECHNOLOGY",
        ),
    ],
)
def test_primary_transaction_mutations(text: str, expected_type: str, expected_family: str):
    model = analyze_primary_transaction(text)
    resolution, domain, routed_type = resolve_domain_selection("sale", text, parse_contract_clauses(text))
    review_ids = [row.item_id for row in get_checklist_for_resolution(routed_type, domain)]

    assert model["type"] == expected_type
    assert model["dominant_family"] == expected_family
    assert resolution.selected_domain == expected_family
    if expected_type == "GOODS_SALE":
        assert routed_type == "sale"
        assert any(item_id.startswith("S-") for item_id in review_ids)
        assert not any(item_id.startswith("TC-") for item_id in review_ids)
    elif expected_type == "TECHNOLOGY_LICENSE":
        assert routed_type == "technology"
        assert any(item_id.startswith("TC-") for item_id in review_ids)
        assert not any(item_id.startswith("S-") for item_id in review_ids)
    else:
        assert routed_type == "technology"
        assert domain.transaction_structure == "MIXED_SALE_TECHNOLOGY_LICENSE"
        assert any(item_id.startswith("TC-") for item_id in review_ids)
        assert any(item_id.startswith("S-") for item_id in review_ids)
