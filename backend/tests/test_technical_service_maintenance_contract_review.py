import json
from pathlib import Path

import pytest
from docx import Document

from review_v2.engine import (
    COVERAGE_NOT_APPLICABLE,
    _technology_third_party_context,
    review_contract_v2,
)
from review_v2.user_profile import UserRiskProfile


ROOT = Path(__file__).resolve().parents[2]
FIXTURE = ROOT / "test_data" / "17 技术服务.docx"


def _document_text(path: Path = FIXTURE) -> str:
    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = [
        "\t".join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
    ]
    return "\n".join(paragraphs + tables)


@pytest.fixture(scope="module")
def maintenance_review() -> dict:
    return review_contract_v2(
        _document_text(),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方 / 委托人（买方）"),
    )


def test_maintenance_contract_context_and_task_role_are_preserved(maintenance_review):
    assert maintenance_review["contract_context"]["contract_subtype"] == "TECHNICAL_SERVICE"
    assert maintenance_review["perspective_resolution"]["display_role"] == "甲方 / 委托方"
    assert maintenance_review["report_context"]["contract_type_label"] == "技术合同—技术服务合同"
    assert maintenance_review["release_quality_gate"]["status"] == "REPORT_READY"


def test_maintenance_contract_uses_service_only_applicability(maintenance_review):
    items = {row["item_id"]: row for row in maintenance_review["items"]}
    assert items["TC-01-04"]["coverage_status"] == COVERAGE_NOT_APPLICABLE
    assert items["TC-03-04"]["coverage_status"] == COVERAGE_NOT_APPLICABLE
    assert all(items[f"TC-08-0{index}"]["coverage_status"] == COVERAGE_NOT_APPLICABLE for index in range(1, 6))
    assert all(items[f"TC-09-0{index}"]["coverage_status"] == "通过" for index in range(1, 5))
    assert all(items[f"TC-10-0{index}"]["coverage_status"] == COVERAGE_NOT_APPLICABLE for index in range(1, 5))
    assert items["TC-11-01"]["coverage_status"] == "部分覆盖"
    assert items["TC-11-02"]["coverage_status"] == "部分覆盖"
    assert items["TC-11-03"]["coverage_status"] == "缺失"
    assert items["TC-11-04"]["coverage_status"] == "部分覆盖"
    assert all(items[f"TC-12-0{index}"]["coverage_status"] == COVERAGE_NOT_APPLICABLE for index in range(1, 5))
    assert items["TC-13-01"]["coverage_status"] == "通过"
    assert items["TC-13-02"]["coverage_status"] == COVERAGE_NOT_APPLICABLE
    assert items["TC-13-03"]["coverage_status"] == COVERAGE_NOT_APPLICABLE


def test_service_fee_payment_uses_article_four_not_breach_clause(maintenance_review):
    item = next(row for row in maintenance_review["items"] if row["item_id"] == "TC-09-04")
    evidence = item["evidence_clauses"]
    assert evidence
    assert all("第四条" in row["location"] for row in evidence)
    assert not any("第九条" in row["location"] for row in evidence)
    assert "技术服务费总额" in evidence[0]["raw_text"]
    assert "支付" in evidence[1]["raw_text"]


def test_service_confidentiality_has_precise_partial_evidence(maintenance_review):
    items = {row["item_id"]: row for row in maintenance_review["items"]}
    liability = items["TC-11-04"]
    assert liability["coverage_status"] == "部分覆盖"
    assert liability["risk_level"] == "中"
    assert "违反保密义务" in liability["evidence_clauses"][0]["raw_text"]
    assert "实际损失" in liability["evidence_clauses"][0]["raw_text"]
    assert "第五条" in liability["evidence_clauses"][0]["location"]


def test_freeze_closure_uses_subtype_safe_wording_and_direct_evidence(maintenance_review):
    items = {row["item_id"]: row for row in maintenance_review["items"]}

    confidentiality = items["TC-11-02"]
    confidentiality_text = json.dumps(confidentiality, ensure_ascii=False)
    assert "非公开咨询方法" not in confidentiality_text
    assert all(term in confidentiality_text for term in (
        "系统架构", "源代码", "服务器信息", "数据库资料", "接口资料", "业务数据", "客户数据",
    ))

    nature = items["TC-01-01"]
    subtype = items["TC-01-02"]
    assert nature["coverage_status"] == "通过"
    assert subtype["coverage_status"] == "通过"
    assert all("第一条" in row["location"] for row in nature["evidence_clauses"])
    assert all("第一条" in row["location"] for row in subtype["evidence_clauses"])
    assert not any("第九条" in row["location"] for row in [*nature["evidence_clauses"], *subtype["evidence_clauses"]])
    assert all(row["semantic_support_level"] == "DIRECT" for row in nature["evidence_clauses"])
    assert "系统开发" in subtype["evidence_clauses"][0]["raw_text"]
    assert "重大功能新增" in subtype["evidence_clauses"][0]["raw_text"]
    assert "整体系统重构" in subtype["evidence_clauses"][0]["raw_text"]

    attachment = items["TC-13-01"]
    locations = [row["location"] for row in attachment["evidence_clauses"]]
    assert attachment["coverage_status"] == "通过"
    assert locations[:6] == [
        "第十四条 与履行本合同有关的下列技术文件，经双方书面或电子邮件确认后，为本合同的组成部分",
        "第十四条第1项", "第十四条第2项", "第十四条第3项", "第十四条第4项", "第十四条第6项",
    ]
    assert all(row["semantic_support_level"] == "DIRECT" for row in attachment["evidence_clauses"][:6])
    assert all(row["semantic_support_level"] == "SUPPORTING" for row in attachment["evidence_clauses"][6:])


def test_maintenance_contract_summary_is_recomputed_after_applicability(maintenance_review):
    risks = maintenance_review["risk_finding_summary"]
    dispositions = maintenance_review["final_disposition_summary"]
    assert dispositions["total"] == 4
    assert risks["重大风险"] == 0
    assert risks["中等风险"] == 2
    assert risks["低风险"] == 2
    assert dispositions["actionable"] == 2
    assert dispositions["optional_improvement"] == 2
    assert dispositions["manual_review"] == 0


@pytest.mark.parametrize(
    ("clause", "expected_actor", "expected_ip", "expected_status"),
    [
        ("未经甲方同意，不得将服务转委托第三方。", True, False, COVERAGE_NOT_APPLICABLE),
        ("卖方履行服务需要使用第三方数据库组件，并保证已取得合法授权。", False, True, "缺失"),
    ],
)
def test_third_party_actor_does_not_imply_third_party_ip(
    clause,
    expected_actor,
    expected_ip,
    expected_status,
):
    text = f"""技术服务合同
甲方委托乙方提供系统维护服务。
服务内容为故障排查、运行监控和技术支持。
服务费为人民币10000元，验收后支付。
{clause}
"""
    context = _technology_third_party_context(text)
    assert context["third_party_actor"] is expected_actor
    assert context["third_party_ip_trigger"] is expected_ip

    output = review_contract_v2(
        text,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方 / 委托人"),
    )
    items = {row["item_id"]: row for row in output["items"]}
    for review_id in ("TC-03-04", "TC-12-02", "TC-13-02"):
        assert items[review_id]["coverage_status"] == expected_status
