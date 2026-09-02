from functools import lru_cache
from pathlib import Path

from docx import Document

from review_v2.engine import COVERAGE_NOT_APPLICABLE, review_contract_v2
from review_v2.technology_service_analysis import (
    analyze_payment_schedule,
    analyze_technical_service_core_facts,
    analyze_technical_service_risks,
    analyze_timeline,
    classify_technology_transaction,
)
from review_v2.user_profile import UserRiskProfile


ROOT = Path(__file__).resolve().parents[2]
FIXTURE = ROOT / "test_data" / "32 技术服务.docx"


@lru_cache(maxsize=1)
def _text() -> str:
    document = Document(str(FIXTURE))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = [
        "\t".join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
    ]
    return "\n".join(paragraphs + tables)


@lru_cache(maxsize=1)
def _review() -> dict:
    return review_contract_v2(
        _text(),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方", contract_scenario="技术服务"),
    )


def _items() -> dict:
    return {row["item_id"]: row for row in _review()["items"]}


def test_service_route_does_not_activate_entrusted_development_from_incidental_words():
    classification = classify_technology_transaction(_text())
    assert classification["primary_subtype"] == "TECHNICAL_SERVICE"
    assert "COMMISSIONED_DEVELOPMENT" not in classification.get("secondary_modules", [])

    items = _items()
    assert all(items[f"TC-04-0{index}"]["coverage_status"] == COVERAGE_NOT_APPLICABLE for index in range(1, 5))


def test_competition_and_improvement_facts_are_shared_across_rules():
    items = _items()
    competition = items["TC-02-01"]
    boundary = items["TC-10-04"]
    assert competition["coverage_status"] == "风险条款"
    assert boundary["coverage_status"] == "风险条款"
    assert any("不得自行开发" in (row.get("text") or "") for row in competition["evidence_clauses"])
    assert any("不得采购" in (row.get("text") or "") for row in boundary["evidence_clauses"])
    assert {"TC-02-01", "TC-10-04"} <= {
        review_id
        for finding in _review()["risk_findings"]
        for review_id in finding["source_review_ids"]
    }


def test_third_party_claim_uses_article_facts_and_only_marks_real_gaps():
    item = _items()["TC-12-02"]
    statuses = {row["element_id"]: row["status"] for row in item["coverage_elements"]}
    for element_id in (
        "claim_handling_responsibility",
        "defense_cost_allocation",
        "third_party_compensation",
        "replacement_remedy",
        "license_cure",
    ):
        assert statuses[element_id] == "COVERED"
    for element_id in ("notification_deadline", "lawyer_selection", "settlement_authority"):
        assert statuses[element_id] == "MISSING"
    assert statuses["defense_control"] == "PARTIALLY_COVERED"
    defense = next(row for row in item["coverage_elements"] if row["element_id"] == "defense_control")
    assert "负责处理全部争议" in defense["proposition"]
    assert "基础处理主导安排" in defense["proposition"]
    assert item["coverage_status"] == "部分覆盖"
    evidence_text = " ".join(row.get("text") or "" for row in item["evidence_clauses"])
    assert "由乙方负责处理全部争议" in evidence_text
    assert "提供替代方案" in evidence_text


def test_acceptance_separates_test_control_authority_from_concrete_parameters():
    item = _items()["TC-13-03"]
    statuses = {row["element_id"]: row["status"] for row in item["coverage_elements"]}
    assert statuses["quantitative_performance_targets"] == "COVERED"
    assert statuses["test_environment"] == "MISSING"
    assert statuses["test_equipment"] == "MISSING"
    assert statuses["test_dataset"] == "MISSING"
    assert statuses["test_method"] == "MISSING"
    assert statuses["acceptance_dispute_retest"] == "COVERED"
    assert statuses["actual_test_attachment"] == "REFERENCED_BUT_NOT_PROVIDED"
    assert item["coverage_status"] == "部分覆盖"
    semantics = item["decision_trace"]["test_parameter_semantics"]
    assert semantics["determination_authority"] == "甲方"
    assert semantics["authority_statement_is_not_parameter_content"] is True
    assert not any(semantics["specific_content"].values())
    evidence_text = " ".join(row.get("text") or "" for row in item["evidence_clauses"])
    assert "SOC估算平均误差不高于3%" in evidence_text
    assert "测试数据、测试设备、测试环境和测试方法" in evidence_text
    assert "第三方检测机构进行复测" in evidence_text


def test_triggered_improvement_notice_rule_cannot_remain_not_applicable():
    item = _items()["TC-10-03"]
    assert item["applicability"]["applicable"] is True
    assert item["applicability"]["status"] == "APPLICABLE"
    assert item["coverage_status"] != COVERAGE_NOT_APPLICABLE
    assert item["coverage_status"] == "缺失"
    assert item["decision_trace"]["applicability_status_invariant"]["status"] == "REPAIRED"
    assert "未触发对应适用条件" not in item["risk_point_description"]
    contradictions = [
        row["item_id"] for row in _review()["items"]
        if row["coverage_status"] == COVERAGE_NOT_APPLICABLE
        and row.get("applicability", {}).get("applicable") is True
        and str(row.get("applicability", {}).get("status") or "").upper() in {
            "APPLICABLE", "CONDITIONALLY_APPLICABLE", "CONDITIONAL", "TRIGGERED",
        }
    ]
    assert contradictions == []


def test_confidentiality_is_split_by_obligated_party_for_recipient_review():
    items = _items()
    period = items["TC-11-02"]
    liability = items["TC-11-04"]
    assert period["coverage_status"] == "风险条款"
    assert "终止后" in period["risk_point_description"] and "甲方" in period["risk_point_description"]
    assert liability["coverage_status"] == "通过"
    assert "第九条一般违约责任" in liability["risk_point_description"]
    assert "违约金及不足部分继续赔偿" in liability["risk_point_description"]
    elements = {row["element_id"]: row for row in liability["coverage_elements"]}
    assert elements["party_a_breach_liability"]["status"] == "COVERED"
    assert elements["party_b_breach_liability"]["status"] == "COVERED"
    trace = liability["decision_trace"]["cross_clause_obligation_resolver"]
    assert trace["status"] == "MERGED"
    assert trace["superseded_vague_special_clause"] is True
    assert "liquidated_damages" in trace["party_b"]["remedy_types"]
    assert "damage_compensation" in trace["party_b"]["remedy_types"]
    evidence_text = " ".join(row.get("text") or "" for row in liability["evidence_clauses"])
    assert "乙方违反保密、知识产权或者数据使用约定" in evidence_text
    assert not any("TC-11-04" in row["source_review_ids"] for row in _review()["risk_findings"])


def test_party_direction_and_finding_consistency_are_enforced():
    output = _review()
    acceptance = next(
        row for row in output["risk_findings"]
        if row["risk_group"] == "technical_acceptance_payment_control"
    )
    assert acceptance["severity"] != "HIGH"
    assert "乙方完成服务后仍可能无法验收和收款" not in acceptance["issue"]
    assert "甲方" in acceptance["issue"] and any(term in acceptance["issue"] for term in ("可执行", "争议"))

    subject = _items()["TC-01-03"]
    assert subject["coverage_status"] == "通过"
    assert not any("TC-01-03" in row["source_review_ids"] for row in output["risk_findings"])
    assert all(row["recommendation"].strip() for row in output["risk_findings"])


def test_timeline_and_deliverables_are_shared_deterministic_facts():
    timeline = analyze_timeline(_text())
    assert timeline["service_start_date"] == "2026-10-01"
    assert timeline["service_end_date"] == "2027-06-30"
    assert any(row["type"] == "MILESTONE_AFTER_PERFORMANCE_END" for row in timeline["conflicts"])

    facts = analyze_technical_service_core_facts(_text())
    assert {
        "技术优化报告", "算法参数文件", "测试记录", "现场技术服务", "现场技术培训", "技术咨询支持",
    } <= set(facts["technical_materials"]["actual_deliverables"])
    material = _items()["TC-13-01"]
    assert all(name in material["risk_point_description"] for name in ("技术优化报告", "算法参数文件", "测试记录"))


def test_report_statistics_are_recomputed_from_final_findings_and_themes():
    output = _review()
    findings = output["risk_findings"]
    themes = output.get("risk_themes") or []
    assert output["finding_count"] == len(findings)
    assert output["risk_theme_count"] == len(themes)
    theme_counts = output["risk_theme_severity_summary"]
    assert sum(theme_counts[level] for level in ("major", "medium", "low")) == len(themes)
    assert output["risk_theme_summary"]["total"] == len(themes)
    assert output["manual_review_summary"]["manual_review_count"] == len(output["manual_review_items"])
    assert output["release_quality_gate"]["status"] == "REPORT_READY"
    assert output["observability"]["production_token_scan"]["forbidden_internal_tokens_found"] == []


def test_amount_only_service_payment_nodes_remain_deterministic():
    payment = analyze_payment_schedule(_text())
    assert payment["contract_total"] == 1_200_000
    assert payment["amount_total"] == 1_400_000
    assert payment["amount_delta"] == 200_000
    assert payment["percentage_total"] == 116.66666666666667
    assert payment["status"] == "INCONSISTENT"
    assert [row["amount"] for row in payment["nodes"]] == [500_000, 400_000, 300_000, 200_000]


def test_counterparty_restrictions_are_projected_as_current_party_enforceability_risk():
    items = _items()
    for review_id in ("TC-02-01", "TC-10-01", "TC-10-02", "TC-10-04"):
        item = items[review_id]
        trace = item["decision_trace"]["risk_to_current_party"]
        assert trace["review_role"] == "SERVICE_RECIPIENT"
        assert trace["current_party_directly_restricted"] is False
        assert trace["risk_scope"] == "CURRENT_PARTY_ENFORCEABILITY_AND_VALIDITY"
        combined = " ".join(filter(None, (
            item.get("risk_point_description"),
            item.get("legal_consequence_if_missing"),
            item.get("suggested_clause_example"),
        )))
        assert "甲方被限制" not in combined
        assert "被强制转让方" not in combined
        assert any(term in combined for term in ("可执行", "效力", "争议"))
        assert item["decision_trace"]["recommendation_direction_check"]["preserves_current_party_control"] is True


def test_positive_evidence_suppresses_conflicting_absence_search():
    output = _review()
    competition = _items()["TC-02-01"]
    positive_ids = {
        row["element_id"]
        for row in competition["coverage_elements"]
        if row["status"] == "ADVERSE" and row.get("evidence_ids")
    }
    assert {
        "independent_research_restriction",
        "alternative_technology_restriction",
        "post_termination_competition_term",
    } <= positive_ids
    finding = next(row for row in output["risk_findings"] if "TC-02-01" in row["source_review_ids"])
    bundle = finding.get("evidence_bundle") or {}
    search_text = " ".join(
        str(value)
        for row in bundle.get("items") or []
        if row.get("source_type") == "SEARCH_AUDIT"
        for value in row.get("missing_elements") or []
    )
    assert "自行开发相同或类似" not in search_text
    assert "采购或使用替代技术" not in search_text
    assert "终止后长期持续" not in search_text


def test_third_party_claim_consequence_and_advice_only_consume_residual_gaps():
    item = _items()["TC-12-02"]
    residual = item["decision_trace"]["residual_uncovered_elements"]
    assert set(residual["element_ids"]) == {
        "notification_requirement", "notification_deadline", "defense_control",
        "lawyer_selection", "settlement_authority", "termination_remedy", "refund_remedy",
    }
    advice = item["suggested_clause_example"]
    assert "通知期限" in advice and "抗辩控制" in advice and "解除权" in advice
    for already_covered in ("必要协助", "责任承担", "替代方案", "取得第三方授权"):
        assert already_covered not in advice


def test_referenced_but_unprovided_attachment_is_not_an_ordinary_missing_clause():
    item = _items()["TC-13-03"]
    attachment = next(
        row for row in item["coverage_elements"]
        if row["element_id"] == "actual_test_attachment"
    )
    assert attachment["status"] == "REFERENCED_BUT_NOT_PROVIDED"
    assert attachment["document_state"] == "REFERENCED_BUT_NOT_PROVIDED"
    assert attachment["external_material_required"] is True
    assert item["decision_trace"]["attachment_material_state"] == {
        "arranged_in_contract": True,
        "provided_in_review_materials": False,
        "status": "REFERENCED_BUT_NOT_PROVIDED",
        "ordinary_missing_clause": False,
    }
    assert "补充提供并核验已约定的验收测试附件正文" in item["suggested_clause_example"]
    assert "第三方复测机构" not in item["suggested_clause_example"]


def test_mixed_deliverables_have_typed_facts_and_type_specific_advice():
    facts = analyze_technical_service_core_facts(_text())
    deliverables = {
        row["name"]: row for row in facts["technical_materials"]["deliverable_facts"]
    }
    assert deliverables["技术优化报告"]["type"] == "DOCUMENT_OR_REPORT"
    assert deliverables["算法参数文件"]["type"] == "DATA_OR_PARAMETER_FILE"
    assert deliverables["测试记录"]["type"] == "TEST_RECORD"
    assert deliverables["现场技术服务"]["type"] == "ONSITE_SERVICE"
    assert deliverables["现场技术培训"]["type"] == "TRAINING"
    assert deliverables["技术咨询支持"]["type"] == "SUPPORT_OR_CONSULTING"
    assert deliverables["技术优化报告"]["document_title"] == "电池管理系统技术优化报告"
    assert deliverables["现场技术培训"]["quantity"] == "两次"
    assert deliverables["技术咨询支持"]["support_duration"] == "项目结束后六个月"
    assert "江苏省苏州市" in deliverables["现场技术服务"]["service_location"]
    advice = _items()["TC-13-01"]["suggested_clause_example"]
    assert "现场技术服务" in advice and "服务人员、工时或持续时间" in advice
    assert "现场技术培训" in advice and "培训对象、培训时长" in advice
    assert "技术咨询支持" in advice and "响应时间、服务渠道" in advice
    assert "培训次数" not in advice
    assert "支持期限" not in advice
    assert "文件名称" not in advice
    assert "服务地点" not in advice
    elements = {row["element_id"]: row for row in _items()["TC-13-01"]["coverage_elements"]}
    assert "培训次数" in elements["training_completion_details"]["covered_fields"]
    assert "支持期限" in elements["support_consulting_details"]["covered_fields"]
    assert "服务地点" in elements["onsite_service_completion_details"]["covered_fields"]
    assert "文件名称" in elements["document_delivery_details"]["covered_fields"]
    assert "现场技术服务的文件名称、版本" not in advice
    assert "现场技术培训的文件名称、版本" not in advice


def test_na_reason_preserves_confirmed_confidential_information_fact():
    item = _items()["TC-12-03"]
    assert item["coverage_status"] == COVERAGE_NOT_APPLICABLE
    reason = item["applicability"]["reason"]
    assert "虽涉及技术秘密或保密资料" in reason
    assert "未形成可识别的技术秘密范围" not in reason
    assert item["coverage_elements"] == []
    assert item["decision_trace"]["applicability_consistency"]["status"] == "PASS"


def test_improvement_rules_have_distinct_atomic_responsibilities_but_may_aggregate():
    items = _items()
    ownership = items["TC-10-01"]
    sharing = items["TC-10-02"]
    boundary = items["TC-10-04"]
    assert len({ownership["title"], sharing["title"], boundary["title"]}) == 3
    ownership_ids = {row["element_id"] for row in ownership["coverage_elements"]}
    sharing_ids = {row["element_id"] for row in sharing["coverage_elements"]}
    boundary_ids = {row["element_id"] for row in boundary["coverage_elements"]}
    assert "compulsory_free_improvement_assignment" in ownership_ids
    assert "improvement_use_scope" in sharing_ids
    assert "alternative_technology_restriction" in boundary_ids
    assert not (ownership_ids == sharing_ids == boundary_ids)
    finding = next(
        row for row in _review()["risk_findings"]
        if {"TC-10-01", "TC-10-02", "TC-10-04"} <= set(row["source_review_ids"])
    )
    assert finding["risk_group"] == "improvement_no_unlawful_restriction"


def test_report_actual_exposure_uses_party_aware_finding_consequence():
    report = _review()["result"]
    assert "该条款对当前审查方形成不利或限制性安排" not in report
    assert "过宽限制可能在争议中被限缩适用或认定部分条款难以执行" in report
    assert "影响甲方最终取得权利的稳定性和可执行性" in report


def test_attachment_finding_only_describes_the_unprovided_referenced_material():
    finding = next(
        row for row in _review()["risk_findings"]
        if "TC-13-03" in row["source_review_ids"]
    )
    assert "附件正文未随本次审查材料提供" in finding["issue"]
    assert "所引用材料未纳入本次审查范围" in finding["consequence"]
    assert "缺少可直接审查的材料基础" in finding["consequence"]
    assert "第三方复测机构" not in finding["consequence"]
    assert "补充提供并核验已约定的验收测试附件正文" in finding["recommendation"]


def test_structural_pass_is_linked_to_higher_priority_substantive_risk():
    items = _items()
    ownership = items["TC-09-03"]
    payment = items["TC-09-04"]
    assert ownership["coverage_status"] == "通过"
    assert payment["coverage_status"] == "通过"
    assert "结构已覆盖" in ownership["action_suggestion"]
    assert "TC-10-01" in ownership["action_suggestion"]
    assert "TC-12-04" not in ownership["action_suggestion"]
    assert "无需修改" not in ownership["action_suggestion"]
    assert "结构已覆盖" in payment["action_suggestion"]
    assert "GEN-PAY-01" in payment["action_suggestion"]
    assert "无需修改" not in payment["action_suggestion"]


def test_ip_liability_and_general_remedies_jointly_trigger_tc1204():
    item = _items()["TC-12-04"]
    assert item["applicability"]["status"] == "APPLICABLE"
    trace = item["decision_trace"]["service_remedy_coordination_applicability"]
    assert trace["ip_specific_liability"] is True
    assert trace["general_contract_remedies"] is True
    statuses = {row["element_id"]: row["status"] for row in item["coverage_elements"]}
    assert statuses["infringement_specific_liability"] == "COVERED"
    assert statuses["general_contract_remedies"] == "COVERED"
    assert statuses["special_general_liability_coordination"] == "PARTIALLY_COVERED"
    assert item["risk_level"] == "低"
    assert item["action_priority"] == "C_OPTIMIZATION"
    trace = item["decision_trace"]["service_remedy_coordination_applicability"]
    assert trace["coordination_grade"] == "IMPLIED_PARTIAL_COORDINATION"
    assert trace["actual_conflict"] is False
    finding = next(row for row in _review()["risk_findings"] if "TC-12-04" in row["source_review_ids"])
    assert finding["severity"] == "LOW"
    assert finding["priority"] == "C_OPTIMIZATION"


def test_major_technology_restrictions_include_specific_legal_authorities():
    items = _items()
    for review_id in ("TC-02-01", "TC-10-01", "TC-10-04"):
        basis = items[review_id]["legal_basis_index"]
        assert "民法典》第850条" in basis
        assert "技术合同纠纷案件适用法律若干问题的解释》第10条" in basis
    report = _review()["result"]
    assert "《中华人民共和国民法典》第850条" in report
    assert "《最高人民法院关于审理技术合同纠纷案件适用法律若干问题的解释》第10条" in report


def test_gross_negligence_liability_cap_triggers_validity_review_without_voiding_contract():
    signals = analyze_technical_service_risks(_text())
    validity = signals["liability_cap_validity"]
    assert validity["present"] is True
    assert validity["analysis"]["gross_negligence_within_cap"] is True
    assert validity["analysis"]["intent_within_cap"] is False
    assert "重大过失" in validity["analysis"]["contract_fact"]
    assert "故意" not in validity["analysis"]["contract_fact"]
    assert validity["analysis"]["legal_effect_is_not_deterministically_invalid"] is True

    item = _items()["GEN-LIAB-02"]
    assert item["coverage_status"] == "风险条款"
    assert "重大过失" in item["risk_point_description"]
    assert "故意" not in item["risk_point_description"]
    assert "民法典》第506条" in item["legal_basis_index"]
    layers = item["decision_trace"]["contract_fact_legal_evaluation"]
    assert "故意" not in layers["contract_fact"]
    assert "故意或者重大过失" in layers["legal_rule"]
    finding = next(row for row in _review()["risk_findings"] if "GEN-LIAB-02" in row["source_review_ids"])
    assert "不当然影响整份合同效力" in finding["consequence"]
    assert "故意、重大过失" in finding["recommendation"]


def test_excessive_liquidated_damages_are_flagged_as_adjustment_risk_not_invalidity():
    signals = analyze_technical_service_risks(_text())
    penalty = signals["potential_excessive_liquidated_damages"]
    assert penalty["present"] is True
    analysis = penalty["analysis"]
    assert any(row["rate_percent_per_day"] == 0.5 for row in analysis["high_daily_penalties"])
    assert any(row["amount"] == 3_000_000 for row in analysis["high_fixed_penalties"])
    assert analysis["contract_total"] == 1_200_000
    assert analysis["contract_price_ratio_is_not_actual_loss_ratio"] is True

    item = _items()["GEN-LIAB-03"]
    assert "每日0.5%" in item["risk_point_description"]
    assert "300万元" in item["risk_point_description"]
    assert "实际损失比例" in item["risk_point_description"]
    assert "民法典》第585条" in item["legal_basis_index"]
    assert "合同编通则若干问题的解释》第65条" in item["legal_basis_index"]
    finding = next(row for row in _review()["risk_findings"] if "GEN-LIAB-03" in row["source_review_ids"])
    combined = " ".join((finding["issue"], finding["consequence"], finding["recommendation"]))
    assert "司法调减" in combined
    assert "必然无效" not in combined


def test_legal_basis_is_routed_by_service_scene_and_propagated_consistently():
    output = _review()
    items = _items()
    assert "第875条" not in items["TC-10-01"]["legal_basis_index"]
    assert "第885条" in items["TC-10-01"]["legal_basis_index"]
    assert "第875条" not in items["TC-10-03"]["legal_basis_index"]
    assert "第874条" not in items["TC-12-02"]["legal_basis_index"]
    assert "第874条" not in items["TC-12-04"]["legal_basis_index"]
    assert all(value in items["TC-09-02"]["legal_basis_index"] for value in ("第845条", "第509条"))
    assert "第850条" in items["TC-10-04"]["legal_basis_index"]
    assert "第10条" in items["TC-10-04"]["legal_basis_index"]

    penalty_finding = next(
        row for row in output["risk_findings"]
        if row["risk_group"] == "generic_potential_excessive_liquidated_damages"
    )
    basis = "；".join(row["text"] for row in penalty_finding["legal_basis"])
    assert "第585条" in basis and "第65条" in basis
    assert "《中华人民共和国民法典》第585条" in output["result"]
    assert "合同编通则若干问题的解释》第65条" in output["result"]


def test_excessive_penalties_preserve_obligor_and_review_party_direction():
    item = _items()["GEN-LIAB-03"]
    direction = item["decision_trace"]["liquidated_damages_direction"]
    current_rows = direction["current_party_obligations"]
    counterparty_rows = direction["counterparty_obligations"]
    assert any(row["obligor"] == "甲方" and row["amount"] == 2_000_000 for row in current_rows)
    assert any(row["obligor"] == "乙方" and row.get("amount") == 3_000_000 for row in counterparty_rows)
    assert all(row["direction_to_review_party"] == "DIRECT_ECONOMIC_EXPOSURE" for row in current_rows)
    assert all(row["direction_to_review_party"] == "FAVORABLE_REMEDY_WITH_ENFORCEABILITY_RISK" for row in counterparty_rows)
    assert "甲方" in item["risk_point_description"] and "200万元" in item["risk_point_description"]
    assert "乙方" in item["risk_point_description"] and "300万元" in item["risk_point_description"]
    assert "增强当前审查方救济" in item["legal_consequence_if_missing"]
    assert "不等同于实际损失比例" in item["risk_point_description"]


def test_valid_cap_exception_and_ordinary_penalty_do_not_trigger_validity_risks():
    text = (
        "合同总价为人民币100万元。"
        "任何一方承担的赔偿责任最高不超过合同总价，但故意、重大过失及依法不得限责的情形不适用该上限。"
        "一方迟延履行的，每逾期一日按应付未付金额的0.05%支付违约金。"
        "发生其他一般违约时支付人民币10万元违约金。"
    )
    signals = analyze_technical_service_risks(text)
    assert signals["liability_cap_validity"]["present"] is False
    assert signals["potential_excessive_liquidated_damages"]["present"] is False
