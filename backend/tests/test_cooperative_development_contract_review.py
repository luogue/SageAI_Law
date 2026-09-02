from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from review_v2 import engine
from review_v2.engine import review_contract_v2
from review_v2.technology_service_analysis import classify_technology_transaction
from review_v2.user_profile import UserRiskProfile


ROOT = Path(__file__).resolve().parents[2]
FIXTURE = ROOT / "test_data" / "11、合作.docx"
MINIMAL_FIXTURE = ROOT / "test_data" / "15 合作开发.docx"


def _fixture_text() -> str:
    document = Document(str(FIXTURE))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = ["\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    return "\n".join(paragraphs + tables)


def _document_text(path: Path) -> str:
    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = ["\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    return "\n".join(paragraphs + tables)


@pytest.fixture(scope="module")
def cooperative_review() -> dict:
    return review_contract_v2(
        _fixture_text(),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/合作开发方"),
    )


def test_cooperative_development_classification(cooperative_review: dict):
    classification = cooperative_review["deterministic_analysis"]["technology_classification"]

    assert classification["primary_category"] == "TECHNICAL_CONTRACT"
    assert classification["primary_subtype"] == "COOPERATIVE_DEVELOPMENT"
    assert classification["development_mode"] == "COOPERATIVE_DEVELOPMENT"
    assert classification["confidence"] == "high"
    assert set(classification["classification_evidence"]) >= {
        "bilateral_research_investment",
        "bilateral_research_tasks",
        "joint_development_results",
        "joint_development_risk",
    }
    assert cooperative_review["report_context"]["contract_type_label"] == "技术合同—技术开发合同—合作开发合同"
    assert cooperative_review["contract_context"]["user_role"] == "cooperative_developer"
    assert cooperative_review["contract_context"]["counterparty_role"] == "cooperative_developer"
    assert cooperative_review["contract_context"]["perspective_resolution"]["display_role"] == "甲方 / 合作开发方"

    by_id = {row["item_id"]: row for row in cooperative_review["items"]}
    assert all(by_id[f"TC-05-0{index}"]["coverage_status"] != "不适用" for index in range(1, 5))
    assert all(by_id[f"TC-04-0{index}"]["coverage_status"] == "不适用" for index in range(1, 5))
    assert cooperative_review["release_quality_gate"]["status"] == "REPORT_READY"
    assert cooperative_review["result"].startswith("# 合作开发合同审查报告")


def test_commissioned_vs_cooperative_development():
    commissioned = classify_technology_transaction(
        "技术开发合同。甲方委托乙方研究开发智能平台，甲方支付研究开发经费100万元，"
        "乙方负责全部研究开发工作并向甲方交付成果。"
    )
    cooperative = classify_technology_transaction(
        "甲方投入研发人员和研发费用并负责算法模块开发；乙方投入研发人员和研发费用并负责接口模块开发。"
        "双方共同组织验收，共同享有开发成果，并各自承担正常技术风险下已经发生的研发费用。"
    )

    assert commissioned["development_mode"] == "ENTRUSTED_DEVELOPMENT"
    assert cooperative["development_mode"] == "COOPERATIVE_DEVELOPMENT"
    assert "bilateral_research_tasks" in cooperative["classification_evidence"]


def test_cooperative_role_rejects_stale_patent_transfer_identity():
    output = review_contract_v2(
        _fixture_text(),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方 / 专利权受让方"),
    )
    resolution = output["perspective_resolution"]

    assert resolution["contract_subtype"] == "COOPERATIVE_DEVELOPMENT"
    assert resolution["canonical_role"] == "COOPERATIVE_DEVELOPER"
    assert resolution["party_side"] == "PARTY_A"
    assert resolution["review_party_identity"] == "PARTY_A"
    assert resolution["review_party_label"] == "甲方"
    assert resolution["transaction_role"] == "COOPERATIVE_DEVELOPER"
    assert resolution["display_role"] == "甲方 / 合作开发方"
    assert output["contract_context"]["review_party_identity"] == "PARTY_A"
    assert output["contract_context"]["review_party_label"] == "甲方"
    assert output["contract_context"]["transaction_role"] == "COOPERATIVE_DEVELOPER"
    assert "PATENT_TRANSFEREE" in resolution["rejected_roles"]
    assert "审查立场：甲方 / 合作开发方" in output["result"]
    assert "审查立场：甲方 / 专利权受让方" not in output["result"]


def test_asymmetric_liability_party_effect_changes_with_review_party(cooperative_review: dict):
    party_a = next(
        row for row in cooperative_review["risk_findings"]
        if row["risk_group"] == "asymmetric_breach_liability"
    )
    assert party_a["issue_type"] == "ASYMMETRIC"
    assert party_a["party_effect"] == "FAVORABLE"
    assert all(term in party_a["consequence"] for term in ("当前审查方为甲方", "100000元", "300000元", "相对有利于甲方"))
    assert "不应无依据扩大当前审查方责任" in party_a["recommendation"]

    report = cooperative_review["result"]
    section_start = report.index("双方违约责任明显不对称（TECH-GEN-01）")
    section_end = report.find("\n## ", section_start)
    asymmetry_section = report[section_start:section_end if section_end >= 0 else None]
    assert all(term in asymmetry_section for term in ("当前审查方为甲方", "100000元", "300000元", "相对有利于甲方"))
    assert "该风险分配对当前审查方不利" not in asymmetry_section
    assert "具体利弊取决于当前审查方" not in asymmetry_section

    restriction = next(
        row for row in cooperative_review["risk_findings"]
        if row["risk_group"] == "technology_restrictive_covenant"
    )
    assert restriction["party_effect"] == "MIXED"
    restriction_start = report.index("竞业限制范围和责任偏重（TECH-GEN-03）")
    restriction_end = report.find("\n## ", restriction_start)
    restriction_section = report[restriction_start:restriction_end if restriction_end >= 0 else None]
    assert "同时作用于双方" in restriction_section
    assert "仅对当前审查方单方面不利" in restriction_section

    party_b_output = review_contract_v2(
        _fixture_text(),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="乙方"),
    )
    party_b = next(
        row for row in party_b_output["risk_findings"]
        if row["risk_group"] == "asymmetric_breach_liability"
    )
    assert party_b_output["perspective_resolution"]["party_side"] == "PARTY_B"
    assert party_b_output["perspective_resolution"]["review_party_identity"] == "PARTY_B"
    assert party_b_output["perspective_resolution"]["review_party_label"] == "乙方"
    assert party_b_output["perspective_resolution"]["transaction_role"] == "COOPERATIVE_DEVELOPER"
    assert "审查立场：乙方 / 合作开发方" in party_b_output["result"]
    assert party_b["issue_type"] == "ASYMMETRIC"
    assert party_b["party_effect"] == "ADVERSE"
    assert all(term in party_b["consequence"] for term in ("当前审查方为乙方", "300000元", "100000元", "更高的直接金额暴露"))


def test_unknown_party_is_the_only_unknown_party_effect_case():
    output = review_contract_v2(
        _fixture_text(),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="合作开发方"),
    )
    resolution = output["perspective_resolution"]
    finding = next(
        row for row in output["risk_findings"]
        if row["risk_group"] == "asymmetric_breach_liability"
    )

    assert resolution["review_party_identity"] == "UNKNOWN"
    assert resolution["transaction_role"] == "COOPERATIVE_DEVELOPER"
    assert finding["party_effect"] == "UNKNOWN"
    assert "具体利弊需要结合审查方确定" in output["result"]
    assert "该风险分配对当前审查方不利" not in output["result"]


def test_missing_element_recommendation_is_natural_without_rewriting_source():
    source = "未明确双方能否独立使用共同改进成果及是否需要对方同意"

    assert engine._recommendation_from_missing_elements([source]) == (
        "建议明确双方能否独立使用共同改进成果及是否需要对方同意。"
    )
    assert source == "未明确双方能否独立使用共同改进成果及是否需要对方同意"


def test_single_manual_review_does_not_abort_report(monkeypatch):
    original = engine._applicability_result_for_checkpoint

    def one_uncertain_item(checkpoint, clauses, applicable, reason, canonical_claims=None):
        result = original(checkpoint, clauses, applicable, reason, canonical_claims)
        if checkpoint.item_id == "TC-08-01":
            result.update({
                "status": "UNCERTAIN",
                "applicable": None,
                "reason": "合同未附咨询任务材料，无法确认该局部事项是否适用。",
                "machine_reason": "Missing consulting task attachment.",
                "user_reason": "合同未附咨询任务材料，需确认是否另有技术咨询任务及对应附件。",
                "depends_on": ["consulting_task_attachment"],
            })
        return result

    monkeypatch.setattr(engine, "_applicability_result_for_checkpoint", one_uncertain_item)
    output = review_contract_v2(
        _fixture_text(),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/合作开发方"),
    )

    assert output["release_quality_gate"]["status"] == "REPORT_READY"
    assert output["manual_review_summary"]["manual_review_count"] == 1
    assert output["manual_review_items"][0]["item_id"] == "TC-08-01"
    assert "需确认" in output["manual_review_items"][0]["reason"] or "无法确认" in output["manual_review_items"][0]["reason"]
    assert "正式报告已停止生成" not in output["result"]
    assert "人工复核事项" in output["result"]
    assert all(
        "TC-08-01" not in finding["source_review_ids"]
        for finding in output["risk_findings"]
    )


def test_manual_review_fact_contradiction_is_local_not_fatal():
    manual_item = SimpleNamespace(
        item_id="TC-XX-01",
        coverage_status=engine.INSUFFICIENT_CONTEXT,
        needs_manual_review=True,
        decision_trace={"fact_coverage_contradiction": {"status": "UNRESOLVED"}},
    )

    assert engine._recall_gate_failures(
        [manual_item],
        {"truncated": False, "parsed_character_ratio": 1.0},
    ) == []


def test_non_fatal_quality_gate_does_not_stop_report_generation():
    output = review_contract_v2(
        "培训服务合同。课程名称：数学培训。课程名称：英语培训。总课时：20节。"
        "上课时间：每周六上午。开课日期：2026年9月10日。预计结课日期：2027年6月30日。",
        "service",
        use_llm=False,
    )

    assert output["release_quality_gate"]["status"] == "BLOCKED"
    assert output["report_generation_decision"]["status"] == "GENERATED_WITH_QUALITY_WARNINGS"
    assert output["report_generation_decision"]["report_generated"] is True
    assert output["system_metadata"]["technology_cooperative_development"] == "FROZEN_REGRESSION"
    assert output["fatal_report_generation_gate"]["status"] == "PASS"
    assert output["result"].startswith("# ")
    assert "正式报告已停止生成" not in output["result"]


def test_only_global_input_failure_is_fatal():
    output = review_contract_v2("", "technology", use_llm=False)
    gate = output["fatal_report_generation_gate"]

    assert gate["status"] == "FATAL"
    assert gate["fatal_rule_ids"] == ["EMPTY_CONTRACT_BODY"]
    assert output["report_generation_decision"]["status"] == "STOPPED_FATAL"
    assert output["report_generation_decision"]["report_generated"] is False


def test_cooperative_rules_use_operational_facts_and_direct_evidence(cooperative_review: dict):
    by_id = {row["item_id"]: row for row in cooperative_review["items"]}

    assert by_id["TC-01-01"]["coverage_status"] == "通过"
    assert by_id["TC-01-04"]["coverage_status"] == "不适用"

    investment = by_id["TC-05-01"]
    assert investment["coverage_status"] == "通过"
    assert {row["element_id"] for row in investment["coverage_elements"]} == {
        "bilateral_research_investment",
        "investment_form_and_amount",
        "research_task_allocation",
        "budget_overrun_allocation",
        "technical_coordination",
        "major_decision_process",
    }
    investment_text = "\n".join(str(row.get("text") or "") for row in investment["evidence_clauses"])
    assert "120万元" in investment_text
    assert "100万元" in investment_text
    assert "各自超出的成本由各自承担" in investment_text

    exit_item = by_id["TC-05-04"]
    assert exit_item["coverage_status"] != "通过"
    exit_text = "\n".join(str(row.get("text") or "") for row in exit_item["evidence_clauses"])
    assert "源代码" in exit_text
    assert "自行决定是否" in exit_text
    assert "无权强制要求" in exit_text

    improvement = by_id["TC-10-02"]
    assert improvement["coverage_status"] == "部分覆盖"
    assert {row["element_id"] for row in improvement["coverage_elements"]} == {
        "improvement_ownership_boundary",
        "independent_use_right",
        "external_license_transfer_sublicense",
        "license_terms_and_revenue",
    }
    improvement_risk_text = " ".join((
        improvement["risk_point_description"],
        improvement["legal_consequence_if_missing"],
    ))
    assert "独立使用" in improvement_risk_text
    assert "许可收益" in improvement_risk_text
    assert "适用范围、履行标准和责任安排不清" not in improvement_risk_text
    assert all("第十四条" in str(row.get("location") or "") for row in improvement["evidence_clauses"])
    assert all("净利润" not in str(row.get("text") or "") for row in improvement["evidence_clauses"])

    notification = by_id["TC-10-03"]
    assert notification["coverage_status"] == "缺失"
    assert notification["needs_manual_review"] is False
    assert notification["system_confidence"]["level"] == "high"
    assert "双方未约定" in str(notification["evidence_clauses"][0].get("text") or "")
    improvement_boundary = by_id["TC-10-04"]
    assert improvement_boundary["coverage_status"] == "不适用"
    assert improvement_boundary["related_risks"] == []

    confidentiality = by_id["TC-11-04"]
    assert confidentiality["coverage_status"] == "部分覆盖"
    assert {str(row.get("location") or "").split()[0] for row in confidentiality["evidence_clauses"]} == {
        "第十条",
        "第十五条",
    }
    confidentiality_quotes = "\n".join(str(row.get("quote") or "") for row in confidentiality["evidence_clauses"])
    assert "不得向第三方披露" in confidentiality_quotes
    assert "10万元违约金" in confidentiality_quotes
    assert "30万元违约金" in confidentiality_quotes
    confidentiality_support = next(
        row["support_summary"] for row in confidentiality["evidence_clauses"]
        if row.get("evidence_role") == "CONFIDENTIALITY_OBLIGATION"
    )
    assert "保密期限" not in confidentiality_support

    claim_response = by_id["TC-12-02"]
    assert claim_response["coverage_status"] == "部分覆盖"
    assert claim_response["issue_type"] == "PARTIAL"
    assert {row["element_id"] for row in claim_response["coverage_elements"]} == {
        "claim_trigger_and_responsibility",
        "claim_notice",
        "defense_and_settlement_control",
        "defense_cost_and_compensation_scope",
        "continuity_and_replacement_remedy",
    }
    assert all(term in claim_response["risk_point_description"] for term in ("通知", "抗辩", "和解", "替代"))

    infringement = by_id["TC-12-04"]
    assert infringement["coverage_status"] == "部分覆盖"
    infringement_locations = {str(row.get("location") or "") for row in infringement["evidence_clauses"]}
    assert any("第九条" in location for location in infringement_locations)
    assert any("第十五条" in location for location in infringement_locations)

    assert by_id["TC-13-02"]["coverage_status"] == "部分覆盖"
    assert by_id["TC-13-03"]["coverage_status"] == "部分覆盖"
    assert cooperative_review["release_quality_gate"]["status"] == "REPORT_READY"


def test_cross_contract_risks_are_general_and_canonical(cooperative_review: dict):
    findings = cooperative_review["risk_findings"]
    by_finding_id = {row["finding_id"]: row for row in findings}

    expected = {
        "RF-LIABILITY-ASYMMETRY-001": "TECH-GEN-01",
        "RF-DISPUTE-CONFLICT-001": "TECH-GEN-02",
        "RF-TECH-RESTRICT-001": "TECH-GEN-03",
        "RF-TECH-OSS-001": "TECH-GEN-04",
        "RF-TECH-EXIT-001": "TC-05-04",
    }
    for finding_id, review_id in expected.items():
        assert finding_id in by_finding_id
        assert by_finding_id[finding_id]["source_review_ids"] == [review_id]
        assert sum(row["finding_id"] == finding_id for row in findings) == 1

    liability_text = str(by_finding_id["RF-LIABILITY-ASYMMETRY-001"])
    assert "10万元" in liability_text
    assert "30万元" in liability_text

    dispute_text = str(by_finding_id["RF-DISPUTE-CONFLICT-001"])
    assert "南京仲裁委员会" in dispute_text
    assert "甲方所在地有管辖权的人民法院" in dispute_text

    restriction_text = str(by_finding_id["RF-TECH-RESTRICT-001"])
    assert "3年" in restriction_text
    assert "200万元" in restriction_text

    open_source_text = str(by_finding_id["RF-TECH-OSS-001"])
    assert "SBOM" in open_source_text


def test_final_evidence_covers_each_cooperative_risk_proposition(cooperative_review: dict):
    findings = {row["risk_group"]: row for row in cooperative_review["risk_findings"]}

    liability = {
        row["fact_type"]: row["quoted_text"]
        for row in findings["asymmetric_breach_liability"]["evidence_bundle"]["items"]
    }
    assert "10万元" in liability["PARTY_A_COMPARABLE_LIABILITY"]
    assert "30万元" not in liability["PARTY_A_COMPARABLE_LIABILITY"]
    assert "30万元" in liability["PARTY_B_COMPARABLE_LIABILITY"]
    assert "10万元" not in liability["PARTY_B_COMPARABLE_LIABILITY"]
    assert "ratio=3" in liability["LIABILITY_RATIO_COMPARISON"]

    dispute = findings["dispute_resolution_choice_conflict"]["evidence_bundle"]["items"]
    assert any("南京仲裁委员会" in row["quoted_text"] for row in dispute)
    assert any("人民法院提起诉讼" in row["quoted_text"] for row in dispute)
    assert any(row["source_type"] == "SEARCH_AUDIT" for row in dispute)

    restriction = findings["technology_restrictive_covenant"]["evidence_bundle"]["items"]
    restriction_text = "\n".join(row["quoted_text"] for row in restriction)
    assert all(term in restriction_text for term in ("3年", "相同或类似功能", "200万元"))
    assert any(row["source_type"] == "SEARCH_AUDIT" and "地域范围" in row["quoted_text"] for row in restriction)

    components = findings["open_source_component_governance_gap"]["evidence_bundle"]["items"]
    component_text = "\n".join(row["quoted_text"] for row in components)
    assert all(term in component_text for term in (
        "自行选择", "无需事先取得", "开源组件登记清单", "许可证审查流程", "SBOM",
        "知识产权索赔", "全部责任由实际引入该组件的一方承担",
    ))

    exit_evidence = findings["cooperation_exit_continuity_risk"]["evidence_bundle"]["items"]
    exit_text = "\n".join(row["quoted_text"] for row in exit_evidence)
    assert all(term in exit_text for term in ("源代码", "自行决定是否", "无权强制要求"))

    claims = findings["third_party_claim_defense"]["evidence_bundle"]["items"]
    assert any("知识产权索赔" in row["quoted_text"] and "全部责任" in row["quoted_text"] for row in claims)
    assert any(row["source_type"] == "SEARCH_AUDIT" and "抗辩" in row["quoted_text"] for row in claims)

    liability_cap = findings["indemnity_and_liability_cap"]["evidence_bundle"]["items"]
    assert any("商业使用受到限制" in row["quoted_text"] and "全部责任" in row["quoted_text"] for row in liability_cap)
    assert any(
        row["source_type"] == "SEARCH_AUDIT"
        and "适用顺序" in row["quoted_text"]
        and "重复赔偿" in row["quoted_text"]
        for row in liability_cap
    )
    assert all("责任总额上限" not in row["quoted_text"] for row in liability_cap)

    provenance = findings["technology_ip_rights_provenance"]["evidence_bundle"]["items"]
    provenance_text = "\n".join(row["quoted_text"] for row in provenance)
    assert all(term in provenance_text for term in ("50%", "共同著作权人", "专利申请"))
    assert any(row["source_type"] == "SEARCH_AUDIT" and "人员成果归集" in row["quoted_text"] for row in provenance)

    audit = cooperative_review["report_integrity_gate"]["evidence_proposition_audit"]
    assert audit["status"] == "PASS"
    assert audit["checked_risk_count"] == 4
    assert {row["risk_group"] for row in audit["rows"]} == {
        "asymmetric_breach_liability",
        "dispute_resolution_choice_conflict",
        "technology_restrictive_covenant",
        "open_source_component_governance_gap",
    }
    assert all(row["direct_evidence_ids"] for row in audit["rows"])


def test_cross_rule_and_partial_narratives_are_semantically_consistent(cooperative_review: dict):
    by_id = {row["item_id"]: row for row in cooperative_review["items"]}
    findings = {row["risk_group"]: row for row in cooperative_review["risk_findings"]}

    for review_id in ("TC-03-04", "TC-13-02"):
        item = by_id[review_id]
        assert item["coverage_status"] == "部分覆盖"
        assert item["issue_type"] == "PARTIAL"
        assert item["related_risks"] == [findings["technology_ip_rights_provenance"]["finding_id"]]
    assert set(findings["technology_ip_rights_provenance"]["source_review_ids"]) == {
        "TC-03-04", "TC-13-02",
    }

    acceptance = findings["acceptance_test_attachment"]
    assert acceptance["covered_element_ids"] == ["body_acceptance_metrics", "test_method_and_report"]
    assert acceptance["missing_element_ids"] == ["independent_test_attachment_details"]
    assert all(term in acceptance["recommendation"] for term in ("固定测试数据集", "设备环境", "操作步骤"))
    assert all(term not in acceptance["recommendation"] for term in ("量化验收指标", "测试报告", "第三方检测"))
    assert all(term not in acceptance["issue"] for term in ("量化验收指标", "测试方法和测试报告要求"))
    assert all(term not in acceptance["consequence"] for term in ("风险转移", "质量异议", "付款条件"))

    notification = by_id["TC-10-03"]
    notification_consequence = notification["narrative_state"]["consequence"]
    assert all(term in notification_consequence for term in ("改进成果", "合作开发成果", "共同权利", "通知"))
    assert all(term not in notification_consequence for term in ("验收节点", "质量异议", "付款条件", "风险转移"))

    confidentiality = by_id["TC-11-04"]["narrative_state"]["consequence"]
    claims = by_id["TC-12-02"]["narrative_state"]["consequence"]
    assert "改进成果" not in confidentiality
    assert "保密" not in claims

    exit_item = by_id["TC-05-04"]
    assert exit_item["legal_basis_index"].startswith("《民法典》第856条")
    assert "第509条" in exit_item["legal_basis_index"]
    assert "第577条" in exit_item["legal_basis_index"]
    assert "第563条" not in exit_item["legal_basis_index"]

    gate = cooperative_review["report_semantic_consistency_gate"]
    assert gate["status"] == "PASS"
    assert gate["non_fatal"] is True
    assert gate["report_blocked"] is False
    assert gate["warning_count"] == 0


def test_minimal_cooperative_development_rules_and_report_are_consistent():
    output = review_contract_v2(
        _document_text(MINIMAL_FIXTURE),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方/合作开发方"),
    )
    by_id = {row["item_id"]: row for row in output["items"]}

    assert output["report_context"]["contract_type_label"] == "技术合同—技术开发合同—合作开发合同"
    assert output["perspective_resolution"]["display_role"] == "甲方 / 合作开发方"
    assert output["release_quality_gate"]["status"] == "REPORT_READY"
    assert output["report_generation_decision"]["report_generated"] is True
    # A hypothetical third-party infringement allocation is not affirmative
    # evidence that the project has actually infringed another person's
    # technical achievement; the former TC-02-02 false positive is excluded.
    assert len(output["risk_findings"]) == 12
    assert output["summary"] == {"高": 0, "中": 8, "低": 8}
    assert sum(finding["severity"] == "HIGH" for finding in output["risk_findings"]) == 0
    assert sum(finding["severity"] == "MEDIUM" for finding in output["risk_findings"]) == 8
    assert sum(finding["severity"] == "LOW" for finding in output["risk_findings"]) == 4
    assert output["final_disposition_summary"]["actionable"] == 9
    assert output["final_disposition_summary"]["optional_improvement"] == 3
    assert output["manual_review_summary"]["manual_review_count"] == 0

    assert by_id["TC-01-03"]["coverage_status"] == "通过"
    subject_evidence = " ".join(str(row.get("text") or "") for row in by_id["TC-01-03"]["evidence_clauses"])
    assert "星云智能企业数据分析平台 V2.0" in subject_evidence

    assert {review_id: by_id[review_id]["coverage_status"] for review_id in (
        "TC-05-01", "TC-05-02", "TC-05-04",
    )} == {"TC-05-01": "部分覆盖", "TC-05-02": "部分覆盖", "TC-05-04": "部分覆盖"}
    assert by_id["TC-05-02"]["risk_level"] == "低"
    assert "《合同法》第340条" in by_id["TC-05-02"]["legal_basis_index"]
    assert "《民法典》第860条" in by_id["TC-05-02"]["legal_basis_index"]
    assert all(term not in by_id["TC-05-02"]["legal_consequence_if_missing"] for term in ("第三方权利主张", "停用", "替换"))
    secret_result = by_id["TC-05-03"]
    assert secret_result["coverage_status"] == "部分覆盖"
    assert secret_result["risk_level"] == "中"
    assert secret_result["decision_trace"]["severity_trigger"]["core_value_exposure"] is False
    assert {row["element_id"]: row["status"] for row in secret_result["coverage_elements"]} == {
        "joint_source_code_ownership": "COVERED",
        "basic_confidentiality_protection": "COVERED",
        "statutory_default_use_and_transfer": "COVERED",
        "authorized_internal_users": "MISSING",
        "independent_use_right": "MISSING",
        "external_license_transfer_disclosure": "MISSING",
        "external_revenue_allocation": "MISSING",
        "post_termination_use_boundary": "MISSING",
        "material_copy_retention_and_disposal": "MISSING",
    }
    assert "《合同法》第341条" in secret_result["legal_basis_index"]
    assert "《民法典》第861条" in secret_result["legal_basis_index"]
    assert "最高人民法院" in secret_result["legal_basis_index"]
    assert "第20条" in secret_result["legal_basis_index"]
    assert "默认规则" in secret_result["risk_point_description"]
    assert "权利并非完全无法确定" in secret_result["legal_consequence_if_missing"]
    assert "主动明确或改变法定默认安排" in secret_result["suggested_clause_example"]
    secret_section = output["result"].split("（TC-05-03）", 1)[1].split("\n---", 1)[0]
    assert "合作结束后，如未明确各合作方" in secret_section
    assert all(term not in secret_section for term in (
        "项目结束后如未明确开发方", "开发方可留存", "委托方", "受托方", "委托开发方",
    ))
    assert "适用范围、履行标准和责任安排" not in " ".join((
        secret_result["risk_point_description"], secret_result["legal_consequence_if_missing"],
        secret_result["suggested_clause_example"],
    ))
    exit_elements = {row["element_id"]: row["status"] for row in by_id["TC-05-04"]["coverage_elements"]}
    assert exit_elements["agreed_and_objective_termination"] == "COVERED"
    assert all(exit_elements[element_id] == "MISSING" for element_id in (
        "delay_and_nonperformance_liability", "unilateral_early_exit", "source_and_material_handover",
        "unfinished_result_ownership", "cost_settlement", "continued_development_right",
    ))
    assert by_id["TC-05-04"]["needs_manual_review"] is False
    assert all(term not in output["result"] for term in ("委托方", "受托方", "委托开发方"))

    assert all(by_id[review_id]["coverage_status"] == "不适用" for review_id in (
        "TC-10-01", "TC-10-02", "TC-10-03", "TC-10-04",
    ))
    assert all(
        not set(finding["source_review_ids"]) & {"TC-10-01", "TC-10-02", "TC-10-03", "TC-10-04"}
        for finding in output["risk_findings"]
    )

    assert {review_id: by_id[review_id]["coverage_status"] for review_id in (
        "TC-11-01", "TC-11-02", "TC-11-04",
    )} == {"TC-11-01": "部分覆盖", "TC-11-02": "部分覆盖", "TC-11-04": "部分覆盖"}
    confidentiality_clause_ids = [
        {str(row.get("clause_id") or "") for row in by_id[review_id]["evidence_clauses"] if row.get("clause_id")}
        for review_id in ("TC-11-01", "TC-11-02", "TC-11-04")
    ]
    assert set.intersection(*confidentiality_clause_ids)
    confidentiality_quotes = {
        review_id: " ".join(
            str(row.get("quote") or row.get("original_excerpt") or row.get("text") or "")
            for row in by_id[review_id]["evidence_clauses"]
        )
        for review_id in ("TC-11-01", "TC-11-02", "TC-11-04")
    }
    assert all("禁止合作方泄露本协议所涉及的相关商业秘密" in quote for quote in confidentiality_quotes.values())
    assert all("私自以团体名义进行业务活动" not in quote for quote in confidentiality_quotes.values())
    assert "如合作人违反上述各条，应按实际损失赔偿" in confidentiality_quotes["TC-11-04"]
    confidentiality_elements = {row["element_id"]: row["status"] for row in by_id["TC-11-04"]["coverage_elements"]}
    assert confidentiality_elements["confidentiality_obligation"] == "COVERED"
    assert confidentiality_elements["breach_damage_liability"] == "COVERED"
    assert all(confidentiality_elements[element_id] == "MISSING" for element_id in (
        "stop_disclosure", "stop_use", "mitigation_obligation", "return_or_destroy", "disposition_confirmation",
    ))

    claim_response = by_id["TC-12-02"]
    assert claim_response["coverage_status"] == "部分覆盖"
    assert claim_response["risk_level"] == "中"
    assert claim_response["decision_trace"]["severity_trigger"]["major_ip_exposure"] is False
    assert {row["element_id"]: row["status"] for row in claim_response["coverage_elements"]} == {
        "infringement_responsibility": "COVERED",
        "claim_notice": "MISSING",
        "defense_control": "MISSING",
        "lawyer_selection": "MISSING",
        "settlement_authority": "MISSING",
        "defense_cost_allocation": "MISSING",
        "continuity_remedies": "MISSING",
        "cooperation_mechanism": "MISSING",
    }
    claim_evidence = " ".join(
        str(row.get("quote") or row.get("original_excerpt") or row.get("text") or "")
        for row in claim_response["evidence_clauses"]
    )
    assert "不得有侵犯他人知识产权的行为" in claim_evidence
    assert "承担全部侵权责任" in claim_evidence

    infringement = by_id["TC-12-04"]
    assert infringement["title"] == "侵权责任与一般救济衔接"
    assert infringement["coverage_status"] == "部分覆盖"
    infringement_text = " ".join((
        infringement["risk_point_description"], infringement["legal_consequence_if_missing"],
        str(infringement["decision_trace"].get("specialized_search_audit") or ""),
    ))
    assert "重复赔偿" in infringement_text
    assert "责任总额上限" not in infringement_text
    infringement_elements = {row["element_id"]: row["status"] for row in infringement["coverage_elements"]}
    assert infringement_elements["infringement_specific_liability"] == "COVERED"
    assert infringement_elements["general_loss_compensation"] == "COVERED"
    assert infringement_elements["contractual_penalty"] == "NOT_APPLICABLE"
    assert infringement_elements["refund_remedy"] == "NOT_APPLICABLE"
    assert infringement_elements["continued_performance_remedy"] == "NOT_APPLICABLE"
    asserted_remedies = " ".join((
        infringement["reasoning_summary"], infringement["risk_point_description"],
        infringement["legal_consequence_if_missing"],
    ))
    assert "部分一般损失赔偿机制" in asserted_remedies
    assert "已有一般损失赔偿、违约金、退款或继续履行" not in asserted_remedies
    infringement_section = output["result"].split("（TC-12-04）", 1)[1].split("\n---", 1)[0]
    assert "合同中已经存在的一般损失赔偿" in infringement_section
    assert "依法或其他合同条款可能适用的其他救济" in infringement_section
    assert all(term not in infringement_section for term in (
        "侵权专项责任与一般损失赔偿、违约金、退款和继续履行等救济",
        "合同已约定退款", "合同已约定继续履行", "合同已约定专项违约金", "责任上限",
    ))
    assert all(term not in output["result"] for term in (
        "委托方", "受托方", "委托开发方", "未明确开发方", "开发方可留存",
    ))

    documentation = by_id["TC-13-01"]
    assert documentation["coverage_status"] == "部分覆盖"
    documentation_elements = {
        row["element_id"]: row["status"] for row in documentation["coverage_elements"]
    }
    assert documentation_elements["body_development_scope"] == "COVERED"
    assert documentation_elements["body_research_work_allocation"] == "COVERED"
    assert all(documentation_elements[element_id] == "MISSING" for element_id in (
        "requirements_specification", "detailed_module_schedule", "interface_documentation",
        "source_code_version_list", "deployment_and_operation_documentation", "test_materials",
        "research_records", "stage_result_list",
    ))
    assert "合同正文已经明确基础研发范围、主要开发模块和双方基本研发分工" in documentation["risk_point_description"]
    assert "未明确研发分工" not in documentation["risk_point_description"]
    assert "需求规格" in documentation["suggested_clause_example"]
    assert "阶段成果清单" in documentation["suggested_clause_example"]
    assert "适用范围、履行标准和责任安排" not in " ".join((
        documentation["risk_point_description"], documentation["legal_consequence_if_missing"],
        documentation["suggested_clause_example"],
    ))

    forbidden_fallback = "适用范围、履行标准和责任安排"
    assert all(forbidden_fallback not in str(finding) for finding in output["risk_findings"])
