from pathlib import Path

from docx import Document

from review_v2.engine import (
    COVERAGE_NOT_APPLICABLE,
    ReviewItemResult,
    _apply_factual_grounding_gate,
    review_contract_v2,
)
from review_v2.technology_service_analysis import (
    analyze_technical_service_core_facts,
    analyze_technology_relationship_risks,
)
from review_v2.user_profile import UserRiskProfile


ROOT = Path(__file__).resolve().parents[2]
FIXTURE = ROOT / "test_data" / "23 技术服务.docx"


def _fixture_text() -> str:
    document = Document(str(FIXTURE))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = ["\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    return "\n".join(paragraphs + tables)


def _review() -> dict:
    return review_contract_v2(
        _fixture_text(),
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity="甲方", contract_scenario="软件技术服务"),
    )


def test_contract_defined_service_roles_override_party_letter_defaults():
    output = _review()
    context = output["contract_context"]
    perspective = context["perspective_resolution"]

    assert context["contract_subtype"] == "TECHNICAL_SERVICE"
    assert context["transaction_role"] == "SERVICE_PROVIDER"
    assert context["user_role"] == "service_provider"
    assert perspective["party_side"] == "PARTY_A"
    assert perspective["legal_role"] == "SERVICE_PROVIDER"
    assert perspective["contract_display_role"] == "服务提供方"
    assert perspective["display_role"] == "甲方 / 服务提供方"
    assert "审查立场：甲方 / 服务提供方" in output["result"]
    assert "甲方 / 委托方" not in output["result"]


def test_service_acceptance_and_new_result_core_facts_are_not_keyword_passes():
    text = _fixture_text()
    facts = analyze_technical_service_core_facts(text)
    assert facts["acceptance"]["mechanism_present"] is True
    assert facts["acceptance"]["standard_deferred"] is True
    assert facts["acceptance"]["deadline_defined"] is True
    assert facts["acceptance"]["written_objection"] is True
    assert facts["acceptance"]["deemed_acceptance"] is True
    assert facts["acceptance"]["remedy_present"] is True
    assert facts["acceptance"]["remedy_deadline_specific"] is False
    assert facts["service_generated_result"]["scope_defined"] is True
    assert facts["service_generated_result"]["ownership_deferred"] is True
    assert facts["service_generated_result"]["ownership_finalized"] is False

    output = _review()
    by_id = {item["item_id"]: item for item in output["items"]}
    acceptance = by_id["TC-09-02"]
    ownership = by_id["TC-09-03"]

    assert acceptance["coverage_status"] == "部分覆盖"
    acceptance_elements = {row["element_id"]: row["status"] for row in acceptance["coverage_elements"]}
    assert acceptance_elements["acceptance_process"] == "COVERED"
    assert acceptance_elements["acceptance_deadline"] == "COVERED"
    assert acceptance_elements["written_objection_mechanism"] == "COVERED"
    assert acceptance_elements["deemed_acceptance_arrangement"] == "COVERED"
    assert acceptance_elements["acceptance_standard_finality"] == "PARTIALLY_COVERED"
    assert acceptance_elements["rectification_deadline"] == "PARTIALLY_COVERED"

    assert ownership["coverage_status"] == "部分覆盖"
    ownership_elements = {row["element_id"]: row["status"] for row in ownership["coverage_elements"]}
    assert ownership_elements["service_result_scope"] == "COVERED"
    assert ownership_elements["service_result_ownership_finality"] == "PARTIALLY_COVERED"
    assert ownership_elements["interim_result_use"] == "COVERED"
    assert "另行协商" in ownership["risk_point_description"]
    assert "已覆盖，无需修改" not in ownership["action_suggestion"]


def test_service_ip_proof_and_personal_information_rules_are_fact_triggered():
    output = _review()
    by_id = {item["item_id"]: item for item in output["items"]}

    assert by_id["TC-03-04"]["coverage_status"] == COVERAGE_NOT_APPLICABLE
    assert by_id["TC-13-02"]["coverage_status"] == COVERAGE_NOT_APPLICABLE
    assert not any(row["risk_group"] == "technology_ip_rights_provenance" for row in output["risk_findings"])
    assert not any(row["risk_group"] == "personal_information_processing_gap" for row in output["risk_findings"])

    generic = analyze_technology_relationship_risks(
        "服务方可访问业务数据、系统数据、接口数据及系统账号；项目结束后删除或返还。"
    )
    assert generic["personal_information_processing_gap"]["present"] is False

    explicit = analyze_technology_relationship_risks(
        "服务方受托处理员工姓名、手机号和身份证号，并在项目结束后删除个人信息。"
    )
    assert explicit["personal_information_processing_gap"]["present"] is True
    assert {"姓名", "手机号", "身份证"} <= set(
        explicit["personal_information_processing_gap"]["analysis"]["matched_categories"]
    )


def test_final_report_contains_no_ungrounded_service_facts():
    output = _review()
    report = output["result"]
    forbidden = (
        "当前技术服务仅以SLA",
        "月度服务记录",
        "最终书面验收确认履行质量",
        "姓名、工号、电话、邮箱及系统日志",
        "专利证书、专利申请文件",
    )
    assert all(token not in report for token in forbidden)

    test_item = next(item for item in output["items"] if item["item_id"] == "TC-13-03")
    assert test_item["coverage_status"] == COVERAGE_NOT_APPLICABLE
    reason = str((test_item.get("applicability") or {}).get("reason") or "")
    assert "SLA" not in reason
    assert "响应时间" not in reason
    assert "月度服务记录" not in reason


def test_factual_grounding_gate_blocks_an_unbound_concrete_contract_fact():
    item = ReviewItemResult(
        item_id="TC-TEST",
        node_id="test",
        title="测试审查项",
        risk_level="中",
        status="待修改",
        risk_point_description="合同已经约定SLA和月度服务记录。",
        suggested_clause_example="建议沿用现有SLA。",
        legal_basis_index="",
        legal_consequence_if_missing="现有SLA可能引发争议。",
        applicable_if="",
        matched_terms=[],
        coverage_status="部分覆盖",
        issue_type="PARTIAL",
        evidence_clauses=[],
        evidence_location="",
        reasoning_summary="合同存在SLA。",
    )
    item.decision_trace["factual_grounding_claims"] = [{
        "claim_id": "UNGROUNDED-SLA",
        "required_term_groups": [["SLA"], ["月度服务记录"]],
        "evidence_clause_ids": [],
        "fields": [
            "risk_point_description", "reasoning_summary", "legal_consequence_if_missing",
            "suggested_clause_example",
        ],
    }]

    checked = _apply_factual_grounding_gate([item], "合同仅约定一般技术支持。")
    assert checked[0].status == "UNKNOWN"
    assert checked[0].issue_type == "MANUAL_REVIEW"
    assert checked[0].needs_manual_review is True
    assert checked[0].decision_trace["factual_grounding_gate"]["status"] == "NEEDS_REVIEW"
    assert "SLA" not in checked[0].risk_point_description


def test_provider_side_secondary_facts_are_recorded_without_forcing_risks():
    facts = analyze_technical_service_core_facts(_fixture_text())
    assert facts["liability"]["third_party_platform_exception"] is True
    assert facts["liability"]["indirect_loss_exclusion"] is True
    assert facts["liability"]["provider_liability_cap"] is True
    assert facts["change_control"]["principle_only"] is True
    assert facts["termination_settlement"]["actual_workload_basis"] is True
    assert facts["termination_settlement"]["workload_determination_unclear"] is True

    output = _review()
    assert output["deterministic_analysis"]["technical_service_core_facts"] == facts


PERSPECTIVE_CASE_TEXT = """
软件技术服务合同
甲方（服务提供方）为乙方（服务接受方）提供软件维护、故障排查和技术支持服务。
服务期限一年，服务费人民币十万元。服务成果按双方确认的标准书面验收；不合格的，甲方负责整改。
甲方保证其提供的程序代码不侵犯第三方知识产权。因甲方提供的程序代码侵犯第三方知识产权并导致乙方遭受损失的，由甲方承担相应责任。
因乙方提供的软件、数据、图片、文档或商标导致知识产权纠纷的，由乙方自行承担责任。
双方对履约中知悉的商业秘密承担保密义务。
"""


def _perspective_review(party: str) -> dict:
    return review_contract_v2(
        PERSPECTIVE_CASE_TEXT,
        "technology",
        use_llm=False,
        profile=UserRiskProfile(user_identity=party, contract_scenario="软件技术服务"),
    )


def test_third_party_claim_direction_changes_between_provider_and_recipient():
    provider = _perspective_review("甲方")
    recipient = _perspective_review("乙方")
    provider_item = next(item for item in provider["items"] if item["item_id"] == "TC-12-02")
    recipient_item = next(item for item in recipient["items"] if item["item_id"] == "TC-12-02")
    provider_finding = next(
        row for row in provider["risk_findings"] if "TC-12-02" in row["source_review_ids"]
    )
    recipient_finding = next(
        row for row in recipient["risk_findings"] if "TC-12-02" in row["source_review_ids"]
    )
    provider_disposition = next(
        row for row in provider["final_dispositions"] if "TC-12-02" in row["review_ids"]
    )
    recipient_disposition = next(
        row for row in recipient["final_dispositions"] if "TC-12-02" in row["review_ids"]
    )

    # Coverage is a contract fact and remains the same; risk and drafting
    # direction change with the reviewed party.
    assert provider_item["coverage_status"] == recipient_item["coverage_status"] == "部分覆盖"
    assert provider_item["risk_level"] == "低"
    assert provider_finding["severity"] == "LOW"
    assert provider_finding["priority"] == "C_OPTIMIZATION"
    assert provider_item["decision_trace"]["risk_to_current_party"]["risk_scope"] == "PROCEDURAL_CONTROL_ONLY"
    assert provider_item["decision_trace"]["recommendation_direction_check"]["expands_current_party_liability"] is False
    assert "未经服务提供方书面同意" in provider_finding["recommendation"]
    assert "退款或赔偿救济" not in provider_finding["recommendation"]
    assert provider_disposition["client_effect"] == "FAVORABLE"
    assert provider_disposition["final_action"] == "OPTIONAL_IMPROVEMENT"
    assert provider_disposition["blocked_recommendation_ids"] == []
    assert "未经服务提供方书面同意" in provider_disposition["recommendation"]

    assert recipient_item["risk_level"] == "中"
    assert recipient_finding["severity"] == "MEDIUM"
    assert recipient_finding["priority"] == "B_PRACTICAL"
    assert recipient_item["decision_trace"]["risk_to_current_party"]["risk_scope"] == "BASE_PROTECTION_AND_REMEDY"
    assert "解除、退款或替代救济" in recipient_finding["recommendation"]
    assert recipient_disposition["client_effect"] == "FAVORABLE"
    assert recipient_disposition["final_action"] == "SHOULD_CHANGE"


def test_service_result_consequence_survives_finding_aggregation():
    output = _review()
    finding = next(
        row for row in output["risk_findings"] if "TC-09-03" in row["source_review_ids"]
    )
    assert "成果权属" in finding["consequence"]
    assert "交付使用范围" in finding["consequence"]
    assert "停用、替换成本" not in finding["consequence"]
    assert "项目定制成果的知识产权归属" in finding["recommendation"]


def test_service_search_audit_and_material_attachment_follow_current_subtype_facts():
    output = _review()
    items = {item["item_id"]: item for item in output["items"]}
    liability_projection = str({
        "item": items["TC-12-04"],
        "findings": [
            row for row in output["risk_findings"] if "TC-12-04" in row["source_review_ids"]
        ],
    })
    assert "被许可方" not in liability_projection
    assert "超范围实施" not in liability_projection

    attachment = items["TC-13-01"]
    assert attachment["coverage_status"] == "部分覆盖"
    attachment_elements = {row["element_id"]: row["status"] for row in attachment["coverage_elements"]}
    assert attachment_elements["attachment_formation_mechanism"] == "COVERED"
    assert attachment_elements["actual_attachment_schedule"] == "MISSING"
    assert "未显示附件级技术资料清单已经实际形成" in attachment["risk_point_description"]


def test_service_confidentiality_period_and_scope_bind_direct_confidentiality_clauses():
    output = _review()
    item = next(row for row in output["items"] if row["item_id"] == "TC-11-02")
    evidence_text = "\n".join(
        str(row.get("text") or row.get("quote") or "") for row in item["evidence_clauses"]
    )
    elements = {row["element_id"]: row["status"] for row in item["coverage_elements"]}

    assert item["coverage_status"] == "通过"
    assert item["action_priority"] == "NONE"
    assert all(elements[element_id] == "COVERED" for element_id in (
        "confidential_information_scope",
        "mutual_confidentiality_responsibility",
        "third_party_disclosure_restriction",
        "statutory_disclosure_exception",
        "post_termination_survival",
        "confidentiality_period",
    ))
    assert "商业秘密、技术资料、客户信息、经营信息及其他未公开信息" in evidence_text
    assert "任何一方不得向第三方披露" in evidence_text
    assert "司法机关或行政机关依法要求披露的除外" in evidence_text
    assert "合同终止后继续有效" in evidence_text
    assert "双方保密义务长期有效" in evidence_text
    assert "仅可在履行本合同所必需的范围内使用乙方提供的数据" not in evidence_text
    assert not any("TC-11-02" in row["source_review_ids"] for row in output["risk_findings"])


def test_service_provider_confidentiality_breach_narrative_excludes_liability_expansion():
    output = _review()
    item = next(row for row in output["items"] if row["item_id"] == "TC-11-04")
    finding = next(row for row in output["risk_findings"] if "TC-11-04" in row["source_review_ids"])
    disposition = next(row for row in output["final_dispositions"] if "TC-11-04" in row["review_ids"])
    elements = {row["element_id"]: row for row in item["coverage_elements"]}
    report = output["result"]
    section_start = report.index("TC-11-04）")
    section_end = report.index("TC-12-01）", section_start)
    section = report[section_start:section_end]

    # Coverage extraction retains the absent compensation fields for internal
    # audit, while party-direction filtering keeps them out of the provider's
    # user-facing risk and recommendation.
    for element_id in ("direct_loss", "lost_profit", "attorney_fee", "litigation_cost"):
        assert elements[element_id]["status"] == "MISSING"
        assert elements[element_id]["risk_relevant"] is False
        assert elements[element_id]["direction"] == "NEUTRAL"
    for element_id in ("stop_disclosure", "stop_use", "mitigation_obligation"):
        assert elements[element_id]["status"] == "MISSING"
        assert elements[element_id]["risk_relevant"] is True

    assert item["decision_trace"]["risk_to_current_party"]["risk_scope"] == "INCIDENT_CONTAINMENT_ONLY"
    assert item["decision_trace"]["recommendation_direction_check"]["expands_current_party_liability"] is False
    assert "停止披露" in finding["issue"]
    assert "停止使用" in finding["issue"]
    assert "防止损失或披露继续扩大" in finding["issue"]
    assert "可得利益损失" not in finding["issue"]
    assert "合理律师费" not in finding["issue"]
    assert "诉讼费用" not in finding["issue"]
    assert "无需重复创设同一赔偿义务" in disposition["recommendation"]
    assert "赔偿范围包括可得利益损失" not in section
    assert "赔偿范围包括合理律师费" not in section
    assert "赔偿范围包括诉讼费用" not in section
