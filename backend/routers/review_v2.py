"""审查引擎第二版 API。

该路由独立于裁判引擎和旧版 LLM 合同审查主流程。
"""

from __future__ import annotations

import html
import hashlib
import inspect
import io
import logging
import os
import re
from datetime import datetime
from typing import Any, Dict, List
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from pydantic import BaseModel, Field
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from auth.routes import get_current_user
from layers.input_layer import InputLayer
from review_v2.engine import (
    COVERAGE_NOT_APPLICABLE,
    EVALUATOR_VERSION,
    ISSUE_NOT_APPLICABLE,
    PIPELINE_VERSION,
    REPORT_FORMAT_NOTICE,
    RULE_VERSION,
    TECHNICAL_CONSULTING_EVIDENCE_IMPLEMENTATION,
    review_contract_v2,
)
from review_v2.contract_registry import get_checklist_for_resolution
from review_v2.technology_service_analysis import (
    TECHNOLOGY_CLASSIFIER_VERSION,
    classify_technology_transaction,
)
from review_v2.llm_semantic_reviewer import get_llm_review_status
from review_v2.user_profile import UserRiskProfile
from review_v2.semantic_reviewer import (
    COVERAGE_MISSING,
    COVERAGE_PARTIAL,
    COVERAGE_PASS,
    COVERAGE_RISKY,
    ISSUE_CLAUSE_DEFECT,
    ISSUE_MISSING,
    ISSUE_PASS,
    ISSUE_RISKY,
    ISSUE_TEMPLATE_BLANK,
    has_blank_value,
)


router = APIRouter(prefix="/review-v2", tags=["contract_review_v2"])
input_layer = InputLayer()
logger = logging.getLogger(__name__)


def _runtime_source_identity(target: Any) -> Dict[str, Any]:
    path = inspect.getsourcefile(target) or ""
    payload = b""
    mtime = None
    if path and os.path.exists(path):
        with open(path, "rb") as source_file:
            payload = source_file.read()
        mtime = os.path.getmtime(path)
    return {
        "module": getattr(target, "__module__", ""),
        "function": getattr(target, "__qualname__", getattr(target, "__name__", "")),
        "file": path,
        "first_line": getattr(getattr(target, "__code__", None), "co_firstlineno", None),
        "mtime": mtime,
        "sha256": hashlib.sha256(payload).hexdigest() if payload else "",
    }


_LOADED_DEPENDENCY_SOURCES = {
    "engine": _runtime_source_identity(review_contract_v2),
    "classifier": _runtime_source_identity(classify_technology_transaction),
    "registry": _runtime_source_identity(get_checklist_for_resolution),
}


def _runtime_source_drift_status() -> Dict[str, Any]:
    """Fail closed when the running worker has older code than the workspace."""
    targets = {
        "engine": review_contract_v2,
        "classifier": classify_technology_transaction,
        "registry": get_checklist_for_resolution,
    }
    current = {name: _runtime_source_identity(target) for name, target in targets.items()}
    changed = [
        name for name, identity in current.items()
        if str(identity.get("sha256") or "")
        != str((_LOADED_DEPENDENCY_SOURCES.get(name) or {}).get("sha256") or "")
    ]
    return {
        "status": "STALE_RUNTIME" if changed else "CURRENT",
        "changed_dependencies": changed,
        "loaded": _LOADED_DEPENDENCY_SOURCES,
        "current": current,
    }


class ReviewV2Input(BaseModel):
    contract_text: str = Field(..., min_length=1, description="合同文本")
    contract_type: str = Field("house_lease", description="合同类型，当前支持 house_lease、sale、gift、deposit、work_contract、transportation、agency、loan、service、guarantee、finance_lease、warehouse、construction、technology、commission_agency、utility_supply、brokerage")
    include_low_risk: bool = Field(True, description="是否返回低风险/已覆盖检查点")
    use_llm: bool = Field(True, description="是否启用大模型语义复核")
    profile: Dict[str, str] = Field(default_factory=dict, description="用户风险画像：{user_identity, contract_scenario, risk_appetite, bargaining_position}")


class ReviewV2Result(BaseModel):
    success: bool
    contract_type: str
    selected_contract_type: str = ""
    contract_type_name: str
    contract_review_trace_id: str = ""
    message: str
    items: List[Dict[str, Any]] = Field(default_factory=list)
    all_items_count: int = 0
    summary: Dict[str, int] = Field(default_factory=dict)
    risk_finding_summary: Dict[str, int] = Field(default_factory=dict)
    coverage_summary: Dict[str, int] = Field(default_factory=dict)
    issue_type_summary: Dict[str, int] = Field(default_factory=dict)
    judge_source_summary: Dict[str, int] = Field(default_factory=dict)
    llm_enabled: bool = True
    llm_status: Dict[str, Any] = Field(default_factory=dict)
    global_fallback_reason: str = ""
    report_notice: str = ""
    result: str = ""
    filename: str = ""
    extracted_text_length: int = 0
    clause_count: int = 0
    formal_clause_count: int = 0
    contract_facts: Dict[str, Any] = Field(default_factory=dict)
    deterministic_analysis: Dict[str, Any] = Field(default_factory=dict)
    document_integrity: Dict[str, Any] = Field(default_factory=dict)
    quality_gate: Dict[str, Any] = Field(default_factory=dict)
    review_results_4_0: List[Dict[str, Any]] = Field(default_factory=list)
    risk_findings: List[Dict[str, Any]] = Field(default_factory=list)
    risk_themes: List[Dict[str, Any]] = Field(default_factory=list)
    risk_theme_count: int = 0
    aggregated_risk_count: int = 0
    finding_count: int = 0
    risk_theme_summary: Dict[str, int] = Field(default_factory=dict)
    risk_theme_severity_summary: Dict[str, int] = Field(default_factory=dict)
    recommendations: List[Dict[str, Any]] = Field(default_factory=list)
    final_dispositions: List[Dict[str, Any]] = Field(default_factory=list)
    manual_review_items: List[Dict[str, Any]] = Field(default_factory=list)
    manual_review_summary: Dict[str, Any] = Field(default_factory=dict)
    contract_context: Dict[str, Any] = Field(default_factory=dict)
    report_context: Dict[str, Any] = Field(default_factory=dict)
    perspective_resolution: Dict[str, Any] = Field(default_factory=dict)
    overall_assessment: Dict[str, Any] = Field(default_factory=dict)
    legal_reference_issues: List[Dict[str, Any]] = Field(default_factory=list)
    legal_citation_claims: List[Dict[str, Any]] = Field(default_factory=list)
    legal_authority_validation: Dict[str, Any] = Field(default_factory=dict)
    benchmark_metrics: Dict[str, Any] = Field(default_factory=dict)
    release_quality_gate: Dict[str, Any] = Field(default_factory=dict)
    terminology_consistency_gate: Dict[str, Any] = Field(default_factory=dict)
    formal_recommendation_validation: Dict[str, Any] = Field(default_factory=dict)
    technology_transaction_consistency_gate: Dict[str, Any] = Field(default_factory=dict)
    fatal_report_generation_gate: Dict[str, Any] = Field(default_factory=dict)
    report_generation_decision: Dict[str, Any] = Field(default_factory=dict)
    quality4_gate: Dict[str, Any] = Field(default_factory=dict)
    architecture_status: List[Dict[str, Any]] = Field(default_factory=list)
    system_metadata: Dict[str, Any] = Field(default_factory=dict)
    observability: Dict[str, Any] = Field(default_factory=dict)
    production_path_trace: Dict[str, Any] = Field(default_factory=dict)
    profile: Dict[str, Any] = Field(default_factory=dict)
    litigation: Dict[str, Any] = Field(default_factory=dict)
    litigation_result: str = ""
    negotiation: Dict[str, Any] = Field(default_factory=dict)
    negotiation_result: str = ""
    jurisdiction: Dict[str, Any] = Field(default_factory=dict)
    jurisdiction_result: str = ""


class ReviewV2PdfInput(BaseModel):
    filename: str = "合同审查报告"
    contract_type_name: str = "房屋租赁合同"
    result: str = ""
    summary: Dict[str, int] = Field(default_factory=dict)
    risk_finding_summary: Dict[str, int] = Field(default_factory=dict)
    coverage_summary: Dict[str, int] = Field(default_factory=dict)
    issue_type_summary: Dict[str, int] = Field(default_factory=dict)
    judge_source_summary: Dict[str, int] = Field(default_factory=dict)
    llm_status: Dict[str, Any] = Field(default_factory=dict)
    global_fallback_reason: str = ""
    report_notice: str = ""
    risk_findings: List[Dict[str, Any]] = Field(default_factory=list)
    recommendations: List[Dict[str, Any]] = Field(default_factory=list)
    overall_assessment: Dict[str, Any] = Field(default_factory=dict)
    benchmark_metrics: Dict[str, Any] = Field(default_factory=dict)
    release_quality_gate: Dict[str, Any] = Field(default_factory=dict)
    quality4_gate: Dict[str, Any] = Field(default_factory=dict)
    system_metadata: Dict[str, Any] = Field(default_factory=dict)
    items: List[Dict[str, Any]] = Field(default_factory=list)


class ReviewV2LlmStatus(BaseModel):
    enabled_by_env: bool
    has_model_key: bool
    available: bool
    configured_models: List[str] = Field(default_factory=list)
    priority: str = ""
    message: str = ""


def _register_cjk_font() -> str:
    candidates = [
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
    ]
    for path in candidates:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont("SageCJK", path))
                return "SageCJK"
            except Exception:
                continue
    return "Helvetica"


def _p(text: Any) -> str:
    return html.escape(str(text or "")).replace("\n", "<br/>")


def _item_evidence_text(item: Dict[str, Any]) -> str:
    return "\n".join(
        str(evidence.get("text") or "")
        for evidence in (item.get("evidence_clauses") or [])
        if isinstance(evidence, dict)
    )


def _item_issue_type(item: Dict[str, Any]) -> str:
    issue_type = str(item.get("issue_type") or "").strip()
    if issue_type:
        return issue_type

    coverage_status = str(item.get("coverage_status") or "").strip()
    evidence_text = _item_evidence_text(item)
    if coverage_status in {COVERAGE_PASS, COVERAGE_PARTIAL} and has_blank_value(evidence_text):
        return ISSUE_TEMPLATE_BLANK
    if coverage_status == COVERAGE_PASS:
        return ISSUE_PASS
    if coverage_status == COVERAGE_RISKY:
        return ISSUE_RISKY
    if coverage_status == COVERAGE_MISSING:
        return ISSUE_MISSING
    return ISSUE_CLAUSE_DEFECT


def _item_legal_consequence(item: Dict[str, Any]) -> str:
    if item.get("coverage_status") == COVERAGE_NOT_APPLICABLE:
        applicability = item.get("applicability") if isinstance(item.get("applicability"), dict) else {}
        return str(applicability.get("user_reason") or applicability.get("reason") or "本项不适用于当前合同。")
    if item.get("coverage_status") == COVERAGE_PASS:
        related = item.get("related_risks") or []
        if related:
            return f"关联风险 {', '.join(map(str, related))}：同一条款或事实仍存在需要处理的经济或法律风险。"
        return "无明显残余风险"
    narrative = item.get("narrative_state") if isinstance(item.get("narrative_state"), dict) else {}
    return (
        str(narrative.get("consequence") or "").strip()
        if item.get("coverage_status") == COVERAGE_PARTIAL else ""
    ) or (
        str(item.get("legal_consequence_if_missing") or "").strip()
        or str(item.get("risk_point_description") or "").strip()
        or "未补全该审查点会增加合同履行和争议处理的不确定性，建议签署前补充明确约定。"
    )


def _item_production_narrative(item: Dict[str, Any], field: str, fallback: str = "") -> str:
    narrative = item.get("narrative_state") if isinstance(item.get("narrative_state"), dict) else {}
    if item.get("coverage_status") == COVERAGE_PARTIAL and narrative.get(field):
        return str(narrative[field])
    if item.get("coverage_status") == COVERAGE_NOT_APPLICABLE:
        applicability = item.get("applicability") if isinstance(item.get("applicability"), dict) else {}
        return str(applicability.get("user_reason") or applicability.get("reason") or fallback)
    return str(fallback or "")


def _item_legal_label(item: Dict[str, Any]) -> str:
    if item.get("coverage_status") == COVERAGE_NOT_APPLICABLE:
        return "适用说明"
    if item.get("coverage_status") == COVERAGE_PASS:
        if item.get("related_risks"):
            return "关联风险"
        return "残余风险"
    return "缺失的法律后果"


def _item_suggested_label(item: Dict[str, Any]) -> str:
    coverage = item.get("coverage_status")
    if coverage in {COVERAGE_NOT_APPLICABLE, COVERAGE_PASS}:
        return ""
    if coverage == COVERAGE_PARTIAL:
        return "建议补充"
    if coverage == COVERAGE_RISKY:
        return "建议修改原条款"
    return "建议新增"


def _item_citation(item: Dict[str, Any]) -> Dict[str, str]:
    citation = item.get("citation") if isinstance(item.get("citation"), dict) else {}
    return {
        "clause_number": str(citation.get("clause_number") or "未召回正式条款"),
        "original_excerpt": str(citation.get("original_excerpt") or "无正式条款证据"),
        "position_description": str(citation.get("position_description") or "合同正式条款及附件中未定位到可验证条款"),
    }


def _user_facing_text(value: Any) -> str:
    text = str(value or "")
    replacements = {
        "统一事实图谱提取": "根据合同原文可确认",
        "统一事实图谱识别": "根据合同原文可识别",
        "统一事实图谱显示": "根据合同原文显示",
        "统一事实图谱": "合同事实",
        "Fact Graph": "合同事实",
        "Scenario Engine": "金额测算",
        "Final Quality Gate": "最终校验",
        "Cross Review Consistency Verifier": "一致性校验",
        "rule_fallback": "自动审查",
        "LLM": "语义复核",
        "建议人工复核": "建议重点核对",
        "需要人工复核": "需要重点核对",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def _priority_label(priority: str) -> str:
    return {
        "A_CORE": "优先处理",
        "B_PRACTICAL": "建议处理",
        "C_OPTIMIZATION": "可选优化",
        "N/A": "不适用",
    }.get(priority or "", priority or "未标注")


def _finding_title_map(payload: "ReviewV2PdfInput") -> Dict[str, str]:
    return {finding.get("finding_id", ""): finding.get("title", "关联风险") for finding in (payload.risk_findings or [])}


def _related_risk_titles(item: Dict[str, Any], title_map: Dict[str, str]) -> str:
    ids = item.get("related_risks") or []
    if not ids:
        return "无"
    return "、".join(title_map.get(risk_id, "关联风险") for risk_id in ids)


def _issue_type_summary_from_items(items: List[Dict[str, Any]]) -> Dict[str, int]:
    summary = {
        ISSUE_PASS: 0,
        ISSUE_TEMPLATE_BLANK: 0,
        ISSUE_CLAUSE_DEFECT: 0,
        ISSUE_MISSING: 0,
        ISSUE_RISKY: 0,
        ISSUE_NOT_APPLICABLE: 0,
    }
    for item in items:
        issue_type = _item_issue_type(item)
        summary[issue_type] = summary.get(issue_type, 0) + 1
    return summary


def _production_markdown_or_raise(payload: ReviewV2PdfInput) -> str:
    markdown = (payload.result or "").strip()
    if not markdown:
        raise ValueError("正式导出必须使用生产 Markdown result，禁止由导出器重建业务报告。")
    return markdown


def _markdown_blocks(markdown: str) -> List[Dict[str, str]]:
    blocks: List[Dict[str, str]] = []
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line == "---":
            blocks.append({"type": "separator", "text": ""})
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            blocks.append({"type": f"h{len(heading.group(1))}", "text": heading.group(2)})
            continue
        if line.startswith("- "):
            blocks.append({"type": "bullet", "text": line[2:]})
            continue
        if line.startswith("> "):
            blocks.append({"type": "quote", "text": line[2:]})
            continue
        blocks.append({"type": "paragraph", "text": line})
    return blocks


def _strip_markdown_inline(text: str) -> str:
    cleaned = re.sub(r"\*\*(.*?)\*\*", r"\1", text or "")
    cleaned = re.sub(r"`([^`]*)`", r"\1", cleaned)
    return cleaned.strip()


def _build_pdf(payload: ReviewV2PdfInput) -> bytes:
    markdown = _production_markdown_or_raise(payload)
    font_name = _register_cjk_font()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=1.6 * cm,
        leftMargin=1.6 * cm,
        topMargin=1.6 * cm,
        bottomMargin=1.6 * cm,
    )
    styles = getSampleStyleSheet()
    base = ParagraphStyle("SageBase", parent=styles["Normal"], fontName=font_name, fontSize=9, leading=14, alignment=TA_LEFT)
    title = ParagraphStyle("SageTitle", parent=base, fontSize=18, leading=24, spaceAfter=12)
    heading = ParagraphStyle("SageHeading", parent=base, fontSize=12, leading=18, spaceBefore=10, spaceAfter=6)
    subheading = ParagraphStyle("SageSubHeading", parent=base, fontSize=10, leading=15, spaceBefore=6, spaceAfter=4)
    quote = ParagraphStyle("SageQuote", parent=base, textColor=colors.HexColor("#4b5563"), leftIndent=10, leading=14)
    story = []
    for block in _markdown_blocks(markdown):
        text = _p(_strip_markdown_inline(block["text"]))
        if block["type"] == "separator":
            story.append(Spacer(1, 6))
        elif block["type"] == "h1":
            story.append(Paragraph(text, title))
        elif block["type"] in {"h2", "h3"}:
            story.append(Paragraph(text, heading))
        elif block["type"].startswith("h"):
            story.append(Paragraph(text, subheading))
        elif block["type"] == "bullet":
            story.append(Paragraph(f"• {text}", base))
        elif block["type"] == "quote":
            story.append(Paragraph(text, quote))
        else:
            story.append(Paragraph(text, base))
    doc.build(story)
    return buffer.getvalue()


def _build_debug_pdf(payload: ReviewV2PdfInput) -> bytes:
    font_name = _register_cjk_font()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=1.6 * cm,
        leftMargin=1.6 * cm,
        topMargin=1.6 * cm,
        bottomMargin=1.6 * cm,
    )

    styles = getSampleStyleSheet()
    base = ParagraphStyle(
        "SageBase",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=9,
        leading=14,
        alignment=TA_LEFT,
    )
    title = ParagraphStyle(
        "SageTitle",
        parent=base,
        fontSize=18,
        leading=24,
        spaceAfter=12,
    )
    heading = ParagraphStyle(
        "SageHeading",
        parent=base,
        fontSize=12,
        leading=18,
        spaceBefore=10,
        spaceAfter=6,
    )

    summary = payload.summary or {}
    risk_finding_summary = payload.risk_finding_summary or {}
    coverage_summary = payload.coverage_summary or {}
    issue_type_summary = _issue_type_summary_from_items(payload.items)
    judge_source_summary = payload.judge_source_summary or {}
    finding_titles = _finding_title_map(payload)
    story = [
        Paragraph("合同审查报告（第二版）", title),
        Paragraph(f"合同类型：{_p(payload.contract_type_name)}", base),
        Paragraph(f"文件名称：{_p(payload.filename)}", base),
        Paragraph(f"总体可签署性：{_p((payload.overall_assessment or {}).get('signability') or '未记录')}", base),
        Paragraph(f"模型状态：{_p((payload.llm_status or {}).get('message') or '未记录')}", base),
        Paragraph(f"报告提示：{_p(payload.report_notice or REPORT_FORMAT_NOTICE)}", base),
        Paragraph(f"全局回退原因：{_p(payload.global_fallback_reason or '无')}", base),
        Paragraph(f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", base),
        Spacer(1, 10),
        Table(
            [
                ["重大风险", "中等风险", "低风险", "人工复核"],
                [
                    str(risk_finding_summary.get("重大风险", 0)),
                    str(risk_finding_summary.get("中等风险", 0)),
                    str(risk_finding_summary.get("低风险", 0)),
                    str(risk_finding_summary.get("人工复核", 0)),
                ],
            ],
            colWidths=[4 * cm, 4 * cm, 4 * cm, 4 * cm],
            style=TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), font_name),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f6fb")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d8dee9")),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ]
            ),
        ),
        Spacer(1, 10),
        Table(
            [
                ["高风险", "中风险", "低风险", "不适用", "审查项总数"],
                [
                    str(summary.get("高", 0)),
                    str(summary.get("中", 0)),
                    str(summary.get("低", 0)),
                    str(coverage_summary.get(COVERAGE_NOT_APPLICABLE, 0)),
                    str(len(payload.items)),
                ],
            ],
            colWidths=[3.2 * cm, 3.2 * cm, 3.2 * cm, 3.2 * cm, 3.2 * cm],
            style=TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), font_name),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f6fb")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d8dee9")),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ]
            ),
        ),
        Spacer(1, 12),
    ]

    if coverage_summary:
        story.extend(
            [
                Paragraph("覆盖状态统计", heading),
                Table(
                    [
                        ["通过", "部分覆盖", "缺失", "风险条款", "不适用"],
                        [
                            str(coverage_summary.get("通过", 0)),
                            str(coverage_summary.get("部分覆盖", 0)),
                            str(coverage_summary.get("缺失", 0)),
                            str(coverage_summary.get("风险条款", 0)),
                            str(coverage_summary.get(COVERAGE_NOT_APPLICABLE, 0)),
                        ],
                    ],
                    colWidths=[3.2 * cm, 3.2 * cm, 3.2 * cm, 3.2 * cm, 3.2 * cm],
                    style=TableStyle(
                        [
                            ("FONTNAME", (0, 0), (-1, -1), font_name),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f6fb")),
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d8dee9")),
                            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                            ("TOPPADDING", (0, 0), (-1, -1), 8),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                        ]
                    ),
                ),
                Spacer(1, 12),
            ]
        )

    if issue_type_summary:
        story.extend(
            [
                Paragraph("问题类型统计", heading),
                Table(
                    [
                        ["通过", "模板空白", "条款缺陷", "完全缺失", "风险条款", "不适用"],
                        [
                            str(issue_type_summary.get("通过", 0)),
                            str(issue_type_summary.get("模板空白", 0)),
                            str(issue_type_summary.get("条款缺陷", 0)),
                            str(issue_type_summary.get("完全缺失", 0)),
                            str(issue_type_summary.get("风险条款", 0)),
                            str(issue_type_summary.get(ISSUE_NOT_APPLICABLE, 0)),
                        ],
                    ],
                    colWidths=[2.65 * cm, 2.65 * cm, 2.65 * cm, 2.65 * cm, 2.65 * cm, 2.65 * cm],
                    style=TableStyle(
                        [
                            ("FONTNAME", (0, 0), (-1, -1), font_name),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f6fb")),
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d8dee9")),
                            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                            ("TOPPADDING", (0, 0), (-1, -1), 8),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                        ]
                    ),
                ),
                Spacer(1, 12),
            ]
        )

    for index, item in enumerate(payload.items, start=1):
        issue_type = _item_issue_type(item)
        legal_consequence = _item_legal_consequence(item)
        legal_label = _item_legal_label(item)
        suggested_label = _item_suggested_label(item)
        citation = _item_citation(item)
        reasoning = str(item.get("reasoning") or "").strip()
        story.extend(
            [
                Paragraph(f"{index}. {_p(item.get('title'))}（{_p(item.get('item_id'))}）", heading),
                Paragraph(
                    f"覆盖状态：{_p(item.get('coverage_status'))}　"
                    f"审查优先级：{_p(_priority_label(item.get('review_priority') or 'N/A'))}　"
                    f"关联风险发现：{_p(_related_risk_titles(item, finding_titles))}　"
                    f"问题类型：{_p(issue_type)}　"
                    f"复核状态：{_p(item.get('review_status') or '待处理')}",
                    base,
                ),
                Paragraph(
                    f"证据来源：{_p('语义复核' if item.get('judge_source') == 'llm' else '合同原文/结构化事实')}　"
                    f"模型置信度：{_p(item.get('llm_confidence') or '无')}",
                    base,
                ),
                Paragraph(f"原文位置：{_p(item.get('evidence_location') or '未召回正式条款')}", base),
                Paragraph(f"系统判断：{_p(_user_facing_text(_item_production_narrative(item, 'judgment', item.get('reasoning_summary'))))}", base),
                Paragraph(f"【风险点描述】：{_p(_user_facing_text(_item_production_narrative(item, 'issue', item.get('risk_point_description'))))}", base),
                Paragraph(f"【{_p(legal_label)}】：{_p(_user_facing_text(legal_consequence))}", base),
                Paragraph(f"【法律依据索引】：{_p(item.get('legal_basis_index'))}", base),
                Paragraph("命中原文：", base),
                *[
                    Paragraph(
                        f"{_p(evidence.get('location'))}：{_p(evidence.get('text'))}",
                        base,
                    )
                    for evidence in (item.get("evidence_clauses") or [])[:3]
                ],
                *(
                    [
                        Paragraph("【推理过程】", heading),
                        *[Paragraph(_p(line), base) for line in reasoning.splitlines()[1:] if line.strip()],
                        Paragraph("【证据定位】", heading),
                        Paragraph(f"条款编号：{_p(citation['clause_number'])}", base),
                        Paragraph(f"原文片段：{_p(citation['original_excerpt'])}", base),
                        Paragraph(f"位置描述：{_p(citation['position_description'])}", base),
                    ]
                    if reasoning
                    else []
                ),
                Paragraph(f"人工备注：{_p(item.get('reviewer_note'))}", base),
                Spacer(1, 8),
            ]
        )
        if suggested_label:
            suggested_value = _item_production_narrative(item, "recommendation", item.get("suggested_clause_example"))
            story.insert(-1, Paragraph(f"【{_p(suggested_label)}】：{_p(suggested_value)}", base))

    doc.build(story)
    return buffer.getvalue()


def _set_docx_run_font(run, *, bold: bool = False, size: int = 10) -> None:
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    try:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    except Exception:
        pass


def _add_docx_line(doc: DocxDocument, label: str, value: Any = "", *, bold_label: bool = True) -> None:
    paragraph = doc.add_paragraph()
    label_run = paragraph.add_run(label)
    _set_docx_run_font(label_run, bold=bold_label)
    if value not in (None, ""):
        value_run = paragraph.add_run(str(value))
        _set_docx_run_font(value_run)


def _add_docx_table(doc: DocxDocument, headers: List[str], values: List[Any]) -> None:
    table = doc.add_table(rows=2, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
        table.cell(1, index).text = str(values[index])


def _build_debug_docx(payload: ReviewV2PdfInput) -> bytes:
    doc = DocxDocument()
    finding_titles = _finding_title_map(payload)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10)

    title = doc.add_heading("合同审查报告（第二版）", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _add_docx_line(doc, "合同类型：", payload.contract_type_name)
    _add_docx_line(doc, "文件名称：", payload.filename)
    _add_docx_line(doc, "总体可签署性：", (payload.overall_assessment or {}).get("signability") or "未记录")
    _add_docx_line(doc, "模型状态：", (payload.llm_status or {}).get("message") or "未记录")
    _add_docx_line(doc, "报告提示：", payload.report_notice or REPORT_FORMAT_NOTICE)
    _add_docx_line(doc, "全局回退原因：", payload.global_fallback_reason or "无")
    _add_docx_line(doc, "导出时间：", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

    summary = payload.summary or {}
    risk_finding_summary = payload.risk_finding_summary or {}
    coverage_summary = payload.coverage_summary or {}
    doc.add_heading("风险发现统计", level=2)
    _add_docx_table(
        doc,
        ["重大风险", "中等风险", "低风险", "人工复核"],
        [
            risk_finding_summary.get("重大风险", 0),
            risk_finding_summary.get("中等风险", 0),
            risk_finding_summary.get("低风险", 0),
            risk_finding_summary.get("人工复核", 0),
        ],
    )
    if coverage_summary:
        doc.add_heading("覆盖状态统计", level=2)
        _add_docx_table(
            doc,
            ["通过", "部分覆盖", "缺失", "风险条款", "不适用"],
            [
                coverage_summary.get("通过", 0),
                coverage_summary.get("部分覆盖", 0),
                coverage_summary.get("缺失", 0),
                coverage_summary.get("风险条款", 0),
                coverage_summary.get(COVERAGE_NOT_APPLICABLE, 0),
            ],
        )

    issue_type_summary = _issue_type_summary_from_items(payload.items)
    if issue_type_summary:
        doc.add_heading("问题类型统计", level=2)
        _add_docx_table(
            doc,
            ["通过", "模板空白", "条款缺陷", "完全缺失", "风险条款", "不适用"],
            [
                issue_type_summary.get("通过", 0),
                issue_type_summary.get("模板空白", 0),
                issue_type_summary.get("条款缺陷", 0),
                issue_type_summary.get("完全缺失", 0),
                issue_type_summary.get("风险条款", 0),
                issue_type_summary.get(ISSUE_NOT_APPLICABLE, 0),
            ],
        )

    for index, item in enumerate(payload.items, start=1):
        doc.add_heading(f"{index}. {item.get('title')}（{item.get('item_id')}）", level=2)
        issue_type = _item_issue_type(item)
        legal_consequence = _item_legal_consequence(item)
        legal_label = _item_legal_label(item)
        suggested_label = _item_suggested_label(item)
        citation = _item_citation(item)
        reasoning = str(item.get("reasoning") or "").strip()
        source = "语义复核" if item.get("judge_source") == "llm" else "合同原文/结构化事实"

        _add_docx_line(
            doc,
            "基本信息：",
            f"覆盖状态：{item.get('coverage_status') or ''}；审查优先级：{_priority_label(item.get('review_priority') or 'N/A')}；"
            f"关联风险发现：{_related_risk_titles(item, finding_titles)}；"
            f"问题类型：{issue_type}；复核状态：{item.get('review_status') or '待处理'}",
        )
        _add_docx_line(doc, "证据来源：", f"{source}；模型置信度：{item.get('llm_confidence') or '无'}")
        _add_docx_line(doc, "原文位置：", item.get("evidence_location") or "未召回正式条款")
        _add_docx_line(doc, "系统判断：", _user_facing_text(_item_production_narrative(item, "judgment", item.get("reasoning_summary"))))
        _add_docx_line(doc, "【风险点描述】：", _user_facing_text(_item_production_narrative(item, "issue", item.get("risk_point_description"))))
        _add_docx_line(doc, f"【{legal_label}】：", _user_facing_text(legal_consequence))
        if suggested_label:
            _add_docx_line(doc, f"【{suggested_label}】：", _item_production_narrative(item, "recommendation", item.get("suggested_clause_example")))
        _add_docx_line(doc, "【法律依据索引】：", item.get("legal_basis_index") or "")

        evidences = item.get("evidence_clauses") or []
        _add_docx_line(doc, "命中原文：")
        if evidences:
            for evidence in evidences[:3]:
                _add_docx_line(doc, "", f"{evidence.get('location') or ''}：{evidence.get('text') or ''}", bold_label=False)
        else:
            _add_docx_line(doc, "", "无正式条款证据", bold_label=False)

        if reasoning:
            _add_docx_line(doc, "【推理过程】：")
            for line in reasoning.splitlines()[1:]:
                if line.strip():
                    _add_docx_line(doc, "", line, bold_label=False)
            _add_docx_line(doc, "【证据定位】：")
            _add_docx_line(doc, "- 条款编号：", citation["clause_number"])
            _add_docx_line(doc, "- 原文片段：", citation["original_excerpt"])
            _add_docx_line(doc, "- 位置描述：", citation["position_description"])
        _add_docx_line(doc, "人工备注：", item.get("reviewer_note") or "")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _build_docx(payload: ReviewV2PdfInput) -> bytes:
    markdown = _production_markdown_or_raise(payload)
    doc = DocxDocument()
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10)
    for block in _markdown_blocks(markdown):
        text = _strip_markdown_inline(block["text"])
        if block["type"] == "separator":
            doc.add_paragraph("")
            continue
        if block["type"] == "h1":
            paragraph = doc.add_heading(text, level=1)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if block["type"] == "h2":
            doc.add_heading(text, level=2)
            continue
        if block["type"] == "h3":
            doc.add_heading(text, level=3)
            continue
        if block["type"].startswith("h"):
            doc.add_heading(text, level=4)
            continue
        paragraph = doc.add_paragraph()
        if block["type"] == "bullet":
            paragraph.style = "List Bullet"
        run = paragraph.add_run(text)
        _set_docx_run_font(run)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@router.get("/llm-status", response_model=ReviewV2LlmStatus, summary="合同审查第二版：大模型语义复核状态")
async def review_v2_llm_status(
    current_user: dict = Depends(get_current_user),
):
    """返回审查第二版当前是否可以调用大模型语义复核。"""

    return get_llm_review_status()


@router.post("/litigation-analysis", summary="合同审查第二版：诉讼策略分析")
async def review_v2_litigation_analysis(
    payload: ReviewV2Input,
    current_user: dict = Depends(get_current_user),
):
    """仅对已标记高/中风险的条款执行诉讼策略分析（证据缺口+法院裁量+成本收益）。"""

    from review_v2.litigation_analyzer import analyze_litigation_strategy, format_litigation_markdown

    profile = UserRiskProfile.from_dict(payload.profile) if payload.profile else UserRiskProfile.default()

    # 先运行审查获取 items
    review_result = review_contract_v2(
        payload.contract_text,
        payload.contract_type,
        include_low_risk=payload.include_low_risk,
        use_llm=payload.use_llm,
        profile=profile,
    )

    if not review_result.get("success"):
        return {"success": False, "message": review_result.get("message", "审查失败"), "items": [], "summary": {}}

    report = analyze_litigation_strategy(review_result["items"], profile)
    return {
        "success": True,
        "disclaimer": report.disclaimer,
        "items": [item.to_public_dict() for item in report.items],
        "summary": report.summary,
        "result": format_litigation_markdown(report),
        "profile": profile.to_dict(),
    }


@router.post("/negotiation-analysis", summary="合同审查第二版：商业谈判分析")
async def review_v2_negotiation_analysis(
    payload: ReviewV2Input,
    current_user: dict = Depends(get_current_user),
):
    """对已标记高/中风险的条款生成谈判筹码转换方案（谈判属性+对方意图+替代方案+交换筹码+话术）。"""

    from review_v2.negotiation_analyzer import analyze_negotiation_strategy, format_negotiation_markdown

    profile = UserRiskProfile.from_dict(payload.profile) if payload.profile else UserRiskProfile.default()

    review_result = review_contract_v2(
        payload.contract_text,
        payload.contract_type,
        include_low_risk=payload.include_low_risk,
        use_llm=payload.use_llm,
        profile=profile,
    )

    if not review_result.get("success"):
        return {"success": False, "message": review_result.get("message", "审查失败"), "items": [], "overall_strategy": ""}

    report = analyze_negotiation_strategy(review_result["items"], profile)
    return {
        "success": True,
        "disclaimer": report.disclaimer,
        "items": [item.to_public_dict() for item in report.items],
        "must_fight": report.must_fight,
        "core_negotiate": report.core_negotiate,
        "flexible": report.flexible,
        "concession_order": report.concession_order,
        "overall_strategy": report.overall_strategy,
        "result": format_negotiation_markdown(report),
        "profile": profile.to_dict(),
    }


@router.post("/jurisdiction-prediction", summary="合同审查第二版：跨地域裁判预判")
async def review_v2_jurisdiction_prediction(
    payload: ReviewV2Input,
    current_user: dict = Depends(get_current_user),
):
    """结合用户所在地法院的司法实践，对高/中风险条款进行地域裁判预判。"""

    from review_v2.jurisdiction_analyzer import analyze_jurisdiction_prediction, format_jurisdiction_markdown

    profile = UserRiskProfile.from_dict(payload.profile) if payload.profile else UserRiskProfile.default()

    review_result = review_contract_v2(
        payload.contract_text,
        payload.contract_type,
        include_low_risk=payload.include_low_risk,
        use_llm=payload.use_llm,
        profile=profile,
    )

    if not review_result.get("success"):
        return {"success": False, "message": review_result.get("message", "审查失败"), "items": [], "environment_assessment": ""}

    report = analyze_jurisdiction_prediction(review_result["items"], profile, payload.contract_text)
    return {
        "success": True,
        "disclaimer": report.disclaimer,
        "jurisdiction": report.jurisdiction,
        "items": [item.to_public_dict() for item in report.items],
        "environment_assessment": report.environment_assessment,
        "jurisdiction_strategy_summary": report.jurisdiction_strategy_summary,
        "result": format_jurisdiction_markdown(report),
        "profile": profile.to_dict(),
    }


@router.post("/checklist", response_model=ReviewV2Result, summary="合同审查第二版：清单式审查")
async def review_v2_checklist(
    payload: ReviewV2Input,
    current_user: dict = Depends(get_current_user),
):
    """按文本运行独立审查引擎，供调试和后续批量迁移使用。"""

    profile = UserRiskProfile.from_dict(payload.profile) if payload.profile else UserRiskProfile.default()

    return review_contract_v2(
        payload.contract_text,
        payload.contract_type,
        include_low_risk=payload.include_low_risk,
        use_llm=payload.use_llm,
        profile=profile,
    )


@router.post("/file", response_model=ReviewV2Result, summary="合同审查第二版：文件上传审查")
async def review_v2_file(
    file: UploadFile = File(...),
    contract_type: str = Form("house_lease"),
    include_low_risk: bool = Form(True),
    use_llm: bool = Form(True),
    profile_json: str = Form("{}"),
    current_user: dict = Depends(get_current_user),
):
    """接收 Word/PDF 合同文件，抽取文本后运行清单式审查。"""

    source_drift = _runtime_source_drift_status()
    if source_drift["status"] != "CURRENT":
        logger.critical(
            "[STALE_REVIEW_RUNTIME] dependencies=%s process_id=%s",
            source_drift["changed_dependencies"],
            os.getpid(),
        )
        raise HTTPException(
            status_code=503,
            detail="审查后端已加载旧代码，请重启本地服务后重新审查。",
        )

    filename = file.filename or ""
    if not filename.lower().endswith((".doc", ".docx", ".pdf")):
        raise HTTPException(status_code=400, detail="审查第二版当前支持 .doc、.docx 和 .pdf 文件")

    import json as _json
    try:
        profile_dict = _json.loads(profile_json) if profile_json else {}
    except _json.JSONDecodeError:
        profile_dict = {}
    profile = UserRiskProfile.from_dict(profile_dict)

    try:
        contract_text = input_layer.extract_text_from_file(file)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"合同文本读取失败：{exc}") from exc

    result = review_contract_v2(
        contract_text,
        contract_type,
        include_low_risk=include_low_risk,
        use_llm=use_llm,
        profile=profile,
    )
    actual_pipeline = str((result.get("system_metadata") or {}).get("pipeline_version") or "")
    if actual_pipeline != PIPELINE_VERSION:
        logger.critical(
            "[PIPELINE_ASSERTION_FAILED] expected=%s actual=%s entrypoint=review_v2_file",
            PIPELINE_VERSION,
            actual_pipeline,
        )
        raise HTTPException(status_code=500, detail="技术合同审查未进入 canonical review-v2 pipeline")
    result["filename"] = filename
    result["extracted_text_length"] = len(contract_text)
    result["production_path_trace"] = {
        "entrypoint": "routers.review_v2.review_v2_file",
        "entrypoint_file": inspect.getsourcefile(review_v2_file) or __file__,
        "entrypoint_line": review_v2_file.__code__.co_firstlineno,
        "engine_function": "review_v2.engine.review_contract_v2",
        "engine_file": inspect.getsourcefile(review_contract_v2) or "",
        "engine_line": review_contract_v2.__code__.co_firstlineno,
        "process_id": os.getpid(),
        "evidence_implementation": TECHNICAL_CONSULTING_EVIDENCE_IMPLEMENTATION,
        "pipeline_id": PIPELINE_VERSION,
        "engine_version": (result.get("system_metadata") or {}).get("system_version"),
        "classifier_version": TECHNOLOGY_CLASSIFIER_VERSION,
        "evaluator_version": EVALUATOR_VERSION,
        "ruleset_version": RULE_VERSION,
        "runtime_sources": {
            "entrypoint": _runtime_source_identity(review_v2_file),
            "engine": _runtime_source_identity(review_contract_v2),
            "classifier": _runtime_source_identity(classify_technology_transaction),
            "registry": _runtime_source_identity(get_checklist_for_resolution),
        },
        "runtime_source_drift": source_drift,
        "document_hash": hashlib.sha256(contract_text.encode("utf-8")).hexdigest(),
        "cache": {"status": "MISS", "reason": "review-v2/file has no rule, evidence, or report result cache"},
        "input_visibility": {
            "normalized_text_length": len(contract_text),
            "confidentiality_terms_present": any(term in contract_text for term in (
                "保密", "商业秘密", "技术秘密", "不得泄露", "不得披露", "持续有效",
            )),
        },
        "classification_lifecycle": ((result.get("observability") or {}).get("classification_lifecycle") or {}),
        "classification_runtime": ((result.get("observability") or {}).get("classification_runtime") or {}),
        "registry_runtime": ((result.get("observability") or {}).get("review_registry_debug") or {}),
        "payment_risk_lifecycle": ((result.get("observability") or {}).get("payment_risk_lifecycle") or {}),
    }
    payment_events = result["production_path_trace"]["payment_risk_lifecycle"].get("events") or []
    for event in payment_events:
        fact_clause_ids = [row.get("clause_id") for row in event.get("source_clauses") or []]
        logger.warning(
            "%s review_id=%s filename=%s candidate_id=%s source_rule=%s source_clause_ids=%s "
            "scenario=%s applicability=%s status=%s filter_reason=%s invocations=%s arithmetic_status=%s "
            "percentage_total=%s amount_total=%s merged=%s finding_id=%s final_present=%s rendered=%s",
            event.get("event"),
            result.get("contract_review_trace_id"),
            filename,
            event.get("candidate_id", ""),
            event.get("source_rule", ""),
            event.get("source_clause_ids") or fact_clause_ids,
            event.get("scenario", ""),
            event.get("applicability", {}),
            event.get("status", ""),
            event.get("filter_reason") or event.get("reason", ""),
            event.get("invocation_count", ""),
            (event.get("output") or {}).get("status", ""),
            event.get("structured_percentages") or (event.get("output") or {}).get("percentage_total", ""),
            event.get("structured_amounts") or (event.get("output") or {}).get("amount_total", ""),
            event.get("merged", ""),
            event.get("finding_id", ""),
            event.get("present", ""),
            event.get("rendered", ""),
        )
    logger.warning(
        "[TRACE_REVIEW_V2] pid=%s pipeline=%s classifier=%s evaluator=%s ruleset=%s "
        "entry=%s:%s engine=%s:%s subtype=%s classifier_executions=%s registry_id=%s "
        "evidence_impl=%s doc_hash=%s confidentiality=%s",
        os.getpid(),
        PIPELINE_VERSION,
        TECHNOLOGY_CLASSIFIER_VERSION,
        EVALUATOR_VERSION,
        RULE_VERSION,
        result["production_path_trace"]["entrypoint_file"],
        result["production_path_trace"]["entrypoint_line"],
        result["production_path_trace"]["engine_file"],
        result["production_path_trace"]["engine_line"],
        ((result["production_path_trace"].get("classification_lifecycle") or {}).get("renderer") or {}).get("contract_subtype"),
        (result["production_path_trace"].get("classification_runtime") or {}).get("actual_execution_count"),
        (result["production_path_trace"].get("registry_runtime") or {}).get("registry_object_id"),
        TECHNICAL_CONSULTING_EVIDENCE_IMPLEMENTATION,
        result["production_path_trace"]["document_hash"],
        result["production_path_trace"]["input_visibility"]["confidentiality_terms_present"],
    )
    return result


@router.post("/export-pdf", summary="合同审查第二版：导出 PDF 报告")
async def review_v2_export_pdf(
    payload: ReviewV2PdfInput,
    current_user: dict = Depends(get_current_user),
):
    """将当前结构化审查报告导出为 PDF。"""

    try:
        content = _build_pdf(payload)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    base_name = (payload.filename or "合同审查报告").rsplit(".", 1)[0]
    download_name = f"{base_name}_审查第二版报告.pdf"
    encoded_name = quote(download_name)
    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_name}"},
    )


@router.post("/export-docx", summary="合同审查第二版：导出 DOCX 报告")
async def review_v2_export_docx(
    payload: ReviewV2PdfInput,
    current_user: dict = Depends(get_current_user),
):
    """将当前结构化审查报告导出为 Word 文档。"""

    try:
        content = _build_docx(payload)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    base_name = (payload.filename or "合同审查报告").rsplit(".", 1)[0]
    download_name = f"{base_name}_审查第二版报告.docx"
    encoded_name = quote(download_name)
    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_name}"},
    )
