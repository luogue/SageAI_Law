"""输出封装层：将校验后的报告打包成统一格式，添加元数据，支持多种输出格式"""
import json
import time
import io
import re
from typing import Dict, Optional, Tuple

class OutputLayer:
    SUPPORTED_FORMATS = ["md", "docx", "pdf"]
    
    def __init__(self):
        pass
    
    def _filter_internal_content(self, raw_report: str) -> str:
        """
        过滤报告中的内部推理步骤，只保留对最终用户有用的内容
        
        需要隐藏的内容：
        - 从"前置步骤：关键前提确认"到"4. 自检"之间的所有结构化推理步骤
        - "1. 自上而下建逻辑结构"整个章节
        - "2. 逻辑推理"整个章节
        - "3. 风险评估与商业可行性评估"章节
        - "4. 自检"章节
        - 任何"层级标注"、"模型补充"等内部标记
        
        用户应该看到：
        - 5. 结论与行动建议（包含结论、行动建议、声明）
        - 二、对用户额外要求的专项回应（如果有）
        """
        lines = raw_report.split('\n')
        filtered_lines = []
        skip_section = False
        
        # 内部推理章节的起始标记（支持markdown格式化版本）
        internal_section_starts = [
            r'^\*{0,2}\s*\d+\.\s*自上而下建逻辑结构',
            r'^\*{0,2}\s*\d+\.\s*逻辑推理(?!步骤)',
            r'^\*{0,2}\s*\d+\.\s*风险评估与商业可行性评估',
            r'^\*{0,2}\s*\d+\.\s*自检',
            r'^\*{0,2}\s*[iv]+\.\s*自检',
            r'^\*{0,2}\s*前置步骤',
            r'^\*{0,2}\s*【新增】信息来源',
            r'^\*{0,2}\s*【优化】风险/问题概率评估',
            r'^\*{0,2}\s*【新增】商业可行性评估',
            r'^\*{0,2}\s*###?\s*b\.\s*第二步',
            r'^\*{0,2}\s*####?\s*i\.\s*约束合规链式推理',
        ]
        
        # 用户可见章节的起始标记
        user_visible_starts = [
            r'^\*{0,2}\s*\d+\.\s*结论与行动建议',
            r'^\*{0,2}\s*##\s*二、对用户额外要求',
            r'^\*{0,2}\s*二、对用户额外要求',
        ]
        
        for line in lines:
            stripped = line.strip()
            
            # 检查是否是内部推理章节开头
            is_internal_start = False
            for pattern in internal_section_starts:
                if re.match(pattern, stripped, re.IGNORECASE):
                    is_internal_start = True
                    skip_section = True
                    break
            
            if is_internal_start:
                continue
            
            # 检查是否是用户可见章节开头（退出跳过模式）
            if skip_section:
                for pattern in user_visible_starts:
                    if re.match(pattern, stripped, re.IGNORECASE):
                        skip_section = False
                        break
            
            # 如果在跳过模式中，继续跳过
            if skip_section:
                continue
            
            # 过滤单行内部标记
            if '层级标注' in line or '模型补充' in line:
                continue
            
            filtered_lines.append(line)
        
        return '\n'.join(filtered_lines)
    
    def _generate_metadata(self, validation_result: Dict) -> Dict:
        """生成元数据"""
        metadata = {
            "timestamp": time.strftime("%Y-%m-%d %H:%M:%S", time.localtime()),
            "model_version": "qwen-plus",
            "confidence_score": validation_result.get('confidence_check', {}).get('confidence_score', 0),
            "validation_status": validation_result.get('status', 'unknown'),
            "has_warnings": validation_result.get('has_warnings', False),
            "warnings": validation_result.get('warnings', []),
            "format_check_status": validation_result.get('format_check', {}).get('status', 'unknown'),
            "confidence_check_status": validation_result.get('confidence_check', {}).get('status', 'unknown'),
            "source_check_status": validation_result.get('source_check', {}).get('status', 'unknown')
        }
        return metadata
    
    def _format_markdown(self, raw_report: str) -> str:
        """格式化Markdown输出"""
        # 过滤内部推理步骤，只保留对最终用户有用的内容
        filtered_report = self._filter_internal_content(raw_report)
        
        markdown = f"""# 合同审查报告

---

{filtered_report}

---

*报告生成时间：{time.strftime("%Y-%m-%d %H:%M:%S", time.localtime())}*
*模型版本：qwen-plus*
"""
        return markdown
    
    def generate_report_file(self, report_markdown: str, format: str = "md", filename: str = "合同审查报告") -> Tuple[bytes, str]:
        """
        根据格式生成报告文件
        
        :param report_markdown: Markdown格式的报告内容
        :param format: 输出格式 ("md", "docx", "pdf")
        :param filename: 文件名（不含扩展名）
        :return: (二进制文件内容, MIME类型)
        """
        if format == "md":
            return self._generate_markdown(report_markdown, filename)
        elif format == "docx":
            return self._generate_word(report_markdown, filename)
        elif format == "pdf":
            return self._generate_pdf(report_markdown, filename)
        else:
            raise ValueError(f"不支持的格式: {format}，支持的格式: {self.SUPPORTED_FORMATS}")
    
    def _generate_markdown(self, report_markdown: str, filename: str) -> Tuple[bytes, str]:
        """生成Markdown文件"""
        content = report_markdown.encode('utf-8')
        return content, "text/markdown"
    
    def _generate_word(self, report_markdown: str, filename: str) -> Tuple[bytes, str]:
        """将Markdown转换为Word文档"""
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        doc = Document()
        
        for line in report_markdown.split('\n'):
            line = line.rstrip()
            
            if not line:
                doc.add_paragraph()
                continue
            
            if line.startswith('# '):
                heading = doc.add_heading(line[2:], level=1)
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('#### '):
                doc.add_heading(line[5:], level=4)
            elif line.startswith('- ') or line.startswith('* '):
                paragraph = doc.add_paragraph(line[2:], style='List Bullet')
            elif line.startswith('| '):
                pass
            elif line.startswith('---'):
                pass
            elif line.startswith('*') and line.endswith('*'):
                paragraph = doc.add_paragraph(line[1:-1])
                for run in paragraph.runs:
                    run.italic = True
            else:
                formatted_line = self._parse_inline_formatting(line)
                paragraph = doc.add_paragraph()
                for text, bold in formatted_line:
                    run = paragraph.add_run(text)
                    run.bold = bold
        
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        content = buffer.getvalue()
        
        return content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    
    def _parse_inline_formatting(self, line: str) -> list:
        """解析行内格式（加粗等）"""
        result = []
        pattern = r'\*\*([^*]+)\*\*|\*([^*]+)\*'
        last_end = 0
        
        for match in re.finditer(pattern, line):
            if match.start() > last_end:
                result.append((line[last_end:match.start()], False))
            
            if match.group(1):
                result.append((match.group(1), True))
            elif match.group(2):
                result.append((match.group(2), False))
            
            last_end = match.end()
        
        if last_end < len(line):
            result.append((line[last_end:], False))
        
        if not result:
            result.append((line, False))
        
        return result
    
    def _generate_pdf(self, report_markdown: str, filename: str) -> Tuple[bytes, str]:
        """将Markdown转换为PDF"""
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=2*cm,
            leftMargin=2*cm,
            topMargin=2*cm,
            bottomMargin=2*cm
        )
        
        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            alignment=TA_CENTER,
            spaceAfter=20,
            spaceBefore=20
        )
        
        heading2_style = ParagraphStyle(
            'CustomHeading2',
            parent=styles['Heading2'],
            fontSize=14,
            spaceAfter=12,
            spaceBefore=12
        )
        
        heading3_style = ParagraphStyle(
            'CustomHeading3',
            parent=styles['Heading3'],
            fontSize=12,
            spaceAfter=8,
            spaceBefore=8
        )
        
        normal_style = styles['Normal']
        bullet_style = ParagraphStyle(
            'Bullet',
            parent=styles['Normal'],
            leftIndent=20,
            bulletIndent=10,
            spaceAfter=4
        )
        
        story = []
        
        for line in report_markdown.split('\n'):
            line = line.rstrip()
            
            if not line:
                story.append(Spacer(1, 6))
                continue
            
            if line.startswith('# '):
                story.append(Paragraph(line[2:], title_style))
            elif line.startswith('## '):
                story.append(Paragraph(line[3:], heading2_style))
            elif line.startswith('### '):
                story.append(Paragraph(line[4:], heading3_style))
            elif line.startswith('- ') or line.startswith('* '):
                text = line[2:]
                text = self._strip_markdown_formatting(text)
                story.append(Paragraph(f"• {text}", bullet_style))
            elif line.startswith('| '):
                pass
            elif line.startswith('---'):
                story.append(Spacer(1, 10))
            elif line.startswith('*') and line.endswith('*'):
                text = line[1:-1]
                text = self._strip_markdown_formatting(text)
                story.append(Paragraph(text, normal_style))
            else:
                text = self._strip_markdown_formatting(line)
                story.append(Paragraph(text, normal_style))
        
        doc.build(story)
        buffer.seek(0)
        content = buffer.getvalue()
        
        return content, "application/pdf"
    
    def _strip_markdown_formatting(self, text: str) -> str:
        """去除Markdown格式符号"""
        text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
        text = re.sub(r'\*([^*]+)\*', r'\1', text)
        text = re.sub(r'`([^`]+)`', r'\1', text)
        return text
    
    def generate_html_preview(self, report_markdown: str) -> str:
        """
        生成HTML预览版本
        
        :param report_markdown: Markdown格式的报告内容
        :return: HTML字符串
        """
        import markdown
        
        html_content = markdown.markdown(
            report_markdown,
            extensions=['tables', 'fenced_code', 'codehilite']
        )
        
        html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>合同审查报告预览</title>
    <style>
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
            max-width: 900px;
            margin: 0 auto;
            padding: 40px 20px;
            color: #333;
            line-height: 1.6;
        }}
        h1 {{
            text-align: center;
            color: #2c3e50;
            border-bottom: 2px solid #409EFF;
            padding-bottom: 10px;
        }}
        h2 {{
            color: #34495e;
            margin-top: 30px;
            border-left: 4px solid #409EFF;
            padding-left: 10px;
        }}
        h3 {{
            color: #546e7a;
        }}
        table {{
            border-collapse: collapse;
            width: 100%;
            margin: 20px 0;
        }}
        th, td {{
            border: 1px solid #ddd;
            padding: 12px;
            text-align: left;
        }}
        th {{
            background-color: #f5f7fa;
        }}
        code {{
            background-color: #f5f7fa;
            padding: 2px 6px;
            border-radius: 3px;
            font-family: Consolas, Monaco, monospace;
        }}
        pre {{
            background-color: #f5f7fa;
            padding: 15px;
            border-radius: 5px;
            overflow-x: auto;
        }}
        blockquote {{
            border-left: 4px solid #409EFF;
            margin: 20px 0;
            padding: 10px 20px;
            background-color: #f5f7fa;
        }}
        .warning {{
            background-color: #fff3cd;
            border: 1px solid #ffc107;
            padding: 15px;
            border-radius: 5px;
            margin: 20px 0;
        }}
        .footer {{
            text-align: center;
            color: #999;
            margin-top: 40px;
            padding-top: 20px;
            border-top: 1px solid #eee;
        }}
    </style>
</head>
<body>
    {html_content}
    <div class="footer">
        <p>报告生成时间：{time.strftime("%Y-%m-%d %H:%M:%S", time.localtime())}</p>
    </div>
</body>
</html>"""
        
        return html
    
    def package_output(self, raw_report: str, validation_result: Dict, input_data: Dict = None) -> Dict:
        """打包输出"""
        metadata = self._generate_metadata(validation_result)
        markdown_content = self._format_markdown(raw_report)
        
        # 如果有 input_data，将合同文本添加到 metadata 中（用于反馈功能）
        if input_data and 'contract_text' in input_data:
            metadata['contract_text'] = input_data['contract_text']
        
        key_info = self._extract_key_info(raw_report)
        
        output = {
            "status": "success",
            "metadata": metadata,
            "content": {
                "markdown": markdown_content,
                "html_preview": self.generate_html_preview(markdown_content),
                "summary": key_info.get('summary', ''),
                "risk_level": key_info.get('risk_level', '未知'),
                "risk_points": key_info.get('risk_points', []),
                "suggestions": key_info.get('suggestions', [])
            },
            "validation": {
                "has_warnings": validation_result.get('has_warnings', False),
                "warnings": validation_result.get('warnings', []),
                "confidence_score": validation_result.get('confidence_check', {}).get('confidence_score', 0)
            },
            "raw_report": raw_report
        }
        
        return output
    
    def _extract_key_info(self, report: str) -> Dict:
        """从报告中提取关键信息"""
        risk_level_pattern = r'风险等级[：:]?\s*(低|中|高|极高)'
        risk_level_match = re.search(risk_level_pattern, report)
        risk_level = risk_level_match.group(1) if risk_level_match else '未知'
        
        risk_pattern = r'[一二三四五六七八九十]+\.\s*([^\n]+)'
        risk_points = re.findall(risk_pattern, report)[:10]
        
        suggestion_pattern = r'(建议|修改建议)[：:]?\s*([^\n]+)'
        suggestions = re.findall(suggestion_pattern, report)[:5]
        suggestions = [s[1] if len(s) > 1 else s[0] for s in suggestions]
        
        conclusion_pattern = r'(结论与建议)[：:]?\s*([\s\S]*?)(?=\n### |$)'
        conclusion_match = re.search(conclusion_pattern, report)
        summary = conclusion_match.group(2).strip() if conclusion_match else ''
        
        return {
            "risk_level": risk_level,
            "risk_points": risk_points,
            "suggestions": suggestions,
            "summary": summary
        }
    
    def format_for_frontend(self, output: Dict) -> Dict:
        """格式化输出供前端使用"""
        frontend_output = {
            "success": True,
            "message": "审查完成",
            "result": output['content']['markdown'],
            "html_preview": output['content'].get('html_preview', ''),
            "metadata": output['metadata'],
            "validation": output['validation'],
            "summary": {
                "risk_level": output['content']['risk_level'],
                "risk_count": len(output['content']['risk_points']),
                "suggestion_count": len(output['content']['suggestions']),
                "confidence": output['validation']['confidence_score']
            }
        }
        
        return frontend_output